"""库存库位快照分组、筛选、数量格式和 Excel 导出测试。"""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from deeplinkerp_branding.services import inventory_detail_service as service
from deeplinkerp_branding.services.inventory_detail_service import (
	CATEGORY_DEFINITIONS,
	build_categorized_inventory_payload,
	build_categorized_inventory_xlsx,
	build_inventory_location_payload,
	build_inventory_location_xlsx,
	format_quantity,
	get_category_definition,
)

ROWS = [
	{
		"snapshot_key": "YWFM-2026-09-29",
		"snapshot_date": "2026-09-29",
		"company": "YUEWEI MX",
		"item_code": "FL002917",
		"item_name": "PET片材 / PET SHEET",
		"warehouse": "IML 仓库 - YWFM",
		"original_location": "AI-11-A02",
		"location_qty": 1865,
		"stock_uom": "张：hoja",
		"item_group": "FL Suministros Auxiliares辅料",
		"dpci": "",
		"external_code": "EXT-2917",
		"original_identifier_alias": "PET 2917",
	},
	{
		"snapshot_key": "YWFM-2026-09-29",
		"snapshot_date": "2026-09-29",
		"company": "YUEWEI MX",
		"item_code": "FL002917",
		"item_name": "PET片材 / PET SHEET",
		"warehouse": "IML 仓库 - YWFM",
		"original_location": "AI-12-T02",
		"location_qty": 695,
		"stock_uom": "张：hoja",
		"item_group": "FL Suministros Auxiliares辅料",
		"dpci": "",
		"external_code": "EXT-2917",
		"original_identifier_alias": "PET 2917",
	},
	{
		"snapshot_key": "YWFM-2026-09-29",
		"snapshot_date": "2026-09-29",
		"company": "YUEWEI MX",
		"item_code": "FL007979",
		"item_name": "色母粒 / MASTERBATCH AZUL",
		"warehouse": "综合仓库 - YWFM",
		"original_location": "AI-4-C01",
		"location_qty": 19.6,
		"stock_uom": "kg",
		"item_group": "FL Suministros Auxiliares辅料",
		"dpci": "",
		"external_code": "FL000164",
		"original_identifier_alias": "FL000164",
	},
]


def test_payload_groups_locations_and_keeps_complete_warehouse_total_when_location_filtered() -> None:
	payload = build_inventory_location_payload(
		ROWS,
		{
			"company": "YUEWEI MX",
			"original_location": "AI-11-A02",
		},
	)

	assert payload["snapshot_key"] == "YWFM-2026-09-29"
	assert len(payload["groups"]) == 1
	group = payload["groups"][0]
	assert group["item_code"] == "FL002917"
	assert group["total_qty"] == 2560
	assert group["locations"] == [{"original_location": "AI-11-A02", "location_qty": 1865}]


def test_payload_filters_keyword_warehouse_and_item_group_before_grouping() -> None:
	payload = build_inventory_location_payload(
		ROWS,
		{
			"keyword": "masterbatch",
			"warehouse": "综合仓库 - YWFM",
			"item_group": "FL Suministros Auxiliares辅料",
		},
	)

	assert [row["item_code"] for row in payload["groups"]] == ["FL007979"]
	assert payload["groups"][0]["total_qty"] == 19.6


def test_quantity_format_uses_integer_count_units_and_two_decimal_continuous_units() -> None:
	assert format_quantity(2560, "张：hoja") == "2,560"
	assert format_quantity(1.5, "个：pieza") == "1.50"
	assert format_quantity(19.6, "kg") == "19.60"
	assert format_quantity(19, "kg") == "19.00"


def test_company_document_permission_is_checked_before_raw_snapshot_queries(monkeypatch) -> None:
	checked: list[str] = []

	class FakeCompany:
		@staticmethod
		def check_permission(permission_type: str) -> None:
			checked.append(permission_type)

	class FakeFrappe:
		@staticmethod
		def get_doc(doctype: str, name: str):
			assert (doctype, name) == (
				"Company",
				"YUEWEI MX",
			)
			return FakeCompany()

	monkeypatch.setattr(service, "frappe", FakeFrappe())
	service._require_company_permission("YUEWEI MX")

	assert checked == ["read"]


def test_snapshot_sql_is_limited_to_permission_filtered_items_and_warehouses(monkeypatch) -> None:
	queries: list[tuple[str, dict]] = []

	class FakeDB:
		@staticmethod
		def sql(query: str, values: dict, as_dict: bool):
			assert as_dict is True
			queries.append((query, values))
			return []

	class FakeFrappe:
		db = FakeDB()

	monkeypatch.setattr(service, "frappe", FakeFrappe())
	monkeypatch.setattr(service, "_accessible_items", lambda: ("ITEM-ALLOWED",))
	monkeypatch.setattr(
		service,
		"_accessible_warehouses",
		lambda company: ("WAREHOUSE-ALLOWED",) if company == "ACME" else (),
	)
	service._load_snapshot_rows({"company": "ACME", "snapshot_key": "SNAPSHOT-1"})

	query, values = queries[0]
	assert "snapshot.item_code IN %(items)s" in query
	assert "snapshot.warehouse IN %(warehouses)s" in query
	assert values["items"] == ("ITEM-ALLOWED",)
	assert values["warehouses"] == ("WAREHOUSE-ALLOWED",)


def test_category_sql_is_limited_to_permission_filtered_items(monkeypatch) -> None:
	queries: list[tuple[str, dict]] = []

	class FakeDB:
		@staticmethod
		def sql(query: str, values: dict, as_dict: bool):
			assert as_dict is True
			queries.append((query, values))
			return []

	class FakeFrappe:
		db = FakeDB()

	monkeypatch.setattr(service, "frappe", FakeFrappe())
	context = {
		"company": "ACME",
		"warehouses": ("WAREHOUSE-ALLOWED",),
		"items": ("ITEM-ALLOWED",),
		"group_lft": 1,
		"group_rgt": 2,
	}
	service._load_category_stock_rows(context)

	query, values = queries[0]
	assert "item.name IN %(items)s" in query
	assert values["items"] == ("ITEM-ALLOWED",)


def test_xlsx_export_merges_group_fields_but_keeps_location_rows_separate() -> None:
	payload = build_inventory_location_payload(ROWS, {})
	workbook = load_workbook(BytesIO(build_inventory_location_xlsx(payload)))
	sheet = workbook.active

	assert sheet.title == "物料库存明细"
	assert [sheet.cell(1, column).value for column in range(1, 12)] == [
		"正式物料编码",
		"物料名称（双语）",
		"仓库",
		"原始库位",
		"库位数量",
		"总库存",
		"库存单位",
		"物料组",
		"DPCI",
		"外部编码",
		"原始标识/别名",
	]
	# FL002917 has two locations; shared fields are true merged cells.
	merged = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
	for column in ("A", "B", "C", "F", "G", "H", "I", "J", "K"):
		assert f"{column}2:{column}3" in merged
	assert sheet["D2"].value == "AI-11-A02"
	assert sheet["D3"].value == "AI-12-T02"
	assert sheet["E2"].value == 1865
	assert sheet["E3"].value == 695
	assert sheet["F2"].value == 2560
	assert sheet["E2"].number_format == "#,##0"
	assert sheet["E4"].number_format == "#,##0.00"


CATEGORY_STOCK_ROWS = [
	{
		"item_code": "NSEMI-001",
		"item_name": "注塑半成品",
		"warehouse": "半成品仓 - YWFM",
		"actual_qty": 12,
		"stock_uom": "个：pieza",
		"item_group": "注塑半成品",
		"dpci": "D-001",
		"external_code": "EXT-S1",
		"original_identifier_alias": "SEMI-1",
	},
	{
		"item_code": "NSEMI-001",
		"item_name": "注塑半成品",
		"warehouse": "待检仓 - YWFM",
		"actual_qty": -2,
		"stock_uom": "个：pieza",
		"item_group": "注塑半成品",
		"dpci": "D-001",
		"external_code": "EXT-S1",
		"original_identifier_alias": "SEMI-1",
	},
	{
		"item_code": "NSEMI-002",
		"item_name": "喷油半成品",
		"warehouse": "",
		"actual_qty": 0,
		"stock_uom": "个：pieza",
		"item_group": "喷油半成品",
		"dpci": "",
		"external_code": "",
		"original_identifier_alias": "",
	},
]

CATEGORY_SNAPSHOT_ROWS = [
	{
		**CATEGORY_STOCK_ROWS[0],
		"snapshot_key": "YWFM-2026-09-29",
		"snapshot_date": "2026-09-29",
		"original_location": "A-01",
		"location_qty": 5,
	},
	{
		**CATEGORY_STOCK_ROWS[0],
		"snapshot_key": "YWFM-2026-09-29",
		"snapshot_date": "2026-09-29",
		"original_location": "A-02",
		"location_qty": 5,
	},
]


def test_category_definitions_are_server_whitelisted_and_use_erp_item_group_roots() -> None:
	assert set(CATEGORY_DEFINITIONS) == {"semi_finished", "finished_goods", "mold"}
	assert get_category_definition("semi_finished")["root_item_group"] == "半成品Semiterminado"
	assert get_category_definition("finished_goods")["root_item_group"] == "成品Producto terminado"
	assert get_category_definition("mold")["root_item_group"] == "模具Moldes"

	try:
		get_category_definition("N-prefix")
	except ValueError as error:
		assert "不支持的库存分类" in str(error)
	else:  # pragma: no cover - invalid category must never be accepted
		raise AssertionError("invalid category was accepted")


def test_missing_category_root_returns_an_empty_page_instead_of_a_server_error(monkeypatch) -> None:
	missing_group = "半成品Semiterminado"
	monkeypatch.setattr(service, "_require_categorized_inventory_read_permission", lambda: None)
	monkeypatch.setattr(service, "_require_company_permission", lambda _company: None)
	monkeypatch.setattr(
		service,
		"_category_query_context",
		lambda _category, _filters: {
			"company": "Yuewei",
			"missing_item_group": missing_group,
			"warehouses": (),
			"snapshot_key": "",
		},
	)
	monkeypatch.setattr(
		service,
		"_load_category_stock_rows",
		lambda _context: (_ for _ in ()).throw(AssertionError("stock query should be skipped")),
	)
	monkeypatch.setattr(
		service,
		"_load_category_snapshot_rows",
		lambda _context: (_ for _ in ()).throw(AssertionError("snapshot query should be skipped")),
	)

	payload = service.get_categorized_inventory_detail(
		"semi_finished",
		filters={"company": "Yuewei"},
		start=0,
		page_length=100,
	)

	assert payload["total_count"] == 0
	assert payload["groups"] == []
	assert payload["item_group_options"] == []
	assert payload["warning"] == f"ERP 未维护分类物料组：{missing_group}"


def test_categorized_payload_keeps_real_time_qty_and_marks_snapshot_difference_and_missing_location() -> None:
	payload = build_categorized_inventory_payload(
		CATEGORY_STOCK_ROWS,
		CATEGORY_SNAPSHOT_ROWS,
		category="semi_finished",
		filters={},
		start=0,
		page_length=100,
	)

	assert payload["total_count"] == 3
	assert payload["page_count"] == 3
	assert payload["has_next"] is False
	groups = {(row["item_code"], row["warehouse"]): row for row in payload["groups"]}

	current = groups[("NSEMI-001", "半成品仓 - YWFM")]
	assert current["actual_qty"] == 12
	assert current["snapshot_qty"] == 10
	assert current["difference_qty"] == 2
	assert current["inventory_status"] == "库存差异"
	assert current["locations"] == [
		{"reference_location": "A-01", "snapshot_location_qty": 5, "snapshot_date": "2026-09-29"},
		{"reference_location": "A-02", "snapshot_location_qty": 5, "snapshot_date": "2026-09-29"},
	]

	missing = groups[("NSEMI-002", "")]
	assert missing["actual_qty"] == 0
	assert missing["snapshot_qty"] is None
	assert missing["difference_qty"] is None
	assert missing["inventory_status"] == "库位待维护"
	assert missing["locations"][0]["reference_location"] == "库位待维护"


def test_categorized_payload_filters_locations_stock_and_paginates_groups() -> None:
	location_payload = build_categorized_inventory_payload(
		CATEGORY_STOCK_ROWS,
		CATEGORY_SNAPSHOT_ROWS,
		category="semi_finished",
		filters={"reference_location": "A-02"},
		start=0,
		page_length=100,
	)
	assert location_payload["total_count"] == 1
	assert location_payload["groups"][0]["actual_qty"] == 12
	assert location_payload["groups"][0]["locations"] == [
		{"reference_location": "A-02", "snapshot_location_qty": 5, "snapshot_date": "2026-09-29"}
	]

	stock_payload = build_categorized_inventory_payload(
		CATEGORY_STOCK_ROWS,
		CATEGORY_SNAPSHOT_ROWS,
		category="semi_finished",
		filters={"only_with_stock": 1},
		start=1,
		page_length=1,
	)
	assert stock_payload["total_count"] == 2
	assert stock_payload["page_count"] == 1
	assert stock_payload["has_previous"] is True
	assert stock_payload["has_next"] is False
	assert stock_payload["groups"][0]["actual_qty"] == -2


def test_selected_parent_item_group_keeps_rows_already_limited_to_its_descendants() -> None:
	payload = build_categorized_inventory_payload(
		CATEGORY_STOCK_ROWS,
		CATEGORY_SNAPSHOT_ROWS,
		category="semi_finished",
		filters={"item_group": "半成品Semiterminado"},
		start=0,
		page_length=100,
	)

	assert payload["total_count"] == 3


def test_category_pagination_parsing_falls_back_safely() -> None:
	assert service._parse_start("not-a-number") == 0
	assert service._parse_start(-10) == 0
	assert service._parse_start("200") == 200
	assert service.parse_page_length(None) == 200
	assert service.parse_page_length(500) == 500
	with pytest.raises(ValueError):
		service.parse_page_length("invalid")


def test_categorized_inventory_checks_snapshot_item_bin_and_warehouse_permissions(monkeypatch) -> None:
	checked: list[tuple[str, str]] = []

	class FakeFrappe:
		PermissionError = PermissionError

		@staticmethod
		def has_permission(doctype: str, permission_type: str) -> bool:
			checked.append((doctype, permission_type))
			return True

	monkeypatch.setattr(service, "frappe", FakeFrappe())
	service._require_categorized_inventory_read_permission()

	assert checked == [
		(service.SNAPSHOT_DOCTYPE, "read"),
		("Item", "read"),
		("Bin", "read"),
		("Warehouse", "read"),
	]


def test_real_time_query_uses_bin_and_item_group_bounds_not_code_prefixes(monkeypatch) -> None:
	captured: list[tuple[str, dict]] = []

	class FakeDB:
		@staticmethod
		def sql(query: str, params: dict, as_dict: bool):
			assert as_dict is True
			captured.append((query, params))
			return []

	monkeypatch.setattr(service, "frappe", SimpleNamespace(db=FakeDB()))
	context = {
		"company": "YUEWEI MX",
		"group_lft": 10,
		"group_rgt": 20,
		"warehouses": ("半成品仓 - YWFM",),
		"items": ("NSEMI-001",),
		"snapshot_key": "YWFM-2026-09-29",
	}
	service._load_category_stock_rows(context)

	query, params = captured[0]
	assert "`tabBin`" in query and "bin.actual_qty" in query
	assert "item.disabled = 0" in query and "item.is_stock_item = 1" in query
	assert "item_group.lft >= %(group_lft)s" in query
	assert "item_group.rgt <= %(group_rgt)s" in query
	assert "LIKE 'M%'" not in query and "LIKE 'N%'" not in query
	assert params["warehouses"] == ("半成品仓 - YWFM",)


def test_snapshot_only_warehouse_replaces_synthetic_zero_warehouse() -> None:
	snapshot = {
		**CATEGORY_STOCK_ROWS[2],
		"warehouse": "半成品仓 - YWFM",
		"snapshot_key": "YWFM-2026-09-29",
		"snapshot_date": "2026-09-29",
		"original_location": "B-01",
		"location_qty": 0,
	}
	payload = build_categorized_inventory_payload(
		[CATEGORY_STOCK_ROWS[2]],
		[snapshot],
		category="semi_finished",
		filters={},
		start=0,
		page_length=100,
	)

	assert [(row["warehouse"], row["actual_qty"]) for row in payload["groups"]] == [("半成品仓 - YWFM", 0)]
	assert payload["groups"][0]["inventory_status"] == "库存一致"


def test_categorized_xlsx_exports_all_reference_locations_and_real_time_status() -> None:
	payload = build_categorized_inventory_payload(
		CATEGORY_STOCK_ROWS,
		CATEGORY_SNAPSHOT_ROWS,
		category="semi_finished",
		filters={},
		start=0,
		page_length=None,
	)
	workbook = load_workbook(BytesIO(build_categorized_inventory_xlsx(payload)))
	sheet = workbook.active

	assert sheet.title == "半成品库存明细"
	assert [sheet.cell(1, column).value for column in range(1, 15)] == [
		"正式物料编码",
		"物料名称（双语）",
		"仓库",
		"参考库位",
		"快照库位数量",
		"实时库存",
		"库存差异",
		"库存状态",
		"快照日期",
		"库存单位",
		"物料组",
		"DPCI",
		"外部编码",
		"原始标识/别名",
	]
	assert sheet["D2"].value == "A-01"
	assert sheet["D3"].value == "A-02"
	assert sheet["F2"].value == 12
	assert sheet["G2"].value == 2
	assert sheet["H2"].value == "库存差异"
