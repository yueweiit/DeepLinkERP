"""Optional link privacy and availability contracts, without a database."""
import unittest
from datetime import date, datetime
from io import BytesIO
from unittest.mock import patch
from types import SimpleNamespace

from openpyxl import load_workbook

import frappe
import frappe.permissions
from deeplinkerp_branding.services import purchase_payment_service as service
from deeplinkerp_branding.services import purchase_document_actions as actions


def native(doctype, name, **changes):
    values = {"doctype": doctype, "name": name, "company": "C", "supplier": "S", "currency": "CNY", "grand_total": 100,
              "docstatus": 1, "status": "To Receive and Bill", "modified": "v1", "per_billed": 0, "items": [], **changes}
    return SimpleNamespace(**values, get=lambda field, default=None: values.get(field, default), check_permission=lambda permission: None)


class PurchaseLinkTests(unittest.TestCase):
    def test_procurement_payment_rejects_order_sources_before_reads(self):
        def reject(message, exception=frappe.ValidationError):
            raise exception(message)
        with patch.object(frappe, 'throw', side_effect=reject), patch.object(service, '_source', side_effect=AssertionError('PO source read before rejection')) as read:
            for invoke in (lambda: service.payment_target('Purchase Order', 'PO'),
                           lambda: service.create_payment_draft.__wrapped__.__wrapped__('Purchase Order', 'PO'),
                           lambda: service.preview_payment_batch.__wrapped__.__wrapped__('Purchase Order', [{'name': 'PO'}])):
                with self.subTest(invoke=invoke), self.assertRaises(frappe.ValidationError):
                    invoke()
            read.assert_not_called()

    def test_receipt_eligibility_uses_projected_order_state_without_reads(self):
        for changes, allowed in (({}, True), ({'docstatus': 0}, False), ({'status': 'Closed'}, False),
                                 ({'status': 'Completed'}, False), ({'per_received': 100}, False),
                                 ({'per_received': 'NaN'}, False), ({'per_received': 'invalid'}, False)):
            with self.subTest(changes=changes):
                result = service.receipt_eligibility({'docstatus': 1, 'status': 'To Receive and Bill', 'per_received': 0, **changes}, can_create=True)
                self.assertEqual(result['allowed'], allowed)
                self.assertEqual(bool(result['reason']), not allowed)
        self.assertFalse(service.receipt_eligibility({'docstatus': 1, 'per_received': 0}, can_create=False)['allowed'])
        self.assertFalse(service.receipt_eligibility({'docstatus': 1, 'per_received': 0}, can_create=True, reversal_pending=True)['allowed'])

    def setUp(self):
        self.flags_patch = patch.object(frappe, 'flags', frappe._dict())
        self.flags_patch.start()
        self.addCleanup(self.flags_patch.stop)
        query = patch.object(service, '_query_fields', return_value=set(service.RECEIPT_FIELDS) | {'owner', 'creation', 'modified'})
        query.start()
        self.addCleanup(query.stop)
        normalize = patch.object(service, '_normalize_filter', side_effect=lambda dt, field, operator, value: [dt, field, operator, value])
        normalize.start()
        self.addCleanup(normalize.stop)

    def test_caught_framework_throw_does_not_queue_private_messages(self):
        frappe.local.message_log = []
        frappe.flags.mute_messages = False
        def private_error(*args):
            frappe.throw('private-target does not exist', frappe.DoesNotExistError)
        with patch.object(service, '_read', side_effect=private_error):
            warnings = []
            self.assertIsNone(service._related('Purchase Order', 'private-target', warnings))
        self.assertFalse(frappe.local.message_log)
        self.assertFalse(frappe.flags.mute_messages)
        self.assertEqual(warnings, [service.LINK_WARNING])

    def test_message_flag_is_restored_after_unexpected_error(self):
        frappe.flags.mute_messages = False
        with patch.object(service, '_read', side_effect=RuntimeError):
            with self.assertRaises(RuntimeError):
                service._related('Purchase Order', 'target', [])
        self.assertFalse(frappe.flags.mute_messages)

    def test_missing_and_denied_targets_have_the_same_public_result(self):
        for error in (frappe.DoesNotExistError, frappe.PermissionError):
            with self.subTest(error=error), patch.object(service, '_read', side_effect=error):
                warnings = []
                self.assertIsNone(service._related('Purchase Order', 'private-target', warnings))
                self.assertEqual(warnings, [service.LINK_WARNING])
                self.assertNotIn('private-target', str(warnings))

    def test_cross_company_and_supplier_links_are_not_returned(self):
        for doc in (frappe._dict(company='Other', supplier='S'), frappe._dict(company='C', supplier='Other')):
            with self.subTest(doc=doc), patch.object(service, '_read', return_value=doc):
                warnings = []
                self.assertIsNone(service._related('Purchase Order', 'private-target', warnings, company='C', supplier='S'))
                self.assertEqual(warnings, [service.LINK_WARNING])

    def test_available_link_retains_the_existing_document(self):
        doc = frappe._dict(company='C', supplier='S')
        with patch.object(service, '_read', return_value=doc):
            warnings = []
            self.assertIs(service._related('Purchase Order', 'valid', warnings, company='C', supplier='S'), doc)
            self.assertFalse(warnings)

    def test_programming_and_amount_errors_are_not_swallowed_as_missing_links(self):
        with patch.object(service, '_read', side_effect=ValueError('invalid amount')):
            with self.assertRaises(ValueError):
                service._related('Purchase Order', 'target', [])

    def test_mixed_links_keep_readable_targets_and_one_generic_warning(self):
        source = SimpleNamespace(company='C', supplier='S', items=[
            frappe._dict(purchase_order=n) for n in ('valid', 'missing', 'denied', 'valid')])
        def read(doctype, name, fields):
            if name == 'missing':
                raise frappe.DoesNotExistError
            if name == 'denied':
                raise frappe.PermissionError
            return frappe._dict(company='C', supplier='S')
        with patch.object(service, '_read', side_effect=read):
            warnings = []
            self.assertEqual(service._source_links(source, 'Purchase Order', 'purchase_order', warnings), ['valid'])
            self.assertEqual(warnings, [service.LINK_WARNING])

    def test_unreadable_receipt_row_does_not_break_other_rows_or_show_balance(self):
        rows = [frappe._dict(name=name, company='C', currency='CNY', grand_total=10) for name in ('unreadable', 'visible')]
        def chain(doctype, name, include_payments):
            if name == 'unreadable':
                raise frappe.PermissionError
            return {'orders': [], 'balances': [{'outstanding': 10, 'settled': 2}],
                    'can_create': True, 'can_create_invoice': False, 'draft_invoices': [], 'draft_orders': [], 'reason': '', 'warnings': [], 'incomplete_links': False, 'invoices': []}
        with patch.object(service, '_require_fields'), patch.object(service._RecordReader, 'preload'), patch.object(service, '_read'), patch.object(frappe, 'get_list', return_value=rows), patch.object(service, 'get_purchase_chain', side_effect=chain):
            result = service.get_receipt_list(page_length=2)
        self.assertEqual(result['total_count'], 2)
        self.assertFalse(result['rows'][0]['can_create'])
        self.assertEqual(result['rows'][0]['balances'], [])
        self.assertTrue(result['rows'][1]['can_create'])

    def test_denied_receipt_source_is_excluded_from_count_totals_and_rows(self):
        rows = [frappe._dict(name='private', company='C', currency='CNY', grand_total=100)]
        with patch.object(service, '_require_fields'), patch.object(service._RecordReader, 'preload'), patch.object(service, '_read', side_effect=frappe.PermissionError), patch.object(frappe, 'get_list', return_value=rows):
            result = service.get_receipt_list()
        self.assertEqual(result, {'rows': [], 'total_count': 0, 'totals': []})


class PurchasePayableScopeTests(unittest.TestCase):
    def setUp(self):
        flags = patch.object(frappe, 'flags', frappe._dict())
        flags.start()
        self.addCleanup(flags.stop)

    def invoice(self, name='PI', **changes):
        values = dict(name=name, docstatus=1, is_return=0, company='C', supplier='S',
                      currency='CNY', party_account_currency='CNY', grand_total=1200,
                      rounded_total=0, disable_rounded_total=1, outstanding_amount=300,
                      posting_date='2026-10-05', status='Partly Paid',
                      items=[frappe._dict(purchase_order='PO', purchase_receipt='PR')])
        values.update(changes)
        return SimpleNamespace(**values, get=lambda field: values.get(field), invoice_is_blocked=lambda: False)

    def scope(self, docs, sources, **args):
        with patch.object(service, '_require_fields'), patch.object(service, '_read', side_effect=lambda dt, name, fields: docs[name]), \
             patch.object(service._RecordReader, 'preload'), \
             patch.object(service, '_source_links', side_effect=lambda doc, dt, field, warnings: sources.get((doc.name, dt), [])), \
             patch.object(frappe, 'has_permission', return_value=True), \
             patch.object(frappe, 'get_list', return_value=[frappe._dict(name=name) for name in [*docs, *docs]]):
            return service.get_purchase_payables(**args)

    def test_po_only_pr_only_both_share_one_native_invoice_scope_contract(self):
        for orders, receipts in ((['PO'], []), ([], ['PR']), (['PO'], ['PR'])):
            with self.subTest(orders=orders, receipts=receipts):
                row = self.scope({'PI': self.invoice()}, {('PI', 'Purchase Order'): orders, ('PI', 'Purchase Receipt'): receipts})['rows'][0]
                self.assertEqual(row['orders'], orders)
                self.assertEqual(row['receipts'], receipts)
                self.assertEqual(row['outstanding'], 300)

    def test_operating_invoice_and_invoice_without_readable_source_do_not_enter_count_or_paging(self):
        docs = {name: self.invoice(name) for name in ('OPERATING', 'HIDDEN-SOURCE', 'VISIBLE')}
        result = self.scope(docs, {('VISIBLE', 'Purchase Receipt'): ['PR']}, start=0, page_length=1)
        self.assertEqual(result['total_count'], 1)
        self.assertEqual([row['name'] for row in result['rows']], ['VISIBLE'])
        self.assertNotIn('HIDDEN', str(result))

    def test_mixed_invoice_keeps_whole_native_total_without_inventing_procurement_amount(self):
        doc = self.invoice(items=[frappe._dict(purchase_order='PO'), frappe._dict(purchase_order=None, purchase_receipt=None)])
        result = self.scope({'PI': doc}, {('PI', 'Purchase Order'): ['PO']})
        self.assertEqual(len(result['rows']), 1)
        self.assertTrue(result['rows'][0]['shared'])
        self.assertEqual(result['rows'][0]['grand_total'], 1200)
        self.assertEqual(result['rows'][0]['outstanding'], 300)
        self.assertNotIn('procurement_amount', result['rows'][0])
        self.assertIn('整张应付单', result['notice'])

    def test_draft_cancelled_and_return_keep_native_state_and_disable_quick_payment(self):
        for status, is_return in ((0, 0), (2, 0), (1, 1)):
            with self.subTest(status=status, is_return=is_return):
                row = self.scope({'PI': self.invoice(docstatus=status, is_return=is_return)}, {('PI', 'Purchase Order'): ['PO']})['rows'][0]
                self.assertEqual(row['docstatus'], status)
                self.assertEqual(row['is_return'], bool(is_return))
                self.assertFalse(row['can_pay'])
                if status != 1:
                    self.assertNotIn('outstanding', row)

    def test_native_doctype_permission_and_field_rejection_are_not_suppressed(self):
        with patch.object(frappe, 'has_permission', return_value=True), patch.object(service, '_require_fields', side_effect=frappe.PermissionError):
            with self.assertRaises(frappe.PermissionError):
                service.get_purchase_payables()
        with patch.object(frappe, 'has_permission', return_value=True), patch.object(service, '_require_fields'), patch.object(frappe, 'get_list', side_effect=frappe.PermissionError):
            with self.assertRaises(frappe.PermissionError):
                service.get_purchase_payables()

    def test_payable_totals_use_all_readable_filtered_invoices_before_pagination_and_account_currency(self):
        docs = {'PI-1': self.invoice('PI-1'),
                'PI-2': self.invoice('PI-2', currency='USD', party_account_currency='CNY', base_grand_total=2400, outstanding_amount=400),
                'DRAFT': self.invoice('DRAFT', docstatus=0),
                'RETURN': self.invoice('RETURN', is_return=1),
                'PRIVATE': self.invoice('PRIVATE')}
        sources = {(name, 'Purchase Receipt'): ['PR'] for name in docs if name != 'PRIVATE'}
        with patch.object(frappe, 'get_cached_value', return_value='CNY'):
            result = self.scope(docs, sources, page_length=1)
        self.assertEqual(len(result['rows']), 1)
        self.assertEqual(result['totals'], [{'currency': 'CNY', 'total': 3600, 'settled': 2900, 'outstanding': 700}])
        self.assertEqual(result['total_count'], 4)
        self.assertNotIn('PRIVATE', str(result))

    def test_payable_export_uses_all_authorized_search_matches_beyond_page_and_skips_missing_sources(self):
        docs = {f'PI-{index}': self.invoice(f'PI-{index}') for index in range(120)}
        docs.update(OTHER=self.invoice('OTHER'), PRIVATE=self.invoice('PRIVATE'))
        sources = {(name, 'Purchase Receipt'): ['PR'] for name in docs if name != 'PRIVATE'}
        with patch.object(service, '_export', return_value=None) as export:
            result = self.scope(docs, sources, search='PI-', start=50, page_length=20,
                                export_format='xlsx', columns=['name', 'total', 'settled', 'outstanding'])
        self.assertIsNone(result)
        self.assertEqual(export.call_args.args[0], 'Purchase Invoice')
        rows = export.call_args.args[1]
        self.assertEqual(len(rows), 120)
        self.assertEqual(len({row['name'] for row in rows}), 120)
        self.assertNotIn('OTHER', str(rows))
        self.assertNotIn('PRIVATE', str(rows))


class CrossborderInvoiceChainTests(unittest.TestCase):
    def setUp(self):
        flags = patch.object(frappe, "flags", frappe._dict())
        flags.start(); self.addCleanup(flags.stop)
        fields = patch.object(service, "_require_fields")
        fields.start(); self.addCleanup(fields.stop)

    def test_later_receipt_discovers_earlier_order_invoice_via_explicit_po_links(self):
        receipt = native("Purchase Receipt", "PR", items=[frappe._dict(purchase_order="PO")])
        with patch.object(service, "_source", return_value=receipt), patch.object(service, "_source_links", return_value=["PO"]), \
             patch.object(frappe, "get_all", side_effect=[["PR-PI"], ["PO-PI", "PR-PI"]]) as query:
            self.assertEqual(service._invoice_names("Purchase Receipt", "PR"), ["PR-PI", "PO-PI"])
        self.assertEqual(query.call_args_list[1].kwargs["filters"], {"purchase_order": ["in", ["PO"]]})

    def test_unreadable_order_link_remains_a_warning_during_invoice_discovery(self):
        receipt = native("Purchase Receipt", "PR", items=[frappe._dict(purchase_order="PRIVATE")])
        def links(doc, dt, field, warnings):
            warnings.append(service.LINK_WARNING)
            return []
        with patch.object(service, "_source", return_value=receipt), patch.object(service, "_source_links", side_effect=links), \
             patch.object(frappe, "get_all", return_value=[]):
            warnings = []
            self.assertEqual(service._invoice_names("Purchase Receipt", "PR", warnings=warnings), [])
        self.assertEqual(warnings, [service.LINK_WARNING])

    def test_order_invoice_capability_keeps_historical_drafts_without_enabling_procurement_actions(self):
        order = native("Purchase Order", "PO", per_billed=100, items=[frappe._dict(name="A", qty=10)])
        for names in ([], ["PI-DRAFT"]):
            with self.subTest(names=names), patch.object(service, "_source", return_value=order), \
                 patch.object(service, "_invoice_names", return_value=names), \
                 patch.object(service, "_related", return_value=native("Purchase Invoice", "PI-DRAFT", docstatus=0)), \
                 patch.object(service, "_invoice_row", return_value={"name": "PI-DRAFT", "docstatus": 0, "can_pay": False, "shared": False}), \
                 patch.object(service, "_order_progress", return_value=[]), \
                 patch.object(frappe, "has_permission", side_effect=lambda dt, perm: dt != "Payment Entry"), \
                 patch.object(frappe, "db", SimpleNamespace(get_values=lambda *args, **kwargs: [], get_single_value=lambda *args: "No")), \
                 patch.object(actions, "_mapping_fields"), patch.object(service, "order_execution_reason", return_value=""), \
                 patch.object(actions, "_fields"), patch.object(actions, "_native") as mapper:
                chain = service.get_purchase_chain("Purchase Order", "PO", include_payments=False)
            self.assertFalse(chain["can_create_invoice"])
            self.assertFalse(chain["can_create"])
            self.assertEqual(chain["draft_invoices"], [{"name": name, "docstatus": 0} for name in names])
            self.assertEqual(chain["invoice_reason"], "已有应付草稿，请选择继续编辑" if names else "")
            self.assertEqual(chain["source_modified"], "v1")
            mapper.assert_not_called()

    def test_direct_order_invoice_on_receipt_is_shared_whole_balance(self):
        invoice = native("Purchase Invoice", "PI", is_return=0, party_account_currency="CNY", disable_rounded_total=1,
            rounded_total=0, outstanding_amount=70, items=[frappe._dict(purchase_order="PO", purchase_receipt=None)])
        invoice.invoice_is_blocked = lambda: False
        row = service._invoice_row(invoice, "Purchase Receipt", "PR")
        self.assertTrue(row["shared"])
        self.assertEqual(row["outstanding"], 70)
        self.assertEqual(row["scope_label"], "共享应付整单余额")
        self.assertNotIn("receipt_paid", row)

    def test_completed_order_allows_existing_invoice_payment_but_closed_order_does_not(self):
        invoice = native("Purchase Invoice", "PI")
        invoice_row = {"name": "PI", "docstatus": 1, "is_return": False, "shared": True, "can_pay": True,
                       "currency": "CNY", "total": 70, "settled": 25, "outstanding": 45}
        for status in ("Closed", "Completed"):
            order = native("Purchase Order", "PO", status=status, per_billed=100)
            for source_type in ("Purchase Order", "Purchase Receipt"):
                source = order if source_type == "Purchase Order" else native("Purchase Receipt", "PR")
                with self.subTest(status=status, source_type=source_type), \
                     patch.object(service, "_source", return_value=source), patch.object(service, "_invoice_names", return_value=["PI"]), \
                     patch.object(service, "_related", side_effect=lambda dt, *args: order if dt == "Purchase Order" else invoice), \
                     patch.object(service, "_invoice_row", return_value=invoice_row), \
                     patch.object(service, "_source_links", side_effect=lambda doc, *args: ["PO"] if doc.doctype == "Purchase Receipt" else []), \
                     patch.object(service, "_order_progress", return_value=[]), \
                     patch.object(actions, "_mapping_fields"), patch.object(frappe, "has_permission", return_value=True):
                    chain = service.get_purchase_chain(source_type, source.name, include_payments=False)
                    self.assertEqual(chain["can_create"], status == "Completed" and source_type == "Purchase Receipt")
                    self.assertEqual(chain["balances"], [{"currency": "CNY", "total": 70, "settled": 25, "outstanding": 45}])
                    self.assertFalse(chain["can_create_invoice"])
                    self.assertFalse(chain["can_prepay"])
                    self.assertTrue(service.order_execution_reason(order))
                    if status == "Completed":
                        self.assertEqual(chain["reason"], "")
                        self.assertIn("完成", chain["invoice_reason"])
                    else:
                        self.assertIn("关闭", chain["reason"])

    def test_receipt_payment_filter_includes_invoice_refs_to_its_verified_order(self):
        receipt = native("Purchase Receipt", "PR")
        payment = native("Payment Entry", "PE", references=[], party="S")
        row = {"name": "PE", "payment_type": "Pay", "docstatus": 1, "currency": "CNY", "amount": 30,
               "references": [{"doctype": "Purchase Invoice", "name": "PI", "orders": ["PO"], "receipts": [], "allocated": 30}]}
        with patch.object(service, "_source", return_value=receipt), patch.object(service, "_source_links", return_value=["PO"]), \
             patch.object(service, "_query_fields", return_value={"name", "party", "party_type", "payment_type", "posting_date", "creation"}), \
             patch.object(service, "_normalize_filter", side_effect=lambda dt, field, op, value: [dt, field, op, value]), \
             patch.object(frappe, "get_list", return_value=[frappe._dict(name="PE", party="S")]), \
             patch.object(service._RecordReader, "preload"), patch.object(service._RecordReader, "doc", return_value=payment), \
             patch.object(service, "_payment_row", return_value=row), patch.object(service, "_vouchers", return_value=[]), \
             patch.object(service, "_decorate_payments"):
            # Preloading is isolated above; preserve the native reader's presence test.
            reader = service._RecordReader(); reader.docs["Payment Entry", "PE"] = payment
            with patch.object(service, "_RecordReader", return_value=reader):
                result = service.get_payment_records(purchase_receipt="PR")
        self.assertEqual(result["total_count"], 1)
        self.assertEqual(result["rows"][0]["references"][0]["name"], "PI")


class PaymentRecordReaderTests(unittest.TestCase):
    def setUp(self):
        flags = patch.object(frappe, 'flags', frappe._dict())
        flags.start()
        self.addCleanup(flags.stop)
        query = patch.object(service, '_query_fields', return_value={'name', 'party', 'party_type', 'payment_type', 'posting_date', 'creation', 'modified'})
        query.start()
        self.addCleanup(query.stop)
        normalize = patch.object(service, '_normalize_filter', side_effect=lambda dt, field, operator, value: [dt, field, operator, value])
        normalize.start()
        self.addCleanup(normalize.stop)

    def test_bulk_hydration_preserves_children_and_masks_parent(self):
        reader = service._RecordReader()
        rows = [frappe._dict(name='PE-1', company='C')]
        children = [frappe._dict(name='REF-1', parent='PE-1', idx=1)]
        doc = SimpleNamespace(mask_fields=unittest.mock.Mock())
        meta = SimpleNamespace(get_table_fields=lambda: [frappe._dict(fieldname='references', options='Payment Entry Reference')])
        values = unittest.mock.Mock(side_effect=[rows, children])
        with patch.object(frappe, 'db', SimpleNamespace(get_values=values)), patch.object(frappe, 'get_meta', return_value=meta), patch.object(frappe, 'get_doc', return_value=doc) as get_doc:
            reader.preload('Payment Entry', ['PE-1', 'missing', 'PE-1'])
            reader.preload('Payment Entry', ['PE-1', 'missing'])
        self.assertEqual(values.call_count, 2)
        self.assertEqual(values.call_args_list[1].args[1], {'parent': ['in', ['PE-1']], 'parenttype': 'Payment Entry', 'parentfield': 'references'})
        self.assertEqual(get_doc.call_args.args[0]['references'], children)
        doc.mask_fields.assert_called_once()
        self.assertIs(reader.doc('Payment Entry', 'PE-1'), doc)
        with self.assertRaises(frappe.DoesNotExistError):
            reader.doc('Payment Entry', 'missing')

    def test_denied_permission_is_never_cached_as_success(self):
        reader = service._RecordReader()
        doc = SimpleNamespace(doctype='Payment Entry', name='PE-1', check_permission=unittest.mock.Mock(side_effect=frappe.PermissionError))
        for _ in range(2):
            with self.assertRaises(frappe.PermissionError): reader.check(doc)
        self.assertEqual(doc.check_permission.call_count, 2)
        self.assertFalse(reader.checked)

    def test_context_is_restored_after_unexpected_error(self):
        outer = service._RecordReader()
        token = service._record_reader.set(outer)
        try:
            with patch.object(service, '_payment_records', side_effect=RuntimeError):
                with self.assertRaises(RuntimeError): service.get_payment_records()
            self.assertIs(service._record_reader.get(), outer)
        finally:
            service._record_reader.reset(token)

    def test_empty_page_does_not_read_vouchers(self):
        with patch.object(service, '_require_fields'), patch.object(frappe, 'get_list', return_value=[]), patch.object(service, '_vouchers') as vouchers:
            result = service.get_payment_records(search='no match')
        self.assertEqual(result['total_count'], 0)
        self.assertEqual(result['rows'], [])
        vouchers.assert_not_called()


class ProcurementExportTests(unittest.TestCase):
    """Read native XLSX files; isolate document access and the site-only default format."""

    def export(self, doctype, rows, columns=None, can_export=True, check_read=None,
               generic_export=True, owner_export=False, owner='buyer@example.test'):
        response = {}
        doc = SimpleNamespace(check_permission=check_read or (lambda permission: None),
                              get=lambda field: owner if field == 'owner' else None)
        allowed = service.RECEIPT_COLUMNS if doctype == 'Purchase Receipt' else service.PAYABLE_COLUMNS if doctype == 'Purchase Invoice' else service.PAYMENT_COLUMNS
        from frappe.utils.xlsxutils import XLSXStyleBuilder
        capability = lambda doctype, is_owner=False: owner_export if is_owner else can_export
        def throw_without_site(message, exception):
            raise exception(message)
        with patch.object(frappe, 'response', response), patch.object(frappe, 'has_permission', return_value=generic_export), patch.object(frappe.permissions, 'can_export', side_effect=capability), patch.object(frappe, 'session', SimpleNamespace(user='buyer@example.test')), patch.object(frappe, 'throw', side_effect=throw_without_site), patch.object(service, '_read_doc', return_value=doc), patch.object(XLSXStyleBuilder, 'get_datetime_format', return_value='yyyy-mm-dd hh:mm:ss'):
            service._export(doctype, rows, columns, allowed)
        return load_workbook(BytesIO(response['filecontent']))

    def test_native_capability_allows_export_when_generic_and_document_export_are_false(self):
        def readable(permission):
            if permission == 'export':
                raise frappe.PermissionError
        for doctype in ('Purchase Receipt', 'Payment Entry'):
            with self.subTest(doctype=doctype):
                workbook = self.export(doctype, [{'name': 'DOC-1', 'references': []}], ['name'],
                                       generic_export=False, check_read=readable)
                self.assertEqual(workbook.active.cell(2, 1).value, 'DOC-1')

    def test_owner_only_export_allows_exact_user_and_rejects_other_or_missing_owners(self):
        for doctype in ('Purchase Receipt', 'Payment Entry'):
            for owner in ('buyer@example.test', 'other@example.test', None):
                with self.subTest(doctype=doctype, owner=owner):
                    def export():
                        return self.export(doctype, [{'name': 'DOC-1', 'references': []}], ['name'],
                                           can_export=False, owner_export=True, owner=owner)
                    if owner == 'buyer@example.test':
                        self.assertEqual(export().active.cell(2, 1).value, 'DOC-1')
                    else:
                        with self.assertRaises(frappe.PermissionError):
                            export()

    def test_selected_amount_columns_keep_order_and_append_distinct_currencies(self):
        cases = (
            ('Purchase Receipt', 'grand_total', ['入库日期', '入库金额', '采购入库单号', '入库币种'], ['USD']),
            ('Payment Entry', 'amount', ['付款日期', '金额', '付款单', '付款币种', '核销币种'], ['USD', 'CNY']),
        )
        for doctype, amount_field, headers, currencies in cases:
            with self.subTest(doctype=doctype):
                row = {'name': 'DOC-1', 'posting_date': '2026-10-04', amount_field: 4000.123456,
                       'currency': 'USD', 'references': [{'currency': 'CNY'}]}
                selected = ['posting_date', amount_field, 'name']
                workbook = self.export(doctype, [row], selected)
                self.assertEqual([cell.value for cell in workbook.active[1]], headers)
                self.assertEqual([cell.value for cell in workbook.active[2]][3:], currencies)
                self.assertEqual(selected, ['posting_date', amount_field, 'name'])
                self.assertEqual(workbook.active.cell(2, 2).value, 4000.123456)
                self.assertEqual(workbook.active.cell(2, 2).number_format, '#,##0.00')

    def test_payable_export_keeps_invoice_and_account_currencies_distinct_without_duplicating_columns(self):
        row = {'name': 'PI', 'grand_total': 100.125, 'invoice_currency': 'USD',
               'total': 600.75, 'settled': 450.25, 'outstanding': 150.5, 'currency': 'CNY'}
        for selected in (['name', 'grand_total', 'total', 'settled', 'outstanding'],
                         ['name', 'grand_total', 'invoice_currency', 'total', 'settled', 'outstanding', 'currency']):
            workbook = self.export('Purchase Invoice', [row], selected)
            labels = [cell.value for cell in workbook.active[1]]
            self.assertEqual(labels.count(service.PAYABLE_COLUMNS['currency']), 1)
            self.assertEqual(labels.count(service.PAYABLE_COLUMNS['invoice_currency']), 1)
            values = dict(zip(labels, [cell.value for cell in workbook.active[2]]))
            self.assertEqual(values[service.PAYABLE_COLUMNS['invoice_currency']], 'USD')
            self.assertEqual(values[service.PAYABLE_COLUMNS['currency']], 'CNY')
            for field in ('grand_total', 'total', 'settled', 'outstanding'):
                index = labels.index(service.PAYABLE_COLUMNS[field]) + 1
                self.assertEqual(workbook.active.cell(2, index).value, row[field])
                self.assertEqual(workbook.active.cell(2, index).number_format, '#,##0.00')

    def test_dates_are_typed_and_formatted_without_filling_missing_dates(self):
        for doctype in ('Purchase Receipt', 'Payment Entry'):
            for source_date in (date(2026, 10, 4), datetime(2026, 10, 4, 12, 30), '2026-10-04', None):
                with self.subTest(doctype=doctype, source_date=source_date):
                    row = {'name': 'DOC-1', 'posting_date': source_date, 'references': []}
                    workbook = self.export(doctype, [row], ['name', 'posting_date'])
                    cell = workbook.active.cell(2, 2)
                    self.assertEqual(cell.value, datetime(2026, 10, 4) if source_date else None)
                    if source_date:
                        self.assertEqual(cell.number_format, 'yyyy-mm-dd')
                        self.assertTrue(cell.is_date)
                    else:
                        # Native XLSXWriter omits empty cells; Excel inherits the column format.
                        self.assertEqual(workbook.active.column_dimensions['B'].number_format, 'yyyy-mm-dd')

    def test_default_and_explicit_currency_columns_are_not_duplicated(self):
        expected_receipt_labels = {
            'name': '采购入库单号', 'supplier_name': '供应商名称', 'supplier': '供应商编码',
            'posting_date': '入库日期', 'status': '入库状态', 'company': '公司',
            'currency': '入库币种', 'grand_total': '入库金额', 'docstatus': '单据状态', 'is_return': '是否退货',
        }
        self.assertNotIn('modified', service.RECEIPT_COLUMNS)
        self.assertNotIn('modified', service.RECEIPT_FIELDS)
        for doctype, amount_field in (('Purchase Receipt', 'grand_total'), ('Payment Entry', 'amount')):
            allowed = service.RECEIPT_COLUMNS if doctype == 'Purchase Receipt' else service.PAYMENT_COLUMNS
            row = {'name': 'DOC-1', amount_field: 0, 'currency': 'CNY', 'references': []}
            for selected in (None, ['currency', amount_field, 'name']):
                with self.subTest(doctype=doctype, selected=selected):
                    workbook = self.export(doctype, [row], selected)
                    labels = [cell.value for cell in workbook.active[1]]
                    self.assertEqual(labels.count(allowed['currency']), 1)
                    if selected:
                        self.assertEqual(labels[:3], [allowed[column] for column in selected])
                    elif doctype == 'Purchase Receipt':
                        self.assertEqual(labels[:len(service.RECEIPT_FIELDS)],
                                         [expected_receipt_labels[column] for column in service.RECEIPT_FIELDS])
                    if doctype == 'Payment Entry':
                        self.assertEqual(labels.count('核销币种'), 1)

    def test_export_permissions_and_formula_protection_remain_native(self):
        for prefix in ('=', '+', '-', '@'):
            with self.subTest(prefix=prefix):
                row = {'name': 'DOC-1', 'supplier_name': '<span>' + prefix + 'unsafe</span>'}
                workbook = self.export('Purchase Receipt', [row], ['supplier_name'])
                cell = workbook.active.cell(2, 1)
                self.assertEqual(cell.value, "'" + prefix + 'unsafe')
                self.assertEqual(cell.data_type, 's')
        # Pure CI has no site-bound message flags; the native QA script tests real throws.
        def throw_without_site(message, exception):
            raise exception(message)
        with patch.object(frappe, 'throw', side_effect=throw_without_site), self.assertRaises(frappe.PermissionError):
            self.export('Purchase Receipt', [], ['name'], can_export=False)

    def test_each_exported_document_still_requires_read_permission(self):
        checked = []
        def denied(permission):
            checked.append(permission)
            if permission == 'read':
                raise frappe.PermissionError
        for doctype in ('Purchase Receipt', 'Payment Entry'):
            with self.subTest(doctype=doctype), self.assertRaises(frappe.PermissionError):
                self.export(doctype, [{'name': 'PRIVATE', 'references': []}], ['name'], check_read=denied)
        self.assertEqual(checked, ['read', 'read'])


if __name__ == '__main__':
    unittest.main()
