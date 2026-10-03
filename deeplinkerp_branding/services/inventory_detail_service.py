"""四类库存明细查询、导出及未保存物料移动单准备服务。"""

from __future__ import annotations

import base64
import json
import math
from collections import OrderedDict
from collections.abc import Iterable
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
	import frappe
except Exception:  # pragma: no cover - 本地单测环境不安装 Frappe
	frappe = None


SNAPSHOT_DOCTYPE = "Inventory Original Location Snapshot"
COUNT_UOM_TOKENS = (
	"个",
	"件",
	"套",
	"卷",
	"包",
	"张",
	"片",
	"支",
	"条",
	"台",
	"pieza",
	"conjunto",
	"rollo",
	"paquete",
	"hoja",
	"ramo",
	"barra",
	"unidad",
)
COMMON_EXPORT_COLUMNS = (1, 2, 3, 6, 7, 8, 9, 10, 11)
EXPORT_HEADERS = (
	"正式物料编码",
	"物料名称（双语）",
	"仓库",
	"原始库位",
	"库位数量",
	"总库存",
	"库存单位",
	"物料组",
	"DPCI",
	"外部编码",
	"原始标识/别名",
)
CATEGORY_DEFINITIONS = {
	"semi_finished": {
		"title": "半成品库存明细",
		"root_item_group": "半成品Semiterminado",
		"route": "semi-finished-inventory-detail",
		"file_prefix": "semi-finished-inventory-detail",
	},
	"finished_goods": {
		"title": "成品库存明细",
		"root_item_group": "成品Producto terminado",
		"route": "finished-goods-inventory-detail",
		"file_prefix": "finished-goods-inventory-detail",
	},
	"mold": {
		"title": "模具库存明细",
		"root_item_group": "模具Moldes",
		"route": "mold-inventory-detail",
		"file_prefix": "mold-inventory-detail",
	},
}
CATEGORY_EXPORT_HEADERS = (
	"正式物料编码",
	"物料名称（双语）",
	"仓库",
	"参考库位",
	"快照库位数量",
	"实时库存",
	"库存差异",
	"库存状态",
	"快照日期",
	"库存单位",
	"物料组",
	"DPCI",
	"外部编码",
	"原始标识/别名",
)
CATEGORY_EXPORT_MERGE_COLUMNS = (1, 2, 3, 6, 7, 8, 10, 11, 12, 13, 14)
DEFAULT_PAGE_LENGTH = 200
MIN_PAGE_LENGTH = 1
MAX_PAGE_LENGTH = 2500
SIMPLE_STOCK_ENTRY_PURPOSES = {
	"Material Receipt",
	"Material Issue",
	"Material Transfer",
}


def _whitelist(function):
	if frappe is None:
		return function
	return frappe.whitelist()(function)


def is_count_uom(stock_uom: str | None) -> bool:
	normalized = str(stock_uom or "").strip().lower()
	return any(token in normalized for token in COUNT_UOM_TOKENS)


def quantity_precision(value: Any, stock_uom: str | None) -> int:
	number = float(value or 0)
	if not is_count_uom(stock_uom):
		return 2
	return 0 if abs(number - round(number)) <= 1e-9 else 2


def format_quantity(value: Any, stock_uom: str | None) -> str:
	precision = quantity_precision(value, stock_uom)
	return f"{float(value or 0):,.{precision}f}"


def get_category_definition(category: Any) -> dict[str, str]:
	key = _text(category)
	definition = CATEGORY_DEFINITIONS.get(key)
	if not definition:
		raise ValueError(f"不支持的库存分类：{key or '空'}")
	return dict(definition)


def _text(value: Any) -> str:
	return str(value or "").strip()


def parse_page_length(value: Any = None) -> int:
	"""Return a validated user-entered page length."""

	if value is None or (isinstance(value, str) and not value.strip()):
		return DEFAULT_PAGE_LENGTH
	if isinstance(value, bool):
		raise ValueError("每页行数必须是 1 到 2500 的整数。")
	try:
		parsed = int(value)
	except (TypeError, ValueError) as error:
		raise ValueError("每页行数必须是 1 到 2500 的整数。") from error
	if str(value).strip() not in {str(parsed), f"+{parsed}"}:
		raise ValueError("每页行数必须是 1 到 2500 的整数。")
	if not MIN_PAGE_LENGTH <= parsed <= MAX_PAGE_LENGTH:
		raise ValueError("每页行数必须是 1 到 2500 的整数。")
	return parsed


def _contains(value: Any, keyword: Any) -> bool:
	return _text(keyword).casefold() in _text(value).casefold()


def _matches_group_filters(row: dict[str, Any], filters: dict[str, Any]) -> bool:
	for fieldname in ("company", "warehouse", "item_group", "snapshot_key"):
		selected = _text(filters.get(fieldname))
		if selected and _text(row.get(fieldname)) != selected:
			return False

	snapshot_date = _text(filters.get("snapshot_date"))
	if snapshot_date and _text(row.get("snapshot_date")) != snapshot_date:
		return False

	keyword = _text(filters.get("keyword") or filters.get("item_code"))
	if keyword and not any(
		_contains(row.get(fieldname), keyword)
		for fieldname in (
			"item_code",
			"item_name",
			"dpci",
			"external_code",
			"original_identifier_alias",
		)
	):
		return False
	return True


def build_inventory_location_payload(
	rows: Iterable[dict[str, Any]],
	filters: dict[str, Any] | None = None,
	*,
	start: int = 0,
	page_length: int | None = DEFAULT_PAGE_LENGTH,
) -> dict[str, Any]:
	"""按物料和仓库分组；库位筛选不改变完整仓库小计。"""

	filters = dict(filters or {})
	all_rows = [dict(row) for row in rows]
	group_rows = [row for row in all_rows if _matches_group_filters(row, filters)]
	groups: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()

	for row in sorted(
		group_rows,
		key=lambda value: (
			_text(value.get("warehouse")),
			_text(value.get("item_code")),
			_text(value.get("original_location")),
		),
	):
		key = (_text(row.get("item_code")), _text(row.get("warehouse")))
		group = groups.setdefault(
			key,
			{
				fieldname: row.get(fieldname) or ""
				for fieldname in (
					"item_code",
					"item_name",
					"warehouse",
					"stock_uom",
					"item_group",
					"dpci",
					"external_code",
					"original_identifier_alias",
				)
			}
			| {"total_qty": 0.0, "locations": []},
		)
		group["total_qty"] += float(row.get("location_qty") or 0)
		group["locations"].append(
			{
				"original_location": row.get("original_location") or "",
				"location_qty": float(row.get("location_qty") or 0),
			}
		)

	location_filter = _text(filters.get("original_location"))
	visible_groups: list[dict[str, Any]] = []
	for group in groups.values():
		if location_filter:
			visible_locations = [
				location
				for location in group["locations"]
				if _contains(location["original_location"], location_filter)
			]
			if not visible_locations:
				continue
			group["locations"] = visible_locations
		visible_groups.append(group)

	total_count = len(visible_groups)
	start = _parse_start(start)
	page_groups = (
		visible_groups[start : start + page_length] if page_length is not None else visible_groups[start:]
	)
	snapshot_source = next(
		(row for row in group_rows if row.get("snapshot_key")),
		next((row for row in all_rows if row.get("snapshot_key")), {}),
	)
	return {
		"snapshot_key": snapshot_source.get("snapshot_key") or "",
		"snapshot_date": snapshot_source.get("snapshot_date") or "",
		"company": filters.get("company") or snapshot_source.get("company") or "",
		"groups": page_groups,
		"group_count": total_count,
		"total_count": total_count,
		"page_count": len(page_groups),
		"location_count": sum(len(group["locations"]) for group in page_groups),
		"start": start,
		"page_length": page_length,
		"has_previous": start > 0,
		"has_next": bool(page_length is not None and start + len(page_groups) < total_count),
	}


def _number_format(stock_uom: str | None, value: Any) -> str:
	return "#,##0" if quantity_precision(value, stock_uom) == 0 else "#,##0.00"


def build_inventory_location_xlsx(payload: dict[str, Any]) -> bytes:
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "物料库存明细"
	sheet.freeze_panes = "A2"
	sheet.append(EXPORT_HEADERS)

	header_fill = PatternFill("solid", fgColor="075985")
	for cell in sheet[1]:
		cell.fill = header_fill
		cell.font = Font(color="FFFFFF", bold=True)
		cell.alignment = Alignment(horizontal="center", vertical="center")

	current_row = 2
	for group in payload.get("groups") or []:
		locations = list(group.get("locations") or []) or [{"original_location": "", "location_qty": 0}]
		start_row = current_row
		for index, location in enumerate(locations):
			if index == 0:
				values = [
					group.get("item_code") or "",
					group.get("item_name") or "",
					group.get("warehouse") or "",
					location.get("original_location") or "",
					float(location.get("location_qty") or 0),
					float(group.get("total_qty") or 0),
					group.get("stock_uom") or "",
					group.get("item_group") or "",
					group.get("dpci") or "",
					group.get("external_code") or "",
					group.get("original_identifier_alias") or "",
				]
			else:
				values = [
					"",
					"",
					"",
					location.get("original_location") or "",
					float(location.get("location_qty") or 0),
					"",
					"",
					"",
					"",
					"",
					"",
				]
			sheet.append(values)
			sheet.cell(current_row, 5).number_format = _number_format(
				group.get("stock_uom"), location.get("location_qty")
			)
			if index == 0:
				sheet.cell(current_row, 6).number_format = _number_format(
					group.get("stock_uom"), group.get("total_qty")
				)
			current_row += 1

		end_row = current_row - 1
		if end_row > start_row:
			for column in COMMON_EXPORT_COLUMNS:
				sheet.merge_cells(
					start_row=start_row,
					start_column=column,
					end_row=end_row,
					end_column=column,
				)
		for row_number in range(start_row, end_row + 1):
			for cell in sheet[row_number]:
				cell.alignment = Alignment(vertical="center", wrap_text=True)

	widths = (18, 32, 24, 18, 15, 15, 15, 28, 18, 20, 28)
	for index, width in enumerate(widths, start=1):
		sheet.column_dimensions[chr(64 + index)].width = width

	output = BytesIO()
	workbook.save(output)
	return output.getvalue()


def _category_group_metadata(row: dict[str, Any]) -> dict[str, Any]:
	return {
		fieldname: row.get(fieldname) or ""
		for fieldname in (
			"item_code",
			"item_name",
			"warehouse",
			"stock_uom",
			"item_group",
			"dpci",
			"external_code",
			"original_identifier_alias",
		)
	} | {"actual_qty": float(row.get("actual_qty") or 0), "_snapshot_locations": []}


def _truthy(value: Any) -> bool:
	return _text(value).casefold() in {"1", "true", "yes", "on"}


def _parse_start(value: Any) -> int:
	try:
		return max(0, int(value or 0))
	except (TypeError, ValueError):
		return 0


def _matches_category_filters(group: dict[str, Any], filters: dict[str, Any]) -> bool:
	warehouse = _text(filters.get("warehouse"))
	if warehouse and _text(group.get("warehouse")) != warehouse:
		return False
	if _truthy(filters.get("only_with_stock")) and abs(float(group.get("actual_qty") or 0)) <= 1e-9:
		return False
	keyword = _text(filters.get("keyword") or filters.get("item_code"))
	if keyword and not any(
		_contains(group.get(fieldname), keyword)
		for fieldname in (
			"item_code",
			"item_name",
			"dpci",
			"external_code",
			"original_identifier_alias",
		)
	):
		return False
	return True


def build_categorized_inventory_payload(
	stock_rows: Iterable[dict[str, Any]],
	snapshot_rows: Iterable[dict[str, Any]],
	*,
	category: str,
	filters: dict[str, Any] | None = None,
	start: int = 0,
	page_length: int | None = DEFAULT_PAGE_LENGTH,
) -> dict[str, Any]:
	"""合并实时 Bin 与最新库位快照；快照只作参考，不覆盖实时库存。"""

	definition = get_category_definition(category)
	filters = dict(filters or {})
	groups: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()
	item_metadata: dict[str, dict[str, Any]] = {}

	for source in stock_rows:
		row = dict(source)
		item_code = _text(row.get("item_code"))
		if not item_code:
			continue
		key = (item_code, _text(row.get("warehouse")))
		group = _category_group_metadata(row)
		groups[key] = group
		item_metadata.setdefault(item_code, group)

	snapshot_source: dict[str, Any] = {}
	for source in sorted(
		(dict(row) for row in snapshot_rows),
		key=lambda value: (
			_text(value.get("item_code")),
			_text(value.get("warehouse")),
			_text(value.get("original_location")),
		),
	):
		item_code = _text(source.get("item_code"))
		if not item_code or item_code not in item_metadata:
			continue
		warehouse = _text(source.get("warehouse"))
		key = (item_code, warehouse)
		if key not in groups:
			metadata = dict(item_metadata[item_code])
			metadata["warehouse"] = warehouse
			metadata["actual_qty"] = 0.0
			metadata["_snapshot_locations"] = []
			groups[key] = metadata
		groups[key]["_snapshot_locations"].append(
			{
				"reference_location": source.get("original_location") or "库位待维护",
				"snapshot_location_qty": float(source.get("location_qty") or 0),
				"snapshot_date": source.get("snapshot_date") or "",
			}
		)
		if source.get("snapshot_key"):
			snapshot_source = source

	# Item 的空仓库行只是“从未产生 Bin”的占位；一旦快照提供仓库就不再保留占位行。
	item_has_warehouse = {item_code for item_code, warehouse in groups if warehouse}
	for key in list(groups):
		if not key[1] and key[0] in item_has_warehouse:
			del groups[key]

	location_filter = _text(filters.get("reference_location") or filters.get("original_location"))
	visible: list[dict[str, Any]] = []
	for group in groups.values():
		if not _matches_category_filters(group, filters):
			continue
		all_locations = list(group.pop("_snapshot_locations", []))
		snapshot_qty = (
			sum(float(location.get("snapshot_location_qty") or 0) for location in all_locations)
			if all_locations
			else None
		)
		if location_filter:
			locations = [
				location
				for location in all_locations
				if _contains(location.get("reference_location"), location_filter)
			]
			if not locations:
				continue
		else:
			locations = all_locations

		actual_qty = float(group.get("actual_qty") or 0)
		has_reference_location = any(
			_text(location.get("reference_location")) not in {"", "库位待维护"} for location in all_locations
		)
		difference_qty = actual_qty - snapshot_qty if snapshot_qty is not None else None
		if not has_reference_location:
			inventory_status = "库位待维护"
		elif abs(float(difference_qty or 0)) <= 1e-9:
			inventory_status = "库存一致"
		else:
			inventory_status = "库存差异"
		group.update(
			{
				"snapshot_qty": snapshot_qty,
				"difference_qty": difference_qty,
				"inventory_status": inventory_status,
				"locations": locations
				or [
					{
						"reference_location": "库位待维护",
						"snapshot_location_qty": None,
						"snapshot_date": "",
					}
				],
			}
		)
		visible.append(group)

	visible.sort(
		key=lambda row: (
			abs(float(row.get("actual_qty") or 0)) <= 1e-9,
			_text(row.get("item_code")),
			_text(row.get("warehouse")),
		)
	)
	total_count = len(visible)
	start = _parse_start(start)
	page_groups = visible[start : start + page_length] if page_length is not None else visible[start:]
	return {
		"category": category,
		"title": definition["title"],
		"root_item_group": definition["root_item_group"],
		"snapshot_key": snapshot_source.get("snapshot_key") or "",
		"snapshot_date": snapshot_source.get("snapshot_date") or "",
		"groups": page_groups,
		"total_count": total_count,
		"page_count": len(page_groups),
		"location_count": sum(len(row["locations"]) for row in page_groups),
		"start": start,
		"page_length": page_length,
		"has_previous": start > 0,
		"has_next": bool(page_length is not None and start + len(page_groups) < total_count),
	}


def build_categorized_inventory_xlsx(payload: dict[str, Any]) -> bytes:
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = _text(payload.get("title")) or "分类库存明细"
	sheet.freeze_panes = "A2"
	sheet.append(CATEGORY_EXPORT_HEADERS)

	header_fill = PatternFill("solid", fgColor="075985")
	for cell in sheet[1]:
		cell.fill = header_fill
		cell.font = Font(color="FFFFFF", bold=True)
		cell.alignment = Alignment(horizontal="center", vertical="center")

	current_row = 2
	for group in payload.get("groups") or []:
		locations = list(group.get("locations") or []) or [
			{
				"reference_location": "库位待维护",
				"snapshot_location_qty": None,
				"snapshot_date": "",
			}
		]
		start_row = current_row
		for index, location in enumerate(locations):
			shared = index == 0
			sheet.append(
				[
					group.get("item_code") or "" if shared else "",
					group.get("item_name") or "" if shared else "",
					group.get("warehouse") or "" if shared else "",
					location.get("reference_location") or "库位待维护",
					location.get("snapshot_location_qty"),
					float(group.get("actual_qty") or 0) if shared else "",
					group.get("difference_qty") if shared else "",
					group.get("inventory_status") or "" if shared else "",
					location.get("snapshot_date") or "",
					group.get("stock_uom") or "" if shared else "",
					group.get("item_group") or "" if shared else "",
					group.get("dpci") or "" if shared else "",
					group.get("external_code") or "" if shared else "",
					group.get("original_identifier_alias") or "" if shared else "",
				]
			)
			if location.get("snapshot_location_qty") is not None:
				sheet.cell(current_row, 5).number_format = _number_format(
					group.get("stock_uom"), location.get("snapshot_location_qty")
				)
			if shared:
				sheet.cell(current_row, 6).number_format = _number_format(
					group.get("stock_uom"), group.get("actual_qty")
				)
				if group.get("difference_qty") is not None:
					sheet.cell(current_row, 7).number_format = _number_format(
						group.get("stock_uom"), group.get("difference_qty")
					)
			current_row += 1

		end_row = current_row - 1
		if end_row > start_row:
			for column in CATEGORY_EXPORT_MERGE_COLUMNS:
				sheet.merge_cells(
					start_row=start_row,
					start_column=column,
					end_row=end_row,
					end_column=column,
				)
		for row_number in range(start_row, end_row + 1):
			for cell in sheet[row_number]:
				cell.alignment = Alignment(vertical="center", wrap_text=True)

	widths = (18, 32, 24, 18, 15, 15, 15, 14, 14, 15, 28, 18, 20, 28)
	for index, width in enumerate(widths, start=1):
		sheet.column_dimensions[chr(64 + index)].width = width

	output = BytesIO()
	workbook.save(output)
	return output.getvalue()


def _parse_filters(filters: Any = None, **kwargs: Any) -> dict[str, Any]:
	if isinstance(filters, str):
		filters = json.loads(filters or "{}")
	parsed = dict(filters or {})
	parsed.update({key: value for key, value in kwargs.items() if value is not None})
	return parsed


def _require_read_permission() -> None:
	if frappe is None:
		return
	if not frappe.has_permission(SNAPSHOT_DOCTYPE, "read"):
		frappe.throw("没有权限查看库存库位快照。", frappe.PermissionError)


def _require_categorized_inventory_read_permission() -> None:
	if frappe is None:
		return
	_require_read_permission()
	for doctype in ("Item", "Bin", "Warehouse"):
		if not frappe.has_permission(doctype, "read"):
			frappe.throw(f"没有权限查看{doctype}。", frappe.PermissionError)


def _require_company_permission(company: str) -> None:
	"""Raw SQL 读取前仍须遵守用户对 Company 的文档权限。"""

	if frappe is None:
		return
	frappe.get_doc("Company", company).check_permission("read")


def _resolve_company(company: Any = None) -> str:
	"""Use an explicit company or the current user's Frappe default."""

	resolved = _text(company)
	if resolved:
		return resolved
	if frappe is None:
		raise ValueError("请先选择公司或设置默认公司。")
	resolved = _text(frappe.defaults.get_user_default("Company"))
	if not resolved:
		frappe.throw("请先选择公司或设置默认公司。", frappe.ValidationError)
	return resolved


def _accessible_items() -> tuple[str, ...]:
	rows = frappe.get_list(
		"Item",
		fields=["name"],
		limit_page_length=0,
	)
	return tuple(
		_text(row.get("name") if isinstance(row, dict) else getattr(row, "name", row))
		for row in rows
		if _text(row.get("name") if isinstance(row, dict) else getattr(row, "name", row))
	)


def _latest_snapshot_key(company: str) -> str:
	rows = frappe.db.sql(
		f"""
        SELECT snapshot_key
        FROM `tab{SNAPSHOT_DOCTYPE}`
        WHERE company = %(company)s
        ORDER BY snapshot_date DESC, creation DESC
        LIMIT 1
        """,
		{"company": company},
		as_dict=True,
	)
	return rows[0]["snapshot_key"] if rows else ""


def _load_snapshot_rows(filters: dict[str, Any]) -> list[dict[str, Any]]:
	company = _resolve_company(filters.get("company"))
	snapshot_key = _text(filters.get("snapshot_key")) or _latest_snapshot_key(company)
	if not snapshot_key:
		return []
	items = _accessible_items()
	warehouses = _accessible_warehouses(company)
	selected_warehouse = _text(filters.get("warehouse"))
	if selected_warehouse and selected_warehouse not in warehouses:
		frappe.throw("没有权限查看所选仓库。", frappe.PermissionError)
	if not items or not warehouses:
		return []
	filters["company"] = company
	filters["snapshot_key"] = snapshot_key
	return frappe.db.sql(
		f"""
        SELECT
            snapshot.snapshot_key,
            snapshot.snapshot_date,
            snapshot.company,
            snapshot.item_code,
            item.item_name,
            snapshot.warehouse,
            snapshot.original_location,
            snapshot.location_qty,
            snapshot.stock_uom,
            item.item_group,
            COALESCE(item.custom_dpci, '') AS dpci,
            COALESCE(item.custom_external_code, '') AS external_code,
            COALESCE(item.custom_original_identifier_alias, '') AS original_identifier_alias
        FROM `tab{SNAPSHOT_DOCTYPE}` snapshot
        INNER JOIN `tabItem` item ON item.name = snapshot.item_code
        WHERE snapshot.snapshot_key = %(snapshot_key)s
          AND snapshot.company = %(company)s
		  AND snapshot.item_code IN %(items)s
		  AND snapshot.warehouse IN %(warehouses)s
        ORDER BY snapshot.warehouse, snapshot.item_code, snapshot.original_location
        """,
		{
			"snapshot_key": snapshot_key,
			"company": company,
			"items": items,
			"warehouses": warehouses,
		},
		as_dict=True,
	)


def _list_snapshot_options(company: str) -> list[dict[str, Any]]:
	return frappe.db.sql(
		f"""
        SELECT snapshot_key, MAX(snapshot_date) AS snapshot_date
        FROM `tab{SNAPSHOT_DOCTYPE}`
        WHERE company = %(company)s
        GROUP BY snapshot_key
        ORDER BY snapshot_date DESC, snapshot_key DESC
        """,
		{"company": company},
		as_dict=True,
	)


def _category_group_bounds(category: str, selected_group: str = "") -> tuple[int, int] | None:
	definition = get_category_definition(category)
	root_name = definition["root_item_group"]
	root = frappe.db.get_value("Item Group", root_name, ["lft", "rgt"], as_dict=True)
	if not root:
		return None
	frappe.get_doc("Item Group", root_name).check_permission("read")
	selected_group = _text(selected_group)
	if not selected_group:
		return int(root["lft"]), int(root["rgt"])

	selected = frappe.db.get_value("Item Group", selected_group, ["lft", "rgt"], as_dict=True)
	if not selected:
		raise ValueError(f"物料组不存在：{selected_group}")
	if int(selected["lft"]) < int(root["lft"]) or int(selected["rgt"]) > int(root["rgt"]):
		raise ValueError(f"物料组不属于{root_name}：{selected_group}")
	frappe.get_doc("Item Group", selected_group).check_permission("read")
	return int(selected["lft"]), int(selected["rgt"])


def _accessible_warehouses(company: str) -> tuple[str, ...]:
	rows = frappe.get_list(
		"Warehouse",
		filters={"company": company, "is_group": 0},
		fields=["name"],
		limit_page_length=0,
	)
	return tuple(
		_text(row.get("name") if isinstance(row, dict) else getattr(row, "name", row))
		for row in rows
		if _text(row.get("name") if isinstance(row, dict) else getattr(row, "name", row))
	)


def _category_query_context(category: str, filters: dict[str, Any]) -> dict[str, Any]:
	company = _resolve_company(filters.get("company"))
	bounds = _category_group_bounds(category, _text(filters.get("item_group")))
	if bounds is None:
		return {
			"company": company,
			"missing_item_group": get_category_definition(category)["root_item_group"],
			"warehouses": (),
			"snapshot_key": "",
		}
	group_lft, group_rgt = bounds
	warehouses = _accessible_warehouses(company)
	items = _accessible_items()
	selected_warehouse = _text(filters.get("warehouse"))
	if selected_warehouse and selected_warehouse not in warehouses:
		if frappe is not None:
			frappe.throw("没有权限查看所选仓库。", frappe.PermissionError)
		raise PermissionError("没有权限查看所选仓库。")
	return {
		"company": company,
		"group_lft": group_lft,
		"group_rgt": group_rgt,
		"warehouses": warehouses,
		"items": items,
		"snapshot_key": _latest_snapshot_key(company),
	}


def _load_category_stock_rows(context: dict[str, Any]) -> list[dict[str, Any]]:
	warehouses = context["warehouses"]
	items = context["items"]
	if not items:
		return []
	if warehouses:
		stock_join = """
        LEFT JOIN (
            SELECT bin.item_code, bin.warehouse, bin.actual_qty
            FROM `tabBin` bin
            INNER JOIN `tabWarehouse` warehouse ON warehouse.name = bin.warehouse
            WHERE warehouse.company = %(company)s
              AND warehouse.is_group = 0
              AND warehouse.name IN %(warehouses)s
        ) stock ON stock.item_code = item.name
        """
	else:
		stock_join = """
        LEFT JOIN (
            SELECT NULL AS item_code, NULL AS warehouse, 0 AS actual_qty
            WHERE 1 = 0
        ) stock ON stock.item_code = item.name
        """
	return frappe.db.sql(
		f"""
        SELECT
            item.name AS item_code,
            item.item_name,
            COALESCE(stock.warehouse, '') AS warehouse,
            COALESCE(stock.actual_qty, 0) AS actual_qty,
            item.stock_uom,
            item.item_group,
            COALESCE(item.custom_dpci, '') AS dpci,
            COALESCE(item.custom_external_code, '') AS external_code,
            COALESCE(item.custom_original_identifier_alias, '') AS original_identifier_alias
        FROM `tabItem` item
        INNER JOIN `tabItem Group` item_group ON item_group.name = item.item_group
        {stock_join}
        WHERE item.disabled = 0
          AND item.is_stock_item = 1
		  AND item.name IN %(items)s
          AND item_group.lft >= %(group_lft)s
          AND item_group.rgt <= %(group_rgt)s
        ORDER BY item.name, stock.warehouse
        """,
		{
			"company": context["company"],
			"warehouses": warehouses or ("",),
			"items": items,
			"group_lft": context["group_lft"],
			"group_rgt": context["group_rgt"],
		},
		as_dict=True,
	)


def _load_category_snapshot_rows(context: dict[str, Any]) -> list[dict[str, Any]]:
	if not context["snapshot_key"] or not context["warehouses"] or not context["items"]:
		return []
	return frappe.db.sql(
		f"""
        SELECT
            snapshot.snapshot_key,
            snapshot.snapshot_date,
            snapshot.item_code,
            item.item_name,
            snapshot.warehouse,
            snapshot.original_location,
            snapshot.location_qty,
            snapshot.stock_uom,
            item.item_group,
            COALESCE(item.custom_dpci, '') AS dpci,
            COALESCE(item.custom_external_code, '') AS external_code,
            COALESCE(item.custom_original_identifier_alias, '') AS original_identifier_alias
        FROM `tab{SNAPSHOT_DOCTYPE}` snapshot
        INNER JOIN `tabItem` item ON item.name = snapshot.item_code
        INNER JOIN `tabItem Group` item_group ON item_group.name = item.item_group
        WHERE snapshot.snapshot_key = %(snapshot_key)s
          AND snapshot.company = %(company)s
          AND snapshot.warehouse IN %(warehouses)s
		  AND snapshot.item_code IN %(items)s
          AND item_group.lft >= %(group_lft)s
          AND item_group.rgt <= %(group_rgt)s
        ORDER BY snapshot.item_code, snapshot.warehouse, snapshot.original_location
        """,
		{
			"snapshot_key": context["snapshot_key"],
			"company": context["company"],
			"warehouses": context["warehouses"],
			"items": context["items"],
			"group_lft": context["group_lft"],
			"group_rgt": context["group_rgt"],
		},
		as_dict=True,
	)


def _list_category_item_groups(context: dict[str, Any]) -> list[str]:
	rows = frappe.db.sql(
		"""
        SELECT name
        FROM `tabItem Group`
        WHERE lft >= %(group_lft)s
          AND rgt <= %(group_rgt)s
        ORDER BY lft
        """,
		{
			"group_lft": context["group_lft"],
			"group_rgt": context["group_rgt"],
		},
		as_dict=True,
	)
	return [_text(row.get("name")) for row in rows if _text(row.get("name"))]


def _movement_permission_payload() -> dict[str, Any]:
	if frappe is None:
		return {"can_create_stock_entry": True, "movement_disabled_reason": ""}
	allowed = bool(frappe.has_permission("Stock Entry", "create"))
	return {
		"can_create_stock_entry": allowed,
		"movement_disabled_reason": "" if allowed else "没有创建物料移动单的权限。",
	}


def _parse_list(value: Any, label: str) -> list[dict[str, Any]]:
	if isinstance(value, str):
		try:
			value = json.loads(value or "[]")
		except json.JSONDecodeError as error:
			raise ValueError(f"{label}格式不正确。") from error
	if not isinstance(value, list) or not value:
		raise ValueError(f"{label}不能为空。")
	if len(value) > MAX_PAGE_LENGTH:
		raise ValueError(f"{label}一次不能超过 {MAX_PAGE_LENGTH} 项。")
	if not all(isinstance(row, dict) for row in value):
		raise ValueError(f"{label}格式不正确。")
	return [dict(row) for row in value]


def _require_stock_entry_create_permission() -> None:
	if frappe is None:
		return
	if not frappe.has_permission("Stock Entry", "create"):
		frappe.throw("没有权限创建物料移动单。", frappe.PermissionError)
	for doctype in ("Item", "Bin", "Warehouse", "Stock Entry Type"):
		if not frappe.has_permission(doctype, "read"):
			frappe.throw(f"没有权限读取{doctype}。", frappe.PermissionError)


def _request_page_length(value: Any) -> int:
	try:
		return parse_page_length(value)
	except ValueError as error:
		if frappe is not None:
			frappe.throw(str(error), frappe.ValidationError)
		raise


def _validate_item(item_code: str):
	item_code = _text(item_code)
	if not item_code:
		raise ValueError("物料编码不能为空。")
	item = frappe.get_doc("Item", item_code)
	item.check_permission("read")
	if getattr(item, "disabled", 0) or not getattr(item, "is_stock_item", 0):
		raise ValueError(f"物料不可用于库存移动：{item_code}")
	return item


def _validate_warehouse(company: str, warehouse: str):
	warehouse = _text(warehouse)
	if not warehouse:
		raise ValueError("仓库不能为空。")
	doc = frappe.get_doc("Warehouse", warehouse)
	doc.check_permission("read")
	if _text(getattr(doc, "company", "")) != company:
		raise ValueError(f"仓库不属于当前公司：{warehouse}")
	if int(getattr(doc, "is_group", 0) or 0):
		raise ValueError(f"不能使用分组仓库：{warehouse}")
	return doc


def _actual_qty(item_code: str, warehouse: str) -> float:
	if frappe is None or not warehouse:
		return 0.0
	value = frappe.db.get_value(
		"Bin",
		{"item_code": item_code, "warehouse": warehouse},
		"actual_qty",
	)
	return float(value or 0)


def build_movement_item_spec(
	*,
	purpose: str,
	item_code: str,
	source_warehouse: str,
	target_warehouse: str,
	quantity: Any,
	actual_qty: Any,
) -> dict[str, Any]:
	"""Build simple-purpose warehouses without writing an ERP document."""

	purpose = _text(purpose)
	item_code = _text(item_code)
	source_warehouse = _text(source_warehouse)
	target_warehouse = _text(target_warehouse)
	try:
		quantity = float(quantity)
	except (TypeError, ValueError) as error:
		raise ValueError(f"{item_code or '物料'}的移动数量必须是数字。") from error
	if not math.isfinite(quantity):
		raise ValueError(f"{item_code or '物料'}的移动数量必须是有限数字。")
	if quantity <= 0:
		raise ValueError(f"{item_code or '物料'}的移动数量必须大于零。")

	if purpose in {"Material Issue", "Material Transfer"}:
		if not source_warehouse:
			raise ValueError(f"{item_code}缺少来源仓库。")
		if quantity - float(actual_qty or 0) > 1e-9:
			raise ValueError(f"{item_code}的移动数量超过来源仓实时库存 {float(actual_qty or 0):g}。")
	if purpose in {"Material Receipt", "Material Transfer"} and not target_warehouse:
		raise ValueError("请选择目标仓库。")
	if purpose == "Material Transfer" and source_warehouse == target_warehouse:
		raise ValueError(f"{item_code}的来源仓库和目标仓库不能相同。")

	result = {
		"item_code": item_code,
		"qty": quantity,
		"s_warehouse": "",
		"t_warehouse": "",
	}
	if purpose == "Material Receipt":
		result["t_warehouse"] = target_warehouse
	elif purpose == "Material Issue":
		result["s_warehouse"] = source_warehouse
	elif purpose == "Material Transfer":
		result["s_warehouse"] = source_warehouse
		result["t_warehouse"] = target_warehouse
	return result


def _stock_entry_types() -> list[dict[str, str]]:
	rows = frappe.get_list(
		"Stock Entry Type",
		fields=["name", "purpose"],
		order_by="name asc",
		limit_page_length=0,
	)
	return [
		{
			"name": _text(row.get("name") if isinstance(row, dict) else row.name),
			"purpose": _text(row.get("purpose") if isinstance(row, dict) else row.purpose),
		}
		for row in rows
	]


@_whitelist
def get_inventory_movement_context(company: str, selections: Any) -> dict[str, Any]:
	"""Re-read selected item/warehouse groups from live Bin quantities."""

	_require_stock_entry_create_permission()
	company = _resolve_company(company)
	_require_company_permission(company)
	selected = _parse_list(selections, "已选物料")
	seen: set[tuple[str, str]] = set()
	items: list[dict[str, Any]] = []
	for selection in selected:
		item_code = _text(selection.get("item_code"))
		source_warehouse = _text(selection.get("source_warehouse"))
		key = (item_code, source_warehouse)
		if key in seen:
			raise ValueError(f"重复选择物料仓库组：{item_code} / {source_warehouse}")
		seen.add(key)
		item = _validate_item(item_code)
		_validate_warehouse(company, source_warehouse)
		items.append(
			{
				"item_code": item_code,
				"item_name": _text(getattr(item, "item_name", "")),
				"source_warehouse": source_warehouse,
				"actual_qty": _actual_qty(item_code, source_warehouse),
				"stock_uom": _text(getattr(item, "stock_uom", "")),
			}
		)
	return {
		"company": company,
		"items": items,
		"stock_entry_types": _stock_entry_types(),
	}


@_whitelist
def prepare_inventory_stock_entry(
	company: str,
	stock_entry_type: str,
	target_warehouse: str = "",
	items: Any = None,
) -> dict[str, Any]:
	"""Return a complete local Stock Entry without inserting or saving it."""

	_require_stock_entry_create_permission()
	company = _resolve_company(company)
	_require_company_permission(company)
	stock_entry_type = _text(stock_entry_type)
	if not stock_entry_type:
		raise ValueError("请选择移动类型。")
	type_doc = frappe.get_doc("Stock Entry Type", stock_entry_type)
	type_doc.check_permission("read")
	purpose = _text(getattr(type_doc, "purpose", ""))
	if not purpose:
		raise ValueError(f"移动类型未配置用途：{stock_entry_type}")

	target_warehouse = _text(target_warehouse)
	if purpose in {"Material Receipt", "Material Transfer"}:
		_validate_warehouse(company, target_warehouse)

	selected = _parse_list(items, "移动物料")
	stock_entry = frappe.new_doc("Stock Entry")
	stock_entry.company = company
	stock_entry.stock_entry_type = stock_entry_type
	stock_entry.purpose = purpose
	if purpose in {"Material Receipt", "Material Transfer"}:
		stock_entry.to_warehouse = target_warehouse

	seen: set[tuple[str, str]] = set()
	for selected_item in selected:
		item_code = _text(selected_item.get("item_code"))
		source_warehouse = _text(selected_item.get("source_warehouse"))
		key = (item_code, source_warehouse)
		if key in seen:
			raise ValueError(f"重复移动物料仓库组：{item_code} / {source_warehouse}")
		seen.add(key)
		item = _validate_item(item_code)
		actual_qty = 0.0
		if purpose in {"Material Issue", "Material Transfer"}:
			_validate_warehouse(company, source_warehouse)
			actual_qty = _actual_qty(item_code, source_warehouse)
		spec = build_movement_item_spec(
			purpose=purpose,
			item_code=item_code,
			source_warehouse=source_warehouse,
			target_warehouse=target_warehouse,
			quantity=selected_item.get("qty"),
			actual_qty=actual_qty,
		)
		warehouse = spec.get("s_warehouse") or spec.get("t_warehouse") or ""
		details = stock_entry.get_item_details(
			frappe._dict(
				{
					"item_code": item_code,
					"company": company,
					"qty": spec["qty"],
					"warehouse": warehouse,
				}
			)
		)
		row = dict(details or {})
		row.update(spec)
		row.setdefault("uom", _text(getattr(item, "stock_uom", "")))
		row.setdefault("stock_uom", _text(getattr(item, "stock_uom", "")))
		row.setdefault("conversion_factor", 1)
		stock_entry.append("items", row)

	document = stock_entry.as_dict()
	document["__islocal"] = 1
	return {
		"stock_entry": document,
		"purpose": purpose,
		"requires_completion": purpose not in SIMPLE_STOCK_ENTRY_PURPOSES,
		"completion_message": (
			"该移动类型还需要在标准物料移动表单中补充工单、BOM、批次、序列号或物料角色等字段。"
			if purpose not in SIMPLE_STOCK_ENTRY_PURPOSES
			else ""
		),
	}


@_whitelist
def get_inventory_location_detail(
	filters: Any = None,
	start: Any = 0,
	page_length: Any = DEFAULT_PAGE_LENGTH,
	**kwargs: Any,
) -> dict[str, Any]:
	_require_read_permission()
	parsed = _parse_filters(filters, **kwargs)
	company = _resolve_company(parsed.get("company"))
	parsed["company"] = company
	_require_company_permission(company)
	payload = build_inventory_location_payload(
		_load_snapshot_rows(parsed),
		parsed,
		start=_parse_start(start),
		page_length=_request_page_length(page_length),
	)
	payload["snapshot_options"] = _list_snapshot_options(company)
	payload.update(_movement_permission_payload())
	return payload


@_whitelist
def export_inventory_location_detail(filters: Any = None, **kwargs: Any) -> dict[str, Any]:
	_require_read_permission()
	parsed = _parse_filters(filters, **kwargs)
	company = _resolve_company(parsed.get("company"))
	parsed["company"] = company
	_require_company_permission(company)
	payload = build_inventory_location_payload(_load_snapshot_rows(parsed), parsed, start=0, page_length=None)
	content = build_inventory_location_xlsx(payload)
	snapshot = payload.get("snapshot_date") or "latest"
	return {
		"file_name": f"material-inventory-detail-{snapshot}.xlsx",
		"mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
		"content_base64": base64.b64encode(content).decode("ascii"),
	}


@_whitelist
def get_categorized_inventory_detail(
	category: str,
	filters: Any = None,
	start: Any = 0,
	page_length: Any = DEFAULT_PAGE_LENGTH,
	**kwargs: Any,
) -> dict[str, Any]:
	_require_categorized_inventory_read_permission()
	parsed = _parse_filters(filters, **kwargs)
	company = _resolve_company(parsed.get("company"))
	parsed["company"] = company
	_require_company_permission(company)
	get_category_definition(category)
	context = _category_query_context(category, parsed)
	missing_item_group = _text(context.get("missing_item_group"))
	payload = build_categorized_inventory_payload(
		[] if missing_item_group else _load_category_stock_rows(context),
		[] if missing_item_group else _load_category_snapshot_rows(context),
		category=category,
		filters=parsed,
		start=_parse_start(start),
		page_length=_request_page_length(page_length),
	)
	payload["company"] = company
	payload["item_group_options"] = [] if missing_item_group else _list_category_item_groups(context)
	if missing_item_group:
		payload["warning"] = f"ERP 未维护分类物料组：{missing_item_group}"
	payload.update(_movement_permission_payload())
	return payload


@_whitelist
def export_categorized_inventory_detail(category: str, filters: Any = None, **kwargs: Any) -> dict[str, Any]:
	_require_categorized_inventory_read_permission()
	parsed = _parse_filters(filters, **kwargs)
	company = _resolve_company(parsed.get("company"))
	parsed["company"] = company
	_require_company_permission(company)
	definition = get_category_definition(category)
	context = _category_query_context(category, parsed)
	missing_item_group = _text(context.get("missing_item_group"))
	payload = build_categorized_inventory_payload(
		[] if missing_item_group else _load_category_stock_rows(context),
		[] if missing_item_group else _load_category_snapshot_rows(context),
		category=category,
		filters=parsed,
		start=0,
		page_length=None,
	)
	content = build_categorized_inventory_xlsx(payload)
	return {
		"file_name": f"{definition['file_prefix']}.xlsx",
		"mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
		"content_base64": base64.b64encode(content).decode("ascii"),
	}
