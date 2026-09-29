"""库存库位快照分组、筛选、数量格式和 Excel 导出测试。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from overseas_costing.services import inventory_location_service as service
from overseas_costing.services.inventory_location_service import (
    build_inventory_location_payload,
    build_inventory_location_xlsx,
    format_quantity,
)


ROWS = [
    {
        "snapshot_key": "YWFM-2026-09-29",
        "snapshot_date": "2026-09-29",
        "company": "YW Fabricación MX 核心制造",
        "item_code": "FL002917",
        "item_name": "PET片材 / PET SHEET",
        "warehouse": "IML 仓库 - YWFM",
        "original_location": "AI-11-A02",
        "location_qty": 1865,
        "stock_uom": "张：hoja",
        "item_group": "FL Suministros Auxiliares辅料",
        "dpci": "",
        "external_code": "EXT-2917",
        "original_identifier_alias": "PET 2917",
    },
    {
        "snapshot_key": "YWFM-2026-09-29",
        "snapshot_date": "2026-09-29",
        "company": "YW Fabricación MX 核心制造",
        "item_code": "FL002917",
        "item_name": "PET片材 / PET SHEET",
        "warehouse": "IML 仓库 - YWFM",
        "original_location": "AI-12-T02",
        "location_qty": 695,
        "stock_uom": "张：hoja",
        "item_group": "FL Suministros Auxiliares辅料",
        "dpci": "",
        "external_code": "EXT-2917",
        "original_identifier_alias": "PET 2917",
    },
    {
        "snapshot_key": "YWFM-2026-09-29",
        "snapshot_date": "2026-09-29",
        "company": "YW Fabricación MX 核心制造",
        "item_code": "FL007979",
        "item_name": "色母粒 / MASTERBATCH AZUL",
        "warehouse": "综合仓库 - YWFM",
        "original_location": "AI-4-C01",
        "location_qty": 19.6,
        "stock_uom": "kg",
        "item_group": "FL Suministros Auxiliares辅料",
        "dpci": "",
        "external_code": "FL000164",
        "original_identifier_alias": "FL000164",
    },
]


def test_payload_groups_locations_and_keeps_complete_warehouse_total_when_location_filtered() -> None:
    payload = build_inventory_location_payload(
        ROWS,
        {
            "company": "YW Fabricación MX 核心制造",
            "original_location": "AI-11-A02",
        },
    )

    assert payload["snapshot_key"] == "YWFM-2026-09-29"
    assert len(payload["groups"]) == 1
    group = payload["groups"][0]
    assert group["item_code"] == "FL002917"
    assert group["total_qty"] == 2560
    assert group["locations"] == [
        {"original_location": "AI-11-A02", "location_qty": 1865}
    ]


def test_payload_filters_keyword_warehouse_and_item_group_before_grouping() -> None:
    payload = build_inventory_location_payload(
        ROWS,
        {
            "keyword": "masterbatch",
            "warehouse": "综合仓库 - YWFM",
            "item_group": "FL Suministros Auxiliares辅料",
        },
    )

    assert [row["item_code"] for row in payload["groups"]] == ["FL007979"]
    assert payload["groups"][0]["total_qty"] == 19.6


def test_quantity_format_uses_integer_count_units_and_two_decimal_continuous_units() -> None:
    assert format_quantity(2560, "张：hoja") == "2,560"
    assert format_quantity(1.5, "个：pieza") == "1.50"
    assert format_quantity(19.6, "kg") == "19.60"
    assert format_quantity(19, "kg") == "19.00"


def test_company_document_permission_is_checked_before_raw_snapshot_queries(monkeypatch) -> None:
    checked: list[str] = []

    class FakeCompany:
        @staticmethod
        def check_permission(permission_type: str) -> None:
            checked.append(permission_type)

    class FakeFrappe:
        @staticmethod
        def get_doc(doctype: str, name: str):
            assert (doctype, name) == (
                "Company",
                "YW Fabricación MX 核心制造",
            )
            return FakeCompany()

    monkeypatch.setattr(service, "frappe", FakeFrappe())
    service._require_company_permission("YW Fabricación MX 核心制造")

    assert checked == ["read"]


def test_xlsx_export_merges_group_fields_but_keeps_location_rows_separate() -> None:
    payload = build_inventory_location_payload(ROWS, {})
    workbook = load_workbook(BytesIO(build_inventory_location_xlsx(payload)))
    sheet = workbook.active

    assert sheet.title == "库存库位明细"
    assert [sheet.cell(1, column).value for column in range(1, 12)] == [
        "正式物料编码",
        "物料名称（双语）",
        "仓库",
        "原始库位",
        "库位数量",
        "总库存",
        "库存单位",
        "物料组",
        "DPCI",
        "外部编码",
        "原始标识/别名",
    ]
    # FL002917 has two locations; shared fields are true merged cells.
    merged = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    for column in ("A", "B", "C", "F", "G", "H", "I", "J", "K"):
        assert f"{column}2:{column}3" in merged
    assert sheet["D2"].value == "AI-11-A02"
    assert sheet["D3"].value == "AI-12-T02"
    assert sheet["E2"].value == 1865
    assert sheet["E3"].value == 695
    assert sheet["F2"].value == 2560
    assert sheet["E2"].number_format == "#,##0"
    assert sheet["E4"].number_format == "#,##0.00"
