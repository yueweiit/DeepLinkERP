"""库存原始库位快照查询、分组与导出服务。"""

from __future__ import annotations

import base64
import json
from collections import OrderedDict
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    import frappe
except Exception:  # pragma: no cover - 本地单测环境不安装 Frappe
    frappe = None


SNAPSHOT_DOCTYPE = "Inventory Original Location Snapshot"
DEFAULT_COMPANY = "YW Fabricación MX 核心制造"
COUNT_UOM_TOKENS = (
    "个",
    "件",
    "套",
    "卷",
    "包",
    "张",
    "片",
    "支",
    "条",
    "台",
    "pieza",
    "conjunto",
    "rollo",
    "paquete",
    "hoja",
    "ramo",
    "barra",
    "unidad",
)
COMMON_EXPORT_COLUMNS = (1, 2, 3, 6, 7, 8, 9, 10, 11)
EXPORT_HEADERS = (
    "正式物料编码",
    "物料名称（双语）",
    "仓库",
    "原始库位",
    "库位数量",
    "总库存",
    "库存单位",
    "物料组",
    "DPCI",
    "外部编码",
    "原始标识/别名",
)


def _whitelist(function):
    if frappe is None:
        return function
    return frappe.whitelist()(function)


def is_count_uom(stock_uom: str | None) -> bool:
    normalized = str(stock_uom or "").strip().lower()
    return any(token in normalized for token in COUNT_UOM_TOKENS)


def quantity_precision(value: Any, stock_uom: str | None) -> int:
    number = float(value or 0)
    if not is_count_uom(stock_uom):
        return 2
    return 0 if abs(number - round(number)) <= 1e-9 else 2


def format_quantity(value: Any, stock_uom: str | None) -> str:
    precision = quantity_precision(value, stock_uom)
    return f"{float(value or 0):,.{precision}f}"


def _text(value: Any) -> str:
    return str(value or "").strip()


def _contains(value: Any, keyword: Any) -> bool:
    return _text(keyword).casefold() in _text(value).casefold()


def _matches_group_filters(row: dict[str, Any], filters: dict[str, Any]) -> bool:
    for fieldname in ("company", "warehouse", "item_group", "snapshot_key"):
        selected = _text(filters.get(fieldname))
        if selected and _text(row.get(fieldname)) != selected:
            return False

    snapshot_date = _text(filters.get("snapshot_date"))
    if snapshot_date and _text(row.get("snapshot_date")) != snapshot_date:
        return False

    keyword = _text(filters.get("keyword") or filters.get("item_code"))
    if keyword and not any(
        _contains(row.get(fieldname), keyword)
        for fieldname in (
            "item_code",
            "item_name",
            "dpci",
            "external_code",
            "original_identifier_alias",
        )
    ):
        return False
    return True


def build_inventory_location_payload(
    rows: Iterable[dict[str, Any]], filters: dict[str, Any] | None = None
) -> dict[str, Any]:
    """按物料和仓库分组；库位筛选不改变完整仓库小计。"""

    filters = dict(filters or {})
    all_rows = [dict(row) for row in rows]
    group_rows = [row for row in all_rows if _matches_group_filters(row, filters)]
    groups: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()

    for row in sorted(
        group_rows,
        key=lambda value: (
            _text(value.get("warehouse")),
            _text(value.get("item_code")),
            _text(value.get("original_location")),
        ),
    ):
        key = (_text(row.get("item_code")), _text(row.get("warehouse")))
        group = groups.setdefault(
            key,
            {
                fieldname: row.get(fieldname) or ""
                for fieldname in (
                    "item_code",
                    "item_name",
                    "warehouse",
                    "stock_uom",
                    "item_group",
                    "dpci",
                    "external_code",
                    "original_identifier_alias",
                )
            }
            | {"total_qty": 0.0, "locations": []},
        )
        group["total_qty"] += float(row.get("location_qty") or 0)
        group["locations"].append(
            {
                "original_location": row.get("original_location") or "",
                "location_qty": float(row.get("location_qty") or 0),
            }
        )

    location_filter = _text(filters.get("original_location"))
    visible_groups: list[dict[str, Any]] = []
    for group in groups.values():
        if location_filter:
            visible_locations = [
                location
                for location in group["locations"]
                if _contains(location["original_location"], location_filter)
            ]
            if not visible_locations:
                continue
            group["locations"] = visible_locations
        visible_groups.append(group)

    snapshot_source = next(
        (row for row in group_rows if row.get("snapshot_key")),
        next((row for row in all_rows if row.get("snapshot_key")), {}),
    )
    return {
        "snapshot_key": snapshot_source.get("snapshot_key") or "",
        "snapshot_date": snapshot_source.get("snapshot_date") or "",
        "company": filters.get("company")
        or snapshot_source.get("company")
        or DEFAULT_COMPANY,
        "groups": visible_groups,
        "group_count": len(visible_groups),
        "location_count": sum(len(group["locations"]) for group in visible_groups),
    }


def _number_format(stock_uom: str | None, value: Any) -> str:
    return "#,##0" if quantity_precision(value, stock_uom) == 0 else "#,##0.00"


def build_inventory_location_xlsx(payload: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "库存库位明细"
    sheet.freeze_panes = "A2"
    sheet.append(EXPORT_HEADERS)

    header_fill = PatternFill("solid", fgColor="075985")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    current_row = 2
    for group in payload.get("groups") or []:
        locations = list(group.get("locations") or []) or [
            {"original_location": "", "location_qty": 0}
        ]
        start_row = current_row
        for index, location in enumerate(locations):
            if index == 0:
                values = [
                    group.get("item_code") or "",
                    group.get("item_name") or "",
                    group.get("warehouse") or "",
                    location.get("original_location") or "",
                    float(location.get("location_qty") or 0),
                    float(group.get("total_qty") or 0),
                    group.get("stock_uom") or "",
                    group.get("item_group") or "",
                    group.get("dpci") or "",
                    group.get("external_code") or "",
                    group.get("original_identifier_alias") or "",
                ]
            else:
                values = [
                    "",
                    "",
                    "",
                    location.get("original_location") or "",
                    float(location.get("location_qty") or 0),
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            sheet.append(values)
            sheet.cell(current_row, 5).number_format = _number_format(
                group.get("stock_uom"), location.get("location_qty")
            )
            if index == 0:
                sheet.cell(current_row, 6).number_format = _number_format(
                    group.get("stock_uom"), group.get("total_qty")
                )
            current_row += 1

        end_row = current_row - 1
        if end_row > start_row:
            for column in COMMON_EXPORT_COLUMNS:
                sheet.merge_cells(
                    start_row=start_row,
                    start_column=column,
                    end_row=end_row,
                    end_column=column,
                )
        for row_number in range(start_row, end_row + 1):
            for cell in sheet[row_number]:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    widths = (18, 32, 24, 18, 15, 15, 15, 28, 18, 20, 28)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _parse_filters(filters: Any = None, **kwargs: Any) -> dict[str, Any]:
    if isinstance(filters, str):
        filters = json.loads(filters or "{}")
    parsed = dict(filters or {})
    parsed.update({key: value for key, value in kwargs.items() if value is not None})
    return parsed


def _require_read_permission() -> None:
    if frappe is None:
        return
    if not frappe.has_permission(SNAPSHOT_DOCTYPE, "read"):
        frappe.throw("没有权限查看库存库位快照。", frappe.PermissionError)


def _require_company_permission(company: str) -> None:
    """Raw SQL 读取前仍须遵守用户对 Company 的文档权限。"""

    if frappe is None:
        return
    frappe.get_doc("Company", company).check_permission("read")


def _latest_snapshot_key(company: str) -> str:
    rows = frappe.db.sql(
        f"""
        SELECT snapshot_key
        FROM `tab{SNAPSHOT_DOCTYPE}`
        WHERE company = %(company)s
        ORDER BY snapshot_date DESC, creation DESC
        LIMIT 1
        """,
        {"company": company},
        as_dict=True,
    )
    return rows[0]["snapshot_key"] if rows else ""


def _load_snapshot_rows(filters: dict[str, Any]) -> list[dict[str, Any]]:
    company = _text(filters.get("company")) or DEFAULT_COMPANY
    snapshot_key = _text(filters.get("snapshot_key")) or _latest_snapshot_key(company)
    if not snapshot_key:
        return []
    filters["company"] = company
    filters["snapshot_key"] = snapshot_key
    return frappe.db.sql(
        f"""
        SELECT
            snapshot.snapshot_key,
            snapshot.snapshot_date,
            snapshot.company,
            snapshot.item_code,
            item.item_name,
            snapshot.warehouse,
            snapshot.original_location,
            snapshot.location_qty,
            snapshot.stock_uom,
            item.item_group,
            COALESCE(item.custom_dpci, '') AS dpci,
            COALESCE(item.custom_external_code, '') AS external_code,
            COALESCE(item.custom_original_identifier_alias, '') AS original_identifier_alias
        FROM `tab{SNAPSHOT_DOCTYPE}` snapshot
        INNER JOIN `tabItem` item ON item.name = snapshot.item_code
        WHERE snapshot.snapshot_key = %(snapshot_key)s
          AND snapshot.company = %(company)s
        ORDER BY snapshot.warehouse, snapshot.item_code, snapshot.original_location
        """,
        {"snapshot_key": snapshot_key, "company": company},
        as_dict=True,
    )


def _list_snapshot_options(company: str) -> list[dict[str, Any]]:
    return frappe.db.sql(
        f"""
        SELECT snapshot_key, MAX(snapshot_date) AS snapshot_date
        FROM `tab{SNAPSHOT_DOCTYPE}`
        WHERE company = %(company)s
        GROUP BY snapshot_key
        ORDER BY snapshot_date DESC, snapshot_key DESC
        """,
        {"company": company},
        as_dict=True,
    )


@_whitelist
def get_inventory_location_detail(filters: Any = None, **kwargs: Any) -> dict[str, Any]:
    _require_read_permission()
    parsed = _parse_filters(filters, **kwargs)
    company = _text(parsed.get("company")) or DEFAULT_COMPANY
    parsed["company"] = company
    _require_company_permission(company)
    payload = build_inventory_location_payload(_load_snapshot_rows(parsed), parsed)
    payload["snapshot_options"] = _list_snapshot_options(company)
    return payload


@_whitelist
def export_inventory_location_detail(filters: Any = None, **kwargs: Any) -> dict[str, Any]:
    _require_read_permission()
    parsed = _parse_filters(filters, **kwargs)
    company = _text(parsed.get("company")) or DEFAULT_COMPANY
    parsed["company"] = company
    _require_company_permission(company)
    payload = build_inventory_location_payload(_load_snapshot_rows(parsed), parsed)
    content = build_inventory_location_xlsx(payload)
    snapshot = payload.get("snapshot_date") or "latest"
    return {
        "file_name": f"inventory-location-detail-{snapshot}.xlsx",
        "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "content_base64": base64.b64encode(content).decode("ascii"),
    }
