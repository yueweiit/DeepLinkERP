"""Read workbook names only when a user opens a controlled local attachment."""
from . import effective_logistics_source as effective_source
from . import packing_source_service


def list_attachment_sheets(batch_name, source_kind, source_id):
    if source_kind not in {'approval_attachment', 'manual_attachment'}:
        raise ValueError('请选择当前来源的装箱附件。')
    bundle = effective_source.current_source_bundle(batch_name)
    context = (bundle or {}).get('context') or {}
    effective_source.require_readable(context)
    source = packing_source_service._attachment_source_v2(batch_name, source_id)
    effective_source.validate_packing_source(batch_name, source_kind, source_id, attachment=source)
    if not source or not source.get('file_url'):
        raise ValueError('当前附件尚未归档到本地，请先获取资料。')
    path = packing_source_service.resolve_packing_attachment_path(source, bundle)
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError('归档工作簿超过 20 MB 限制。')
    if path.suffix.lower() == '.xls' and context.get('root_kind') == 'expense':
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            sheets = book.sheet_names()
        finally:
            book.release_resources()
    elif path.suffix.lower() in {'.xlsx', '.xlsm'}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            sheets = book.sheetnames
        finally:
            book.close()
    else:
        raise ValueError('该附件不是当前支持的装箱工作簿。')
    latest = effective_source.current_source_bundle(batch_name)
    if context != ((latest or {}).get('context') or {}):
        raise ValueError('读取期间当前来源已变化，请重新获取资料。')
    return {'ok': True, 'source_id': source_id, 'sheets': sheets, 'source_context': context}
