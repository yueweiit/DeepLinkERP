"""中文用途：预览可信装箱来源，并确认不可变、可追溯的当前装箱快照。"""

from __future__ import annotations

import copy
import hashlib
import json
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Callable

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - 仅在读取本地 Excel Sheet 名时需要
    load_workbook = None

try:
    import frappe
except Exception:  # pragma: no cover
    frappe = None

from overseas_costing.services import packing_source_service
from overseas_costing.services import effective_logistics_source as effective_source
from overseas_costing.services.packing_sheet_recommendation import (
    recommend_packing_sheets,
    summarize_packing_preview,
)


Resolver = Callable[..., dict[str, Any]]


def preview_packing_source_v2(
    batch_name: str,
    source_kind: str,
    source_id: str,
    *,
    sheet_name: str | None = None,
    browser_payload: dict[str, Any] | None = None,
    resolver: Resolver | None = None,
) -> dict[str, Any]:
    """浏览器只提交受控 ID；正文、路径、汇总和工作簿 URL 一律重新从服务器读取。"""

    del browser_payload
    kind = packing_source_service.normalize_packing_source_kind(source_kind)
    trusted = (resolver or packing_source_service.resolve_trusted_packing_source)(
        batch_name=str(batch_name),
        source_kind=kind,
        source_id=str(source_id),
        sheet_name=sheet_name,
    )
    source_hash = str(trusted.get("source_hash") or "")
    if len(source_hash) != 64:
        raise ValueError("可信装箱来源缺少 SHA-256。")
    preview = copy.deepcopy(trusted.get("preview") or {})
    source = _public_source(trusted.get("source") or {})
    revision = packing_source_service._encode_revision(
        kind,
        str(source.get("source_id") or source_id),
        source_hash,
        batch_name=str(batch_name),
        version_name=str((trusted.get('source_context') or {}).get('cost_version') or ''),
        source_context=trusted.get('source_context') or {},
    )
    return {
        **preview,
        "ok": True,
        "preview_ready": True,
        "source": source,
        "source_kind": kind,
        "source_id": source.get("source_id") or str(source_id),
        "source_hash": source_hash,
        "source_revision": revision,
        "source_context": trusted.get('source_context') or {},
    }


def confirm_packing_snapshot(
    batch_name: str,
    source_revision: str,
    resolutions_json: str | dict[str, Any] | None,
    *,
    repository: Any | None = None,
    resolver: Resolver | None = None,
) -> dict[str, Any]:
    claims = packing_source_service.get_source_revision_claims(source_revision)
    if not claims or str(claims.get("batch") or "") != str(batch_name):
        return {"ok": False, "source_changed": True, "message": "装箱预览不属于当前批次，请重新预览。"}
    kind = packing_source_service.normalize_packing_source_kind(str(claims.get("kind") or ""))
    source_id = str(claims.get("id") or "")
    expected_hash = str(claims.get("hash") or "")
    resolutions = _parse_resolutions(resolutions_json)
    repo = repository or FrappePackingSnapshotRepository()
    repo.lock_batch(str(batch_name))
    try:
        trusted = (resolver or packing_source_service.resolve_trusted_packing_source)(
            batch_name=str(batch_name),
            source_kind=kind,
            source_id=source_id,
            sheet_name=resolutions.get("sheet_name"),
        )
        actual_hash = str(trusted.get("source_hash") or "")
        source_context = trusted.get('source_context') or {}
        try:
            effective_source.require_available(source_context)
        except ValueError as exc:
            repo.rollback()
            return {'ok': False, 'analysis_only': True, 'message': str(exc)}
        if (claims.get('source_context') or {}) != source_context:
            repo.rollback()
            return {'ok': False, 'source_changed': True, 'message': '采用来源已变化，请重新预览。'}
        if claims.get('version') and claims['version'] != str(source_context.get('cost_version') or ''):
            repo.rollback()
            return {'ok': False, 'source_changed': True, 'message': '成本版本已变化，请重新预览。'}
        if not actual_hash or actual_hash != expected_hash:
            repo.rollback()
            return {"ok": False, "source_changed": True, "message": "装箱来源已更新，请重新预览后确认。"}
        idempotency_key = _idempotency_key(batch_name, source_revision)
        existing = repo.get_by_idempotency_key(idempotency_key)
        if existing:
            stored_resolutions = _parse_resolutions(_record_value(existing, "resolutions_json"))
            repo.rollback()
            if _json(stored_resolutions) != _json(resolutions):
                return {
                    "ok": False,
                    "source_changed": False,
                    "resolution_conflict": True,
                    "message": "该来源版本已经确认，不能用另一组决定覆盖。请刷新来源后重新预览。",
                }
            return {"ok": True, "idempotent": True, "snapshot": _public_snapshot(existing)}
        preview = _apply_resolutions(trusted.get("preview") or {}, resolutions)
        blocking = (preview.get("validation") or {}).get("blocking") or []
        if blocking:
            repo.rollback()
            return {
                "ok": False,
                "needs_group_confirmation": bool((preview.get("validation") or {}).get("needs_group_confirmation")),
                "validation": preview.get("validation"),
                "message": "装箱分组或汇总仍有待确认项目。",
            }
        gross = _positive_decimal(((preview.get("totals") or {}).get("gross_weight_kg") or {}).get("value"))
        volume = _positive_decimal(((preview.get("totals") or {}).get("volume_m3") or {}).get("value"))
        if gross is None or volume is None:
            repo.rollback()
            return {"ok": False, "message": "比较运费前必须确认整票毛重和总体积。"}
        current = repo.get_current(str(batch_name))
        if current:
            repo.supersede(current)
        source = _public_source(trusted.get("source") or {})
        now = datetime.now(timezone.utc).isoformat()
        next_version = repo.next_version(str(batch_name)) if hasattr(repo, "next_version") else 1
        totals = preview.get("totals") or {}
        values = {
            "batch": str(batch_name),
            "cost_version": source_context.get('cost_version') or repo.current_version(str(batch_name)),
            "source_context_json": _json(source_context),
            "version": next_version,
            "idempotency_key": idempotency_key,
            "source_kind": kind,
            "source_id": source_id,
            "source_revision": source_revision,
            "source_label": source.get("source_label") or source_id,
            "source_locator_json": _json(source),
            "source_read_at": now,
            "source_updated_at": source.get("source_updated_at") or None,
            "source_hash": actual_hash,
            "status": "Confirmed",
            "is_current": 1,
            "material_row_count": int(preview.get("material_row_count") or 0),
            "package_count": int(preview.get("package_count") or 0),
            "total_net_weight_kg": _decimal_text(
                _decimal((totals.get("net_weight_kg") or {}).get("value"))
            ),
            "net_total_kind": (totals.get("net_weight_kg") or {}).get("kind") or "",
            "total_gross_weight_kg": _decimal_text(gross),
            "total_volume_m3": _decimal_text(volume),
            "material_rows_json": _json(preview.get("material_rows") or []),
            "package_groups_json": _json(preview.get("groups") or []),
            "totals_json": _json(totals),
            "validation_json": _json(preview.get("validation") or {}),
            "resolutions_json": _json(resolutions),
            "confirmed_by": _session_user(),
            "confirmed_at": now,
        }
        snapshot = repo.insert_snapshot(values)
        repo.write_audit(
            {
                "batch": str(batch_name),
                "version": repo.current_version(str(batch_name)),
                "action_type": "PACKING_CONFIRM",
                "field_name": "packing_snapshot",
                "old_value": _record_value(current, "name") if current else "",
                "new_value": _record_value(snapshot, "name"),
                "operator_name": _session_user(),
                "action_remark": f"确认装箱来源：{values['source_label']}",
            }
        )
        repo.commit()
        return {"ok": True, "idempotent": False, "snapshot": _public_snapshot(snapshot)}
    except Exception:
        repo.rollback()
        raise


def _apply_resolutions(preview: dict[str, Any], resolutions: dict[str, Any]) -> dict[str, Any]:
    resolved = copy.deepcopy(preview)
    group_resolutions = resolutions.get("groups") if isinstance(resolutions.get("groups"), dict) else {}
    unresolved = []
    resolved_groups = []
    for group in resolved.get("groups") or []:
        if not group.get("needs_confirmation"):
            resolved_groups.append(group)
            continue
        decision = group_resolutions.get(str(group.get("group_id") or ""))
        action = str((decision or {}).get("action") or "") if isinstance(decision, dict) else str(decision or "")
        if action in {"confirm_shared", "link", "shared"}:
            group["needs_confirmation"] = False
            for field in ("net_weight_kg", "gross_weight_kg", "volume_m3", "package_count"):
                if isinstance(group.get(field), dict) and group[field].get("value") is not None:
                    group[field]["count_once"] = True
            group.setdefault("evidence", []).append({"kind": "user_confirmed_shared"})
            resolved_groups.append(group)
        elif action in {"split", "partition"}:
            rows = list(group.get("row_numbers") or [])
            partitions = [[row] for row in rows] if action == "split" else (decision or {}).get("partitions")
            normalized = _validate_partitions(rows, partitions)
            if normalized:
                resolved_groups.extend(_partition_group(group, normalized, action))
            else:
                unresolved.append(group.get("group_id"))
                resolved_groups.append(group)
        else:
            unresolved.append(group.get("group_id"))
            resolved_groups.append(group)
    resolved["groups"] = resolved_groups
    resolved["package_group_count"] = len(resolved_groups)
    _recompute_resolved_totals(resolved)
    resolved_total_fields = _apply_total_resolutions(resolved, resolutions)
    validation = resolved.setdefault("validation", {})
    blocking = [
        item
        for item in (validation.get("blocking") or [])
        if item.get("code") != "group_confirmation_required"
        and not (item.get("code") == "conflicting_merge_ranges" and not unresolved)
        and not (item.get("code") == "total_mismatch" and item.get("field") in resolved_total_fields)
    ]
    if unresolved:
        blocking.append(
            {
                "code": "group_confirmation_required",
                "message": f"仍有 {len(unresolved)} 个共享包装候选未确认。",
            }
        )
    validation["blocking"] = blocking
    validation["needs_group_confirmation"] = bool(unresolved)
    return resolved


def _validate_partitions(rows: list[int], partitions: Any) -> list[list[int]] | None:
    if not isinstance(partitions, list) or not partitions:
        return None
    normalized = []
    for partition in partitions:
        if not isinstance(partition, list) or not partition:
            return None
        try:
            values = [int(row) for row in partition]
        except (TypeError, ValueError):
            return None
        if len(values) != len(set(values)):
            return None
        normalized.append(values)
    flattened = [row for partition in normalized for row in partition]
    return normalized if len(flattened) == len(set(flattened)) and sorted(flattened) == sorted(rows) else None


def _partition_group(group: dict[str, Any], partitions: list[list[int]], action: str) -> list[dict[str, Any]]:
    result = []
    for index, partition in enumerate(partitions, start=1):
        subgroup = copy.deepcopy(group)
        subgroup["group_id"] = f"{group.get('group_id')}-part-{index}"
        subgroup["row_numbers"] = partition
        subgroup["needs_confirmation"] = False
        subgroup["suggestion_reason"] = None
        subgroup["merge_conflict"] = False
        subgroup.setdefault("evidence", []).append(
            {"kind": f"user_{action}", "source_group_id": group.get("group_id"), "rows": partition}
        )
        metrics = [
            subgroup.get("net_weight_kg"),
            subgroup.get("gross_weight_kg"),
            subgroup.get("volume_m3"),
            subgroup.get("package_count"),
            *((subgroup.get("dimensions") or {}).values()),
        ]
        for metric in metrics:
            if not isinstance(metric, dict):
                continue
            source_row = metric.get("source_row")
            owns_value = source_row in partition if source_row is not None else index == 1
            if owns_value and metric.get("value") is not None:
                metric["count_once"] = True
            else:
                metric["value"] = None
                metric["count_once"] = False
                metric["source_row"] = None
        result.append(subgroup)
    return result


def _apply_total_resolutions(preview: dict[str, Any], resolutions: dict[str, Any]) -> set[str]:
    decisions = resolutions.get("totals") if isinstance(resolutions.get("totals"), dict) else {}
    resolved_fields: set[str] = set()
    for field, total in (preview.get("totals") or {}).items():
        declared = _decimal((total or {}).get("declared_value"))
        calculated = _decimal((total or {}).get("calculated_value"))
        if declared is None or calculated is None:
            continue
        tolerance = max(Decimal("0.01"), abs(declared) * Decimal("0.001"))
        if abs(declared - calculated) <= tolerance:
            resolved_fields.add(field)
    for field, decision in decisions.items():
        total = (preview.get("totals") or {}).get(field)
        if not isinstance(total, dict):
            continue
        action = str((decision or {}).get("action") or "") if isinstance(decision, dict) else str(decision or "")
        if action == "use_declared" and total.get("declared_value") is not None:
            total["value"] = total["declared_value"]
            total["kind"] = "user_selected_source_total"
            resolved_fields.add(field)
        elif action == "use_calculated" and total.get("calculated_value") is not None:
            total["value"] = total["calculated_value"]
            total["kind"] = "user_selected_calculated"
            resolved_fields.add(field)
    return resolved_fields


def _recompute_resolved_totals(preview: dict[str, Any]) -> None:
    """用用户已确认的 count_once 包装组重新汇总，但保留表内整票声明值。"""

    groups = preview.get("groups") or []
    totals = preview.setdefault("totals", {})
    for field in ("net_weight_kg", "gross_weight_kg", "volume_m3"):
        values = [
            _decimal((group.get(field) or {}).get("value"))
            for group in groups
            if (group.get(field) or {}).get("count_once")
        ]
        calculated = sum((value for value in values if value is not None), Decimal("0")) if any(value is not None for value in values) else None
        total = totals.setdefault(field, {})
        total["calculated_value"] = _decimal_text(calculated)
        if field == "net_weight_kg" or total.get("kind") != "source_total":
            total["value"] = _decimal_text(calculated)
            total["kind"] = "calculated_detail_sum"

    package_values = [
        _decimal((group.get("package_count") or {}).get("value"))
        for group in groups
        if (group.get("package_count") or {}).get("count_once")
    ]
    package_calculated = sum(
        (value for value in package_values if value is not None), Decimal("0")
    ) if any(value is not None for value in package_values) else None
    package_total = totals.setdefault("package_count", {})
    package_total["calculated_value"] = _decimal_text(package_calculated)
    if package_total.get("kind") != "source_total" and package_calculated is not None:
        package_total["value"] = _decimal_text(package_calculated)
        package_total["kind"] = "calculated_group_sum"
    resolved_package_count = _decimal(package_total.get("value"))
    if resolved_package_count is not None and resolved_package_count == resolved_package_count.to_integral_value():
        preview["package_count"] = int(resolved_package_count)


def _parse_resolutions(value: str | dict[str, Any] | None) -> dict[str, Any]:
    if isinstance(value, dict):
        result = value
    else:
        try:
            result = json.loads(value or "{}")
        except (TypeError, ValueError) as error:
            raise ValueError("装箱确认内容不是有效 JSON。") from error
    if not isinstance(result, dict):
        raise ValueError("装箱确认内容必须是对象。")
    return result


def _idempotency_key(batch_name: str, source_revision: str) -> str:
    return hashlib.sha256(f"{batch_name}|{source_revision}".encode("utf-8")).hexdigest()


def _public_source(source: dict[str, Any]) -> dict[str, Any]:
    allowed = {
        "source_kind",
        "source_id",
        "source_label",
        "source_updated_at",
        "sheet_name",
        "workbook_id",
        "sheet_id",
        "instance_id",
        "comment_user",
    }
    return {key: value for key, value in source.items() if key in allowed}


def _public_snapshot(snapshot: Any) -> dict[str, Any]:
    allowed = {
        "name",
        "batch",
        "version",
        "idempotency_key",
        "source_kind",
        "source_id",
        "source_label",
        "source_hash",
        "status",
        "is_current",
        "material_row_count",
        "package_count",
        "total_net_weight_kg",
        "net_total_kind",
        "total_gross_weight_kg",
        "total_volume_m3",
        "confirmed_by",
        "confirmed_at",
        "cost_version",
        "source_context_json",
    }
    if isinstance(snapshot, dict):
        return {key: snapshot.get(key) for key in allowed if key in snapshot}
    return {key: getattr(snapshot, key) for key in allowed if hasattr(snapshot, key)}


def _record_value(record: Any, fieldname: str) -> Any:
    return record.get(fieldname) if isinstance(record, dict) else getattr(record, fieldname, None)


def _decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    return result if result.is_finite() else None


def _positive_decimal(value: Any) -> Decimal | None:
    result = _decimal(value)
    return result if result is not None and result > 0 else None


def _decimal_text(value: Decimal | None) -> str | None:
    return format(value, "f") if value is not None else None


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _session_user() -> str:
    user = getattr(getattr(frappe, "session", None), "user", "") if frappe is not None else ""
    return str(user or "")


class FrappePackingSnapshotRepository:
    def __init__(self) -> None:
        if frappe is None:
            raise RuntimeError("当前环境未连接 Frappe。")

    def lock_batch(self, batch_name: str) -> None:
        effective_source.current_source_bundle(batch_name, lock=True)
        rows = frappe.db.sql(
            "SELECT name FROM `tabOverseas Cost Batch` WHERE name=%s FOR UPDATE",
            (batch_name,),
        )
        if not rows:
            raise ValueError("未找到当前批次。")

    def current_version(self, batch_name: str) -> str:
        return str(frappe.db.get_value("Overseas Cost Batch", batch_name, "current_version") or "")

    def next_version(self, batch_name: str) -> int:
        rows = frappe.db.sql(
            "SELECT COALESCE(MAX(version), 0) + 1 FROM `tabOverseas Packing Snapshot` WHERE batch=%s",
            (batch_name,),
        )
        return int(rows[0][0]) if rows else 1

    def get_by_idempotency_key(self, key: str) -> Any:
        name = frappe.db.get_value("Overseas Packing Snapshot", {"idempotency_key": key}, "name")
        return frappe.get_doc("Overseas Packing Snapshot", name) if name else None

    def get_current(self, batch_name: str) -> Any:
        name = frappe.db.get_value(
            "Overseas Packing Snapshot",
            {"batch": batch_name, "status": "Confirmed", "is_current": 1},
            "name",
        )
        return frappe.get_doc("Overseas Packing Snapshot", name) if name else None

    def supersede(self, snapshot: Any) -> None:
        snapshot.status = "Superseded"
        snapshot.is_current = 0
        snapshot.flags.allow_packing_supersede = True
        snapshot.save(ignore_permissions=True)

    def insert_snapshot(self, values: dict[str, Any]) -> Any:
        return frappe.get_doc({"doctype": "Overseas Packing Snapshot", **values}).insert(ignore_permissions=True)

    def write_audit(self, values: dict[str, Any]) -> None:
        frappe.get_doc({"doctype": "Overseas Cost Audit Log", **values}).insert(ignore_permissions=True)

    def commit(self) -> None:
        frappe.db.commit()

    def rollback(self) -> None:
        frappe.db.rollback()


def get_current_packing_snapshot(batch_name: str, version_name: str | None = None) -> dict[str, Any] | None:
    if frappe is None:
        raise RuntimeError("当前环境未连接 Frappe。")
    name = frappe.db.get_value(
        "Overseas Packing Snapshot",
        {"batch": str(batch_name), "status": "Confirmed", "is_current": 1},
        "name",
    )
    if not name:
        return None
    snapshot = _public_snapshot(frappe.get_doc("Overseas Packing Snapshot", name))
    bundle = effective_source.current_source_bundle(batch_name, version_name)
    if bundle and (bundle['context']['root_kind'] == 'expense' or (bundle['context'].get('packing') or {}).get('selected_source') or effective_source.json_dict(snapshot.get('source_context_json'))):
        context = bundle['context']
        saved = effective_source.json_dict(snapshot.get('source_context_json'))
        if saved.get('fingerprint') != context['fingerprint'] or not context['approved'] or context['invalid'] or not context['available']:
            return None
    return snapshot


def list_packing_sources(batch_name: str, *, approval_detail: dict | None = None, include_wiki: bool = True) -> dict[str, Any]:
    """返回受控来源 ID；不返回服务器路径、对象键、原始审批 JSON 或任何凭据。"""

    if frappe is None:
        raise RuntimeError("当前环境未连接 Frappe。")
    bundle = effective_source.current_source_bundle(batch_name)
    if bundle and (bundle['context'].get('packing') or {}).get('selected_source'):
        selected=selected_packing_ai_sources(batch_name,bundle)
        return {'approval_sources':[{k:v for k,v in row.items() if k not in ('scoped_goods','scoped_text','form_fields')} for row in selected],
                'manual_sources':[],'manual_attachments':[],'wiki_workbooks':[],'source_context':bundle['context']}
    if bundle and bundle['context']['root_kind'] == 'expense':
        sources = _bound_material_sources(batch_name, bundle)
        # The picker receives descriptors, never raw form/comment content.
        return {'approval_sources': [{k: v for k, v in row.items() if k not in {'form_fields', 'approval_decisions', 'comment_text'}}
                    for row in sources if row['source_kind'] not in {'approval_form', 'wiki_sheet'}],
                'manual_sources': [], 'manual_attachments': [],
                'wiki_workbooks': [{'workbook_id': row['source_id'].split(':')[0], 'label': row['source_label'], 'sheets': [row]}
                                   for row in sources if include_wiki and row['source_kind'] == 'wiki_sheet'],
                'source_context': bundle['context']}
    attachment_rows = frappe.get_list(
        "Overseas Cost Attachment",
        filters={"batch": str(batch_name)},
        fields=["name", "batch", "source_type", "oa_attachment_origin", "attachment_type", "file_name", "file_url", "modified", "parse_result_json"],
        limit_page_length=1000,
    )
    manual = []
    approval = []
    approval_by_identity = {}
    approval_by_name = {}
    for row in attachment_rows:
        if not _is_material_ai_attachment(row.get("file_name")):
            continue
        if str(row.get("source_type") or "").upper() == "OA" and packing_source_service._attachment_is_audit_only(row):
            continue
        snapshot = packing_source_service.import_service._json_loads_dict(row.get("parse_result_json"))
        archive = snapshot.get("archive") if isinstance(snapshot.get("archive"), dict) else {}
        download = snapshot.get("download") if isinstance(snapshot.get("download"), dict) else {}
        sheets = _attachment_sheet_names(row)
        item = {
            "source_id": row.get("name"),
            "attachment_name": row.get("name"),
            "source_label": row.get("file_name") or row.get("name"),
            "source_updated_at": row.get("modified"),
            "available": bool(row.get("file_url")),
            "download_required": not bool(row.get("file_url")),
            "supported_for_material_import": str(row.get("file_name") or "").lower().endswith((".xlsx", ".xlsm")),
            "attachment_type": row.get("attachment_type") or "",
            "content_hash": str(download.get("sha256") or archive.get("sha256") or ""),
            "sheets": sheets,
        }
        if str(row.get("source_type") or "").upper() == "OA":
            item.update({
                "source_kind": "approval_attachment",
                "origin": row.get("oa_attachment_origin") or snapshot.get("attachment_origin") or "Form",
                "process_instance_id": str(snapshot.get("process_instance_id") or snapshot.get("instance_id") or ""),
                "file_id": str(snapshot.get("file_id") or ""),
                "archive_status": archive.get("status") or ("archived" if item["available"] else "pending"),
                "can_download": not item["available"] and (archive.get("status") == "archived" or bool(snapshot.get("file_id"))),
                "actor_name": str(snapshot.get("comment_user_name") or ""),
                "occurred_at": str(snapshot.get("comment_time") or row.get("modified") or ""),
            })
            approval.append(item)
            approval_by_name[str(row.get("name") or "")] = item
            if item["process_instance_id"] and item["file_id"]:
                approval_by_identity[(item["process_instance_id"], item["file_id"])] = item
        else:
            manual.append({**item, "source_kind": "manual_attachment"})

    detail = approval_detail if approval_detail is not None else packing_source_service.dingtalk_approval_service.get_batch_dingtalk_approval_detail(str(batch_name))
    comments = []
    for approval_row in [
        detail.get("main_approval"),
        *(detail.get("linked_purchase_approvals") or []),
        *(detail.get("excluded_linked_purchase_approvals") or []),
    ]:
        if not isinstance(approval_row, dict):
            continue
        approval_excluded = bool(approval_row.get("excluded"))
        approval_exclusion_reason = str(
            approval_row.get("exclusion_reason")
            or ("审批已失效，不参与分析。" if approval_excluded else "")
        )
        for attachment in approval_row.get("attachments") or []:
            if not isinstance(attachment, dict) or not _is_material_ai_attachment(attachment.get("file_name")):
                continue
            instance_id = str(approval_row.get("instance_id") or "")
            file_id = str(attachment.get("file_id") or "")
            if not instance_id or not file_id:
                continue
            identity = (instance_id, file_id)
            existing = approval_by_identity.get(identity) or approval_by_name.get(str(attachment.get("attachment_name") or ""))
            available = bool((existing or {}).get("available"))
            archive_status = str(attachment.get("archive_status") or "pending")
            metadata = {
                "process_instance_id": instance_id,
                "file_id": file_id,
                "origin": attachment.get("origin") or "Form",
                "archive_status": archive_status,
                "can_download": not available and archive_status == "archived",
                "download_required": not available,
                "content_hash": str(
                    attachment.get("sha256") or attachment.get("content_sha256") or ""
                ),
                "approval_no": str(
                    approval_row.get("business_id") or approval_row.get("approval_no") or ""
                ),
                "actor_name": str(
                    attachment.get("comment_user_name")
                    or approval_row.get("originator_user_name")
                    or approval_row.get("originator_name")
                    or ""
                ),
                "occurred_at": str(
                    attachment.get("comment_time")
                    or attachment.get("source_updated_at")
                    or approval_row.get("finish_time")
                    or approval_row.get("create_time")
                    or ""
                ),
                "excluded": approval_excluded,
                "exclude_reason": approval_exclusion_reason,
            }
            if existing:
                existing.update(metadata)
                approval_by_identity[identity] = existing
                continue
            source_id = "oa:" + hashlib.sha256(_json(list(identity)).encode("utf-8")).hexdigest()
            item = {
                **metadata,
                "source_kind": "approval_attachment",
                "source_id": source_id,
                "attachment_name": "",
                "source_label": attachment.get("file_name") or file_id,
                "source_updated_at": attachment.get("source_updated_at") or attachment.get("comment_time") or detail.get("source_updated_at"),
                "available": False,
                "supported_for_material_import": str(attachment.get("file_name") or "").lower().endswith((".xlsx", ".xlsm")),
                "attachment_type": "",
                "sheets": [],
            }
            approval.append(item)
            approval_by_identity[identity] = item
        for timeline in approval_row.get("timeline") or []:
            if (
                not isinstance(timeline, dict)
                or not timeline.get("source_id")
                or not str(timeline.get("remark") or "").strip()
            ):
                continue
            comments.append(
                {
                    "source_kind": "approval_comment",
                    "source_id": timeline.get("source_id"),
                    "source_label": f"评论 · {timeline.get('user_name') or timeline.get('user_id') or '未知人员'}",
                    "source_updated_at": timeline.get("operation_time"),
                    "process_instance_id": approval_row.get("instance_id") or "",
                    "approval_no": approval_row.get("business_id") or approval_row.get("approval_no") or "",
                    "actor_name": timeline.get("user_name") or timeline.get("user_id") or "",
                    "occurred_at": timeline.get("operation_time") or "",
                    "available": not approval_excluded,
                    "excluded": approval_excluded,
                    "exclude_reason": approval_exclusion_reason,
                    "remark_preview": str(timeline.get("remark") or "")[:160],
                }
            )

    if not include_wiki:
        return {"approval_sources": [*approval, *comments], "manual_sources": manual, "wiki_workbooks": []}
    wiki = []
    wiki_error = ""
    try:
        from overseas_costing.integrations.dingtalk_packing_source import get_packing_runtime_clients

        clients = get_packing_runtime_clients()
        batch_context = _packing_batch_context(str(batch_name), detail)
        workbook_rows: list[dict[str, Any]] = []
        all_sheet_rows: list[dict[str, Any]] = []
        all_snapshot_states: dict[str, dict[str, Any]] = {}
        for workbook in clients.catalog.list_workbooks():
            workbook_id = str(workbook.get("workbook_id") or "")
            sheets = clients.catalog.list_sheets(workbook_id, limit=500)
            snapshot_states = _packing_snapshot_summaries(clients, workbook_id)
            all_snapshot_states.update(snapshot_states)
            rendered_sheets = []
            for sheet in sheets:
                source_id = f"{sheet.get('workbook_id')}:{sheet.get('sheet_id')}"
                snapshot_state = snapshot_states.get(source_id) or {}
                rendered_sheets.append(
                    {
                        "source_kind": "wiki_sheet",
                        "source_id": source_id,
                        "source_label": sheet.get("sheet_name"),
                        "source_updated_at": sheet.get("source_updated_at"),
                        "indexed_at": sheet.get("indexed_at"),
                        "workbook_id": sheet.get("workbook_id"),
                        "workbook_year": sheet.get("year") or workbook.get("year"),
                        "snapshot_updated_at": snapshot_state.get("snapshot_updated_at"),
                        "snapshot_status": snapshot_state.get("snapshot_status") or "not_cached",
                        "content_hash": snapshot_state.get("content_hash") or "",
                        "available": True,
                    }
                )
            workbook_rows.append(
                {
                    "workbook_id": workbook.get("workbook_id"),
                    "year": workbook.get("year"),
                    "label": workbook.get("label"),
                    "updated_at": workbook.get("updated_at"),
                    "sheets": rendered_sheets,
                }
            )
            all_sheet_rows.extend(rendered_sheets)

        snapshot_summaries = {
            source_id: state.get("summary") or {}
            for source_id, state in all_snapshot_states.items()
            if state.get("summary")
        }
        recommended = recommend_packing_sheets(
            all_sheet_rows,
            batch_context=batch_context,
            snapshot_summaries=snapshot_summaries,
        )
        recommended_by_id = {str(row.get("source_id") or ""): row for row in recommended}
        for workbook in workbook_rows:
            workbook["sheets"] = [
                recommended_by_id[str(row.get("source_id") or "")]
                for row in workbook.get("sheets") or []
                if str(row.get("source_id") or "") in recommended_by_id
            ]
            workbook["sheets"].sort(
                key=lambda row: next(
                    index
                    for index, candidate in enumerate(recommended)
                    if candidate.get("source_id") == row.get("source_id")
                )
            )
        wiki = _pin_recommended_workbook(workbook_rows)
    except Exception:
        wiki_error = "装箱计划表缓存暂不可用，请稍后重试。"
    return {
        "manual_attachments": manual,
        "approval_sources": [*approval, *comments],
        "wiki_workbooks": wiki,
        "wiki_error": wiki_error,
    }


MATERIAL_AI_DOCUMENT_SUFFIXES = (
    ".xlsx",
    ".xlsm",
    ".xls",
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".bmp",
    ".tif",
    ".tiff",
    ".doc",
    ".docx",
    ".txt",
)


def _list_approval_body_ai_sources(batch_name: str, *, detail: dict | None = None) -> list[dict[str, Any]]:
    """Expose only the current logistics approval and its server-verified purchase links."""

    from overseas_costing.services import dingtalk_approval_service

    if detail is None:
        bundle = effective_source.current_source_bundle(batch_name)
        detail = (effective_source.approval_detail_for_bundle(bundle) if bundle and bundle['context']['root_kind'] == 'expense'
                  else dingtalk_approval_service.get_batch_dingtalk_approval_detail(str(batch_name)) or {})
    if not detail.get("ok"):
        return []

    def source(approval: dict, role: str) -> dict[str, Any] | None:
        if not isinstance(approval, dict):
            return None
        instance_id = str(approval.get("instance_id") or "").strip()
        if not instance_id:
            return None
        raw_fields = approval.get("form_fields") or []
        fields = (
            dict(raw_fields)
            if isinstance(raw_fields, dict)
            else {
                str(row.get("label") or ""): row.get("value")
                for row in raw_fields
                if isinstance(row, dict) and str(row.get("label") or "").strip()
            }
        )
        if not fields and not approval.get("excluded"):
            return None
        title = str(approval.get("title") or ("国际物流审批" if role == "international_logistics" else "采购审批"))
        return {
            "source_kind": "approval_form",
            "source_id": f"approval:{instance_id}:form",
            "source_label": f"{title}正文",
            "process_instance_id": instance_id,
            "approval_role": role,
            "excluded": bool(approval.get("excluded")),
            "exclude_reason": str(
                approval.get("exclusion_reason")
                or ("审批已失效，不参与分析。" if approval.get("excluded") else "")
            ),
            "approval_no": str(approval.get("business_id") or approval.get("approval_no") or ""),
            "actor_name": str(
                approval.get("originator_user_name")
                or approval.get("originator_name")
                or approval.get("originator_userid")
                or ""
            ),
            "occurred_at": str(approval.get("finish_time") or approval.get("create_time") or ""),
            "form_fields": fields,
            "approval_decisions": [row for row in approval.get("timeline") or [] if isinstance(row, dict) and row.get("remark")],
            "source_updated_at": str(
                approval.get("finish_time")
                or approval.get("create_time")
                or detail.get("source_updated_at")
                or ""
            ),
        }

    rows = []
    main = source(detail.get("main_approval") or {}, "logistics_expense" if (detail.get('source_context') or {}).get('root_kind') == 'expense' else "international_logistics")
    if main:
        rows.append(main)
    for approval in detail.get("linked_purchase_approvals") or []:
        linked = source(approval, "purchase")
        if linked:
            rows.append(linked)
    for approval in detail.get("excluded_linked_purchase_approvals") or []:
        excluded = source(approval, "purchase")
        if excluded:
            rows.append(excluded)
    return rows


def selected_packing_ai_sources(batch_name, bundle):
    """Selected source content is already shipment-scoped and locally archived."""
    context=bundle['context'];source=bundle['source'];selected=context['packing']['selected_source']
    rows=[{**g,**(g.get('physical') or {}),'source_row':(g.get('evidence') or {}).get('row') or g.get('source_position') or index}
          for index,g in enumerate(source.get('goods') or [],1)]
    return [{'source_id':selected['id'],'logical_source_id':selected['id'],'source_kind':selected['source_kind'],
             'source_label':selected['source_label'],'file_name':(selected.get('evidence') or {}).get('file_name') or selected['source_label'],
             'sheet_name':selected.get('sheet',''),'approval_no':selected['approval_no'],'process_instance_id':context['instance_id'],
             'batch':batch_name,'source_context':context,'source_hash':selected['revision'],'content_hash':selected['revision'],
             'source_updated_at':selected.get('occurred_at',''),'actor_name':selected.get('actor_name',''),
             'available':bool(context['available']),'excluded':not context['available'],'approval_role':'logistics_expense',
             'selected_source':selected,'scoped_packing':True,'scoped_goods':rows,'scoped_text':source.get('scoped_text',''),
             'form_fields':{},'approval_decisions':[],'can_download':False}]


def list_material_ai_sources(batch_name: str, version_name: str | None = None) -> list[dict[str, Any]]:
    """Return a stable manifest of every trusted source the material AI task may read."""

    if frappe is None:
        raise RuntimeError("当前环境未连接 Frappe。")
    bundle = effective_source.current_source_bundle(batch_name, version_name)
    if bundle and (bundle['context'].get('packing') or {}).get('selected_source'):
        return selected_packing_ai_sources(batch_name,bundle)
    if bundle and bundle['context']['root_kind'] == 'expense':
        return _bound_material_sources(batch_name, bundle)
    detail = packing_source_service.dingtalk_approval_service.get_batch_dingtalk_approval_detail(str(batch_name)) or {}
    packing = list_packing_sources(str(batch_name), approval_detail=detail, include_wiki=False)
    comment_index = {str(row.get("source_id") or ""): row for approval in
        [detail.get("main_approval") or {}, *(detail.get("linked_purchase_approvals") or [])]
        for row in approval.get("timeline") or [] if isinstance(row, dict)}
    result: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()

    def append_source(source: dict[str, Any], *, sheet_name: str = "") -> None:
        kind = str(source.get("source_kind") or "")
        source_id = str(source.get("source_id") or source.get("attachment_name") or "")
        process_instance_id = str(source.get("process_instance_id") or "")
        file_id = str(source.get("file_id") or "")
        logical_source_id = (
            f"oa:{process_instance_id}:{file_id}"
            if process_instance_id and file_id
            else source_id
        )
        key = (kind, logical_source_id, str(sheet_name or ""))
        if not kind or not source_id or key in seen:
            return
        seen.add(key)
        public = {
            "source_kind": kind,
            "source_id": source_id,
            "logical_source_id": logical_source_id,
            "source_label": str(source.get("source_label") or source.get("file_name") or source_id),
            "file_name": str(source.get("file_name") or source.get("source_label") or ""),
            "sheet_name": str(sheet_name or source.get("sheet_name") or ""),
            "source_updated_at": str(source.get("source_updated_at") or source.get("modified") or ""),
            "available": bool(source.get("available", True)),
            "can_download": bool(source.get("can_download")),
            "download_required": bool(source.get("download_required")),
            "process_instance_id": str(source.get("process_instance_id") or ""),
            "file_id": str(source.get("file_id") or ""),
            "approval_role": str(source.get("approval_role") or ""),
            "approval_no": str(source.get("approval_no") or source.get("business_id") or ""),
            "actor_name": str(
                source.get("actor_name")
                or source.get("comment_user_name")
                or source.get("comment_user")
                or source.get("user_name")
                or ""
            ),
            "occurred_at": str(source.get("occurred_at") or source.get("source_updated_at") or ""),
            "excluded": bool(source.get("excluded")),
            "exclude_reason": str(source.get("exclude_reason") or source.get("exclusion_reason") or ""),
            "form_fields": source.get("form_fields") if isinstance(source.get("form_fields"), dict) else {},
            "approval_decisions": source.get("approval_decisions") or [],
            "dedicated_packing": bool(source.get("dedicated_packing")),
            "content_hash": str(
                source.get("content_hash")
                or source.get("content_sha256")
                or source.get("source_hash")
                or ""
            ),
        }
        hash_basis = {
            "source_kind": kind,
            "logical_source_id": logical_source_id,
            "sheet_name": public["sheet_name"],
            "source_updated_at": public["source_updated_at"],
            "content_hash": public["content_hash"],
            "content": public["form_fields"],
            "decisions": public["approval_decisions"],
            "comment_text": str((comment_index.get(source_id) or {}).get("remark") or ""),
        }
        if kind == "approval_comment":
            public["comment_text"] = hash_basis["comment_text"]
        public["source_hash"] = hashlib.sha256(_json(hash_basis).encode("utf-8")).hexdigest()
        result.append(public)

    try:
        current_snapshot = get_current_packing_snapshot(str(batch_name)) or {}
    except Exception:
        current_snapshot = {}
    current_wiki_source = (
        str(current_snapshot.get("source_id") or "")
        if str(current_snapshot.get("source_kind") or "") == "wiki_sheet"
        else ""
    )
    for workbook in packing.get("wiki_workbooks") or []:
        for sheet in workbook.get("sheets") or []:
            source_id = str(sheet.get("source_id") or "")
            if source_id == current_wiki_source:
                append_source(
                    {
                        **sheet,
                        "source_hash": sheet.get("source_hash")
                        or current_snapshot.get("source_hash"),
                    }
                )
    if current_wiki_source and not any(
        row.get("source_kind") == "wiki_sheet" and row.get("source_id") == current_wiki_source
        for row in result
    ):
        # Revalidate just the attached Sheet. A confirmed snapshot's old hash
        # cannot prove that the current server archive still has that content.
        try:
            from overseas_costing.integrations.dingtalk_packing_source import get_packing_runtime_clients

            workbook_id, separator, sheet_id = current_wiki_source.partition(":")
            if not separator or not workbook_id or not sheet_id:
                raise ValueError("当前装箱计划来源 ID 不合法。")
            manifest = get_packing_runtime_clients().catalog.get_latest_snapshot(workbook_id, sheet_id) or {}
            content_hash = str(manifest.get("content_sha256") or "")
            if not content_hash:
                raise ValueError("当前装箱计划 Sheet 缺少可验证的归档哈希。")
            append_source({**current_snapshot, "content_hash": content_hash,
                           "source_updated_at": manifest.get("capture_finished_at") or ""})
        except Exception as exc:
            append_source({**current_snapshot, "source_hash": "", "content_hash": "",
                           "available": False, "excluded": True, "exclude_reason": str(exc)})
    try:
        approval_body_sources = _list_approval_body_ai_sources(str(batch_name), detail=detail)
    except Exception:
        approval_body_sources = []
    for source in approval_body_sources:
        append_source(source)
    approval_body_by_instance = {str(source.get("process_instance_id") or ""): source for source in approval_body_sources}
    for source in packing.get("approval_sources") or []:
        owning = approval_body_by_instance.get(str(source.get("process_instance_id") or "")) or {}
        packing_fields = {key: value for key, value in (owning.get("form_fields") or {}).items() if "装箱单附件" in key}
        source = {**source, "approval_role": owning.get("approval_role") or source.get("approval_role"),
                  "dedicated_packing": bool(packing_fields and str(source.get("file_name") or source.get("source_label") or "--") in _json(packing_fields))}
        is_unmaterialized = str(source.get("source_id") or "").startswith("oa:") or not source.get("attachment_name")
        if source.get("source_kind") != "approval_comment" and not is_unmaterialized:
            continue
        if source.get("source_kind") == "approval_comment" or not source.get("sheets"):
            append_source(source)
        else:
            for sheet_name in source.get("sheets") or []:
                append_source(source, sheet_name=str(sheet_name or ""))
    attachment_rows = frappe.get_list(
        "Overseas Cost Attachment",
        filters={"batch": str(batch_name)},
        fields=[
            "name",
            "version",
            "source_type",
            "oa_attachment_origin",
            "file_name",
            "file_url",
            "modified",
            "parse_result_json",
        ],
        limit_page_length=5000,
    )
    approval_body_by_instance = {
        str(source.get("process_instance_id") or ""): source
        for source in approval_body_sources
        if str(source.get("process_instance_id") or "")
    }
    for row in attachment_rows:
        if version_name and row.get("version") and str(row.get("version")) != str(version_name):
            continue
        file_name = str(row.get("file_name") or "")
        if not file_name.lower().endswith(MATERIAL_AI_DOCUMENT_SUFFIXES):
            continue
        audit_only = (
            str(row.get("source_type") or "").upper() == "OA"
            and packing_source_service._attachment_is_audit_only(row)
        )
        snapshot = packing_source_service.import_service._json_loads_dict(row.get("parse_result_json"))
        attachment_instance = str(
            snapshot.get("process_instance_id") or snapshot.get("instance_id") or ""
        )
        excluded_instances = {
            str(source.get("process_instance_id") or "")
            for source in approval_body_sources
            if source.get("excluded") and str(source.get("process_instance_id") or "")
        }
        invalid_approval = bool(attachment_instance and attachment_instance in excluded_instances)
        if str(row.get("source_type") or "").upper() == "OA":
            allowed_instances = {
                str(source.get("process_instance_id") or "")
                for source in [*approval_body_sources, *(packing.get("approval_sources") or [])]
                if str(source.get("process_instance_id") or "")
            }
            if not audit_only and (
                not attachment_instance or attachment_instance not in allowed_instances
            ):
                continue
        owning_approval = approval_body_by_instance.get(attachment_instance) or {}
        packing_fields = {key: value for key, value in (owning_approval.get("form_fields") or {}).items() if "装箱单附件" in key}
        source = {
            "source_kind": (
                "approval_attachment"
                if str(row.get("source_type") or "").upper() == "OA"
                else "manual_attachment"
            ),
            "source_id": row.get("name"),
            "approval_role": owning_approval.get("approval_role") or "",
            "dedicated_packing": bool(packing_fields and file_name in _json(packing_fields)),
            "source_label": file_name or row.get("name"),
            "file_name": file_name,
            "source_updated_at": row.get("modified"),
            "available": bool(row.get("file_url")),
            "download_required": not bool(row.get("file_url")),
            "process_instance_id": str(snapshot.get("process_instance_id") or snapshot.get("instance_id") or ""),
            "file_id": str(snapshot.get("file_id") or ""),
            "approval_no": str(owning_approval.get("approval_no") or ""),
            "actor_name": str(
                snapshot.get("comment_user_name") or owning_approval.get("actor_name") or ""
            ),
            "occurred_at": str(
                snapshot.get("comment_time") or owning_approval.get("occurred_at") or ""
            ),
            "excluded": audit_only or invalid_approval,
            "exclude_reason": (
                "审计专用附件，不参与资料分析。"
                if audit_only
                else "所属审批已失效，不参与资料分析。"
                if invalid_approval
                else ""
            ),
            "content_hash": str(packing_source_service._attachment_hash(row) if row.get("file_url") else (
                (
                    snapshot.get("download")
                    if isinstance(snapshot.get("download"), dict)
                    else {}
                ).get("sha256")
                or packing_source_service._attachment_hash(row)
            )),
        }
        sheets = _attachment_sheet_names(row) if file_name.lower().endswith((".xlsx", ".xlsm")) else []
        if sheets:
            for sheet_name in sheets:
                append_source(source, sheet_name=sheet_name)
        else:
            append_source(source)
    return sorted(
        result,
        key=lambda row: (
            0 if row.get("approval_role") == "international_logistics" and row.get("source_kind") == "approval_form" else
            1 if row.get("dedicated_packing") else 2 if row.get("source_kind") == "approval_form" else 3,
            str(row.get("source_kind") or ""),
            str(row.get("source_id") or ""),
            str(row.get("sheet_name") or ""),
        ),
    )


def _bound_material_sources(batch_name, bundle):
    context = bundle['context']
    detail = effective_source.approval_detail_for_bundle(bundle)
    approval = detail['main_approval']
    result = _list_approval_body_ai_sources(batch_name, detail=detail)
    for comment in approval.get('timeline') or []:
        if comment.get('remark') and comment.get('source_id'):
            result.append({'source_kind': 'approval_comment', 'source_id': comment['source_id'],
                           'source_label': '采购支出评论', 'comment_text': comment['remark'],
                           'process_instance_id': context['instance_id'], 'source_updated_at': comment.get('operation_time'),
                           'approval_role': 'logistics_expense', 'excluded': approval['excluded']})
    rows = frappe.get_list('Overseas Cost Attachment', filters={'batch': str(batch_name)},
        fields=['name', 'version', 'source_type', 'file_name', 'file_url', 'modified', 'parse_result_json'], limit_page_length=5000)
    materialized_documents = set()
    for row in rows:
        if not effective_source.attachment_allowed(row, bundle, for_analysis=True) or not _is_material_ai_attachment(row.get('file_name')):
            continue
        meta = effective_source.json_dict(row.get('parse_result_json'))
        document = meta.get('settlement_document') or {}
        materialized_documents.add(str(document.get('document_id') or ''))
        cached = next((d for d in (bundle.get('source') or {}).get('documents') or []
                       if str(d.get('id')) == str(document.get('document_id'))), {})
        manifest = document.get('manifest') or {}
        excel = str(row.get('file_name') or '').lower().endswith(('.xlsx','.xlsm','.xls'))
        # Listing is local metadata only. Corrupt/unreadable bytes are handled by
        # the individual preview/AI read, never by opening files in the catalogue.
        sheets = sorted({str(t.get('title') or '') for t in cached.get('tables') or [] if t.get('title')}) if excel else []
        failed = excel and cached.get('status') in {'failed','error','invalid'}
        for sheet in sheets or ['']:
            result.append({'source_kind': 'approval_attachment', 'source_id': row['name'], 'attachment_name': row['name'],
                'logical_source_id': f"oa:{context['instance_id']}:{meta.get('file_id') or document.get('document_id')}",
                'source_label': row.get('file_name') or row['name'], 'file_name': row.get('file_name'), 'sheet_name': sheet, 'sheets': sheets,
                'process_instance_id': context['instance_id'], 'file_id': meta.get('file_id'),
                'available': bool(row.get('file_url')) and not failed, 'can_download': False, 'download_required': False,
                'source_updated_at': row.get('modified'), 'content_hash': manifest.get('sha256') or manifest.get('content_sha256') or document.get('fingerprint'),
                'approval_role': 'logistics_expense', 'dedicated_packing': any(t.get('kind') == 'packing' for t in document.get('tables') or []),
                'excluded': approval['excluded'] or failed,
                'exclude_reason':'归档文件读取失败：'+'；'.join(map(str,cached.get('issues') or ['请核对当前资料'])) if failed else '',
                'supported_for_material_import': excel and not failed})
    for document in (bundle.get('source') or {}).get('documents') or []:
        from .logistics_settlement.document_writer import document_retired
        if str(document.get('id') or '') in materialized_documents or document_retired(document) or not _is_material_ai_attachment(document.get('file_name')):
            continue
        result.append({'source_kind': 'approval_attachment', 'source_id': f"pending:{document['id']}",
            'process_instance_id': context['instance_id'], 'source_label': document.get('file_name'),
            'file_name': document.get('file_name'), 'available': False, 'can_download': False,
            'download_required': False, 'excluded': True, 'exclude_reason': '采购支出附件尚未完成本地归档或版本登记，请等待同步。',
            'approval_role': 'logistics_expense', 'content_hash': document.get('fingerprint') or document['id']})
    for source_id in sorted(effective_source.explicit_wiki_sources(bundle.get('source'))):
        cached=packing_source_service.load_bound_wiki_snapshot(context,source_id)
        result.append({'source_kind': 'wiki_sheet', 'source_id': source_id, 'source_label': '采购支出链接的装箱计划表',
                       **packing_source_service.wiki_refresh_status(context,source_id),
                       'process_instance_id': context['instance_id'], 'available': bool(cached) and not approval['excluded'],
                       'content_hash':(cached or {}).get('source_hash'), 'source_updated_at':((cached or {}).get('source') or {}).get('source_updated_at'),
                       'excluded':approval['excluded'] or not cached,
                       'exclude_reason':'' if cached else '当前采购支出工作表尚未获取到本地，请点击刷新或获取资料。',
                       'requires_refresh':not bool(cached)})
    for row in result:
        row['approval_no'] = (bundle.get('source') or {}).get('approval_no') or ''
        row['source_context'] = context
        row['analysis_only'] = bool(not context['approved'] or context['invalid'])
        content = {key:value for key,value in row.items() if key not in {
            'cache_refreshed_at','refresh_last_checked_at','refresh_last_success_at','refresh_error'}}
        row['source_hash'] = hashlib.sha256(_json({'context': context, 'source': content}).encode()).hexdigest()
    return result


def _is_excel_packing_attachment(file_name: Any) -> bool:
    return str(file_name or "").strip().lower().endswith((".xlsx", ".xlsm"))


def _is_material_ai_attachment(file_name: Any) -> bool:
    return str(file_name or "").strip().lower().endswith(MATERIAL_AI_DOCUMENT_SUFFIXES)


def _packing_batch_context(batch_name: str, detail: dict[str, Any]) -> dict[str, Any]:
    batch = frappe.db.get_value(
        "Overseas Cost Batch",
        batch_name,
        [
            "name",
            "batch_no",
            "waybill_no",
            "source_approval_no",
            "source_instance_id",
            "project_collection",
            "creation",
        ],
        as_dict=True,
    ) or {}
    items = frappe.get_list(
        "Overseas Cost Item",
        filters={"batch": str(batch.get("name") or batch_name)},
        fields=["material_code", "product_name", "source_doc_no"],
        limit_page_length=5000,
    )
    item_codes = {str(row.get("material_code") or "").strip() for row in items}
    references = {
        str(batch.get(field) or "").strip()
        for field in ("batch_no", "waybill_no", "source_approval_no", "source_instance_id")
    }
    keywords = {str(batch.get("project_collection") or "").strip()}
    for row in items:
        references.add(str(row.get("source_doc_no") or "").strip())
        keywords.add(str(row.get("product_name") or "").strip())
    approvals = [detail.get("main_approval"), *(detail.get("linked_purchase_approvals") or [])]
    for approval in approvals:
        if not isinstance(approval, dict) or approval.get("excluded"):
            continue
        references.update(
            str(approval.get(field) or "").strip()
            for field in ("business_id", "instance_id")
        )
        keywords.add(str(approval.get("title") or "").strip())
        for field in approval.get("form_fields") or []:
            if not isinstance(field, dict):
                continue
            value = str(field.get("value") or "").strip()
            label = str(field.get("label") or "").lower()
            if any(token in label for token in ("审批", "订单", "单号", "order", "物流")):
                references.add(value)
            elif 2 <= len(value) <= 80:
                keywords.add(value)
    creation = str(batch.get("creation") or "")
    return {
        "item_codes": sorted(value for value in item_codes if value),
        "references": sorted(value for value in references if value),
        "keywords": sorted(value for value in keywords if value),
        "reference_date": creation[:10] if len(creation) >= 10 else None,
    }


def _packing_snapshot_summaries(clients: Any, workbook_id: str) -> dict[str, dict[str, Any]]:
    """读取现有 MinIO 快照；任何单个对象失败都只影响对应 Sheet。"""

    from overseas_costing.services.packing_grid import build_grid_from_dingtalk_snapshot
    from overseas_costing.services.packing_parse_service import parse_packing_grid

    result: dict[str, dict[str, Any]] = {}
    for manifest in clients.catalog.list_latest_snapshots(workbook_id):
        sheet_id = str(manifest.get("sheet_id") or "")
        source_id = f"{workbook_id}:{sheet_id}"
        state = {
            "snapshot_status": "ready",
            "snapshot_updated_at": manifest.get("created_at"),
            "content_hash": str(
                manifest.get("content_sha256") or manifest.get("sha256") or ""
            ),
        }
        try:
            payload = clients.archive.download(manifest)
            preview = parse_packing_grid(build_grid_from_dingtalk_snapshot(payload))
            state["summary"] = summarize_packing_preview(preview)
        except Exception:
            state["snapshot_status"] = "unreadable"
        result[source_id] = state
    return result


def _pin_recommended_workbook(workbooks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """多年度表时把包含全局推荐项的工作簿整体置顶。"""

    return sorted(
        workbooks,
        key=lambda workbook: (
            0 if any(sheet.get("is_recommended") for sheet in workbook.get("sheets") or []) else 1,
        ),
    )


def _attachment_sheet_names(row: dict[str, Any]) -> list[str]:
    """只返回工作表名，不向浏览器暴露站点路径。"""

    file_url = str(row.get("file_url") or "").strip()
    file_name = str(row.get("file_name") or file_url).lower()
    if file_url and file_name.endswith('.xls'):
        import xlrd
        path = packing_source_service.import_service._resolve_excel_file_path(file_url=file_url)
        book=xlrd.open_workbook(str(path),on_demand=True)
        try:
            return [str(name)[:200] for name in book.sheet_names()]
        finally:
            book.release_resources()
    if not file_url or not file_name.endswith((".xlsx", ".xlsm")) or load_workbook is None:
        return []
    try:
        path = packing_source_service.import_service._resolve_excel_file_path(file_url=file_url)
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            return [str(name)[:200] for name in workbook.sheetnames]
        finally:
            workbook.close()
    except (FileNotFoundError, ValueError, OSError):
        return []
