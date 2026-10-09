"""Native drawer actions on one synthetic site, always rolled back."""
import json
import subprocess
import sys
import uuid
import unittest
from functools import wraps
from unittest.mock import patch

import frappe
from erpnext.buying.doctype.purchase_order.purchase_order import make_purchase_receipt
from erpnext.stock.doctype.purchase_receipt.purchase_receipt import (
    make_purchase_invoice,
    make_purchase_return,
)
from frappe.utils import add_days, nowdate

from deeplinkerp_branding.services import purchase_payment_service as service

SITE = "po-grid-qa.localhost"
COMPANY = "QA Second Company"


def execute():
    if frappe.local.site != SITE:
        raise RuntimeError("Synthetic QA site only")
    frappe.set_user("Administrator")
    frappe.flags.in_test = True
    frappe.flags.mute_emails = True
    doctypes = ("Purchase Order", "Purchase Receipt", "Purchase Invoice", "Payment Entry", "GL Entry", "China Accounting Voucher", "China Voucher Sync Issue")
    before = {doctype: frappe.db.count(doctype) for doctype in doctypes}
    results = []
    try:
        po = frappe.get_doc({"doctype": "Purchase Order", "company": COMPANY,
            "supplier": "QA Test Supplier", "currency": "CNY", "schedule_date": add_days(nowdate(), 1),
            "items": [{"item_code": "QA-PO-ITEM", "qty": 10, "rate": 1000,
                "schedule_date": add_days(nowdate(), 1)}]}).insert()
        po.submit()
        pr = make_purchase_receipt(po.name)
        pr.items[0].qty = 4
        pr.insert()
        pr.submit()
        pi = make_purchase_invoice(pr.name)
        pi.items[0].qty = 2
        pi.insert()
        pi.submit()
        chain = service.get_purchase_chain("Purchase Receipt", pr.name, include_payments=False)
        assert chain["can_create_invoice"], "Submitted partial PI must permit native remaining quantity"
        assert chain["draft_invoices"] == [], "Submitted PI must not be a resume action"
        results.append("partial submitted invoice permits remaining native quantity")

        from unittest.mock import patch

        from deeplinkerp_branding.services import purchase_document_actions as actions
        original_fields = service._require_fields
        for denied_doctype, denied_parent in (("Purchase Receipt", None), ("Purchase Taxes and Charges", "Purchase Receipt")):
            def protect_source_tax(doctype, fields, parenttype=None):
                if doctype == denied_doctype and parenttype == denied_parent and (doctype == "Purchase Taxes and Charges" or "taxes" in fields):
                    raise frappe.PermissionError
                return original_fields(doctype, fields, parenttype)
            with patch.object(service, "_require_fields", side_effect=protect_source_tax):
                try:
                    actions.preview_document("Purchase Receipt", pr.name, "Purchase Invoice")
                except frappe.PermissionError:
                    pass
                else:
                    raise AssertionError("Protected source taxes leaked through mapped target totals")
        results.append("protected source header and child tax fields deny mapped preview")
        preview = actions.preview_document("Purchase Receipt", pr.name, "Purchase Invoice")
        item = preview["document"]["items"][0]
        assert item["qty"] == item["max_qty"] == 2
        assert "credit_to" not in preview["document"]
        assert "owner" not in preview["document"]
        request_id = str(uuid.uuid4())
        changes = {"items": [{"key": item["key"], "qty": 1, "rate": 1000}], "bill_no": "QA-DRAWER"}
        first = actions.save_document_draft("Purchase Receipt", pr.name, "Purchase Invoice", changes, request_id)
        retry = actions.save_document_draft("Purchase Receipt", pr.name, "Purchase Invoice", changes, request_id)
        assert first["document"]["name"] == retry["document"]["name"] and retry["reused"]
        draft = first["document"]
        for bad_request, bad_changes in (("bad", changes), (request_id, {"remarks": "changed request"})):
            try:
                actions.save_document_draft("Purchase Receipt", pr.name, "Purchase Invoice", bad_changes, bad_request)
            except frappe.ValidationError:
                pass
            else:
                raise AssertionError("Invalid or changed request identifier accepted")
        assert draft["docstatus"] == 0 and draft["grand_total"] == 1000
        chain = service.get_purchase_chain("Purchase Receipt", pr.name, include_payments=False)
        assert not chain["can_create_invoice"] and chain["draft_invoices"][0]["name"] == draft["name"]
        assert frappe.db.get_value("Purchase Invoice", pi.name, "outstanding_amount") == 2000
        assert not frappe.db.count("GL Entry", {"voucher_type": "Purchase Invoice", "voucher_no": draft["name"]})
        results.append("partial invoice draft resumes; duplicate request has no duplicate GL or invoice")
        for invalid in ({"items": [{"key": item["key"], "qty": qty}]} for qty in (-1, 0, 3)):
            try:
                actions.save_document_draft("Purchase Receipt", pr.name, "Purchase Invoice", invalid, str(uuid.uuid4()),
                    target_name=draft["name"], expected_modified=draft["modified"])
            except frappe.ValidationError:
                pass
            else:
                raise AssertionError("Unsafe drawer mutation accepted")
        try:
            actions.save_document_draft("Purchase Receipt", pr.name, "Purchase Invoice", {"credit_to": "arbitrary"}, str(uuid.uuid4()),
                target_name=draft["name"], expected_modified=draft["modified"])
        except frappe.ValidationError:
            pass
        else:
            raise AssertionError("Unrestricted payload update accepted")
        try:
            actions.save_document_draft("Purchase Receipt", pr.name, "Purchase Invoice", changes, str(uuid.uuid4()),
                target_name=draft["name"], expected_modified="1900-01-01 00:00:00")
        except frappe.ValidationError:
            pass
        else:
            raise AssertionError("Stale version accepted")
        results.append("latest native maximum, strict field whitelist, and version guard")
        receipt_preview = actions.preview_document("Purchase Order", po.name, "Purchase Receipt")
        receipt_item = receipt_preview["document"]["items"][0]
        assert receipt_item["max_qty"] == 6
        receipt_draft = actions.save_document_draft("Purchase Order", po.name, "Purchase Receipt",
            {"items": [{"key": receipt_item["key"], "qty": 2}]}, str(uuid.uuid4()), allow_another_draft=1)["document"]
        assert receipt_draft["docstatus"] == 0 and receipt_draft["grand_total"] == 2000
        assert frappe.db.get_value("Purchase Order", po.name, "per_received") == 40
        assert receipt_draft["allowed_actions"] == ["Submit"]
        received = actions.submit_document("Purchase Receipt", receipt_draft["name"], receipt_draft["modified"], request_id=str(uuid.uuid4()))
        assert received["document"]["docstatus"] == 1
        assert frappe.db.get_value("Purchase Order", po.name, "per_received") == 60
        results.append("explicit native receipt submission consumes remaining quantity once")
        changed = actions.save_document_draft("Purchase Receipt", pr.name, "Purchase Invoice",
            {"items": [{"key": item["key"], "qty": 2}]}, str(uuid.uuid4()),
            target_name=draft["name"], expected_modified=draft["modified"])["document"]
        assert changed["grand_total"] == 2000 and "Submit" in changed["allowed_actions"]
        posted = actions.submit_document("Purchase Invoice", changed["name"], changed["modified"])["document"]
        assert posted["docstatus"] == 1
        assert not service.get_purchase_chain("Purchase Receipt", pr.name, include_payments=False)["can_create_invoice"]
        results.append("explicit native PI submission consumes latest remaining quantity")
        args = dict(source_doctype="Purchase Receipt", source_name=pr.name, purchase_invoice=posted["name"],
            amount_to_pay=1000, bank_account="Cash - QAB", request_id=str(uuid.uuid4()))
        pe_name = service.create_payment_draft(**args)["name"]
        manual_payment = frappe.get_doc("Payment Entry", pe_name)
        manual_reference_date = add_days(nowdate(), -4)
        manual_payment.reference_date = manual_reference_date
        manual_payment.save()
        payment = actions.preview_payment(pe_name)["document"]
        assert payment["amount"] == 1000 and payment["currency"] == "CNY"
        bank_name = payment["bank_account"]
        foreign_company = frappe.db.get_value("Company", {"name": ["!=", COMPANY]}, "name")
        assert foreign_company
        for field, value in (("disabled", 1), ("account_currency", "USD"), ("account_type", "Receivable"), ("company", foreign_company)):
            frappe.db.savepoint("bank_guard")
            try:
                frappe.db.set_value("Account", bank_name, field, value)
                for action in (lambda: actions.preview_payment(pe_name),
                               lambda: actions.submit_document("Payment Entry", pe_name, payment["modified"]),
                               lambda: actions.update_payment_draft(pe_name, {"remarks": "must not save"}, payment["modified"])):
                    try:
                        action()
                    except (frappe.PermissionError, frappe.ValidationError):
                        pass
                    else:
                        raise AssertionError("Changed bank state accepted")
                guarded_record = service.get_payment_records(search=pe_name)["rows"][0]
                assert guarded_record["bank_account"] is None and not guarded_record["can_edit"] and not guarded_record["allowed_actions"]
            finally:
                frappe.db.rollback(save_point="bank_guard")
        original_read = service._read
        def denied_bank(doctype, name, fields=()):
            if doctype == "Account" and name == bank_name:
                frappe.throw("QA-PRIVATE-ACCOUNT " + name, frappe.PermissionError)
            return original_read(doctype, name, fields)
        before_messages = len(frappe.message_log)
        with patch.object(service, "_read", side_effect=denied_bank):
            try:
                actions.preview_payment(pe_name)
            except frappe.PermissionError as exc:
                assert bank_name not in str(exc)
            else:
                raise AssertionError("Revoked Account read accepted")
            hidden_bank = service.get_payment_records(search=pe_name)
            assert bank_name not in json.dumps(hidden_bank, default=str)
        assert all("QA-PRIVATE-ACCOUNT" not in str(message) and bank_name not in str(message) for message in frappe.message_log[before_messages:])
        results.append("bank read/state/type/currency revocation disables preview, submit and record controls without private names")
        for changes in ({"remarks": "QA preserve manual reference date"}, {"amount": 1000},
                        {"bank_account": bank_name}, {"posting_date": add_days(nowdate(), 1)}, {"posting_date": nowdate()}):
            payment = actions.update_payment_draft(pe_name, changes, payment["modified"])["document"]
            assert str(frappe.db.get_value("Payment Entry", pe_name, "reference_date")) == str(manual_reference_date), "Drawer edits must preserve native manual reference date"
        results.append("remarks, amount, account and posting-date edits preserve native manual reference date")
        from unittest.mock import patch
        read_only_payment = frappe.get_doc("Payment Entry", pe_name)
        original_permission = read_only_payment.has_permission
        with patch.object(read_only_payment, "has_permission", side_effect=lambda permission, *a, **kw: False if permission == "write" else original_permission(permission, *a, **kw)):
            assert actions._payment(read_only_payment)["editable_fields"] == [], "Read-only payment must expose no write controls"
        edited = actions.update_payment_draft(pe_name, {"amount": 1500, "remarks": "QA drawer edited"},
            payment["modified"])["document"]
        assert edited["amount"] == 1500 and edited["remarks"] == "QA drawer edited"
        assert frappe.db.get_value("Purchase Invoice", posted["name"], "outstanding_amount") == 2000
        posted_payment = actions.submit_document("Payment Entry", pe_name, edited["modified"])["document"]
        assert posted_payment["docstatus"] == 1
        assert frappe.db.get_value("Purchase Invoice", posted["name"], "outstanding_amount") == 500
        results.append("simple payment edits remain draft until explicit native submit; native balance changes once")
        listing = service.get_receipt_list(filters={"company": COMPANY},
            native_filters=[["name", "=", pr.name], ["owner", "=", "Administrator"]], order_by="`tabPurchase Receipt`.`grand_total` asc", page_length=2500)
        assert [row["name"] for row in listing["rows"]] == [pr.name], "Native filters must constrain authorized rows"
        assert listing["total_count"] == 1 and listing["totals"] == [{"currency": "CNY", "grand_total": 4000.0}]
        for field, operator, operand in (("posting_date", "Timespan", "today"), ("creation", "Timespan", "today"),
                                         ("posting_date", "previous", "1 week"), ("posting_date", "next", "1 week")):
            native_conditions = [["name", "=", pr.name], [field, operator, operand]]
            expected_names = frappe.get_list("Purchase Receipt", filters=native_conditions, pluck="name")
            actual_names = [row["name"] for row in service.get_receipt_list(native_filters=native_conditions)["rows"]]
            assert actual_names == expected_names, "Date operators must keep native Frappe normalization"
        assert "can_create_invoice" in listing["rows"][0] and "draft_invoices" in listing["rows"][0]
        service.get_receipt_list(filters={"company": COMPANY},
            native_filters=[["name", "=", pr.name]], export_format="xlsx", page_length=20,
            columns=["name", "grand_total", "currency"])
        from io import BytesIO

        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(frappe.response["filecontent"]))
        assert workbook.active.max_row == 2
        assert [cell.value for cell in workbook.active[2]] == [pr.name, 4000, "CNY"]
        assert workbook.active.cell(2, 2).number_format == "#,##0.00", "Money cells require two display decimals"
        service._export("Purchase Receipt", [{"name": pr.name, "grand_total": 4000.123456}],
            ["name", "grand_total"], service.RECEIPT_COLUMNS)
        precise_book = load_workbook(BytesIO(frappe.response["filecontent"]))
        assert precise_book.active.cell(2, 2).value == 4000.123456, "Export styles must not round raw values"
        assert precise_book.active.cell(2, 2).number_format == "#,##0.00"
        service.get_receipt_list(native_filters=[["name", "=", pr.name]], export_format="xlsx",
            columns=["name", "settled", "outstanding"])
        balances_book = load_workbook(BytesIO(frappe.response["filecontent"]))
        settled = json.loads(balances_book.active.cell(2, 2).value)
        outstanding = json.loads(balances_book.active.cell(2, 3).value)
        assert settled[0]["currency"] == "CNY" and settled[0]["settled"] == 1500
        assert outstanding[0]["currency"] == "CNY" and outstanding[0]["outstanding"] == 2500
        actual_permission = frappe.has_permission
        assert "System Manager" in frappe.get_roles(), "QA must exercise the installed System Manager export exception"
        with patch.object(frappe, "has_permission", side_effect=lambda doctype, ptype="read", *a, **kw: False if doctype == "Purchase Receipt" and ptype == "export" else actual_permission(doctype, ptype, *a, **kw)):
            assert not frappe.has_permission("Purchase Receipt", "export")
            assert frappe.permissions.can_export("Purchase Receipt"), "Native export capability must remain real"
            service.get_receipt_list(native_filters=[["name", "=", pr.name]], export_format="xlsx",
                columns=["name", "grand_total", "currency"])
            native_book = load_workbook(BytesIO(frappe.response["filecontent"]))
            assert native_book.active.max_row == 2
            assert [cell.value for cell in native_book.active[2]] == [pr.name, 4000, "CNY"]
        with patch.object(frappe.permissions, "can_export", return_value=False) as native_export:
            assert frappe.has_permission("Purchase Receipt", "read"), "Export denial must preserve real read permission"
            try:
                service.get_receipt_list(filters={"company": COMPANY}, export_format="xlsx")
            except frappe.PermissionError:
                pass
            else:
                raise AssertionError("Read permission must not imply export permission")
            native_export.assert_any_call("Purchase Receipt")
            native_export.assert_any_call("Purchase Receipt", is_owner=True)
        records = service.get_payment_records(company=COMPANY, from_date=nowdate(), to_date=nowdate(),
            status=1, search=pe_name, page_length=2500, order_by="`tabPayment Entry`.`posting_date` asc", native_filters=[["owner", "=", "Administrator"]])
        assert records["total_count"] == 1 and records["rows"][0]["name"] == pe_name
        assert "sync_issues" in records["rows"][0] and "can_edit" in records["rows"][0]
        service.create_payment_draft(**dict(args, amount_to_pay=100, remarks="<span>=1+1</span>", request_id=str(uuid.uuid4())))
        grouped = service.get_payment_records(purchase_receipt=pr.name)
        assert {(row["payment_type"], row["docstatus"], row["currency"], row["amount"]) for row in grouped["totals"]} == {("Pay", 0, "CNY", 100.0), ("Pay", 1, "CNY", 1500.0)}
        service.get_payment_records(purchase_receipt=pr.name, search=pe_name, export_format="xlsx",
            columns=["name", "amount", "allocation_currency"])
        payment_book = load_workbook(BytesIO(frappe.response["filecontent"]))
        assert payment_book.active.max_row == 2
        assert [cell.value for cell in payment_book.active[1]] == ["付款单", "金额", "核销币种", "付款币种"]
        assert [cell.value for cell in payment_book.active[2]] == [pe_name, 1500, "CNY", "CNY"]
        assert payment_book.active.cell(2, 2).number_format == "#,##0.00"
        service.get_payment_records(search=pe_name, export_format="xlsx", columns=["name", "amount"])
        appended_currency = load_workbook(BytesIO(frappe.response["filecontent"]))
        assert [cell.value for cell in appended_currency.active[1]] == ["付款单", "金额", "付款币种", "核销币种"]
        assert [cell.value for cell in appended_currency.active[2]] == [pe_name, 1500, "CNY", "CNY"]
        service.get_payment_records(purchase_receipt=pr.name, export_format="xlsx", start=1, page_length=1,
            columns=["name", "amount", "allocation_currency", "remarks"])
        full_book = load_workbook(BytesIO(frappe.response["filecontent"]))
        assert full_book.active.max_row == 3, "Export must include every filtered row regardless pagination"
        assert all(row[3].data_type != "f" for row in list(full_book.active.rows)[1:]), "Source HTML must not become an Excel formula"
        issue = frappe.get_doc({"doctype": "China Voucher Sync Issue", "company": COMPANY,
            "posting_date": nowdate(), "status": "Pending", "source_doctype": "Payment Entry", "source_name": pe_name,
            "issue_key": "QA-DRAWER-" + str(uuid.uuid4()), "last_error": "QA-PRIVATE-ERROR"}).insert()
        issue_records = service.get_payment_records(search=pe_name)
        assert issue_records["rows"][0]["sync_issues"][0]["name"] == issue.name
        assert "QA-PRIVATE-ERROR" not in json.dumps(issue_records, default=str)
        real_require = service._require_fields
        def deny_issue_fields(doctype, fields, parenttype=None):
            if doctype == "China Voucher Sync Issue":
                raise frappe.PermissionError
            return real_require(doctype, fields, parenttype)
        with patch.object(service, "_require_fields", side_effect=deny_issue_fields):
            hidden_issues = service.get_payment_records(search=pe_name)
        assert hidden_issues["rows"][0]["sync_issues"] == [] and issue.name not in json.dumps(hidden_issues, default=str)
        results.append("whole bank totals separate native payment states; ordered payment export and sync privacy")
        results.append("native receipt filters/count/currency totals/export allow and deny; payment record filters/capabilities")
        pi.reload()
        pi.cancel()
        returned = make_purchase_return(pr.name)
        returned.items[0].qty = -1
        returned.insert()
        returned.submit()
        remaining = actions.preview_document("Purchase Receipt", pr.name, "Purchase Invoice")["document"]["items"][0]
        expected_remaining = 2 if frappe.db.get_single_value("Buying Settings", "bill_for_rejected_quantity_in_purchase_invoice") else 1
        assert remaining["qty"] == remaining["max_qty"] == expected_remaining, remaining
        assert service.get_purchase_chain("Purchase Receipt", pr.name, include_payments=False)["can_create_invoice"]
        results.append("cancelled partial invoices ignored and native returns reduce remaining invoice quantity")
        frappe.set_user("qa-po-reader@example.invalid")
        try:
            actions.preview_document("Purchase Receipt", pr.name, "Purchase Invoice", target_name=posted["name"])
        except frappe.PermissionError:
            pass
        else:
            raise AssertionError("Other company drawer leaked")
        frappe.set_user("Administrator")
        results.append("drawer source/company permissions")
        print(json.dumps({"site": SITE, "tests": results, "status": "passed"}))
    finally:
        frappe.db.rollback()
        after = {doctype: frappe.db.count(doctype) for doctype in doctypes}
        assert before == after, (before, after)
        print(json.dumps({"rollback_counts_unchanged": True, "counts": after}))


class NativeAtomicPurchaseTests(unittest.TestCase):
    """The shared drawer/native-form boundary, restricted to the new isolated clone."""
    def setUp(self):
        if frappe.local.site != SITE or frappe.conf.db_name != "qa_procurement_5" or frappe.conf.db_host != "db":
            raise RuntimeError("Isolated synthetic procurement database only")
        self.assertFalse(frappe.in_test, "Native stock reposts must follow the production branch")
        frappe.set_user("Administrator")
        frappe.flags.in_test = True
        frappe.flags.mute_emails = True
        from deeplinkerp_branding.services.purchase_repost_boundary import initialize
        initialize()  # BEFORE this fixture's first synthetic business insert.
        frappe.db.after_commit.reset()  # BEFORE any fixture acquires a lease.
        self.types = ("Purchase Order", "Purchase Receipt", "Purchase Invoice", "Payment Entry", "GL Entry",
            "Payment Ledger Entry", "Stock Ledger Entry", "Bin", "China Accounting Voucher", "China Voucher Sync Issue", "Integration Request", "Item",
            "ToDo", "Version", "Notification Log", "China Cash Flow Assignment", "User", "Warehouse",
            "Workflow", "Workflow State", "Workflow Action Master", "Custom Field", "Advance Payment Ledger Entry", "Account", "Supplier")
        self.types += ("Sales Order", "Stock Reservation Entry", "Submission Queue", "Repost Item Valuation",
            "Material Request", "Purchase Fulfilment Link", "China Cash Equivalent Scope", "Delivery Note", "Customer", "Price List", "Item Price", "Sales Invoice")
        self.types += ("Stock Entry",)
        self.types += ("Stock Entry Type", "Project")  # exact new native default/opaque fixtures
        self.types += ("Product Bundle",)  # exact new native late-packed-items fixture
        self.types += ("Putaway Rule",)  # exact new late warehouse native fixture
        self.types += ("OA Purchase Request", "Comment")  # exact source recheck/write boundary fixture
        self.types += ("MES Integration Log", "MES Material Request Task", "Error Log")  # native C1 side effects
        self.types += ("DocShare", "DefaultValue")  # exact new User fixture side effects, never historical rows
        self.before = {doctype: frappe.db.count(doctype) for doctype in self.types}
        self.initial_names = {doctype: set(frappe.get_all(doctype, pluck="name", limit_page_length=0)) for doctype in self.types}
        self.committed_names = None
        self.addCleanup(self.clean)
        self.item = "QA-ATOMIC-" + uuid.uuid4().hex[:10]
        frappe.get_doc({"doctype": "Item", "item_code": self.item, "item_name": self.item, "item_group": "All Item Groups",
            "stock_uom": "Nos", "is_stock_item": 1}).insert()

    def clean(self):
        frappe.db.rollback()
        for doctype in ("Accounts Settings", "Buying Settings", "System Settings", "Stock Reposting Settings"):
            frappe.clear_document_cache(doctype, doctype)
        frappe.clear_document_cache("Company", COMPANY)
        if self.committed_names:
            # Only exact identities created by this isolated synthetic fixture;
            # never run this cleanup on another site/database or production.
            if frappe.local.site != SITE or frappe.conf.db_name != "qa_procurement_5" or frappe.conf.db_host != "db":
                raise RuntimeError("Synthetic cleanup scope changed")
            for doctype, names in self.committed_names.items():
                for field in frappe.get_meta(doctype).get_table_fields():
                    frappe.db.delete(field.options, {"parenttype": doctype, "parent": ["in", sorted(names)]})
                frappe.db.delete(doctype, {"name": ["in", sorted(names)]})
                if doctype == "User" and names:
                    frappe.db.delete("DefaultValue", {"parent": ["in", sorted(names)]})  # exact newly created users only
            frappe.db.commit()
        self.assertEqual(self.before, {doctype: frappe.db.count(doctype) for doctype in self.types})

    def commit_fixture(self, primary_doctype="Purchase Order", name_prefix="QA-ATOMIC-PO-"):
        frappe.db.commit()
        self.remember_new_names()
        self.assertTrue(self.committed_names[primary_doctype])
        self.assertTrue(all(name.startswith(name_prefix) for name in self.committed_names[primary_doctype]))

    def remember_new_names(self):
        """Immediate exact deltas, including subjects of a failed subTest."""
        if self.committed_names is None:
            self.committed_names = {doctype: set() for doctype in self.types}
        for doctype in self.types:
            self.committed_names[doctype] |= set(frappe.get_all(doctype, pluck="name", limit_page_length=0)) - self.initial_names[doctype]

    def order(self, currency="CNY", conversion_rate=1, supplier="QA Test Supplier", transaction_date=None, qty=10, rate=12.345, items=None, submit=True):
        po = frappe.get_doc({"doctype": "Purchase Order", "company": COMPANY,
            "transaction_date": transaction_date or nowdate(),
            "supplier": supplier, "currency": currency, "conversion_rate": conversion_rate, "schedule_date": add_days(nowdate(), 1),
            "items": items or [{"item_code": self.item, "qty": qty, "rate": rate,
                "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]}).insert(
                    set_name="QA-ATOMIC-PO-" + uuid.uuid4().hex[:16])
        if submit:
            po.submit()
        return po

    def test_batch_receipts_individual_merged_partial_and_durable_replay(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        self.assertTrue(callable(getattr(actions, "record_document_batch", None)), "Explicit native receipt batch is missing")
        for merge in (0, 1):
            with self.subTest(merge=merge):
                with patch("oa_purchase_request.oa_purchase_request.oa_purchase_request.auto_create_purchase_receipt"):
                    orders = [self.order(qty=5, rate=10), self.order(qty=6, rate=10)]
                self.commit_fixture()
                sources = [{"name": po.name, "modified": str(po.modified)} for po in orders]
                preview = actions.preview_document_batch(sources, merge=merge)
                edits = [{"items": [{"key": row["key"], "qty": 2, "warehouse": "Stores - QAB"}
                    for row in projection["document"]["items"]]} for projection in preview["documents"]]
                request = str(uuid.uuid4())
                result = actions.record_document_batch(sources, edits, request, merge=merge, confirm=1)
                self.assertEqual(len(result["documents"]), 1 if merge else 2)
                names = [row["document"]["name"] for row in result["documents"]]
                self.assertTrue(all(row["document"]["docstatus"] == 1 for row in result["documents"]))
                backlinks = {(row.purchase_order, row.purchase_order_item) for name in names
                    for row in frappe.get_doc("Purchase Receipt", name).items}
                self.assertEqual(backlinks, {(po.name, po.items[0].name) for po in orders})
                self.assertEqual([frappe.db.get_value("Purchase Order Item", po.items[0].name, "received_qty") for po in orders], [2, 2])
                self.commit_fixture()
                replay = actions.record_document_batch(sources, edits, request, merge=merge, confirm=1)
                self.assertTrue(replay["reused"])
                self.assertEqual([row["document"]["name"] for row in replay["documents"]], names)
                with self.assertRaises(frappe.ValidationError):
                    actions.record_document_batch(sources, [{}], request, merge=merge, confirm=1)

    def test_batch_second_native_receipt_failure_rolls_back_whole_confirmation(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        self.assertTrue(callable(getattr(actions, "record_document_batch", None)), "Explicit native receipt batch is missing")
        with patch("oa_purchase_request.oa_purchase_request.oa_purchase_request.auto_create_purchase_receipt"):
            orders = [self.order(qty=5), self.order(qty=5)]
        self.commit_fixture()
        sources = [{"name": po.name, "modified": str(po.modified)} for po in orders]
        before = {dt: frappe.db.count(dt) for dt in self.types}
        native_submit = actions._submit_document
        calls = []
        def second_failure(*args, **kwargs):
            result = native_submit(*args, **kwargs)
            calls.append(result)
            if len(calls) == 2:
                raise RuntimeError("QA batch second receipt after native stock/GL")
            return result
        with patch.object(actions, "_submit_document", side_effect=second_failure), self.assertRaisesRegex(RuntimeError, "second receipt"):
            actions.record_document_batch(sources, [{}, {}], str(uuid.uuid4()), confirm=1)
        self.assertEqual(len(calls), 2)
        self.assertEqual(before, {dt: frappe.db.count(dt) for dt in self.types})
        self.assertEqual([frappe.db.get_value("Purchase Order Item", po.items[0].name, "received_qty") for po in orders], [0, 0])
        print(json.dumps({"proof": "batch_second_receipt_rollback", "site": SITE, "orders": [po.name for po in orders],
            "before": before, "after": {dt: frappe.db.count(dt) for dt in self.types}, "received_after": [0, 0]}))

    def test_batch_merged_receipt_invoice_and_partial_payment_exact_shared_reference(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        self.assertTrue(callable(getattr(actions, "record_document_batch", None)), "Explicit native receipt batch is missing")
        with patch("oa_purchase_request.oa_purchase_request.oa_purchase_request.auto_create_purchase_receipt"):
            orders = [self.order(qty=3, rate=10), self.order(qty=4, rate=10)]
        self.commit_fixture()
        result = actions.record_document_batch([{"name": po.name, "modified": str(po.modified)} for po in orders],
            [{}], str(uuid.uuid4()), merge=1, confirm=1)
        receipt = result["documents"][0]["document"]
        self.commit_fixture()
        payable = actions.record_document_batch([{"name": receipt["name"], "modified": str(receipt["modified"])}],
            [{}], str(uuid.uuid4()), source_doctype="Purchase Receipt", target_doctype="Purchase Invoice", confirm=1)
        invoice = payable["documents"][0]["document"]
        self.assertEqual(invoice["docstatus"], 1)
        self.assertFalse(frappe.db.get_value("Purchase Invoice", invoice["name"], "update_stock"))
        self.commit_fixture()
        sources = [{"name": po.name} for po in orders]
        payables = service.preview_payment_batch("Purchase Order", sources)
        self.assertEqual([row["name"] for row in payables["invoices"]], [invoice["name"]])
        payment = actions.record_payment("Purchase Order", orders[0].name, sources=payables["sources"],
            allocations=[{"name": invoice["name"], "amount": 25}], bank_account="Cash - QAB", request_id=str(uuid.uuid4()))
        self.assertEqual(payment["document"]["docstatus"], 1)
        entry = frappe.get_doc("Payment Entry", payment["document"]["name"])
        self.assertEqual({row.reference_name for row in entry.references}, {invoice["name"]})
        self.assertEqual(sum(row.allocated_amount for row in entry.references), 25)
        self.assertEqual(frappe.db.get_value("Purchase Invoice", invoice["name"], "outstanding_amount"), 45)

    def test_batch_payment_gl_failure_retains_prior_payables_and_balances(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        self.assertTrue(callable(getattr(actions, "record_document_batch", None)), "Explicit native receipt batch is missing")
        with patch("oa_purchase_request.oa_purchase_request.oa_purchase_request.auto_create_purchase_receipt"):
            orders = [self.order(qty=3, rate=10), self.order(qty=4, rate=10)]
        self.commit_fixture()
        receipts = actions.record_document_batch([{"name": po.name, "modified": str(po.modified)} for po in orders],
            [{}, {}], str(uuid.uuid4()), confirm=1)["documents"]
        self.commit_fixture()
        invoices = actions.record_document_batch([{"name": row["document"]["name"], "modified": str(row["document"]["modified"])} for row in receipts],
            [{}, {}], str(uuid.uuid4()), source_doctype="Purchase Receipt", target_doctype="Purchase Invoice", confirm=1)["documents"]
        self.commit_fixture()
        before = {dt: frappe.db.count(dt) for dt in self.types}
        outstanding = {row["document"]["name"]: frappe.db.get_value("Purchase Invoice", row["document"]["name"], "outstanding_amount") for row in invoices}
        selected = service.preview_payment_batch("Purchase Order", [{"name": po.name} for po in orders])["sources"]
        native_submit = actions._confirm_payment
        def fail_after_gl(*args, **kwargs):
            native_submit(*args, **kwargs)
            raise RuntimeError("QA merged payment after native GL")
        with patch.object(actions, "_confirm_payment", side_effect=fail_after_gl), self.assertRaisesRegex(RuntimeError, "native GL"):
            actions.record_payment("Purchase Order", orders[0].name, sources=selected,
                allocations=[{"name": name, "amount": 10} for name in outstanding], bank_account="Cash - QAB", request_id=str(uuid.uuid4()))
        self.assertEqual(before, {dt: frappe.db.count(dt) for dt in self.types})
        self.assertEqual(outstanding, {name: frappe.db.get_value("Purchase Invoice", name, "outstanding_amount") for name in outstanding})
        print(json.dumps({"proof": "merged_payment_failure_retains_prior_AP", "site": SITE, "outstanding_before": outstanding,
            "outstanding_after": {name: frappe.db.get_value("Purchase Invoice", name, "outstanding_amount") for name in outstanding},
            "before": before, "after": {dt: frappe.db.count(dt) for dt in self.types}}))

    def test_batch_existing_automatic_receipt_drafts_preserve_manual_values(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        orders = [self.order(qty=3, rate=10), self.order(qty=4, rate=10)]
        drafts = []
        for po in orders:
            name = frappe.db.get_value("Purchase Receipt Item", {"purchase_order": po.name, "docstatus": 0}, "parent")
            draft = frappe.get_doc("Purchase Receipt", name)
            draft.items[0].qty = 1
            draft.items[0].received_qty = 0
            draft.items[0].warehouse = "Work In Progress - QAB"
            draft.remarks = "QA manual receipt"
            draft.save(); drafts.append(draft)
        self.commit_fixture()
        sources = [{"name": po.name, "modified": str(po.modified)} for po in orders]
        preview = actions.preview_document_batch(sources)
        self.assertEqual({row["document"]["name"] for row in preview["documents"]}, {draft.name for draft in drafts})
        self.assertTrue(all(row["document"]["items"][0]["qty"] == 1 for row in preview["documents"]))
        documents = [{"doctype": "Purchase Receipt", "name": row["document"]["name"], "modified": str(row["document"]["modified"])} for row in preview["documents"]]
        edits = [{} if row["document"]["sources"][0]["name"] == orders[0].name else
            {"remarks": "QA explicit edit", "items": [{"key": row["document"]["items"][0]["key"], "qty": 2,
                "warehouse": "Stores - QAB"}]} for row in preview["documents"]]
        result = actions.record_document_batch(sources, edits, str(uuid.uuid4()), documents=documents, confirm=1)
        self.assertEqual({row["document"]["name"] for row in result["documents"]}, {draft.name for draft in drafts})
        by_source = {row["document"]["sources"][0]["name"]: row["document"] for row in result["documents"]}
        self.assertEqual((by_source[orders[0].name]["remarks"], by_source[orders[0].name]["items"][0]["warehouse"]),
            ("QA manual receipt", "Work In Progress - QAB"))
        self.assertEqual((by_source[orders[1].name]["remarks"], by_source[orders[1].name]["items"][0]["warehouse"]),
            ("QA explicit edit", "Stores - QAB"))
        self.assertEqual([frappe.db.get_value("Purchase Order Item", po.items[0].name, "received_qty") for po in orders], [1, 2])

    def test_batch_receipt_rejects_changed_quantity_invalid_warehouse_and_merge_headers(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        with patch("oa_purchase_request.oa_purchase_request.oa_purchase_request.auto_create_purchase_receipt"):
            orders = [self.order(qty=3, rate=10), self.order(qty=4, rate=10)]
        self.commit_fixture()
        sources = [{"name": po.name, "modified": str(po.modified)} for po in orders]
        before = {dt: frappe.db.count(dt) for dt in self.types}
        for changes in ([{"items": [{"key": orders[0].items[0].name, "qty": 99}]}, {}],
                        [{"items": [{"key": orders[0].items[0].name, "warehouse": "All Warehouses - QAB"}]}, {}]):
            # Server groups are sorted by actual document identity.
            edits = sorted(zip(orders, changes), key=lambda pair: pair[0].name)
            with self.subTest(changes=changes):
                failed = actions.record_document_batch(sources, [row[1] for row in edits], str(uuid.uuid4()), confirm=1)
                self.assertTrue(failed.get("failed"))
            self.assertEqual(before, {dt: frappe.db.count(dt) for dt in self.types})
        # Another current receipt consumed the first order after the preview.
        taken = make_purchase_receipt(orders[0].name); taken.items[0].qty = 2; taken.insert().submit()
        self.commit_fixture()
        fresh = actions.preview_document_batch(sources)
        self.assertEqual(next(row["document"]["items"][0]["max_qty"] for row in fresh["documents"]
            if row["document"]["items"][0]["source_name"] == orders[0].name), 1)
        stale = [{"name": po.name, "modified": "1900-01-01" if po == orders[0] else str(po.modified)} for po in orders]
        before_stale = {dt: frappe.db.count(dt) for dt in self.types}
        failed = actions.record_document_batch(stale, [{}, {}], str(uuid.uuid4()), confirm=1)
        self.assertTrue(failed.get("failed")); self.assertIn("改变", failed["error"])
        self.assertEqual(before_stale, {dt: frappe.db.count(dt) for dt in self.types})
        for field, value in (("company", "YUEWEI MX"), ("currency", "USD"), ("conversion_rate", 2), ("discount_amount", 1)):
            with self.subTest(field=field):
                old = orders[1].get(field); orders[1].set(field, value)
                with self.assertRaises(frappe.ValidationError): actions._merge_compatible(orders)
                orders[1].set(field, old)
        orders[1].append("taxes", {"charge_type": "Actual", "account_head": "Stock Received But Not Billed - QAB", "tax_amount": 1})
        with self.assertRaisesRegex(frappe.ValidationError, "固定税费"):
            actions._merge_compatible(orders)

    def test_batch_payment_existing_draft_fresh_outstanding_hold_and_replay(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        with patch("oa_purchase_request.oa_purchase_request.oa_purchase_request.auto_create_purchase_receipt"):
            orders = [self.order(qty=3, rate=10), self.order(qty=4, rate=10)]
        self.commit_fixture()
        receipts = actions.record_document_batch([{"name": po.name, "modified": str(po.modified)} for po in orders],
            [{}, {}], str(uuid.uuid4()), confirm=1)["documents"]
        self.commit_fixture()
        invoices = actions.record_document_batch([{"name": row["document"]["name"], "modified": str(row["document"]["modified"])} for row in receipts],
            [{}, {}], str(uuid.uuid4()), source_doctype="Purchase Receipt", target_doctype="Purchase Invoice", confirm=1)["documents"]
        self.commit_fixture()
        preview = service.preview_payment_batch("Purchase Receipt", [{"name": row["document"]["name"]} for row in receipts])
        args = dict(source_doctype="Purchase Receipt", source_name=preview["sources"][0]["name"], sources=preview["sources"],
            allocations=[{"name": row["name"], "amount": 10} for row in preview["invoices"]], bank_account="Cash - QAB", request_id=str(uuid.uuid4()), confirm=0)
        draft = actions.record_payment(**args)
        self.assertEqual(draft["document"]["amount"], 20)
        self.commit_fixture()
        repeated = actions.record_payment(**args)
        self.assertTrue(repeated["reused"])
        existing = actions.record_payment(**{**args, "request_id": str(uuid.uuid4()), "confirm": 1})
        self.assertTrue(existing["needs_review"])
        self.assertEqual(existing["document"]["name"], draft["document"]["name"])
        self.assertEqual(existing["document"]["docstatus"], 0)
        # Current native hold prohibits this already saved merged draft.
        invoice = frappe.get_doc("Purchase Invoice", preview["invoices"][0]["name"])
        invoice.on_hold = 1; invoice.release_date = None; invoice.db_update()
        self.commit_fixture()
        with self.assertRaises(frappe.ValidationError):
            actions.submit_document("Payment Entry", draft["document"]["name"], draft["document"]["modified"])
        self.assertEqual(frappe.db.get_value("Payment Entry", draft["document"]["name"], "docstatus"), 0)
        invoice.on_hold = 0; invoice.db_update(); self.commit_fixture()
        # A separate real payment consumes the first invoice before confirmation.
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        paid = frappe.db.get_value("Purchase Invoice", invoice.name, "outstanding_amount") - 5
        other = get_payment_entry("Purchase Invoice", invoice.name, bank_account="Cash - QAB", bank_amount=paid)
        other.paid_amount = other.received_amount = paid
        for ref in other.references: ref.allocated_amount = paid
        other.insert().submit(); self.commit_fixture()
        with self.assertRaises(frappe.ValidationError):
            actions.submit_document("Payment Entry", draft["document"]["name"], draft["document"]["modified"])
        self.assertEqual(frappe.db.get_value("Payment Entry", draft["document"]["name"], "docstatus"), 0)

    def test_native_source_same_link_recheck_pending_blocks_evidence_and_comment_then_nested_write_rolls_back(self):
        from deeplinkerp_branding.services import purchase_source_service as sources, purchase_repost_boundary as boundary
        po = self.order(submit=False)
        raw = {"version": "QA-ATOMIC-SOURCE-OLD", "eligible": True, "currency": "CNY", "items": []}
        with sources.managed_write():
            oa = frappe.get_doc({"doctype": "OA Purchase Request", "oa_code": "QA-ATOMIC-OA-" + uuid.uuid4().hex,
                "purchase_order": po.name, "target_company": COMPANY,
                sources.SOURCE_FIELD: json.dumps(raw), sources.BOUND_FIELD: raw["version"]}).insert()
        frappe.db.set_value(po.doctype, po.name, "custom_oa_purchase_expense", oa.name, update_modified=False)
        owner = self.pending_owner()
        frappe.db.set_value(po.doctype, po.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        comments, before = frappe.db.count("Comment"), frappe.get_doc(oa.doctype, oa.name).as_dict()
        fresh = {**raw, "version": "QA-ATOMIC-SOURCE-NEW"}
        with self.assertRaisesRegex(frappe.ValidationError, "相关采购来源待完成"):
            sources._bind(oa, po.reload(), fresh, {}, "QA-ATOMIC-source recheck")
        self.assertEqual(frappe.get_doc(oa.doctype, oa.name).as_dict(), before)
        self.assertEqual(frappe.db.count("Comment"), comments)
        with boundary.execution(), boundary.acquire((boundary.fence_key(),)):
            with boundary.acquire((boundary.lock_key("document", po.doctype, po.name),)):
                frappe.db.set_value(po.doctype, po.name, boundary.POINTER, None, update_modified=False)
        self.commit_fixture()
        modified = frappe.db.get_value(po.doctype, po.name, "modified")
        sources._bind(oa, po.reload(), fresh, {}, "QA-ATOMIC-source recheck")
        self.assertEqual(frappe.db.get_value(po.doctype, po.name, "modified"), modified)
        self.assertEqual(frappe.db.get_value(oa.doctype, oa.name, sources.BOUND_FIELD), fresh["version"])
        self.assertEqual(frappe.db.count("Comment"), comments + 1)
        self.commit_fixture()
        # Changed-link still uses native PO save and its existing audit flow.
        # A later OA failure rolls both nested writes back as one execution.
        frappe.db.set_value(po.doctype, po.name, "custom_oa_purchase_expense", None, update_modified=False)
        self.commit_fixture()
        sql_set = frappe.db.set_value
        def fail_source(doctype, *args, **kwargs):
            if doctype == oa.doctype:
                raise RuntimeError("QA-ATOMIC source evidence failure")
            return sql_set(doctype, *args, **kwargs)
        audits = frappe.db.count("Integration Request")
        with patch.object(frappe.db, "set_value", side_effect=fail_source), self.assertRaisesRegex(RuntimeError, "source evidence failure"):
            sources._bind(oa, po.reload(), {**fresh, "version": "QA-ATOMIC-SOURCE-NEXT"}, {})
        self.assertFalse(frappe.db.get_value(po.doctype, po.name, "custom_oa_purchase_expense"))
        self.assertEqual(frappe.db.count("Integration Request"), audits)
        self.assertEqual(frappe.db.get_value(oa.doctype, oa.name, sources.BOUND_FIELD), fresh["version"])
        sources._bind(oa, po.reload(), fresh, {})
        self.assertEqual(frappe.db.get_value(po.doctype, po.name, "custom_oa_purchase_expense"), oa.name)

    def boundary_database(self):
        """Actual second mysqlclient business adapter, never a lock-only service."""
        from frappe.database import get_db
        database = get_db(socket=frappe.conf.db_socket, host=frappe.conf.db_host, port=frappe.conf.db_port,
            user=frappe.conf.db_user, password=frappe.conf.db_password, cur_db_name=frappe.conf.db_name)
        database.connect()
        self.addCleanup(lambda: database.close())
        return database

    def epoch_receipts(self, database, name):
        """Exact owned QA audit rows; exercise the actual physical seam."""
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        state = boundary.initialize(database)
        case = self
        class Participant:
            def __init__(self):
                self.pending, self.committed, self.resolutions = [], [], []
            def prepare(self, identity):
                prior = database.sql("SELECT data,output FROM `tabIntegration Request` WHERE name=%s FOR UPDATE", (name,))[0]
                receipt = {**identity, "effects": list(self.pending)}
                output = json.loads(prior[1] or "[]") + [receipt]
                database.sql("UPDATE `tabIntegration Request` SET output=%s WHERE name=%s", (json.dumps(output), name))
                return receipt
            def promote(self, receipt):
                self.committed.append(receipt)
                case.assertEqual(state.epoch, receipt["epoch"])
                self.pending.clear()
            def invalidate(self):
                self.pending.clear()
            def resolve(self, identity, candidate):
                case.assertTrue(state.poisoned)
                case.assertFalse(state.connection.open)
                fresh = case.boundary_database()
                case.assertIsNot(fresh._conn, state.connection)
                case.assertNotEqual(fresh._conn.thread_id(), state.connection_id)
                fresh.begin(read_only=True)
                result = fresh.sql("SELECT output FROM `tabIntegration Request` WHERE name=%s", (name,))[0][0]
                self.resolutions.append([row for row in json.loads(result or "[]") if all(row[key] == value for key, value in identity.items())])
                fresh.rollback()
                matches = self.resolutions[-1]
                return json.dumps(matches[0], ensure_ascii=False, sort_keys=True).encode("utf-8") if len(matches) == 1 else None
        state.receipt_participant = Participant()
        return state, state.receipt_participant

    def test_native_epoch_partial_receipts_callbacks_nested_raw_commit_and_outer_fence(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        immutable = frappe.db.get_value("Integration Request", audit.name, "data")
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        state = boundary.initialize(left)
        key = boundary.lock_key("qa-epoch-receipt", self.item)
        with boundary.execution(db=left), boundary.acquire((key,), db=left):
            state, receipts = self.epoch_receipts(left, audit.name)
            receipts.pending.append("first")
            def before():
                left.sql("COMMIT")
                receipts.pending.append("second")
            left.before_commit.add(before)
            left.after_commit.reset()
            left.commit()
            left.after_commit.add(lambda: receipts.pending.append("later"))
            left.commit()
            left.commit()
            self.assertEqual([row["effects"] for row in receipts.committed], [["first"], ["second"], [], ["later"]])
            self.assertEqual([row["epoch"] for row in receipts.committed], list(range(4)))
            self.assertEqual(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0], state.connection_id)
            for command in ("SAVEPOINT own", "ROLLBACK TO SAVEPOINT own", "RELEASE SAVEPOINT own"):
                with self.subTest(command=command), self.assertRaises(frappe.ValidationError): left.sql(command)
            for options in ({"run": False}, {"explain": True}):
                with self.subTest(options=options):
                    left.sql("COMMIT", **options)
                    self.assertEqual(state.epoch, 4)
        self.assertEqual(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0], state.connection_id)
        left.commit()  # supported outer physical boundary retains the owner until this point
        self.assertIsNone(state.receipt_participant)
        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0])
        right.rollback()
        data, output, status = right.sql("SELECT data,output,status FROM `tabIntegration Request` WHERE name=%s", (audit.name,))[0]
        self.assertEqual((data, status), (immutable, "Queued"))
        self.assertEqual(len(json.loads(output)), 5)
        print("C2A_PHYSICAL_EPOCHS=" + json.dumps(json.loads(output)), flush=True)

    def test_native_epoch_receipt_write_failure_full_rollback_and_later_status_commit(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        with boundary.execution(db=left), boundary.acquire((boundary.lock_key("qa-receipt-failure", self.item),), db=left):
            state, receipts = self.epoch_receipts(left, audit.name)
            left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-MUST-ROLLBACK", self.item))
            receipts.pending.append("rolled_back_gl")
            native_sql = state.native_sql
            def fail(query, *args, **kwargs):
                if str(query).startswith("UPDATE `tabIntegration Request` SET output="):
                    raise RuntimeError("QA receipt write failure")
                return native_sql(query, *args, **kwargs)
            with patch.object(state, "native_sql", side_effect=fail), self.assertRaisesRegex(RuntimeError, "QA receipt write failure"):
                left.commit()
            self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], self.item)
            self.assertFalse(receipts.pending)
            left.sql("UPDATE `tabIntegration Request` SET status='Failed' WHERE name=%s", (audit.name,))
            left.commit()
            receipts.pending.append("second_rollback_gl")
            left.sql("ROLLBACK")
            left.commit()
            self.assertEqual([row["effects"] for row in receipts.committed], [[], []])
        left.commit()
        right.rollback()
        output, status = right.sql("SELECT output,status FROM `tabIntegration Request` WHERE name=%s", (audit.name,))[0]
        self.assertEqual(status, "Failed")
        self.assertFalse(any(row["effects"] for row in json.loads(output)))

    def test_native_epoch_commit_response_unknown_resolves_durable_receipt_on_fresh_readonly_connection(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        for close_failure in (False, True):
            with self.subTest(close_failure=close_failure):
                audit = self.pending_owner()
                self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
                left, right = self.boundary_database(), self.boundary_database()
                key = boundary.lock_key("qa-unknown-receipt", self.item)
                with boundary.execution(db=left), boundary.acquire((key,), db=left):
                    state, receipts = self.epoch_receipts(left, audit.name)
                    receipts.pending.append("committed_synthetic_marker")
                    left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-DURABLE", self.item))
                    native_sql, commits, resolutions = state.native_sql, [], []
                    native_resolve = receipts.resolve
                    def resolve(identity, candidate):
                        resolutions.append(identity)
                        return native_resolve(identity, candidate)
                    receipts.resolve = resolve
                    physical = state.connection
                    native_close = type(physical).close
                    def close(connection, *args, **kwargs):
                        # Fault only this actual mysqlclient owner's close;
                        # keep the real native db/SQL/COMMIT implementation.
                        if close_failure and connection is physical: raise RuntimeError("QA physical close failed")
                        return native_close(connection, *args, **kwargs)
                    def unknown(query, *args, **kwargs):
                        result = native_sql(query, *args, **kwargs)
                        if str(query).lower() == "commit":
                            commits.append(query)
                            raise RuntimeError("QA lost commit response")
                        return result
                    with patch.object(type(physical), "close", close), patch.object(state, "native_sql", side_effect=unknown), self.assertRaisesRegex(RuntimeError, "QA lost commit response"):
                        left.commit()
                    self.assertEqual(commits, ["commit"])
                    self.assertEqual(len(resolutions), 0 if close_failure else 1)
                    if not close_failure:
                        self.assertEqual([row["effects"] for row in receipts.resolutions[0]], [["committed_synthetic_marker"]])
                    self.assertFalse(receipts.committed)
                    self.assertFalse(receipts.pending)
                    self.assertEqual(state.receipt_resolution["result"], "unknown" if close_failure else "durable")
                    self.assertEqual(bool(physical.open), close_failure)
                    with self.assertRaises(frappe.ValidationError): left.commit()
                left.close()  # exact own fixture owner after fault is removed
                right.rollback()
                self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], "QA-ATOMIC-DURABLE")
                self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0])

    def test_native_epoch_kill_close_poison_and_full_physical_rollback_discard_pending_markers(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        for ending in ("kill", "close", "physical", "full"):
            with self.subTest(ending=ending):
                left, right = self.boundary_database(), self.boundary_database()
                key = boundary.lock_key("qa-epoch-ending", self.item, ending)
                with boundary.execution(db=left), boundary.acquire((key,), db=left):
                    state, receipts = self.epoch_receipts(left, audit.name)
                    receipts.pending.append("rolled_back_synthetic_marker")
                    left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-ROLLBACK", self.item))
                    if ending == "kill":
                        right.sql("KILL CONNECTION %s", (state.connection_id,))
                        with self.assertRaises(Exception): left.sql("SELECT CONNECTION_ID()")
                    elif ending == "close": left.close()
                    elif ending == "physical": state.physical_rollback()
                    else: left.rollback()
                    self.assertFalse(receipts.pending)
                    self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], self.item)
                    if ending in ("kill", "close"):
                        self.assertTrue(state.poisoned)
                        with self.assertRaises(frappe.ValidationError): left.commit()
                        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0])
                    else:
                        left.commit()
                        self.assertEqual([row["effects"] for row in receipts.committed], [[]])
                if not state.poisoned: left.commit()

    def test_native_epoch_known_commit_promotion_failure_does_not_rollback_committed_business(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        with boundary.execution(db=left), boundary.acquire((boundary.lock_key("qa-epoch-promote", self.item),), db=left):
            state, receipts = self.epoch_receipts(left, audit.name)
            receipts.pending.append("durable_synthetic_marker")
            left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-KNOWN-COMMIT", self.item))
            receipts.promote = lambda receipt: (_ for _ in ()).throw(RuntimeError("QA original promotion failure"))
            with self.assertRaisesRegex(RuntimeError, "QA original promotion failure"): left.commit()
            self.assertTrue(state.poisoned)
        self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], "QA-ATOMIC-KNOWN-COMMIT")
        output = right.sql("SELECT output FROM `tabIntegration Request` WHERE name=%s", (audit.name,))[0][0]
        self.assertEqual([row["effects"] for row in json.loads(output)], [["durable_synthetic_marker"]])

    def test_native_epoch_unknown_response_requires_exact_candidate_not_identity_only(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        for result in ("mismatch", "absent"):
            with self.subTest(result=result):
                left = self.boundary_database()
                with boundary.execution(db=left), boundary.acquire((boundary.lock_key("qa-epoch-resolution", self.item, result),), db=left):
                    state, receipts = self.epoch_receipts(left, audit.name)
                    receipts.pending.append("synthetic_candidate")
                    native_sql = state.native_sql
                    def unknown(query, *args, **kwargs):
                        if str(query).lower() == "commit":
                            output = json.loads(native_sql("SELECT output FROM `tabIntegration Request` WHERE name=%s", (audit.name,))[0][0])
                            if result == "mismatch": output[-1]["effects"] = ["different_same_identity"]
                            else: output = output[:-1]
                            native_sql("UPDATE `tabIntegration Request` SET output=%s WHERE name=%s", (json.dumps(output), audit.name))
                            native_sql(query, *args, **kwargs)
                            raise RuntimeError("QA original unknown response")
                        return native_sql(query, *args, **kwargs)
                    with patch.object(state, "native_sql", side_effect=unknown), self.assertRaisesRegex(RuntimeError, "QA original unknown response"):
                        left.commit()
                    self.assertEqual(state.receipt_resolution["result"], result)
                    self.assertFalse(receipts.pending)
                    self.assertFalse(receipts.committed)
                    self.assertTrue(state.poisoned)

    def test_native_epoch_real_prepare_sql_error_never_commits_business_or_claims_durable(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        with boundary.execution(db=left), boundary.acquire((boundary.lock_key("qa-epoch-prepare-sql", self.item),), db=left):
            state, receipts = self.epoch_receipts(left, audit.name)
            left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-UNCOMMITTED", self.item))
            receipts.pending.append("synthetic_failed_prepare")
            original = receipts.prepare
            def invalid_sql(identity):
                receipt = original(identity)
                left.sql("UPDATE `tabIntegration Request` SET `qa_c2a_nonexistent_column`=1 WHERE name=%s", (audit.name,))
                return receipt
            with patch.object(receipts, "prepare", side_effect=invalid_sql), self.assertRaises(Exception) as caught: left.commit()
            self.assertEqual(caught.exception.args[0], 1054)
            self.assertIsNone(state.receipt_resolution)
            self.assertFalse(state.poisoned)
            self.assertFalse(receipts.pending)
            self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], self.item)
            self.assertIsNone(right.sql("SELECT output FROM `tabIntegration Request` WHERE name=%s", (audit.name,))[0][0])
            left.sql("UPDATE `tabIntegration Request` SET status='Failed' WHERE name=%s", (audit.name,))
            left.commit()
            self.assertEqual([row["effects"] for row in receipts.committed], [[]])
        left.commit()

    def test_native_epoch_after_commit_new_participant_waits_for_its_own_physical_boundary(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        key, later = (boundary.lock_key("qa-epoch-after", self.item, phase) for phase in ("old", "new"))
        with boundary.execution(db=left), boundary.acquire((key,), db=left):
            state, original = self.epoch_receipts(left, audit.name)
            original.pending.append("old_synthetic_marker")
        created = []
        def callback():
            with boundary.execution(db=left), boundary.acquire((later,), db=left):
                _, participant = self.epoch_receipts(left, audit.name)
                participant.pending.append("new_synthetic_marker")
                left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-AFTER-COMMIT", self.item))
                created.append(participant)
            raise RuntimeError("QA original after commit failure")
        left.after_commit.add(callback)
        with self.assertRaisesRegex(RuntimeError, "QA original after commit failure"): left.commit()
        self.assertEqual(original.committed[0]["epoch"], 0)
        self.assertEqual(created[0].pending, ["new_synthetic_marker"])
        self.assertFalse(created[0].committed)
        self.assertIs(state.receipt_participant, created[0])
        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0])
        self.assertEqual(right.sql("SELECT IS_USED_LOCK(%s)", (later,))[0][0], state.connection_id)
        self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], self.item)
        left.commit()
        self.assertEqual([(row["epoch"], row["effects"]) for row in created[0].committed], [(1, ["new_synthetic_marker"])])
        right.rollback()
        self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], "QA-ATOMIC-AFTER-COMMIT")
        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (later,))[0][0])
        self.assertIsNone(state.receipt_participant)

    def test_native_epoch_physical_rollback_callback_commit_preserves_prior_receipt_only(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        with boundary.execution(db=left), boundary.acquire((boundary.lock_key("qa-epoch-rollback-callback", self.item),), db=left):
            state, receipts = self.epoch_receipts(left, audit.name)
            left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-CALLBACK-COMMITTED", self.item))
            receipts.pending.append("old_committed_synthetic_marker")
            def before():
                left.commit()
                left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-CALLBACK-ROLLED-BACK", self.item))
                receipts.pending.append("new_rolled_back_synthetic_marker")
            left.before_rollback.add(before)
            state.physical_rollback()
            self.assertEqual([row["effects"] for row in receipts.committed], [["old_committed_synthetic_marker"]])
            self.assertFalse(receipts.pending)
            self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], "QA-ATOMIC-CALLBACK-COMMITTED")
            left.commit()
            self.assertEqual([row["effects"] for row in receipts.committed], [["old_committed_synthetic_marker"], []])
        left.commit()

    def test_native_epoch_raw_rollback_failure_closes_pending_business_owner(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        with boundary.execution(db=left), boundary.acquire((boundary.lock_key("qa-epoch-raw-rollback", self.item),), db=left):
            state, receipts = self.epoch_receipts(left, audit.name)
            left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-RAW-PENDING", self.item))
            receipts.pending.append("raw_rollback_pending_marker")
            native_sql = state.native_sql
            def failed(query, *args, **kwargs):
                if str(query).lower() == "rollback": raise RuntimeError("QA original raw rollback failure")
                return native_sql(query, *args, **kwargs)
            with patch.object(state, "native_sql", side_effect=failed), self.assertRaisesRegex(RuntimeError, "QA original raw rollback failure"): left.sql("ROLLBACK")
            self.assertTrue(state.poisoned)
            self.assertFalse(state.connection.open)
            self.assertFalse(receipts.pending)
            with self.assertRaises(frappe.ValidationError): left.commit()
            self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], self.item)

    def test_native_epoch_cleanup_close_never_lazily_reconnects_or_advances_old_epoch(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        for rollback in ("raw", "physical"):
            with self.subTest(rollback=rollback):
                left = self.boundary_database()
                with boundary.execution(db=left), boundary.acquire((boundary.lock_key("qa-epoch-cleanup-close", self.item, rollback),), db=left):
                    state, receipts = self.epoch_receipts(left, audit.name)
                    entered = []
                    def close():
                        if not entered:
                            entered.append(True)
                            left.close()
                    receipts.invalidate = close
                    if rollback == "raw":
                        with self.assertRaises(frappe.ValidationError): left.sql("ROLLBACK")
                    else: state.physical_rollback()
                    self.assertTrue(state.poisoned)
                    self.assertIsNone(left._conn)
                    self.assertEqual(state.epoch, 0)

    def test_native_epoch_prepare_failure_caught_then_new_callback_writes_are_rolled_back(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        with boundary.execution(db=left), boundary.acquire((boundary.lock_key("qa-epoch-caught-prepare", self.item),), db=left):
            state, receipts = self.epoch_receipts(left, audit.name)
            prepare = receipts.prepare
            receipts.prepare = lambda identity: (_ for _ in ()).throw(RuntimeError("QA first prepare error"))
            def before():
                with self.assertRaisesRegex(RuntimeError, "QA first prepare error"): left.sql("COMMIT")
                receipts.prepare = prepare
                left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-NEW-EPOCH-PENDING", self.item))
                receipts.pending.append("new_failed_callback_marker")
                raise RuntimeError("QA original later callback error")
            left.before_commit.add(before)
            with self.assertRaisesRegex(RuntimeError, "QA original later callback error"): left.commit()
            self.assertFalse(receipts.pending)
            self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], self.item)
            left.sql("UPDATE `tabIntegration Request` SET status='Failed' WHERE name=%s", (audit.name,))
            left.commit()
            self.assertEqual([row["effects"] for row in receipts.committed], [[]])
        left.commit()

    def test_native_epoch_owned_auto_commit_and_promote_reentry_fail_closed(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        audit = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left = self.boundary_database()
        with boundary.execution(db=left), boundary.acquire((boundary.lock_key("qa-epoch-no-reentry", self.item),), db=left):
            state, receipts = self.epoch_receipts(left, audit.name)
            for query in ("COMMIT", "ROLLBACK"):
                with self.subTest(query=query), self.assertRaises(frappe.ValidationError): left.sql(query, auto_commit=True)
        for phase in ("promote", "invalidate"):
            for action in ("sql", "method", "initialize", "execution", "acquire"):
                with self.subTest(phase=phase, action=action):
                    left = self.boundary_database()
                    key = boundary.lock_key("qa-memory-reentry", self.item, phase, action)
                    with boundary.execution(db=left), boundary.acquire((key,), db=left):
                        state, receipts = self.epoch_receipts(left, audit.name)
                        entered, reentry_queries = [], []
                        native_sql = state.native_sql
                        def observe(query, *args, **kwargs):
                            reentry_queries.append(str(query))
                            return native_sql(query, *args, **kwargs)
                        def reenter(*args):
                            if not entered:
                                entered.append(True)
                                with patch.object(state, "native_sql", side_effect=observe):
                                    if action == "method": left.commit()
                                    elif action == "sql": left.sql("SELECT 1")
                                    elif action == "initialize": boundary.initialize(left)
                                    elif action == "execution":
                                        with boundary.execution(db=left): pass
                                    else:
                                        with boundary.acquire((key,), db=left): pass
                        setattr(receipts, phase, reenter)
                        with self.assertRaises(frappe.ValidationError):
                            if phase == "promote": left.commit()
                            else: left.sql("ROLLBACK")
                        self.assertEqual(reentry_queries, [], "Memory-only callback must not query even IS_USED_LOCK")
                        self.assertTrue(state.poisoned)
                        self.assertFalse(state.connection.open)

    def test_native_session_lease_sql_epochs_callbacks_and_two_connection_authority(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        first = boundary.initialize(left)
        second = boundary.initialize(right)
        self.assertNotEqual(first.connection_id, second.connection_id)
        self.assertEqual(type(first.connection).__module__, "MySQLdb.connections")
        key, later = boundary.lock_key("qa-lease", self.item), boundary.lock_key("qa-later", self.item)
        def assert_owner(owner):
            self.assertEqual(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0], owner)
        for query in ("COMMIT", "ROLLBACK", "BEGIN", "CREATE TABLE qa_preview_only (name int)"):
            for option in ({"run": False}, {"explain": True}):
                for positional in (False, True):
                    for expected_begin in (False, True) if query == "BEGIN" else (False,):
                        with self.subTest(preview=query, option=option, positional=positional, expected_begin=expected_begin):
                            try:
                                with boundary.execution(db=left), boundary.acquire((key,), db=left):
                                    if expected_begin:
                                        left.sql("COMMIT")  # real active boundary, lease retained
                                    else:
                                        left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-PREVIEW-PENDING", self.item))
                                before = (first.epoch, first.begin_expected, dict(first.locks), dict(first.references), left.transaction_writes)
                                result = left.sql(query, (), **option) if positional else left.sql(query=query, values=(), **option)
                                right.rollback()  # fresh peer read, not an old RR snapshot
                                owner = right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0]
                                item_name = right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0]
                                print("SQL_PREVIEW_OBSERVATION", json.dumps({"query": query, "option": option,
                                    "positional": positional, "expected_begin": expected_begin, "epoch_before": before[0],
                                    "epoch_after": first.epoch, "owner": owner, "expected_owner": first.connection_id,
                                    "peer_item_name": item_name, "expected_item_name": self.item}), flush=True)
                                self.assertEqual(result, query if "run" in option else None)
                                self.assertEqual((first.epoch, first.begin_expected, first.locks, first.references, left.transaction_writes), before)
                                self.assertEqual(owner, first.connection_id)
                                self.assertEqual(item_name, self.item)  # preview never commits pending business
                                with self.assertRaises(frappe.ValidationError):
                                    with boundary.execution(db=right), boundary.acquire((key,), db=right): pass
                            finally:
                                right.rollback(); left.rollback()
        with boundary.execution(db=left), boundary.acquire((key,), db=left): pass
        before = (first.epoch, first.begin_expected, dict(first.locks))
        for args, kwargs in ((("COMMIT", (), False), {}), (("COMMIT",), {"unknown_sql_option": True})):
            with self.subTest(native_argument_error=args, kwargs=kwargs), self.assertRaises(TypeError):
                left.sql(*args, **kwargs)
            self.assertEqual((first.epoch, first.begin_expected, first.locks), before)
            assert_owner(first.connection_id)
        left.rollback()
        with boundary.execution(db=left), boundary.acquire((key,), db=left):
            left.commit(); assert_owner(first.connection_id)
            left.rollback(); assert_owner(first.connection_id)
            left.savepoint("scope_test"); left.rollback(save_point="scope_test"); assert_owner(first.connection_id)
            with self.assertRaises(frappe.ValidationError):
                with boundary.execution(db=right), boundary.acquire((key,), db=right): pass
        assert_owner(first.connection_id)
        left.before_commit.add(lambda: (_ for _ in ()).throw(RuntimeError("native early callback")))
        with self.assertRaisesRegex(RuntimeError, "native early callback"): left.commit()
        assert_owner(None)
        left.before_commit.reset()
        with boundary.execution(db=left), boundary.acquire((key,), db=left): pass
        left.before_rollback.add(lambda: (_ for _ in ()).throw(RuntimeError("native rollback callback")))
        with self.assertRaisesRegex(RuntimeError, "native rollback callback"): left.rollback()
        assert_owner(None)
        left.before_rollback.reset()
        with boundary.execution(db=left), boundary.acquire((key,), db=left): pass
        def next_epoch():
            with boundary.execution(db=left), boundary.acquire((later,), db=left): pass
            raise RuntimeError("after physical commit")
        left.after_commit.add(next_epoch)
        with self.assertRaisesRegex(RuntimeError, "after physical commit"): left.commit()
        assert_owner(None)
        self.assertEqual(right.sql("SELECT IS_USED_LOCK(%s)", (later,))[0][0], first.connection_id)
        left.commit()
        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (later,))[0][0])
        with boundary.execution(db=left), boundary.acquire((key,), db=left): pass
        left.close()
        assert_owner(None)
        with self.assertRaises(frappe.ValidationError): boundary.initialize(left)

    def test_native_killed_session_cannot_reconnect_or_reacquire(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        left, right = self.boundary_database(), self.boundary_database()
        state = boundary.initialize(left)
        key = boundary.lock_key("qa-kill", self.item)
        with boundary.execution(db=left), boundary.acquire((key,), db=left):
            right.sql("KILL CONNECTION %s", (state.connection_id,))
            with self.assertRaises(Exception): left.sql("SELECT CONNECTION_ID()")
            self.assertTrue(state.poisoned)
            with self.assertRaises(frappe.ValidationError): boundary.initialize(left)
        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0])

    def test_native_killed_before_first_handshake_poisons_execution_without_reconnect(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        left, right = self.boundary_database(), self.boundary_database()
        physical = left._conn
        right.sql("KILL CONNECTION %s", (physical.thread_id(),))
        with self.assertRaises(Exception): boundary.initialize(left)
        self.assertTrue(left._purchase_session.poisoned)
        self.assertFalse(physical.open)
        with self.assertRaises(frappe.ValidationError): boundary.initialize(left)
        self.assertIs(left._conn, physical)

    def test_nested_native_before_commit_rollback_then_new_writes_fail_physically_cleaned(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        state = boundary.initialize(left)
        key = boundary.lock_key("qa-nested-callback", self.item)
        committed_markers = []
        with boundary.execution(db=left), boundary.acquire((key,), db=left): pass
        def callback():
            left.rollback()
            with boundary.execution(db=left), boundary.acquire((key,), db=left):
                left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-PENDING", self.item))
                left.after_commit.add(lambda: committed_markers.append("rolled-back business"))
            raise RuntimeError("original nested callback failure")
        left.before_commit.add(callback)
        with self.assertRaisesRegex(RuntimeError, "original nested callback failure"):
            left.commit()
        self.assertFalse(state.poisoned)
        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0])
        self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], self.item)
        left.commit()  # genuine later native status/error-path commit
        self.assertEqual(committed_markers, [])

    def test_native_before_rollback_nested_commit_and_after_commit_new_epoch_data(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        left, right = self.boundary_database(), self.boundary_database()
        state = boundary.initialize(left)
        key, later = boundary.lock_key("qa-before-rollback", self.item), boundary.lock_key("qa-after-commit", self.item)
        markers = []
        def before():
            left.commit()  # actually commits prior chunk, cannot undo it
            left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-ROLLBACK-PENDING", self.item))
            left.after_commit.add(lambda: markers.append("rolled-back callback"))
            raise RuntimeError("original before rollback nested failure")
        with boundary.execution(db=left), boundary.acquire((key,), db=left):
            left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-COMMITTED-CHUNK", self.item))
            left.before_rollback.add(before)
            with self.assertRaisesRegex(RuntimeError, "original before rollback nested failure"):
                left.rollback()
            self.assertEqual(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0], state.connection_id)
            self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], "QA-ATOMIC-COMMITTED-CHUNK")
            left.commit()  # native status/error commit, still active lease
            self.assertEqual(markers, [])
        left.commit()
        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0])
        def after():
            with boundary.execution(db=left), boundary.acquire((later,), db=left):
                left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-NEXT-EPOCH", self.item))
            raise RuntimeError("after own real commit")
        left.after_commit.add(after)
        with self.assertRaisesRegex(RuntimeError, "after own real commit"): left.commit()
        self.assertEqual(right.sql("SELECT IS_USED_LOCK(%s)", (later,))[0][0], state.connection_id)
        right.rollback()
        self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], "QA-ATOMIC-COMMITTED-CHUNK")
        left.commit()  # only NEXT actual commit publishes/releases new epoch
        right.rollback()
        self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], "QA-ATOMIC-NEXT-EPOCH")
        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (later,))[0][0])

        for cleanup_failure in (False, True):
            with self.subTest(nested_rollback_cleanup_failure=cleanup_failure):
                left, right = self.boundary_database(), self.boundary_database()
                state = boundary.initialize(left)
                key = boundary.lock_key("qa-before-rollback-nested", self.item, cleanup_failure)
                previous = right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0]
                markers = []
                def nested_rollback():
                    left.rollback()  # nested SQL is not the enclosing rollback's SQL
                    left.sql("UPDATE tabItem SET item_name=%s WHERE name=%s", ("QA-ATOMIC-NESTED-PENDING", self.item))
                    left.after_commit.add(lambda: markers.append("rolled-back business"))
                    if cleanup_failure:
                        left.after_rollback.add(lambda: (_ for _ in ()).throw(RuntimeError("unsafe cleanup")))
                    raise RuntimeError("original nested rollback failure")
                with boundary.execution(db=left), boundary.acquire((key,), db=left):
                    left.before_rollback.add(nested_rollback)
                    with self.assertRaisesRegex(RuntimeError, "original nested rollback failure"):
                        left.rollback()
                    right.rollback()
                    self.assertEqual(right.sql("SELECT item_name FROM tabItem WHERE name=%s", (self.item,))[0][0], previous)
                    if cleanup_failure:
                        self.assertTrue(state.poisoned)
                        self.assertFalse(state.connection.open)
                        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0])
                        with self.assertRaises(frappe.ValidationError): left.commit()
                    else:
                        self.assertFalse(state.poisoned)
                        self.assertEqual(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0], state.connection_id)
                        left.commit()  # allowed native follow-up cannot carry stale callbacks
                    self.assertEqual(markers, [])
                if not cleanup_failure:
                    left.commit()
                    self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (key,))[0][0])

    def test_real_fresh_public_payment_and_source_entries_handshake_before_first_lock(self):
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        pe = get_payment_entry(pi.doctype, pi.name, bank_account="Cash - QAB", bank_amount=10).insert()
        self.commit_fixture()
        for action in ("entry-payment-update", "entry-payment-submit", "entry-source-create"):
            with self.subTest(action=action):
                pe.reload()
                result = self.native_peer({"action": action, "item": self.item, "name": pe.name,
                    "modified": str(pe.modified)})
                self.assertLess(result["identity_query"], result["first_lock"])
                self.assertTrue(result["same_physical"])
                self.assertTrue(result["execution_id"])
                frappe.db.rollback()  # only this fixture's read epoch, no pending business writes

    def pending_owner(self):
        return frappe.get_doc({"doctype": "Integration Request", "integration_request_service": "QA Boundary Pointer",
            "status": "Queued", "data": "{}"}).insert(ignore_permissions=True)

    def native_peer(self, payload):
        process = subprocess.run([sys.executable, __file__, "--boundary-peer", json.dumps(payload)],
            capture_output=True, text=True, timeout=30)
        self.assertEqual(process.returncode, 0, process.stdout + process.stderr)
        result = json.loads(process.stdout.strip().splitlines()[-1])
        # A may still own an old RR snapshot. Register the peer's exact typed
        # committed delta immediately, even if the next assertion fails.
        for doctype, names in result["created"].items():
            self.committed_names.setdefault(doctype, set()).update(names)
        return result

    def test_pending_bin_blocks_real_native_stock_draft_but_unrelated_first_stock_works(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        po = self.order()
        owner = self.pending_owner()
        frappe.db.set_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, boundary.POINTER, owner.name)
        self.commit_fixture()
        blocked = frappe.get_doc({"doctype": "Stock Entry", "purpose": "Material Receipt", "stock_entry_type": "Material Receipt", "company": COMPANY,
            "items": [{"item_code": self.item, "qty": 1, "t_warehouse": "Stores - QAB"}]})
        blocked.flags.purchase_reversal_internal = True
        with self.assertRaises(frappe.ValidationError): blocked.insert(ignore_permissions=True)
        extra = self.scope_item()
        self.assertFalse(frappe.db.exists("Bin", {"item_code": extra, "warehouse": "Stores - QAB"}))
        ordinary = frappe.get_doc({"doctype": "Stock Entry", "purpose": "Material Receipt", "stock_entry_type": "Material Receipt", "company": COMPANY,
            "items": [{"item_code": extra, "qty": 1, "t_warehouse": "Stores - QAB"}]}).insert()
        self.assertEqual(ordinary.docstatus, 0)
        self.assertEqual(ordinary.items[0].stock_uom, "Nos")
        self.assertFalse(frappe.db.exists("Bin", {"item_code": extra, "warehouse": "Stores - QAB"}))

    def test_pointer_rewrite_is_rejected_before_native_ignore_permission_save(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        po, owner = self.order(), self.pending_owner()
        self.commit_fixture()
        po.reload()
        po.set(boundary.POINTER, owner.name)
        po.flags.purchase_reversal_internal = True
        with self.assertRaises(frappe.PermissionError): po.save(ignore_permissions=True)
        self.assertFalse(frappe.db.get_value(po.doctype, po.name, boundary.POINTER))

    def test_pending_bin_unchanged_pointer_blocks_native_qty_api_and_same_pair_insert(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        from frappe.client import set_value
        self.order()
        owner = self.pending_owner()
        name = frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "name")
        frappe.db.set_value("Bin", name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        for action in ("native", "rpc", "insert"):
            with self.subTest(action=action):
                original = frappe.get_doc("Bin", name)
                before = original.actual_qty
                with self.assertRaisesRegex(frappe.ValidationError, "相关库存范围待完成"):
                    if action == "native":
                        original.actual_qty = before + 7
                        original.flags.purchase_reversal_internal = True
                        original.save(ignore_permissions=True)
                    elif action == "rpc":
                        set_value("Bin", name, "actual_qty", before + 7)
                    else:
                        frappe.get_doc({"doctype": "Bin", "item_code": self.item, "warehouse": "Stores - QAB",
                            "actual_qty": 7}).insert(ignore_permissions=True)
                self.assertEqual(frappe.db.get_value("Bin", name, "actual_qty"), before)
                self.assertEqual(frappe.db.count("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}), 1)

    def test_material_request_real_purchase_backlink_pointer_blocks_native_save(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        mr, po = self.material_request_order()
        po.submit()
        owner = self.pending_owner()
        frappe.db.set_value(po.doctype, po.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        mr.reload()
        with self.assertRaises(frappe.ValidationError): mr.save(ignore_permissions=True)

    def test_native_material_request_status_rpc_pending_bin_and_backlinks_then_native_resume(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        from erpnext.stock.doctype.material_request.material_request import update_status
        mr, po = self.material_request_order()
        owner = self.pending_owner()
        bin_name = frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "name")
        self.commit_fixture()
        for doctype, name in (("Bin", bin_name), (po.doctype, po.name)):
            frappe.db.set_value(doctype, name, boundary.POINTER, owner.name, update_modified=False)
            self.commit_fixture()
            before = (frappe.db.get_value(mr.doctype, mr.name, "status"),
                frappe.db.get_value("Bin", bin_name, "indented_qty"),
                frappe.db.get_value("Purchase Order Item", po.items[0].name, "material_request_item"))
            with self.assertRaisesRegex(frappe.ValidationError, "相关.*待完成"):
                update_status(mr.name, "Stopped")
            self.assertEqual(before, (frappe.db.get_value(mr.doctype, mr.name, "status"),
                frappe.db.get_value("Bin", bin_name, "indented_qty"),
                frappe.db.get_value("Purchase Order Item", po.items[0].name, "material_request_item")))
            frappe.db.set_value(doctype, name, boundary.POINTER, None, update_modified=False)
            self.commit_fixture()
        update_status(mr.name, "Stopped")
        self.assertEqual(frappe.db.get_value("Bin", bin_name, "indented_qty"), 0)
        self.commit_fixture()
        update_status(mr.name, "Pending")
        self.assertEqual(frappe.db.get_value("Bin", bin_name, "indented_qty"), 10)
        self.commit_fixture()
        po.reload().submit()
        pr = make_purchase_receipt(po.name).insert()
        frappe.db.set_value(pr.doctype, pr.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        before = (frappe.db.get_value(mr.doctype, mr.name, "status"), frappe.db.get_value("Bin", bin_name, "indented_qty"))
        with self.assertRaisesRegex(frappe.ValidationError, "相关.*待完成"):
            update_status(mr.name, "Stopped")
        self.assertEqual((frappe.db.get_value(mr.doctype, mr.name, "status"),
            frappe.db.get_value("Bin", bin_name, "indented_qty")), before)

    def test_native_purchase_status_rpc_and_bulk_preflight_all_before_any_write(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        first, second = self.order(), self.order()
        receipt = self.receipt(first)
        owner = self.pending_owner()
        self.commit_fixture()
        cases = [("Purchase Order", second.name,
            "erpnext.buying.doctype.purchase_order.purchase_order.update_status", {"name": second.name, "status": "Closed"}),
            ("Purchase Receipt", receipt.name,
            "erpnext.stock.doctype.purchase_receipt.purchase_receipt.update_purchase_receipt_status",
            {"docname": receipt.name, "status": "Closed"})]
        for doctype, name, path, kwargs in cases:
            frappe.db.set_value(doctype, name, boundary.POINTER, owner.name, update_modified=False)
            self.commit_fixture()
            before = frappe.db.get_value(doctype, name, "status")
            with self.assertRaisesRegex(frappe.ValidationError, "相关.*待完成"):
                frappe.get_attr(frappe.override_whitelisted_method(path))(**kwargs)
            self.assertEqual(frappe.db.get_value(doctype, name, "status"), before)
            frappe.db.set_value(doctype, name, boundary.POINTER, None, update_modified=False)
            self.commit_fixture()
        frappe.db.set_value(second.doctype, second.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        path = "erpnext.buying.doctype.purchase_order.purchase_order.close_or_unclose_purchase_orders"
        before = [frappe.db.get_value("Purchase Order", doc.name, "status") for doc in (first, second)]
        with self.assertRaisesRegex(frappe.ValidationError, "相关.*待完成"):
            frappe.get_attr(frappe.override_whitelisted_method(path))(json.dumps([first.name, second.name]), "Closed")
        self.assertEqual([frappe.db.get_value("Purchase Order", doc.name, "status") for doc in (first, second)], before)
        frappe.db.set_value(second.doctype, second.name, boundary.POINTER, None, update_modified=False)
        self.commit_fixture()
        audits = set(frappe.get_all("Integration Request", pluck="name"))
        frappe.get_attr(frappe.override_whitelisted_method(path))(json.dumps([first.name, second.name]), "Closed")
        self.assertEqual([frappe.db.get_value("Purchase Order", doc.name, "status") for doc in (first, second)], ["Closed", "Closed"])
        new_audits = set(frappe.get_all("Integration Request", pluck="name")) - audits
        self.assertEqual(len(new_audits), 1)  # one existing operation/postcheck graph, not per-doc queues
        facts = json.loads(frappe.get_doc("Integration Request", next(iter(new_audits))).data)
        for doc in (first, second):
            self.assertEqual(facts["after"][doc.doctype + ":" + doc.name]["status"], "Closed")

    def test_native_order_update_child_qty_rate_guards_before_delete_or_child_writes_and_validates_parent(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        from erpnext.controllers import accounts_controller as native
        po = self.order(items=[{"item_code": self.item, "qty": qty, "rate": 12.345,
            "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)} for qty in (10, 2)])
        other = self.order()
        owner = self.pending_owner()
        frappe.db.set_value(po.doctype, po.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        path = "erpnext.controllers.accounts_controller.update_child_qty_rate"
        def data(rows, qty=8):
            return json.dumps([{"docname": row.name, "item_code": row.item_code, "qty": qty,
                "rate": row.rate, "uom": row.uom, "conversion_factor": row.conversion_factor,
                "schedule_date": str(row.schedule_date), "description": row.description} for row in rows])
        for rows in (po.items, po.items[:1]):
            with patch.object(native, "validate_and_delete_children", wraps=native.validate_and_delete_children) as deletion:
                with self.assertRaisesRegex(frappe.ValidationError, "相关.*待完成"):
                    frappe.get_attr(frappe.override_whitelisted_method(path))(po.doctype, data(rows), po.name)
                deletion.assert_not_called()  # actual native early delete helper must not run
            self.assertEqual([row.qty for row in frappe.get_doc(po.doctype, po.name).items], [10, 2])
        frappe.db.set_value(po.doctype, po.name, boundary.POINTER, None, update_modified=False)
        self.commit_fixture()
        with patch.object(native, "validate_and_delete_children", wraps=native.validate_and_delete_children) as deletion:
            with self.assertRaisesRegex(frappe.ValidationError, "明细.*所属"):
                frappe.get_attr(frappe.override_whitelisted_method(path))(po.doctype, data(other.items), po.name)
            deletion.assert_not_called()
        self.assertEqual(frappe.db.get_value("Purchase Order Item", other.items[0].name, "qty"), 10)
        frappe.get_attr(frappe.override_whitelisted_method(path))(po.doctype, data(po.items, 8), po.name)
        self.assertEqual([row.qty for row in frappe.get_doc(po.doctype, po.name).items], [8, 8])
        self.commit_fixture()
        po.reload()
        frappe.get_attr(frappe.override_whitelisted_method(path))(po.doctype, data(po.items[:1], 7), po.name)
        self.assertEqual([row.qty for row in frappe.get_doc(po.doctype, po.name).items], [7])

    def test_native_order_update_child_qty_rate_cross_parent_refused_before_native_helper(self):
        from erpnext.controllers import accounts_controller as native
        po, other = self.order(), self.order()
        self.commit_fixture()
        row = other.items[0]
        data = json.dumps([{"docname": row.name, "item_code": row.item_code, "qty": 8, "rate": row.rate,
            "uom": row.uom, "conversion_factor": row.conversion_factor, "schedule_date": str(row.schedule_date),
            "description": row.description}])
        with patch.object(native, "validate_and_delete_children", wraps=native.validate_and_delete_children) as deletion:
            with self.assertRaisesRegex(frappe.ValidationError, "明细.*所属"):
                frappe.get_attr(frappe.override_whitelisted_method("erpnext.controllers.accounts_controller.update_child_qty_rate"))(
                    po.doctype, data, po.name)
            deletion.assert_not_called()

    def test_native_purchase_invoice_hold_db_set_pending_refuses_entire_block_and_clear_works(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        from erpnext.accounts.doctype.purchase_invoice.purchase_invoice import block_invoice, unblock_invoice, change_release_date
        po = self.order()
        pi = make_purchase_invoice(self.receipt(po).name).insert()
        owner = self.pending_owner()
        frappe.db.set_value(pi.doctype, pi.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        before = frappe.db.get_value(pi.doctype, pi.name, ["on_hold", "hold_comment", "release_date"])
        for call in (lambda: block_invoice(pi.name, add_days(nowdate(), 2), "QA hold"),
                lambda: unblock_invoice(pi.name), lambda: change_release_date(pi.name, add_days(nowdate(), 3))):
            with self.assertRaisesRegex(frappe.ValidationError, "相关.*待完成"):
                call()
            self.assertEqual(frappe.db.get_value(pi.doctype, pi.name, ["on_hold", "hold_comment", "release_date"]), before)
        frappe.db.set_value(pi.doctype, pi.name, boundary.POINTER, None, update_modified=False)
        self.commit_fixture()
        block_invoice(pi.name, add_days(nowdate(), 2), "QA hold")
        self.assertEqual(frappe.db.get_value(pi.doctype, pi.name, ["on_hold", "hold_comment"]), (1, "QA hold"))
        self.commit_fixture()
        change_release_date(pi.name, add_days(nowdate(), 3))
        self.assertEqual(str(frappe.db.get_value(pi.doctype, pi.name, "release_date")), add_days(nowdate(), 3))
        unblock_invoice(pi.name)
        self.assertEqual(frappe.db.get_value(pi.doctype, pi.name, ["on_hold", "release_date"]), (0, None))

    def test_native_purchase_invoice_partial_hold_failure_rolls_back_prior_real_fields(self):
        from erpnext.accounts.doctype.purchase_invoice.purchase_invoice import block_invoice, unblock_invoice, change_release_date
        po = self.order()
        pi = make_purchase_invoice(self.receipt(po).name).insert()
        self.commit_fixture()
        fields = ["on_hold", "hold_comment", "release_date", "modified", "modified_by"]
        for action, failed_field, after_write in (("block", "hold_comment", False),
                ("block", "release_date", False), ("unblock", "release_date", False),
                ("change-release", "release_date", True)):
            with self.subTest(action=action, failed_field=failed_field):
                if action == "block":
                    unblock_invoice(pi.name)
                else:
                    block_invoice(pi.name, add_days(nowdate(), 2), "QA original hold")
                self.commit_fixture()  # distinct real request baseline, never inside tested failure
                before = frappe.db.get_value(pi.doctype, pi.name, fields)
                counts = {doctype: frappe.db.count(doctype) for doctype in self.types}
                self.remember_effects()  # shared native stock/financial/source/audit evidence
                written, partial = [], []
                native_set = frappe.db.set_value
                first_error = RuntimeError("QA-ATOMIC first hold failure " + action + " " + failed_field)
                def fail_field(doctype, name, fieldname, *args, **kwargs):
                    target = doctype == pi.doctype and name == pi.name and fieldname in fields[:3]
                    if target and fieldname == failed_field and not after_write:
                        partial.append(frappe.db.get_value(pi.doctype, pi.name, fields))
                        raise first_error
                    result = native_set(doctype, name, fieldname, *args, **kwargs)
                    if target:
                        written.append(fieldname)
                        if fieldname == failed_field:
                            partial.append(frappe.db.get_value(pi.doctype, pi.name, fields))
                            raise first_error
                    return result
                with patch.object(frappe.db, "set_value", side_effect=fail_field), self.assertRaises(RuntimeError) as caught:
                    if action == "block":
                        block_invoice(pi.name, add_days(nowdate(), 4), "QA changed hold")
                    elif action == "unblock":
                        unblock_invoice(pi.name)
                    else:
                        change_release_date(pi.name, add_days(nowdate(), 4))
                self.assertIs(caught.exception, first_error)
                expected = ["on_hold", "hold_comment"] if action == "block" and failed_field == "release_date" else (
                    ["release_date"] if action == "change-release" else ["on_hold"])
                self.assertEqual(written, expected)  # preceding native SQL really ran
                self.assertEqual(len(partial), 1)
                self.assertNotEqual(partial[0][:3], before[:3])
                # No test-side rollback: the actual boundary has restored the
                # whole request, including prior db_set writes and timestamps.
                self.assertEqual(frappe.db.get_value(pi.doctype, pi.name, fields), before)
                self.assertEqual({doctype: frappe.db.count(doctype) for doctype in self.types}, counts)
                self.assert_effects_unchanged()

    def test_material_request_backlink_requires_actual_detail_and_shared_union_budget(self):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        mr, po = self.material_request_order()
        detail = po.items[0].material_request_item
        self.commit_fixture()
        audit_before = frappe.db.count("Integration Request")
        frappe.db.set_value("Purchase Order Item", po.items[0].name, "material_request_item",
            "QA-ATOMIC-DETAIL-MISSING-" + uuid.uuid4().hex, update_modified=False)
        self.commit_fixture()
        with self.assertRaises(frappe.ValidationError) as caught:
            mr.reload().save(ignore_permissions=True)
        self.assertEqual(caught.exception.purchase_error_id, "native_source_identity_mismatch")
        self.assertEqual(frappe.db.count("Integration Request"), audit_before)
        # A parent-only historical string is not a native detail updater edge.
        # Do not repair history, and do not permanently disable an ordinary MR.
        frappe.db.set_value("Purchase Order Item", po.items[0].name, "material_request_item", None, update_modified=False)
        self.commit_fixture()
        mr.reload().save(ignore_permissions=True)
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "indented_qty"), 10)
        frappe.db.set_value("Purchase Order Item", po.items[0].name, "material_request_item", detail, update_modified=False)
        self.commit_fixture()
        with patch.object(scope, "MAX_VOUCHERS", 2):
            mr.reload().save(ignore_permissions=True)  # duplicate old/new MR counts once
        with patch.object(scope, "MAX_VOUCHERS", 1), self.assertRaises(frappe.ValidationError):
            mr.reload().save(ignore_permissions=True)  # real PO backlink is identity two

    def test_current_scope_sees_peer_native_transfer_and_rejects_expansion_before_write(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        po, prior, future = self.future_receipts()
        frappe.db.rollback()  # release only commit_fixture's evidence read locks, BEFORE tested scope
        original, peer = boundary._scopes, {}
        def interleave(roots, *, current, opaque):
            result = original(roots, current=current, opaque=opaque)
            if not current and not peer:
                peer.update(self.native_peer({"action": "transfer", "item": self.item}))
            return result
        audit_before = frappe.db.count("Integration Request")
        with patch.object(boundary, "_scopes", side_effect=interleave):
            with self.assertRaises((frappe.QueryDeadlockError, frappe.ValidationError)) as caught:
                prior.save(ignore_permissions=True)
        if isinstance(caught.exception, frappe.QueryDeadlockError):
            self.assertEqual(caught.exception.args[0].args[0], 1020)
        else:
            self.assertIn("范围已扩大", str(caught.exception))
        # Exact peer-created synthetic identities, collected before any failure
        # rolls A back; no historical record or broad cleanup target is added.
        frappe.db.rollback()
        self.remember_new_names()
        fresh = boundary._scopes([prior], current=True, opaque=False)
        self.assertIn(peer["name"], {row.name for scope in fresh for row in scope.vouchers})
        self.assertEqual(frappe.db.count("Integration Request"), audit_before)

    def test_current_scope_peer_native_source_replacement_refuses_old_snapshot_then_fresh_reads_new_identity(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        first = self.order()
        second = self.order(items=[{"item_code": self.item, "qty": 10, "rate": 12.345,
            "warehouse": "Work In Progress - QAB", "schedule_date": add_days(nowdate(), 1)}])
        receipt = make_purchase_receipt(first.name).insert()
        self.commit_fixture()
        frappe.db.rollback()  # fixture-only evidence epoch, before A's tested initial scope
        original, peer = boundary._scopes, {}
        def interleave(roots, *, current, opaque):
            result = original(roots, current=current, opaque=opaque)
            if not current and not peer:
                peer.update(self.native_peer({"action": "source", "item": self.item, "name": receipt.name,
                    "source": second.name, "detail": second.items[0].name}))
            return result
        with patch.object(boundary, "_scopes", side_effect=interleave):
            with self.assertRaises((frappe.QueryDeadlockError, frappe.ValidationError)):
                receipt.save(ignore_permissions=True)
        frappe.db.rollback()
        current = frappe.get_doc(receipt.doctype, receipt.name)
        scopes = boundary._scopes([current], current=True, opaque=False)
        sources = {(row.doctype, row.name) for scope in scopes for row in scope.sources}
        self.assertIn((second.doctype, second.name), sources)
        self.assertNotIn((first.doctype, first.name), sources)
        self.assertEqual((current.items[0].purchase_order, current.items[0].purchase_order_item,
            current.items[0].warehouse), (second.name, second.items[0].name, "Work In Progress - QAB"))

    def test_real_rpc_payment_reference_and_party_replacement_cannot_hide_old_pending_purchase_invoice(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        from frappe.client import save
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        pe = get_payment_entry(pi.doctype, pi.name, bank_account="Cash - QAB", bank_amount=10).insert()
        owner = self.pending_owner()
        frappe.db.set_value(pi.doctype, pi.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        pe.reload()
        before = [(row.reference_doctype, row.reference_name) for row in pe.references]
        pe.set("references", [])
        pe.party_type, pe.party = "Customer", "QA Purchase Payments Customer"
        pe.flags.ignore_permissions = True
        pe.flags.purchase_reversal_internal = True
        with self.assertRaisesRegex(frappe.ValidationError, "相关采购来源待完成"):
            save(json.dumps(pe.as_dict(), default=str))
        stored = frappe.get_doc(pe.doctype, pe.name)
        self.assertEqual(stored.party_type, "Supplier")
        self.assertEqual([(row.reference_doctype, row.reference_name) for row in stored.references], before)

    def test_native_cached_type_and_item_snapshot_mismatch_refuses_then_new_request_works(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        entry_type = frappe.get_doc({"doctype": "Stock Entry Type", "name": "QA-ATOMIC-TYPE-" + uuid.uuid4().hex[:10],
            "purpose": "Material Receipt", "is_standard": 0}).insert()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        entry_type.purpose = "Material Transfer"
        with self.assertRaises(frappe.CannotChangeConstantError): entry_type.save()
        frappe.db.rollback()  # native immutable-default proof, before A's tested snapshot
        frappe.get_cached_value(entry_type.doctype, entry_type.name, "purpose")
        frappe.db.get_value(entry_type.doctype, entry_type.name, ("name", "purpose"), as_dict=True, cache=True)
        self.assertEqual(frappe.db.get_value("Item", self.item, "stock_uom"), "Nos")
        self.native_peer({"action": "defaults", "item": self.item, "entry_type": entry_type.name})
        draft = frappe.get_doc({"doctype": "Stock Entry", "company": COMPANY, "stock_entry_type": entry_type.name,
            "items": [{"item_code": self.item, "qty": 1, "t_warehouse": "Stores - QAB"}]})
        with self.assertRaises((frappe.QueryDeadlockError, frappe.ValidationError)) as caught: draft.insert()
        if isinstance(caught.exception, frappe.QueryDeadlockError):
            self.assertEqual(caught.exception.args[0].args[0], 1020)
        else:
            self.assertIn("快照", str(caught.exception))
        # A genuinely new native process/physical session, not clearing the old
        # execution identity or changing isolation in-place.
        peer = self.native_peer({"action": "draft", "item": self.item, "entry_type": entry_type.name})
        self.assertEqual((peer["purpose"], peer["stock_uom"]), ("Material Receipt", "Kg"))
        frappe.db.rollback()
        self.remember_new_names()

    def test_all_six_pointer_fields_reject_real_rpc_rewrite_clear_and_db_set_flags(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary, purchase_operation as operation
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        from frappe.client import insert, save
        from frappe.model.document import Document
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pe = get_payment_entry(po.doctype, po.name, bank_account="Cash - QAB", party_amount=10).insert()
        stock_bin = frappe.get_doc("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"})
        riv = frappe.get_doc({"doctype": "Repost Item Valuation", "company": COMPANY,
            "based_on": "Item and Warehouse", "item_code": self.item, "warehouse": "Stores - QAB",
            "posting_date": nowdate(), "posting_time": "00:00:00"}).insert()
        owner = self.pending_owner()
        docs = [riv, stock_bin, po, pr, pi, pe]
        self.assertEqual({doc.doctype for doc in docs}, set(boundary.POINTER_TYPES))
        for doc in docs:
            frappe.db.set_value(doc.doctype, doc.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        names = {doctype: set(frappe.get_all(doctype, pluck="name", limit_page_length=0)) for doctype in self.types}
        owned = {doc.doctype: set(frappe.get_all(doc.doctype, filters={boundary.POINTER: owner.name},
            pluck="name", limit_page_length=0)) for doc in docs}
        native_insert = Document.insert
        def rpc_payload(subject):
            data = subject.as_dict()
            data["flags"] = {"ignore_permissions": True, "purchase_reversal_internal": True}
            return json.dumps(data, default=str)
        def reject_action(doc, action, expected_owner):
            current = frappe.get_doc(doc.doctype, doc.name)
            current.flags.purchase_reversal_internal = True
            current.flags.ignore_permissions = True
            inserted = []
            def observe_insert(subject, *args, **kwargs):
                inserted.append((subject.doctype, subject.name))
                return native_insert(subject, *args, **kwargs)
            try:
                with patch.object(Document, "insert", new=observe_insert), patch.object(
                        operation, "_reserve", wraps=operation._reserve) as reserved:
                    with self.assertRaises(frappe.PermissionError):
                        if action == "native-insert-new":
                            current.insert(set_name="QA-ATOMIC-COPY-" + uuid.uuid4().hex, ignore_permissions=True)
                        elif action == "rpc-insert-old":
                            insert(rpc_payload(current))
                        elif action.endswith(("-zero", "-false")):
                            value = 0 if action.endswith("-zero") else False
                            current.set(boundary.POINTER, value)
                            if action.startswith("rpc-insert"):
                                created = insert(rpc_payload(current))
                                identity = (created["doctype"], created["name"])
                            elif action.startswith("rpc-save"):
                                save(rpc_payload(current))
                                identity = (current.doctype, current.name)
                            elif action.startswith("native-save"):
                                current.save(ignore_permissions=True)
                                identity = (current.doctype, current.name)
                            else:
                                current.db_set(boundary.POINTER, value)
                                identity = (current.doctype, current.name)
                            # RED observes actual persisted values before rollback;
                            # GREEN rejects before reaching any native write.
                            print("FALSY_POINTER_WRITE=" + json.dumps({"doctype": identity[0],
                                "name": identity[1], "action": action, "stored_pointer":
                                frappe.db.get_value(*identity, boundary.POINTER)}))
                        elif action in ("rpc-save-local", "native-save-local"):
                            current.set("__islocal", 1)
                            if action == "rpc-save-local":
                                save(rpc_payload(current))
                            else:
                                current.save(ignore_permissions=True)
                        elif action == "rpc-save-no-name":
                            current.name = None
                            save(rpc_payload(current))
                        elif action == "native-copy-owner":
                            frappe.copy_doc(current).save(ignore_permissions=True)
                        elif action == "rpc-clear":
                            current.set(boundary.POINTER, None)
                            save(rpc_payload(current))
                        elif action == "native-rewrite":
                            current.set(boundary.POINTER, "QA-ATOMIC-CLIENT-" + uuid.uuid4().hex)
                            current.save(ignore_permissions=True)
                        else:
                            current.db_set(boundary.POINTER, None)
                        self.assertEqual(set(frappe.get_all(doc.doctype, filters={boundary.POINTER: owner.name},
                            pluck="name", limit_page_length=0)), owned[doc.doctype], "Owner copied into a new identity")
                    self.assertEqual(inserted, [], "Rejected before native insert")
                    reserved.assert_not_called()  # including _save -> insert, not merely late rejection
            finally:
                self.remember_new_names()
                frappe.db.rollback()
            self.assertEqual(frappe.db.get_value(doc.doctype, doc.name, boundary.POINTER), expected_owner)
            self.assertEqual(names, {doctype: set(frappe.get_all(doctype, pluck="name", limit_page_length=0))
                for doctype in self.types})
            self.assertEqual({subject.doctype: set(frappe.get_all(subject.doctype,
                filters={boundary.POINTER: owner.name}, pluck="name", limit_page_length=0)) for subject in docs}, owned)
        for doc in docs:
            for action in ("native-insert-new", "rpc-insert-old", "rpc-save-local", "native-save-local",
                    "rpc-save-no-name", "native-copy-owner", "rpc-insert-zero", "rpc-insert-false",
                    "rpc-clear", "native-rewrite", "db-set"):
                with self.subTest(doctype=doc.doctype, action=action):
                    reject_action(doc, action, owner.name)
            # Native copy/amend preparation honors metadata no_copy, independent
            # of whether this particular original is currently frozen.
            safe = frappe.copy_doc(frappe.get_doc(doc.doctype, doc.name), ignore_no_copy=False)
            self.assertFalse(safe.get(boundary.POINTER))
            self.assertTrue(safe.get("__islocal"))
            self.assertFalse(safe.name)
            self.assertEqual(frappe.db.get_value(doc.doctype, doc.name, boundary.POINTER), owner.name)
        # The same ordinary-save rule also protects previously unowned stored
        # identities. Only this exact synthetic fixture changes its baseline.
        for doc in docs:
            frappe.db.set_value(doc.doctype, doc.name, boundary.POINTER, None, update_modified=False)
        self.commit_fixture()
        owned = {doc.doctype: set() for doc in docs}
        for doc in docs:
            for action in ("rpc-save-zero", "rpc-save-false", "native-save-zero", "native-save-false",
                    "db-set-zero", "db-set-false"):
                with self.subTest(doctype=doc.doctype, action=action, previous_owner=None):
                    reject_action(doc, action, None)
        for doc in docs:
            frappe.db.set_value(doc.doctype, doc.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        # Existing same-identity RIV saves keep generation ownership; a genuine
        # native no-copy draft gets its own identity without that ownership.
        original = frappe.get_doc(riv.doctype, riv.name)
        original.save(ignore_permissions=True)
        copied = frappe.copy_doc(original, ignore_no_copy=False).insert(ignore_permissions=True)
        self.assertNotEqual(copied.name, original.name)
        self.assertFalse(copied.get(boundary.POINTER))
        self.assertEqual(frappe.db.get_value(original.doctype, original.name, boundary.POINTER), owner.name)
        # Actual ordinary amendment on a distinct unfrozen stock pair still
        # uses native no-copy/naming/validation; no old identity is rewritten.
        code = self.scope_item()
        amended_source = self.order(items=[{"item_code": code, "qty": 2, "rate": 12,
            "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}], submit=False)
        self.assertEqual(amended_source.meta.get_field("transaction_time").fieldtype, "Time")
        amended_source.transaction_time = "00:11:22.123456"
        amended_source.submit()
        with patch.object(frappe, "log_error", wraps=frappe.log_error) as native_log:
            amended_source.cancel()
            self.assertEqual(native_log.call_args.kwargs["reference_doctype"], amended_source.doctype)
            self.assertEqual(native_log.call_args.kwargs["reference_name"], amended_source.name)
            errors = frappe.db.get_values("Error Log", {"reference_doctype": amended_source.doctype,
                "reference_name": amended_source.name}, ["name", "reference_doctype", "reference_name", "method", "error"], as_dict=True)
            self.assertEqual(len(errors), 1)
            self.assertEqual(errors[0].method, "中国会计凭证冲销同步记录创建失败")
            self.assertIn("in _ensure_cancellation_sync_issue", errors[0].error)
            print("NATIVE_POINTER_FIXTURE_ERRORLOG=" + json.dumps(errors, default=str), flush=True)
        amended = frappe.copy_doc(amended_source, ignore_no_copy=False)
        amended.amended_from = amended_source.name
        amended.insert(ignore_permissions=True)
        self.assertNotEqual(amended.name, amended_source.name)
        self.assertFalse(amended.get(boundary.POINTER))
        self.assertEqual(frappe.db.get_value(amended_source.doctype, amended_source.name, "docstatus"), 2)
        self.remember_new_names()  # native MyISAM Error Log survives the fixture's final rollback

    def test_real_rpc_old_source_and_warehouse_replacement_cannot_escape_pending(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        from frappe.client import save
        po = self.order()
        pr = make_purchase_receipt(po.name).insert()
        owner = self.pending_owner()
        frappe.db.set_value(po.doctype, po.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        pr.reload()
        old_pair = (pr.items[0].purchase_order, pr.items[0].warehouse)
        pr.items[0].purchase_order = None
        pr.items[0].purchase_order_item = None
        pr.items[0].warehouse = "Work In Progress - QAB"
        pr.flags.purchase_reversal_internal = True
        pr.flags.ignore_permissions = True
        with self.assertRaisesRegex(frappe.ValidationError, "相关采购来源待完成"):
            save(json.dumps(pr.as_dict(), default=str))
        stored = frappe.get_doc(pr.doctype, pr.name)
        self.assertEqual((stored.items[0].purchase_order, stored.items[0].warehouse), old_pair)
        # The native Material Request really changes requested Bin quantity;
        # both its pair and actual PR backlink must participate, no MR pointer.
        mr, second_po = self.material_request_order()
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "indented_qty"), 10)
        second_po.submit()
        second_pr = make_purchase_receipt(second_po.name).insert()
        self.assertEqual(second_pr.items[0].material_request, mr.name)
        frappe.db.set_value(second_pr.doctype, second_pr.name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        mr.reload()
        mr.items[0].qty += 1
        with self.assertRaisesRegex(frappe.ValidationError, "相关采购来源待完成"):
            mr.save(ignore_permissions=True)
        self.assertEqual(frappe.db.get_value("Material Request Item", mr.items[0].name, "qty"), 10)

    def test_opaque_native_and_stored_mes_zero_pending_fence_blocks_cross_company_publication(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        from deeplinkerp_branding.services.purchase_reversal_scope import collect_document_scope
        project = frappe.get_doc({"doctype": "Project", "project_name": "QA-ATOMIC-PROJECT-" + uuid.uuid4().hex[:10],
            "company": COMPANY}).insert()
        mr = frappe.get_doc({"doctype": "Material Request", "company": "Yuewei", "material_request_type": "Purchase",
            "transaction_date": nowdate(), "schedule_date": add_days(nowdate(), 1),
            "items": [{"item_code": self.item, "qty": 3, "warehouse": "Stores - Y", "schedule_date": add_days(nowdate(), 1)}]}).insert()
        # Exact synthetic stored classification, not an actor flag exemption.
        frappe.db.set_value(mr.doctype, mr.name, "custom_request_source", "MES", update_modified=False)
        owner = self.pending_owner()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        mr.reload()
        mr.flags.mes_integration_request = False
        mr.submit()  # installed performance mixin delegates this ordinary branch
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - Y"}, "indented_qty"), 3)
        with self.assertRaisesRegex(frappe.ValidationError, "MES"):
            collect_document_scope(mr)  # still NOT an async cancellation adapter
        frappe.db.commit()
        # B1 resolver refusal is read-only; no scope/audit/business rollback.
        self.remember_new_names()
        opaque = frappe.get_doc({"doctype": "Stock Entry", "company": COMPANY, "project": project.name,
            "stock_entry_type": "Material Receipt", "items": [{"item_code": self.item, "qty": 1,
                "t_warehouse": "Stores - QAB", "allow_zero_valuation_rate": 1}]}).insert()
        self.assertEqual(opaque.docstatus, 0)  # genuine native no-pending compatibility
        right = self.boundary_database()
        with self.assertRaises(frappe.ValidationError):
            with boundary.execution(db=right), boundary.acquire((boundary.fence_key(),), db=right): pass
        frappe.db.commit()  # method return alone did NOT free the dormant fence
        self.remember_new_names()
        name = frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - Y"}, "name")
        with boundary.execution(db=right), boundary.acquire((boundary.fence_key(),), db=right):
            # Native QB-backed set_value resolves frappe.db, not this peer
            # instance. Explicit SQL is essential to prove B really commits.
            right.sql("UPDATE `tabBin` SET custom_purchase_reversal_operation=%s WHERE name=%s", (owner.name, name))
        right.commit()  # D-style fixture publication, NOT implemented D authority
        self.assertEqual(right.sql("SELECT custom_purchase_reversal_operation FROM `tabBin` WHERE name=%s", (name,))[0][0], owner.name)
        with self.assertRaises((frappe.QueryDeadlockError, frappe.ValidationError)):
            opaque.reload().save(ignore_permissions=True)
        mr.reload()
        self.assertTrue(boundary._opaque(mr), mr.custom_request_source)
        self.assertEqual(frappe.db.get_value("Bin", name, boundary.POINTER), owner.name)
        with self.assertRaisesRegex(frappe.ValidationError, "未核查的原生关联路径"):
            mr.save(ignore_permissions=True)
        with boundary.execution(db=right), boundary.acquire((boundary.fence_key(),), db=right):
            right.sql("UPDATE `tabBin` SET custom_purchase_reversal_operation=NULL WHERE name=%s", (name,))
        right.commit()
        opaque.reload().save(ignore_permissions=True)
        mr.reload().save(ignore_permissions=True)

    def test_mes_parent_commit_with_lost_callback_retains_native_intent(self):
        from types import SimpleNamespace
        from frappe.utils import CallbackManager
        request = SimpleNamespace(after_response=CallbackManager())
        with patch.object(frappe.local, "request", request, create=True):
            mr = frappe.get_doc({"doctype": "Material Request", "company": COMPANY,
                "material_request_type": "Purchase", "transaction_date": nowdate(),
                "schedule_date": add_days(nowdate(), 1), "items": [{"item_code": self.item,
                    "qty": 3, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]})
            mr.flags.mes_integration_request = True
            mr.insert()
            mr.submit()
            request.after_response.reset()  # crash/lost callback after parent commit
            frappe.db.commit()
            self.remember_new_names()
        rows = frappe.db.get_values("Integration Request", {
            "integration_request_service": "DeepLinkERP native MES Bin intent",
            "reference_doctype": "Material Request", "reference_docname": mr.name},
            ["name", "request_id", "status", "data"], as_dict=True, for_update=True)
        self.assertEqual(len(rows), 1, "Parent commit must independently retain native Bin intent")
        facts = json.loads(rows[0].data)
        self.assertEqual(facts["material_request_name"], mr.name)
        self.assertEqual(facts["pairs"], [[self.item, "Stores - QAB"]])
        self.assertEqual(facts["user"], "Administrator")
        self.assertEqual(facts["site"], SITE)
        self.assertEqual(facts["database"], "qa_procurement_5")
        self.assertEqual(rows[0].status, "Queued")

    def test_mes_parent_rollback_surviving_after_response_cannot_enqueue_phantom(self):
        from types import SimpleNamespace
        from frappe.utils import CallbackManager
        import mes_integration.mes_integration.material_request as mes
        request = SimpleNamespace(after_response=CallbackManager())
        with patch.object(frappe.local, "request", request, create=True):
            mr = frappe.get_doc({"doctype": "Material Request", "company": COMPANY,
                "material_request_type": "Purchase", "transaction_date": nowdate(),
                "schedule_date": add_days(nowdate(), 1), "items": [{"item_code": self.item,
                    "qty": 3, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]})
            mr.flags.mes_integration_request = True
            mr.insert()
            mr.submit()
            frozen_name = mr.name
            frappe.db.rollback()
            mr.name = "QA-ATOMIC-PHANTOM-" + uuid.uuid4().hex
            with patch.object(mes, "enqueue_mes_material_request_bin_sync_job") as enqueue:
                request.after_response.run()  # native HTTP closing iterator survives rollback
            self.assertFalse(frappe.db.exists("Material Request", frozen_name))
            enqueue.assert_not_called()

    def test_native_intent_ir_paths_reject_current_and_proposed_forgery_before_commit(self):
        from itertools import product
        from frappe.model.rename_doc import rename_doc
        paths = list(product(("save", "db_set", "update_status", "success", "failure", "rename", "delete", "insert",
            "incoming-prefix-insert", "incoming-prefix-rename", "db_insert", "db_update"), ("Queued", "Completed"), ("DLP-MES-BIN-",)))
        paths += list(product(("incoming-prefix-insert", "incoming-prefix-rename", "db_insert", "save", "db_set",
            "update_status", "success", "failure", "rename", "delete", "db_update"), ("Queued", "Completed"),
            ("dlp-mes-bin-", "D\u0139P-M\u00c9S-B\u00cdN-")))
        for action, initial_status, prefix in paths:
            with self.subTest(action=action, current_status=initial_status, prefix=prefix):
                self.assertEqual(frappe.db.sql("SELECT CAST(%s AS CHAR CHARACTER SET utf8mb4) "
                    "COLLATE utf8mb4_unicode_ci LIKE %s", (prefix + uuid.uuid4().hex, "DLP-MES-BIN-%"))[0][0], 1)
                name = prefix + uuid.uuid4().hex
                frappe.db.sql("INSERT INTO `tabIntegration Request` "
                    "(name,creation,modified,owner,modified_by,integration_request_service,request_id,"
                    "request_description,status,data) VALUES (%s,NOW(6),NOW(6),%s,%s,%s,%s,%s,%s,%s)",
                    (name, "Administrator", "Administrator",
                        "DeepLinkERP native MES Bin intent" if prefix == "DLP-MES-BIN-" else "QA ordinary", uuid.uuid4().hex,
                        "Native MES Bin sync acknowledged" if initial_status == "Completed" else "Native MES Bin sync pending",
                        initial_status, json.dumps({"generation": name})))
                frappe.db.commit()
                self.remember_new_names()
                doc = frappe.get_doc("Integration Request", name, for_update=True)
                if action == "save":
                    doc.integration_request_service = None
                    doc.status = "Completed"
                    doc.request_description = "Native MES Bin sync acknowledged"
                    call = lambda: doc.save(ignore_permissions=True)
                elif action == "db_set":
                    call = lambda: doc.db_set({"integration_request_service": "ordinary", "status": "Completed"})
                elif action == "update_status":
                    call = lambda: doc.update_status({"generation": "forged"}, "Completed")
                elif action == "success":
                    call = lambda: doc.handle_success({"ack": True})
                elif action == "failure":
                    call = lambda: doc.handle_failure({"error": "forged"})
                elif action == "rename":
                    call = lambda: rename_doc("Integration Request", name, "QA-ATOMIC-RENAMED-" + uuid.uuid4().hex,
                        ignore_permissions=True)
                elif action == "delete":
                    call = lambda: doc.delete(ignore_permissions=True)
                elif action == "insert":
                    doc = frappe.get_doc({"doctype": "Integration Request", "integration_request_service":
                        "DeepLinkERP native MES Bin intent", "status": "Completed",
                        "request_description": "Native MES Bin sync acknowledged", "data": "{}"})
                    call = lambda: doc.insert(ignore_permissions=True)
                elif action == "incoming-prefix-insert":
                    doc = frappe.get_doc({"doctype": "Integration Request", "integration_request_service": "QA ordinary",
                        "status": "Queued", "data": "{}"})
                    incoming_name = prefix + uuid.uuid4().hex
                    call = lambda: doc.insert(ignore_permissions=True, set_name=incoming_name)
                elif action == "incoming-prefix-rename":
                    doc = frappe.get_doc({"doctype": "Integration Request", "integration_request_service": "QA ordinary",
                        "status": "Queued", "data": "{}"}).insert(ignore_permissions=True)
                    frappe.db.commit()
                    self.remember_new_names()
                    incoming_name = prefix + uuid.uuid4().hex
                    call = lambda: rename_doc("Integration Request", doc.name, incoming_name,
                        ignore_permissions=True)
                elif action == "db_insert":
                    doc = frappe.get_doc({"doctype": "Integration Request", "name": prefix + uuid.uuid4().hex,
                        "integration_request_service": "QA ordinary", "status": "Queued", "data": "{}"})
                    call = doc.db_insert
                else:
                    doc.integration_request_service = "QA ordinary"
                    call = doc.db_update
                try:
                    with patch.object(frappe.db, "commit", wraps=frappe.db.commit) as commit, \
                            patch.object(frappe.db, "sql", wraps=frappe.db.sql) as sql:
                        error = None
                        try:
                            call()
                        except Exception as caught:
                            error = caught
                        if not isinstance(error, frappe.PermissionError):
                            target = incoming_name if action in ("incoming-prefix-insert", "incoming-prefix-rename") else doc.name
                            print("C1_SPEC_PREFIX_ADMISSION=" + json.dumps({"action": action, "prefix": prefix,
                                "target": target, "persisted": frappe.db.get_value("Integration Request", target,
                                    ["name", "integration_request_service", "status"], as_dict=True, for_update=True),
                                "error_type": type(error).__name__ if error else None}, default=str, sort_keys=True), flush=True)
                        self.assertIsInstance(error, frappe.PermissionError)
                        commit.assert_not_called()
                        self.assertFalse(any(str(entry.args[0]).lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))
                            for entry in sql.call_args_list), "Namespace refusal must precede native writes")
                    self.assertEqual(frappe.db.get_value("Integration Request", name, "status", for_update=True), initial_status)
                finally:
                    self.remember_new_names()  # native update_status may have committed before RED failure
                    frappe.db.rollback()
        ordinary = frappe.get_doc({"doctype": "Integration Request", "integration_request_service": "QA ordinary",
            "status": "Queued", "data": "{}"}).insert(ignore_permissions=True)
        ordinary.data = "{\"ordinary\": true}"
        ordinary.save(ignore_permissions=True)
        ordinary.db_set("request_description", "QA ordinary phase")
        ordinary.update_status({"result": "ordinary"}, "Completed")
        self.remember_new_names()
        ordinary.handle_failure({"error": "ordinary"})
        self.remember_new_names()
        ordinary.handle_success({"result": "ordinary"})
        self.remember_new_names()
        renamed = rename_doc("Integration Request", ordinary.name, "QA-ATOMIC-ORDINARY-" + uuid.uuid4().hex,
            force=True, ignore_permissions=True)
        self.remember_new_names()
        frappe.get_doc("Integration Request", renamed).delete(ignore_permissions=True)

    def test_native_intent_exact_generated_active_index_handles_unknown_terminal_variants(self):
        column = "custom_purchase_native_intent_active"
        rows = frappe.db.sql("SELECT COLUMN_TYPE,EXTRA,GENERATION_EXPRESSION FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='tabIntegration Request' AND COLUMN_NAME=%s", (column,), as_dict=True)
        self.assertEqual(len(rows), 1, "Native intent requires the exact prospective generated active index")
        self.assertEqual(rows[0].COLUMN_TYPE, "tinyint(4)")
        self.assertEqual(rows[0].EXTRA, "VIRTUAL GENERATED")
        for service_value, status, phase, active in (
                ("DeepLinkERP native MES Bin intent", "Completed", "Native MES Bin sync acknowledged", 0),
                ("DeepLinkERP native MES Bin intent", "completed", "Native MES Bin sync acknowledged", 1),
                ("DeepLinkERP native MES Bin intent", "Completed ", "Native MES Bin sync acknowledged", 1),
                ("DeepLinkERP native MES Bin intent", None, "Native MES Bin sync acknowledged", 1),
                ("DeepLinkERP native MES Bin intent", "", "Native MES Bin sync acknowledged", 1),
                ("DeepLinkERP native MES Bin intent", "Failed", "Native MES Bin sync pending", 1),
                ("DeepLinkERP native MES Bin intent", "alien", "Native MES Bin sync pending", 1),
                ("DeepLinkERP native MES Bin intent", "Completed", None, 1),
                ("DeepLinkERP native MES Bin intent", "Completed", "", 1),
                ("DeepLinkERP native MES Bin intent", "Completed", "native mes bin sync acknowledged", 1),
                ("DeepLinkERP native MES Bin intent", "Completed", "Native MES Bin sync acknowledged ", 1),
                ("DeepLinkERP native MES Bin intent ", "Completed", "Native MES Bin sync acknowledged", 1),
                ("deeplinkerp native mes bin intent", "Completed", "Native MES Bin sync acknowledged", 1)):
            with self.subTest(service=service_value, status=status, phase=phase):
                name = "DLP-MES-BIN-" + uuid.uuid4().hex
                frappe.db.sql("INSERT INTO `tabIntegration Request` "
                    "(name,integration_request_service,status,request_description) VALUES (%s,%s,%s,%s)",
                    (name, service_value, status, phase))
                actual = frappe.db.sql("SELECT custom_purchase_native_intent_active FROM `tabIntegration Request` "
                    "WHERE name=%s FOR UPDATE", (name,))[0][0]
                self.assertEqual(actual, active)

    def test_native_intent_installer_conflicting_actual_metadata_rejects_before_any_ddl(self):
        from copy import deepcopy
        from deeplinkerp_branding import purchase_reversal_install as installer
        from deeplinkerp_branding.services import purchase_native_intent as intent
        native = frappe.db.sql
        columns = native("SELECT COLUMN_NAME,COLUMN_TYPE,CHARACTER_SET_NAME,COLLATION_NAME,EXTRA,GENERATION_EXPRESSION "
            "FROM information_schema.COLUMNS WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='tabIntegration Request'", as_dict=True)
        indexes = native("SHOW INDEX FROM `tabIntegration Request`", as_dict=True)
        for kind in ("column_type", "column_charset", "column_collation", "generated_type", "generated_stored",
                "generated_case", "generated_service", "index_order", "index_prefix", "index_unique", "index_ignored", "index_missing"):
            with self.subTest(conflict=kind):
                actual_columns, actual_indexes = deepcopy(columns), deepcopy(indexes)
                target = next(row for row in actual_columns if row.COLUMN_NAME == intent.ACTIVITY_COLUMN)
                source = next(row for row in actual_columns if row.COLUMN_NAME == "integration_request_service")
                index = next(row for row in actual_indexes if row.Key_name == intent.ACTIVITY_INDEX and row.Seq_in_index == 1)
                if kind.startswith("column_"):
                    source[{"column_type": "COLUMN_TYPE", "column_charset": "CHARACTER_SET_NAME", "column_collation": "COLLATION_NAME"}[kind]] = "unknown"
                elif kind == "generated_type": target.COLUMN_TYPE = "int(11)"
                elif kind == "generated_stored": target.EXTRA = "STORED GENERATED"
                elif kind == "generated_case": target.GENERATION_EXPRESSION = target.GENERATION_EXPRESSION.replace("Completed", "completed")
                elif kind == "generated_service": target.GENERATION_EXPRESSION = "CASE WHEN status='Completed' THEN 0 ELSE 1 END"
                elif kind == "index_order": index.Column_name = "name"
                elif kind == "index_prefix": index.Sub_part = 20
                elif kind == "index_unique": index.Non_unique = 0
                elif kind == "index_ignored": index.Ignored = "YES"
                else: actual_indexes.remove(index)
                def observed(query, *args, **kwargs):
                    if "information_schema.COLUMNS" in str(query):
                        return [row for row in actual_columns if row.COLUMN_NAME == args[0][0]] if "COLUMN_NAME=%s" in str(query) else actual_columns
                    if str(query).startswith("SHOW INDEX"):
                        return [row for row in actual_indexes if row.Key_name == args[0][0]] if "WHERE Key_name=%s" in str(query) else actual_indexes
                    return native(query, *args, **kwargs)
                with patch.object(frappe.db, "sql", side_effect=observed), \
                        patch.object(frappe.db, "sql_ddl", side_effect=AssertionError("No fixture DDL permitted")) as ddl, \
                        patch.object(frappe.db, "add_index", side_effect=AssertionError("No fixture index write permitted")) as add, \
                        patch.object(frappe.db, "commit", wraps=frappe.db.commit) as commit:
                    with self.assertRaisesRegex(frappe.ValidationError, "冲突|定义未经"):
                        installer.install_native_intent_index()
                    ddl.assert_not_called(); add.assert_not_called(); commit.assert_not_called()
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        for kind in ("index_upper_missing_generated", "index_upper_conflict", "column_upper_conflict",
                "column_upper_expression_conflict", "compatible_upper_alias", "compatible_all_upper_columns"):
            with self.subTest(native_identifier=kind):
                table = "tabQA_C1_IDENT_" + uuid.uuid4().hex[:12]
                self.assertFalse(native("SELECT TABLE_NAME FROM information_schema.TABLES WHERE TABLE_SCHEMA=DATABASE() "
                    "AND TABLE_NAME=%s", (table,)))
                fields = ("name", "integration_request_service", "status", "request_description")
                physical = tuple(name.upper() for name in fields) if kind == "compatible_all_upper_columns" else fields
                definition = ",".join("`" + name + "` VARCHAR(140)" for name in physical)
                generated = intent.ACTIVITY_COLUMN if kind == "index_upper_conflict" else intent.ACTIVITY_COLUMN.upper()
                if kind != "index_upper_missing_generated":
                    expression = intent.ACTIVITY_EXPRESSION.replace("Completed", "completed") if kind == "column_upper_expression_conflict" else intent.ACTIVITY_EXPRESSION
                    definition += ",`" + generated + "` " + ("INT" if kind == "column_upper_conflict" else
                        "TINYINT AS (" + expression + ") VIRTUAL")
                if kind.startswith("index_upper") or kind.startswith("compatible"):
                    index_fields = (physical[0],) if kind.startswith("index_upper") else (physical[1], generated, physical[0])
                    definition += ",INDEX `" + intent.ACTIVITY_INDEX.upper() + "`(" + ",".join("`" + name + "`" for name in index_fields) + ")"
                native("CREATE TABLE `" + table + "` (" + definition + ") ENGINE=InnoDB ROW_FORMAT=DYNAMIC "
                    "DEFAULT CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci")
                try:
                    actual = native("SELECT COLUMN_NAME,COLUMN_TYPE,CHARACTER_SET_NAME,COLLATION_NAME,EXTRA,GENERATION_EXPRESSION "
                        "FROM information_schema.COLUMNS WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME=%s", (table,), as_dict=True)
                    found = native("SHOW INDEX FROM `" + table + "` WHERE Key_name=%s", (intent.ACTIVITY_INDEX,), as_dict=True)
                    if kind.startswith("index_upper") or kind.startswith("compatible"):
                        self.assertTrue(frappe.db.has_index(table, intent.ACTIVITY_INDEX), "Actual pinned has_index must recognize the uppercase alias")
                        self.assertTrue(found)
                    if kind != "index_upper_missing_generated":
                        self.assertEqual(len(native("SELECT COLUMN_NAME FROM information_schema.COLUMNS WHERE TABLE_SCHEMA=DATABASE() "
                            "AND TABLE_NAME=%s AND COLUMN_NAME=%s", (table, intent.ACTIVITY_COLUMN))), 1)
                        native("SELECT COUNT(`" + intent.ACTIVITY_COLUMN + "`) FROM `" + table + "`")  # actual identifier resolution, not Python comparison
                    first_ddl = []
                    def synthetic_metadata(query, *args, **kwargs):
                        if str(query).lstrip().startswith("ALTER TABLE"):
                            first_ddl.append(str(query))
                            raise AssertionError("C1 first business DDL intercepted; real IR remains unchanged")
                        if "information_schema.COLUMNS" in str(query) or str(query).startswith("SHOW INDEX"):
                            query = str(query).replace("tabIntegration Request", table)
                        return native(query, *args, **kwargs)
                    with patch.object(frappe.db, "sql", side_effect=synthetic_metadata), \
                            patch.object(frappe.db, "add_index", wraps=frappe.db.add_index) as add, \
                            patch.object(frappe.db, "sql_ddl", side_effect=AssertionError("No business sql_ddl permitted")) as ddl, \
                            patch.object(frappe.db, "commit", wraps=frappe.db.commit) as commit:
                        error = None
                        try:
                            installer.install_native_intent_index()
                        except Exception as caught:
                            error = caught
                        print("C1_QUALITY_NATIVE_IDENTIFIER=" + json.dumps({"case": kind, "fixture_table": table,
                            "actual_columns": [dict(row) for row in actual], "actual_filtered_index": [dict(row) for row in found],
                            "intercepted_first_business_ddl": first_ddl, "native_add_index_calls": add.call_count,
                            "error_type": type(error).__name__ if error else None}, default=str, sort_keys=True), flush=True)
                        if kind.startswith("compatible"):
                            self.assertIsNone(error, "Compatible native case aliases must be idempotent")
                        else:
                            self.assertIsInstance(error, frappe.ValidationError)
                            self.assertRegex(str(error), "冲突|未经核查")
                        self.assertEqual(first_ddl, [], "All identifier-equivalent definitions must be checked before first DDL")
                        ddl.assert_not_called(); add.assert_not_called(); commit.assert_not_called()
                finally:
                    # Exact newly created metadata fixture only, never the IR schema/history.
                    self.assertTrue(table.startswith("tabQA_C1_IDENT_"))
                    native("DROP TABLE `" + table + "`")
                    self.assertFalse(native("SELECT TABLE_NAME FROM information_schema.TABLES WHERE TABLE_SCHEMA=DATABASE() "
                        "AND TABLE_NAME=%s", (table,)))
                    print("C1_IDENTIFIER_TABLE_REMOVED=" + table, flush=True)
        self.assertEqual(native("SELECT COLUMN_NAME,COLUMN_TYPE,CHARACTER_SET_NAME,COLLATION_NAME,EXTRA,GENERATION_EXPRESSION "
            "FROM information_schema.COLUMNS WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='tabIntegration Request'", as_dict=True), columns)
        self.assertEqual(native("SHOW INDEX FROM `tabIntegration Request`", as_dict=True), indexes)

    def test_native_intent_collation_namespace_rejects_accent_identity_and_service_escape(self):
        from deeplinkerp_branding.services import purchase_native_intent as intent
        service = "D\u00e9epLinkERP native MES Bin intent"
        self.assertEqual(frappe.db.sql("SELECT CAST(%s AS CHAR CHARACTER SET utf8mb4) "
            "COLLATE utf8mb4_unicode_ci = %s", (service, "DeepLinkERP native MES Bin intent"))[0][0], 1)
        forged = frappe.get_doc({"doctype": "Integration Request", "integration_request_service": service,
            "status": "Completed", "request_description": "Native MES Bin sync acknowledged", "data": "{}"})
        try:
            with self.assertRaises(frappe.PermissionError):
                forged.insert(ignore_permissions=True)
        finally:
            frappe.db.rollback()
        name = "QA-ATOMIC-IR-" + uuid.uuid4().hex
        frappe.db.sql("INSERT INTO `tabIntegration Request` (name,integration_request_service,status,data) "
            "VALUES (%s,%s,'Queued','{}')", (name, service))
        doc = frappe.get_doc("Integration Request", name, for_update=True)
        with self.assertRaises(frappe.PermissionError):
            doc.db_set({"integration_request_service": "ordinary", "status": "Completed"})
        frappe.db.rollback()
        for action in ("insert", "save"):
            with self.subTest(native_bytes=action):
                service_bytes = intent.SERVICE.encode("utf-8")
                self.assertEqual(frappe.db.sql("SELECT CAST(%s AS CHAR CHARACTER SET utf8mb4) "
                    "COLLATE utf8mb4_unicode_ci = %s", (service_bytes, intent.SERVICE))[0][0], 1)
                forged = frappe.get_doc({"doctype": "Integration Request", "integration_request_service":
                    service_bytes if action == "insert" else "QA ordinary", "status": "Queued", "data": "{}"})
                if action == "save":
                    forged.insert(ignore_permissions=True)
                forged.integration_request_service = service_bytes
                forged.status, forged.request_description = "Completed", intent.ACKNOWLEDGED
                try:
                    with self.assertRaises(frappe.PermissionError):
                        forged.insert(ignore_permissions=True) if action == "insert" else forged.save(ignore_permissions=True)
                finally:
                    if forged.name and frappe.db.exists("Integration Request", forged.name):
                        persisted = frappe.db.get_value("Integration Request", forged.name,
                            ["name", "integration_request_service", "status", "request_description", intent.ACTIVITY_COLUMN],
                            as_dict=True, for_update=True)
                        print("C1_BYTES_NAMESPACE_PERSISTED=" + json.dumps(dict(persisted), default=str, sort_keys=True), flush=True)
                    frappe.db.rollback()  # exact synthetic rows remain uncommitted, never a historical repair
        ordinary = frappe.get_doc({"doctype": "Integration Request", "integration_request_service": 0,
            "status": "Queued", "data": "{}"}).insert(ignore_permissions=True)
        ordinary.handle_success({"ordinary": True})
        self.remember_new_names()
        self.assertEqual(frappe.db.get_value("Integration Request", ordinary.name,
            ["integration_request_service", "status"], for_update=True), ("0", "Completed"))
        from MySQLdb import ProgrammingError
        unsupported_name = (intent.PREFIX + uuid.uuid4().hex).encode("utf-8")
        unsupported = frappe.get_doc({"doctype": "Integration Request", "name": unsupported_name,
            "integration_request_service": "QA ordinary", "status": "Queued", "data": "{}"})
        before = frappe.db.count("Integration Request")
        with patch.object(frappe.db, "sql", wraps=frappe.db.sql) as sql:
            with self.assertRaises(ProgrammingError):
                unsupported.db_insert()  # existing QB current-name filter does not support bytes; no bypass/ACK
            self.assertFalse(any(str(call.args[0]).lstrip().upper().startswith("INSERT") for call in sql.call_args_list))
        self.assertEqual(frappe.db.count("Integration Request"), before)
        self.assertEqual(frappe.db.sql("SELECT name FROM `tabIntegration Request` WHERE name=%s", (unsupported_name,)), ())

    def test_native_ir_captured_retention_alias_preserves_protected_and_clears_ordinary(self):
        from frappe.integrations.doctype.integration_request.integration_request import IntegrationRequest
        captured = IntegrationRequest.clear_old_logs
        protected, ordinary = "DLP-MES-BIN-" + uuid.uuid4().hex, "QA-ATOMIC-IR-" + uuid.uuid4().hex
        for name, service in ((protected, None), (ordinary, "QA ordinary")):
            frappe.db.sql("INSERT INTO `tabIntegration Request` (name,creation,integration_request_service,status,data) "
                "VALUES (%s,NOW() - INTERVAL 60 DAY,%s,'Queued','{}')", (name, service))
        captured(days=30)
        self.assertTrue(frappe.db.exists("Integration Request", protected), "Captured native retention bypass must preserve identity")
        self.assertFalse(frappe.db.exists("Integration Request", ordinary), "Ordinary native cutoff behavior must remain")

    def test_native_ir_generic_compaction_alias_refuses_before_first_ddl_even_empty(self):
        from frappe.core.doctype.log_settings.log_settings import clear_log_table
        captured = clear_log_table
        for phase in ("empty", "native-null-creation"):
            with self.subTest(phase=phase):
                if phase != "empty":
                    frappe.db.sql("INSERT INTO `tabIntegration Request` (name,integration_request_service,status,data) "
                        "VALUES (%s,%s,'Queued','{}')", ("DLP-MES-BIN-" + uuid.uuid4().hex,
                            "DeepLinkERP native MES Bin intent"))
                with patch.object(frappe.db, "sql_ddl", side_effect=AssertionError("QA prohibited any real compaction DDL")) as ddl, \
                        patch.object(frappe.db, "commit", wraps=frappe.db.commit) as commit:
                    with self.assertRaises(frappe.PermissionError):
                        captured("Integration Request", days=30)
                    ddl.assert_not_called()
                    commit.assert_not_called()

    def test_mes_loaded_producer_mutation_is_not_accepted_by_disk_hash_marker(self):
        import mes_integration.mes_integration.material_request as mes
        from deeplinkerp_branding.services import purchase_native_intent as intent
        native = mes.enqueue_mes_material_request_bin_sync
        original = native.__code__
        def changed(material_request, mr_item_rows=None):
            return None
        try:
            native.__code__ = changed.__code__
            with self.assertRaisesRegex(frappe.ValidationError, "调用身份"):
                intent.install()
        finally:
            native.__code__ = original

    def _assert_fixed_request_cache_getter_shapes(self):
        from pathlib import Path
        from types import CellType, CodeType, FunctionType
        import hashlib
        import frappe.utils.caching as caching
        import erpnext.controllers.buying_controller as buying
        import erpnext.accounts.general_ledger as gl
        from deeplinkerp_branding.services import purchase_native_intent as intent
        caching_source = Path(caching.__file__).read_bytes()
        self.assertEqual(hashlib.sha256(caching_source).hexdigest(),
            "95260b7ea1bda8133e1ac127a4bd90628f688450af4a7020e469bb57600329a5")
        compiled = compile(caching_source, caching.__file__, "exec", dont_inherit=True)
        parent = next(code for code in compiled.co_consts if isinstance(code, CodeType) and code.co_name == "request_cache")
        wrapper_code = next(code for code in parent.co_consts if isinstance(code, CodeType) and code.co_name == "wrapper")
        caching_namespace, decorator = caching.__dict__, caching.request_cache
        aliases = [(module, name, getattr(module, name), getattr(module, name).__wrapped__)
            for module, name in ((buying, "get_purchase_expense_account"), (gl, "get_cost_center_allocation_data"))]
        sources = {module.__file__: hashlib.sha256(Path(module.__file__).read_bytes()).hexdigest()
            for module in (caching, buying, gl)}
        self.assertEqual(sources, {caching.__file__: "95260b7ea1bda8133e1ac127a4bd90628f688450af4a7020e469bb57600329a5",
            buying.__file__: "09d4da3ce8c43f0886e9bbece5207be71cf79febcdfc42c5d1080172fe0081cd",
            gl.__file__: "c2228231802b3d979043bc91f0819b7d9e884c72c32051659456cbbe32f91b84"})
        print("C2A2_CACHE_BEFORE=" + json.dumps({"sources": sources, "aliases":
            {name: [id(outer), id(inner)] for _, name, outer, inner in aliases}}, sort_keys=True), flush=True)
        try:
            for module, name, outer, inner in aliases:
                expected_code = intent._code_named(compile(Path(module.__file__).read_bytes(), module.__file__,
                    "exec", dont_inherit=True), name)
                expected_namespace = module.__dict__
                def audited(candidate, expected_getter=inner):
                    try:
                        return intent._audited_function(candidate, expected_code, expected_namespace,
                            request_cache_getter=expected_getter)
                    except TypeError as error:
                        self.fail("Fixed request_cache audit is unavailable: " + str(error))
                with self.subTest(cached_getter=name, change="genuine"):
                    self.assertTrue(audited(outer), "The genuine fixed wrapper must audit without executing its getter")
                    self.assertFalse(intent._audited_function(outer, expected_code, expected_namespace))
                clone = FunctionType(expected_code, expected_namespace)
                for change in ("wrapper-code", "wrapper-globals", "wrapper-defaults", "wrapper-kwdefaults",
                        "captured-clone", "wrapped-clone", "missing-wrapped", "extra-cell", "wrong-freevar",
                        "empty-cell", "not-function"):
                    with self.subTest(cached_getter=name, change=change):
                        code = wrapper_code.replace(co_name="unknown") if change == "wrapper-code" else wrapper_code
                        cells = (CellType(clone if change == "captured-clone" else inner),)
                        if change == "extra-cell":
                            code, cells = code.replace(co_freevars=("func", "extra")), cells + (CellType(None),)
                        elif change == "wrong-freevar": code = code.replace(co_freevars=("unknown",))
                        elif change == "empty-cell": cells = (CellType(),)
                        candidate = FunctionType(code, dict(caching_namespace) if change == "wrapper-globals"
                            else caching_namespace, closure=cells)
                        candidate.__wrapped__ = clone if change == "wrapped-clone" else inner
                        if change == "wrapper-defaults": candidate.__defaults__ = (None,)
                        elif change == "wrapper-kwdefaults": candidate.__kwdefaults__ = {"unknown": None}
                        elif change == "missing-wrapped": del candidate.__wrapped__
                        elif change == "not-function": candidate = object()
                        self.assertFalse(audited(candidate))
                for change in ("getter-code", "getter-globals", "getter-defaults", "getter-kwdefaults", "forged-pair"):
                    with self.subTest(cached_getter=name, change=change):
                        # Detached getter identity is frozen before mutation. Native aliases stay untouched.
                        captured = FunctionType(expected_code, dict(expected_namespace) if change == "getter-globals"
                            else expected_namespace)
                        candidate = decorator(captured)
                        if change == "getter-code": captured.__code__ = expected_code.replace(co_name="unknown")
                        elif change == "getter-defaults": captured.__defaults__ = (None,)
                        elif change == "getter-kwdefaults": captured.__kwdefaults__ = {"unknown": None}
                        self.assertFalse(audited(candidate, inner if change == "forged-pair" else captured))
                with self.subTest(cached_getter=name, change="arbitrary-func-cell"):
                    def arbitrary(func):
                        def wrapper(*args, **kwargs):
                            raise AssertionError("Auditing must never execute a wrapper", func)
                        wrapper.__wrapped__ = func
                        return wrapper
                    self.assertFalse(audited(arbitrary(inner)))
                for other in (caching.site_cache, caching.redis_cache, caching.http_cache()):
                    with self.subTest(cached_getter=name, decorator=other.__name__):
                        captured = FunctionType(expected_code, expected_namespace)
                        self.assertFalse(audited(other(captured), captured))
                with self.subTest(cached_getter=name, change="caching-source"):
                    with patch.object(Path, "read_bytes", return_value=caching_source + b"\n"):
                        self.assertFalse(audited(outer))
        finally:
            self.assertIs(frappe.request_cache, decorator)
            self.assertIs(caching.request_cache, decorator)
            self.assertIs(caching.__dict__, caching_namespace)
            for module, name, outer, inner in aliases:
                self.assertIs(getattr(module, name), outer)
                self.assertIs(outer.__wrapped__, inner)
                self.assertIs(outer.__closure__[0].cell_contents, inner)
            after = {path: hashlib.sha256(Path(path).read_bytes()).hexdigest() for path in sources}
            self.assertEqual(after, sources)
            print("C2A2_CACHE_AFTER=" + json.dumps({"sources": after, "aliases":
                {name: [id(outer), id(inner)] for _, name, outer, inner in aliases}}, sort_keys=True), flush=True)

    def test_mes_all_loaded_dependency_code_and_unknown_callable_shape_close_only_protected_capability(self):
        self._assert_fixed_request_cache_getter_shapes()
        import mes_integration.mes_integration.material_request as mes
        from deeplinkerp_branding.services import purchase_native_intent as intent, purchase_repost_boundary as boundary
        frappe.db.commit()
        self.remember_new_names()
        for name in ("enqueue_mes_material_request_bin_sync_job", "get_mes_material_request_item_warehouse_pairs",
                "get_stock_item_warehouse_pairs", "normalize_mes_item_warehouse_pairs", "get_mes_indented_qty_map"):
            with self.subTest(dependency=name):
                native = getattr(mes, name)
                saved = native.__code__
                def unknown(items=None):
                    return []
                try:
                    native.__code__ = unknown.__code__
                    with self.assertRaisesRegex(frappe.ValidationError, "调用身份"):
                        intent.install()
                finally:
                    native.__code__ = saved
        saved = mes.sync_material_request_bins
        try:
            mes.sync_material_request_bins = object()
            boundary.before_execution()  # unrelated ordinary request is compatible
            with self.assertRaisesRegex(frappe.ValidationError, "调用身份"):
                intent.install()
        finally:
            mes.sync_material_request_bins = saved
        import rq.queue as rq_queue
        import frappe.integrations.doctype.integration_request.integration_request as ir
        import frappe.core.doctype.log_settings.log_settings as logs
        callables = [(mes, mes, name, "_dlp_native_intent_original_" + suffix, suffix == "lock", False)
            for name, suffix in (("enqueue_mes_material_request_bin_sync", "producer"), ("lock_mes_material_request_bins", "lock"),
                ("sync_material_request_bins", "sync"), ("release_mes_material_request_locks", "release"))]
        callables += [(rq_queue.Queue, rq_queue, "enqueue_call", "_dlp_native_intent_original_enqueue", False, False),
            (ir.IntegrationRequest, ir, "clear_old_logs", "_dlp_purchase_retention_original", False, True),
            (logs, logs, "clear_log_table", "_dlp_purchase_retention_original", False, False)]
        for holder, module, name, attribute, context, static in callables:
            outer = getattr(holder, name)
            native = outer.__wrapped__ if context else outer
            for change in ("defaults", "kwdefaults", "same-code-globals", "same-code-object"):
                with self.subTest(adapted=name, change=change):
                    from types import FunctionType
                    self.assertTrue(intent.install())  # establish genuine supported state before the unknown mutation
                    defaults, kwdefaults = native.__defaults__, native.__kwdefaults__
                    try:
                        if change == "defaults":
                            native.__defaults__ = ("QA unknown default",)
                        elif change == "kwdefaults":
                            native.__kwdefaults__ = {"unknown": "QA"}
                        else:
                            replacement = FunctionType(native.__code__, dict(native.__globals__) if change == "same-code-globals"
                                else native.__globals__, native.__name__, native.__defaults__)
                            if context:
                                outer.__wrapped__ = replacement
                            else:
                                setattr(holder, name, staticmethod(replacement) if static else replacement)
                        boundary.before_execution()  # unsupported adapters must not break an ordinary request
                        self.assertFalse(intent.install(strict=False))
                        self.assertEqual(frappe.local.purchase_native_intent_capability, "unsupported")
                        self.assertFalse(mes._dlp_native_intent_capability)
                        with self.assertRaisesRegex(frappe.ValidationError, "调用身份"):
                            intent.install()
                    finally:
                        native.__defaults__, native.__kwdefaults__ = defaults, kwdefaults
                        if context:
                            outer.__wrapped__ = native
                        else:
                            setattr(holder, name, staticmethod(outer) if static else outer)
            original = getattr(module, attribute)
            for change in ("globals", "object"):
                with self.subTest(saved=name, change=change):
                    from contextlib import contextmanager
                    from types import FunctionType
                    self.assertTrue(intent.install())
                    function = original.__wrapped__ if context else original
                    replacement = FunctionType(function.__code__, dict(function.__globals__) if change == "globals"
                        else function.__globals__, function.__name__, function.__defaults__)
                    with patch.object(module, attribute, contextmanager(replacement) if context else replacement):
                        with self.assertRaisesRegex(frappe.ValidationError, "调用身份"):
                            intent.install()
        state = boundary.initialize()
        native_sql, acquired = state.native_sql, []
        def observe_acquisition(query, values=(), **kwargs):
            if str(query).strip() == "SELECT GET_LOCK(%s, %s)":
                acquired.append(values)
            return native_sql(query, values, **kwargs)
        for name, value in (("MES_BIN_LOCK_TIMEOUT_SECONDS", 181), ("MES_INWARD_MATERIAL_REQUEST_TYPES", ("Purchase",)),
                ("flt", lambda value: 0)):
            with self.subTest(linked_global=name), patch.object(mes, name, value), \
                    patch.object(state, "native_sql", side_effect=observe_acquisition):
                with self.assertRaisesRegex(frappe.ValidationError, "调用身份"):
                    with mes.lock_mes_material_request_bins({"company": COMPANY,
                            "items": [{"item_code": self.item, "warehouse": "Stores - QAB"}]}):
                        pass
                self.assertEqual(acquired, [], "Unknown linked globals must close capability before native acquisition")

    def test_native_mes_same_source_new_generation_and_scoped_partial_producer(self):
        from types import SimpleNamespace
        from frappe.utils import CallbackManager
        import mes_integration.mes_integration.material_request as mes
        from deeplinkerp_branding.services import purchase_native_intent as intent
        warehouse = frappe.get_doc({"doctype": "Warehouse", "warehouse_name": "QA-ATOMIC-" + uuid.uuid4().hex,
            "company": COMPANY, "parent_warehouse": "All Warehouses - QAB"}).insert().name
        mr, first = self.native_mes_intent((warehouse, "Stores - QAB"))
        request = SimpleNamespace(after_response=CallbackManager())
        with patch.object(frappe.local, "request", request, create=True):
            mr = frappe.get_doc("Material Request", mr.name)
            mr.flags.mes_integration_request = True
            mr.update_requested_qty([mr.items[0].name])
            request.after_response.reset()
            frappe.db.commit()
            self.remember_new_names()
        records = intent.active_set()
        self.assertEqual(len(records), 2)
        later = next(record for record in records if record.name != first.name)
        self.assertNotEqual(later.generation, first.request_id)
        self.assertEqual(later.facts["new_pairs"], [[self.item, warehouse]])
        self.assertEqual(len(later.facts["source_pairs"]), 2, "A scoped native call still freezes the full source identity")
        from redis.exceptions import ConnectionError
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            intent.dispatch(later.name)
        self.remember_new_names()
        self.assertEqual(frappe.db.get_value("Integration Request", later.name, "status", for_update=True), "Completed")
        self.assertEqual(frappe.db.get_value("Integration Request", first.name, "status", for_update=True), "Queued",
            "A callback may only acknowledge its exact generation, not another pending generation")

    def test_native_mes_explicit_same_actor_recovery_group_covers_warehouse_move_and_qty_generation(self):
        from types import SimpleNamespace
        from frappe.utils import CallbackManager
        from redis.exceptions import ConnectionError
        from deeplinkerp_branding.services import purchase_native_intent as intent
        warehouse = frappe.get_doc({"doctype": "Warehouse", "warehouse_name": "QA-ATOMIC-" + uuid.uuid4().hex,
            "company": COMPANY, "parent_warehouse": "All Warehouses - QAB"}).insert().name
        mr, first = self.native_mes_intent()
        child = frappe.get_doc("Material Request Item", mr.items[0].name)
        child.db_set({"warehouse": warehouse, "qty": 7, "stock_qty": 7})
        request = SimpleNamespace(after_response=CallbackManager())
        with patch.object(frappe.local, "request", request, create=True):
            current = frappe.get_doc("Material Request", mr.name)
            current.flags.mes_integration_request = True
            current.update_requested_qty()
            request.after_response.reset()
            frappe.db.commit()
            self.remember_new_names()
        records = intent.active_set()
        self.assertEqual(len(records), 2)
        later = next(record for record in records if record.name != first.name)
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            intent.dispatch(first.name)
        self.remember_new_names()
        self.assertEqual(frappe.db.get_value("Integration Request", first.name, "status", for_update=True), "Queued")
        frappe.db.commit()  # native failure diagnostic only; the first generation remains unfinished
        self.remember_new_names()
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            self.assertEqual(set(intent.recover()), {first.name, later.name})
        self.remember_new_names()
        for record in records:
            actual = frappe.db.get_value("Integration Request", record.name, ["status", "output"], as_dict=True, for_update=True)
            self.assertEqual(actual.status, "Completed")
            receipt = json.loads(actual.output)
            self.assertEqual({(row["item_code"], row["warehouse"]) for row in receipt["bins"]}, set(record.pairs))
            self.assertEqual(receipt["claim_names"], sorted([first.name, later.name]))
            intent.dispatch(record.name)  # each terminal receipt replays its own fixed subset
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"},
            "indented_qty", for_update=True), 0)
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": warehouse},
            "indented_qty", for_update=True), 7)

    def test_unknown_mes_capability_preserves_ordinary_request_but_refuses_flagged_producer(self):
        from deeplinkerp_branding.services import purchase_native_intent as intent, purchase_repost_boundary as boundary
        mr = frappe.get_doc({"doctype": "Material Request", "company": COMPANY,
            "material_request_type": "Purchase", "transaction_date": nowdate(),
            "schedule_date": add_days(nowdate(), 1), "items": [{"item_code": self.item,
                "qty": 3, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]}).insert()
        frappe.db.commit()
        self.remember_new_names()
        native_read = intent.Path.read_bytes
        with patch.object(intent.Path, "read_bytes", lambda path: b"unknown native MES source"
                if path.name == "material_request.py" else native_read(path)):
            ordinary_error = None
            try:
                boundary.before_execution()
                frappe.get_doc("Item", self.item)
            except frappe.ValidationError as error:
                ordinary_error = error
            self.assertIsNone(ordinary_error, "Unknown MES capability must not disable unrelated ordinary requests")
            mr.flags.mes_integration_request = True
            with self.assertRaisesRegex(frappe.ValidationError, "版本未经核查"):
                mr.submit()

    def test_optional_mes_absence_is_explicit_and_preserves_unrelated_request(self):
        from deeplinkerp_branding.services import purchase_native_intent as intent, purchase_repost_boundary as boundary
        original = frappe.get_hooks
        def optional_absent(key=None, *args, **kwargs):
            value = original(key, *args, **kwargs)
            if key == "extend_doctype_class":
                value = dict(value)
                value["Material Request"] = [hook for hook in value.get("Material Request", [])
                    if not hook.startswith("mes_integration.")]
            return value
        with patch.object(frappe, "get_hooks", side_effect=optional_absent):
            boundary.before_execution()
            self.assertFalse(intent.install(strict=False))  # existing session need not reinstall every ordinary request
            self.assertEqual(frappe.local.purchase_native_intent_capability, "optional_absent")
            self.assertEqual(frappe.get_doc("Item", self.item).name, self.item)
            with self.assertRaisesRegex(frappe.ValidationError, "未安装"):
                intent.install()

    def test_mes_late_server_flag_cannot_take_publication_fence_after_pair_lease(self):
        from types import SimpleNamespace
        from frappe.utils import CallbackManager
        import mes_integration.mes_integration.material_request as mes
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        captured = mes.enqueue_mes_material_request_bin_sync
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        for action in ("submit", "update_requested_qty", "captured_producer"):
            with self.subTest(native_entry=action):
                request = SimpleNamespace(after_response=CallbackManager())
                with patch.object(frappe.local, "request", request, create=True), boundary.execution() as state:
                    mr = frappe.get_doc({"doctype": "Material Request", "company": COMPANY,
                        "material_request_type": "Purchase", "transaction_date": nowdate(),
                        "schedule_date": add_days(nowdate(), 1), "items": [{"item_code": self.item,
                            "qty": 3, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]}).insert()
                    mr.flags.mes_integration_request = True
                    self.assertTrue(state.locks)
                    self.assertNotIn(boundary.fence_key(), state.locks)
                    held = dict(state.locks)
                    before = frappe.db.count("Integration Request")
                    callbacks = tuple(request.after_response._functions)
                    call = mr.submit if action == "submit" else (mr.update_requested_qty if action == "update_requested_qty"
                        else lambda: captured(mr, [mr.items[0].name]))
                    with patch.object(frappe.db, "commit", wraps=frappe.db.commit) as commit, \
                            patch.object(frappe.db, "rollback", wraps=frappe.db.rollback) as rollback, \
                            patch.object(boundary, "acquire", wraps=boundary.acquire) as acquire:
                        error = None
                        try:
                            call()
                        except Exception as caught:
                            error = caught
                        print("C1_SPEC_PUBLICATION_ORDER=" + json.dumps({"entry": action, "material_request": mr.name,
                            "prior_lease_count": len(held), "after_lease_count": len(state.locks),
                            "fence_after": boundary.fence_key() in state.locks,
                            "intents_before": before, "intents_after": frappe.db.count("Integration Request"),
                            "callbacks_added": len(request.after_response._functions) - len(callbacks),
                            "error_type": type(error).__name__ if error else None}, sort_keys=True), flush=True)
                        self.assertIsInstance(error, frappe.ValidationError)
                        self.assertIn("发布隔离顺序", str(error))
                        commit.assert_not_called(); rollback.assert_not_called()
                        self.assertEqual(tuple(request.after_response._functions), callbacks)
                        self.assertFalse(any(boundary.fence_key() in tuple(entry.args[0]) for entry in acquire.call_args_list))
                        self.assertEqual(state.locks, held)
                        self.assertEqual(frappe.db.count("Integration Request"), before)
                frappe.db.rollback()  # a distinct request fixture, never a business-view refresh/reorder

    def native_mes_intent(self, warehouses=("Stores - QAB",)):
        from types import SimpleNamespace
        from frappe.utils import CallbackManager
        request = SimpleNamespace(after_response=CallbackManager())
        with patch.object(frappe.local, "request", request, create=True):
            mr = frappe.get_doc({"doctype": "Material Request", "company": COMPANY,
                "material_request_type": "Purchase", "transaction_date": nowdate(),
                "schedule_date": add_days(nowdate(), 1), "items": [{"item_code": self.item,
                    "qty": index + 3, "warehouse": warehouse, "schedule_date": add_days(nowdate(), 1)}
                    for index, warehouse in enumerate(warehouses)]})
            mr.flags.mes_integration_request = True
            mr.insert().submit()
            request.after_response.reset()
            frappe.db.commit()
            self.remember_new_names()
        row = frappe.db.get_values("Integration Request", {"integration_request_service":
            "DeepLinkERP native MES Bin intent", "reference_docname": mr.name}, "*", as_dict=True, for_update=True)[0]
        frappe.db.rollback()  # only fixture evidence read locks, before the tested execution
        return mr, row

    def synthetic_intents(self, row, count, *, pairs=None, padding="", status="Queued", phase=None):
        from deeplinkerp_branding.services import purchase_native_intent as intent
        for index in range(count):
            name, generation = intent.PREFIX + uuid.uuid4().hex, uuid.uuid4().hex
            facts = json.loads(row.data)
            facts["generation"] = generation
            if pairs is not None:
                pair = pairs(index)
                facts.update(old_pairs=[], new_pairs=[pair], pairs=[pair], warehouse_companies={pair[1]: COMPANY})
            if padding:
                facts["padding"] = padding
            frappe.db.sql("INSERT INTO `tabIntegration Request` "
                "(name,integration_request_service,request_id,status,request_description,reference_doctype,reference_docname,data) "
                "VALUES (%s,%s,%s,%s,%s,'Material Request',%s,%s)",
                (name, intent.SERVICE, generation, status, intent.PENDING if phase is None else phase,
                    row.reference_docname, json.dumps(facts, ensure_ascii=False)))

    def native_mes_actor(self):
        user = "qa-atomic-" + uuid.uuid4().hex[:12] + "@example.invalid"
        with patch.object(frappe, "enqueue", return_value=None):  # User contact side job is outside this MR-only queue fixture
            frappe.get_doc({"doctype": "User", "email": user, "first_name": "Atomic MES QA", "send_welcome_email": 0,
                "roles": [{"role": role} for role in ("Purchase Manager", "Stock Manager", "Purchase User", "Stock User")]}).insert()
        try:
            frappe.db.commit()
        finally:
            self.remember_new_names()  # physical commit may precede an after_commit failure
        self.addCleanup(lambda: frappe.clear_cache(user=user))
        return user

    def test_native_mes_real_rq_serializer_fixed_metadata_preserves_actor_and_native_dedup(self):
        from types import SimpleNamespace
        from frappe.utils import CallbackManager
        from frappe.utils.background_jobs import create_job_id, get_queue, get_job
        from rq.job import JobStatus
        from rq.serializers import DefaultSerializer
        from redis.exceptions import ConnectionError
        import mes_integration.mes_integration.material_request as mes
        from deeplinkerp_branding.services import purchase_native_intent as intent
        actor = self.native_mes_actor()
        frappe.set_user(actor)
        try:
            mr, row = self.native_mes_intent()
            self.assertTrue(frappe.has_permission("Material Request", "submit", doc=mr))
            queue, identifier = get_queue("short"), create_job_id("mes-material-request-bin-sync:" + mr.name)
            self.assertIsNone(get_job("mes-material-request-bin-sync:" + mr.name), "Exact synthetic native job ID must be absent before creation")
            self.assertNotIn(identifier, queue.job_ids)
            try:
                alias = mes.sync_material_request_bins
                self.assertIs(DefaultSerializer.loads(DefaultSerializer.dumps(alias)), alias,
                    "The actual captured adapted callable must retain native serializer identity")
                intent.dispatch(row.name)
                job = get_job("mes-material-request-bin-sync:" + mr.name)
                self.assertIsNotNone(job)
                print("C1_NATIVE_RQ", json.dumps({"id": job.id, "origin": job.origin,
                    "serializer": job.serializer.__module__ + "." + job.serializer.__name__,
                    "func": job.func_name, "user": job.kwargs.get("user"), "meta": job.meta}, sort_keys=True), flush=True)
                self.assertEqual(job.id, identifier)
                self.assertEqual(job.origin, queue.name)
                self.assertEqual(job.func_name, "frappe.utils.background_jobs.execute_job")
                self.assertEqual(job.kwargs["user"], actor)
                self.assertEqual(job.kwargs["site"], SITE)
                self.assertEqual(job.kwargs["method"], intent.METHOD)
                self.assertEqual(set(job.kwargs["kwargs"]), {"material_request_name", "item_warehouse_pairs"})
                fixed = job.meta["deeplinkerp_native_bin_intent"]
                self.assertEqual(fixed["user"], actor)
                self.assertEqual(fixed["records"][0][0:2], [row.name, row.request_id])
                self.assertEqual(fixed["pairs"], [[self.item, "Stores - QAB"]])
                before_kwargs, before_meta = DefaultSerializer.dumps(job.kwargs), DefaultSerializer.dumps(job.meta)
                created = [row.name]
                for status in (JobStatus.QUEUED, JobStatus.STARTED):
                    with self.subTest(native_dedup_status=status):
                        job.set_status(status)  # native Job status fixture, not an actual worker execution
                        child = frappe.get_doc("Material Request Item", mr.items[0].name)
                        child.db_set({"qty": len(created) + 6, "stock_qty": len(created) + 6})
                        request = SimpleNamespace(after_response=CallbackManager())
                        with patch.object(frappe.local, "request", request, create=True):
                            current = frappe.get_doc("Material Request", mr.name)
                            current.flags.mes_integration_request = True
                            current.update_requested_qty()
                            request.after_response.reset()
                            frappe.db.commit()
                            self.remember_new_names()
                        latest = next(record for record in intent.active_set() if record.name not in created)
                        created.append(latest.name)
                        intent.dispatch(latest.name)
                        unchanged = get_job("mes-material-request-bin-sync:" + mr.name)
                        self.assertEqual(DefaultSerializer.dumps(unchanged.kwargs), before_kwargs)
                        self.assertEqual(DefaultSerializer.dumps(unchanged.meta), before_meta)
                        self.assertEqual(frappe.db.get_value("Integration Request", latest.name, "status", for_update=True), "Queued")
                queue.remove(identifier)
                job.delete()  # only the exact job created above; no queue/registry sweep
                self.assertIsNone(get_job("mes-material-request-bin-sync:" + mr.name))
                frappe.db.commit()  # finish current-read evidence before an independent SERVICE recovery
                with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
                    self.assertEqual(set(intent.recover()), set(created))
                self.remember_new_names()
                for name in created:
                    output = frappe.db.get_value("Integration Request", name, ["status", "output"], as_dict=True, for_update=True)
                    self.assertEqual(output.status, "Completed")
                    self.assertEqual(json.loads(output.output)["user"], actor)
            finally:
                remaining = get_job("mes-material-request-bin-sync:" + mr.name)
                if remaining is not None:
                    self.assertEqual(remaining.id, identifier)
                    queue.remove(identifier)
                    remaining.delete()
                self.assertNotIn(identifier, queue.job_ids)
        finally:
            frappe.set_user("Administrator")

    def test_native_mes_mixed_actor_recovery_group_holds_before_executor_write(self):
        from types import SimpleNamespace
        from frappe.utils import CallbackManager
        from deeplinkerp_branding.services import purchase_native_intent as intent
        actor = self.native_mes_actor()
        mr, row = self.native_mes_intent()
        request = SimpleNamespace(after_response=CallbackManager())
        try:
            frappe.set_user(actor)
            with patch.object(frappe.local, "request", request, create=True):
                current = frappe.get_doc("Material Request", mr.name)
                current.flags.mes_integration_request = True
                current.update_requested_qty()
                request.after_response.reset()
                frappe.db.commit()
                self.remember_new_names()
            with patch.object(frappe, "enqueue", side_effect=AssertionError("Mixed actor HOLD must precede helper enqueue")) as enqueue:
                with self.assertRaisesRegex(frappe.ValidationError, "不同生产者执行人"):
                    intent.recover()
                enqueue.assert_not_called()
            self.assertEqual(frappe.db.count("Bin", {"item_code": self.item}), 0)
            self.assertEqual(set(frappe.get_all("Integration Request", filters={"reference_docname": mr.name}, pluck="status")), {"Queued"})
        finally:
            frappe.set_user("Administrator")

    def test_native_intent_registration_must_not_overflow_existing_active_set(self):
        from deeplinkerp_branding.services import purchase_native_intent as intent
        warehouse = frappe.get_doc({"doctype": "Warehouse", "warehouse_name": "QA-ATOMIC-" + uuid.uuid4().hex,
            "company": COMPANY, "parent_warehouse": "All Warehouses - QAB"}).insert().name
        mr, row = self.native_mes_intent()
        with patch.object(intent, "MAX_RECORDS", 1):
            with self.assertRaisesRegex(frappe.ValidationError, "记录.*安全上限"):
                other = frappe.get_doc({"doctype": "Material Request", "company": COMPANY,
                    "material_request_type": "Purchase", "transaction_date": nowdate(),
                    "schedule_date": add_days(nowdate(), 1), "items": [{"item_code": self.item,
                        "qty": 2, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]})
                other.flags.mes_integration_request = True
                other.insert().submit()
        for resource in ("actual-5000-records", "actual-2500-union", "actual-data-output-total"):
            with self.subTest(resource=resource):
                if resource == "actual-5000-records":
                    self.synthetic_intents(row, 4999)
                    message = "记录.*安全上限"
                elif resource == "actual-2500-union":
                    self.synthetic_intents(row, 2499, pairs=lambda index: ["QA-BOUND-" + str(index), "Stores - QAB"])
                    message = "联合范围.*安全上限"
                else:
                    self.synthetic_intents(row, 3, padding="x" * (intent.MAX_BODY_BYTES - len(row.data.encode()) - 100))
                    headers = intent._headers()
                    needed = intent.MAX_TOTAL_BYTES - sum(header.body_bytes for header in headers) - 100
                    # Fill the remaining resource budget across actual outputs, each below 4 MiB.
                    for header in headers:
                        amount = min(needed, intent.MAX_BODY_BYTES)
                        frappe.db.sql("UPDATE `tabIntegration Request` SET output=%s WHERE name=%s", ("x" * amount, header.name))
                        needed -= amount
                    self.assertEqual(needed, 0)
                    message = "新增正文总字节.*安全上限"
                with self.assertRaisesRegex(frappe.ValidationError, message):
                    other = frappe.get_doc({"doctype": "Material Request", "company": COMPANY,
                        "material_request_type": "Purchase", "transaction_date": nowdate(),
                        "schedule_date": add_days(nowdate(), 1), "items": [{"item_code": self.item,
                            "qty": 2, "warehouse": warehouse if resource == "actual-2500-union" else "Stores - QAB",
                            "schedule_date": add_days(nowdate(), 1)}]})
                    other.flags.mes_integration_request = True
                    other.insert().submit()
                self.assertEqual(frappe.db.count("Integration Request", {"integration_request_service": intent.SERVICE}), 1)
                self.assertFalse(frappe.db.exists("Material Request", other.name))

    def test_native_intent_active_selector_rejects_actual_record_pair_and_byte_limits_before_partial_proof(self):
        from deeplinkerp_branding.services import purchase_native_intent as intent
        mr, row = self.native_mes_intent()
        cases = (("records", 5000, None, "", "记录"),
            ("pairs", 2500, lambda index: ["QA-BOUND-" + str(index), "Stores - QAB"], "", "联合范围"),
            ("single-bytes", 1, None, "\u6c49" * ((intent.MAX_BODY_BYTES // 3) + 1), "正文缺失或超过"),
            ("total-bytes", 5, None, "x" * (3400 * 1024), "正文总字节"),
            ("output-single-bytes", 0, None, "\u6c49" * ((intent.MAX_BODY_BYTES // 3) + 1), "输出.*安全上限"),
            ("output-total-bytes", 5, None, "x" * (3400 * 1024), "正文总字节"))
        for label, count, pairs, padding, message in cases:
            with self.subTest(boundary=label):
                frappe.db.savepoint("native_intent_budget_fixture")
                self.synthetic_intents(row, count, pairs=pairs, padding="" if label.startswith("output-") else padding)
                if label.startswith("output-"):
                    frappe.db.sql("UPDATE `tabIntegration Request` SET output=%s WHERE integration_request_service=%s",
                        (padding, intent.SERVICE))
                with patch.object(frappe.db, "sql", wraps=frappe.db.sql) as sql:
                    with self.assertRaisesRegex(frappe.ValidationError, message):
                        intent.active_set()
                    if label != "pairs":
                        self.assertFalse(any(str(call.args[0]).startswith("SELECT name,data")
                            for call in sql.call_args_list), "Header actual byte/count caps must precede body reads")
                # The failed execution rolls back its entire synthetic fixture.
                self.assertEqual(frappe.db.count("Integration Request"), 1)

    def test_native_intent_ack_output_budgets_refuse_before_first_ack_and_rollback_native_bins(self):
        from redis.exceptions import ConnectionError
        from deeplinkerp_branding.services import purchase_native_intent as intent
        for budget in ("single", "total"):
            with self.subTest(budget=budget):
                mr, row = self.native_mes_intent()
                self.synthetic_intents(row, 30)
                frappe.db.commit()
                self.remember_new_names()
                headers = intent._headers()
                total = sum(header.body_bytes for header in headers)
                body_limit = max(header.body_bytes for header in headers) + 100 if budget == "single" else intent.MAX_BODY_BYTES
                frappe.db.commit()  # independent recovery starts after fixture evidence, not a refreshed executor
                with patch.object(intent, "MAX_BODY_BYTES", body_limit), \
                        patch.object(intent, "MAX_TOTAL_BYTES", total + 100 if budget == "total" else intent.MAX_TOTAL_BYTES), \
                        patch.object(frappe.db, "set_value", wraps=frappe.db.set_value) as writes, \
                        patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
                    intent.recover()  # original native fallback can swallow the intentional resource refusal
                    self.remember_new_names()  # immediately after the native commit, before any RED assertion
                    ack_writes = [call for call in writes.call_args_list if call.args and
                        call.args[0] == "Integration Request" and isinstance(call.args[2], dict) and
                        call.args[2].get("status") == "Completed"]
                    self.assertFalse(ack_writes, "Resource caps must reject before the first IR ACK write")
                self.remember_new_names()
                self.assertEqual(frappe.db.count("Bin", {"item_code": self.item}), 0,
                    "Original native Bin writes must physically roll back with the rejected ACK")
                self.assertEqual(set(frappe.get_all("Integration Request", filters={"integration_request_service": intent.SERVICE},
                    pluck="status")), {"Queued"})

    def test_native_intent_active_selector_reads_durable_independent_service(self):
        from deeplinkerp_branding.services import purchase_native_intent as intent
        mr, row = self.native_mes_intent()
        selector = getattr(intent, "active_set", None)
        self.assertTrue(callable(selector), "Committed native intent requires independent SERVICE recovery selector")
        active = selector()
        self.assertEqual([(entry.name, entry.generation, entry.material_request_name, entry.pairs) for entry in active],
            [(row.name, row.request_id, mr.name, ((self.item, "Stores - QAB"),))])

    def test_native_intent_history_growing_actual_explain_stays_on_bounded_active_index(self):
        from deeplinkerp_branding.services import purchase_native_intent as intent
        mr, row = self.native_mes_intent()
        metadata = frappe.db.sql("SELECT COLUMN_NAME,COLUMN_TYPE,CHARACTER_SET_NAME,COLLATION_NAME,EXTRA "
            "FROM information_schema.COLUMNS WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='tabIntegration Request' "
            "AND COLUMN_NAME IN ('name','integration_request_service','status','request_description',%s)",
            (intent.ACTIVITY_COLUMN,), as_dict=True)
        print("C1_INDEX_COLUMN_METADATA=" + json.dumps(metadata, default=str, sort_keys=True), flush=True)
        for size in (100, 20000):
            with self.subTest(history=size):
                self.synthetic_intents(row, size, status="Completed", phase=intent.ACKNOWLEDGED)
                with patch.object(frappe.db, "sql", wraps=frappe.db.sql) as sql:
                    headers = intent._headers()
                    query, values = sql.call_args_list[0].args[:2]
                self.assertEqual([header.name for header in headers], [row.name])
                explain = frappe.db.sql("EXPLAIN " + query, values, as_dict=True)
                actual = json.loads(frappe.db.sql("ANALYZE FORMAT=JSON " + query, values)[0][0])
                table = actual["query_block"]["nested_loop"][0]["table"]
                self.assertEqual(explain[0].key, intent.ACTIVITY_INDEX)
                self.assertNotIn("filesort", explain[0].Extra.lower())
                self.assertEqual(table["key"], intent.ACTIVITY_INDEX)
                self.assertEqual(table["r_rows"], 1, "Inactive history must not become a body/header scan")
                print("C1_ACTIVE_EXPLAIN=" + json.dumps({"history": size, "explain": explain, "actual": actual},
                    default=str, sort_keys=True), flush=True)
        accent = "D\u00e9epLinkERP native MES Bin intent"
        frappe.db.sql("INSERT INTO `tabIntegration Request` (name,integration_request_service,status,request_description,data) "
            "VALUES (%s,%s,'Completed',%s,'{}')", ("DLP-MES-BIN-" + uuid.uuid4().hex, accent, intent.ACKNOWLEDGED))
        with self.assertRaisesRegex(frappe.ValidationError, "身份"):
            intent.active_set()  # collation-equivalent noncanonical service remains visible and rejected

    def test_native_intent_zero_active_selector_holds_publication_fence_until_physical_commit(self):
        from deeplinkerp_branding.services import purchase_native_intent as intent, purchase_repost_boundary as boundary
        frappe.db.commit()
        self.remember_new_names()
        self.assertEqual(intent.active_set(), ())
        peer = self.boundary_database()
        with self.assertRaisesRegex(frappe.ValidationError, "正在办理"):
            with boundary.execution(db=peer), boundary.acquire((boundary.fence_key(),), db=peer):
                pass
        frappe.db.commit()
        with boundary.execution(db=peer), boundary.acquire((boundary.fence_key(),), db=peer):
            pass
        peer.rollback()

    def test_native_mes_producer_source_facts_match_real_persisted_controller(self):
        import hashlib
        from deeplinkerp_branding.services import purchase_operation as audit, purchase_native_intent as intent
        mr, row = self.native_mes_intent()
        actual = frappe.get_doc("Material Request", mr.name, for_update=True)
        facts = json.loads(row.data)
        self.assertEqual(facts["source_modified"], str(actual.modified))
        self.assertEqual(facts["source_persisted_version"], hashlib.sha256(audit.encode(actual.as_dict()).encode()).hexdigest(),
            "Real on_submit producer and reloaded persisted MR must have equivalent source evidence")
        self.assertEqual(facts["source_version"], intent._source_identity(actual))

    def test_native_mes_native_po_progression_does_not_strand_producer_intent(self):
        from redis.exceptions import ConnectionError
        from erpnext.stock.doctype.material_request.material_request import make_purchase_order
        from deeplinkerp_branding.services import purchase_native_intent as intent
        mr, row = self.native_mes_intent()
        before = json.loads(row.data)
        po = make_purchase_order(mr.name)
        po.supplier = "QA Test Supplier"
        po.items[0].qty = 1
        po.items[0].rate = 10
        po.insert(set_name="QA-ATOMIC-PO-" + uuid.uuid4().hex[:16]).submit()
        frappe.db.commit()
        self.remember_new_names()
        current = frappe.get_doc("Material Request", mr.name, for_update=True)
        self.assertEqual(current.items[0].ordered_qty, 1)
        self.assertAlmostEqual(current.per_ordered, 100 / 3, places=5)
        self.assertEqual(frappe.db.count("Integration Request", {"integration_request_service": intent.SERVICE,
            "reference_docname": mr.name}), 1, "Native PO progression must not register a new deferred producer")
        self.assertNotEqual(intent._source_version(current), before.get("source_persisted_version", before["source_version"]))
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            intent.dispatch(row.name)
        self.remember_new_names()
        actual = frappe.db.get_value("Integration Request", row.name, ["status", "output"], as_dict=True, for_update=True)
        self.assertEqual(actual.status, "Completed")
        receipt = json.loads(actual.output)
        self.assertEqual(receipt["generation"], row.request_id)
        self.assertEqual(receipt["bins"][0]["indented_qty"], "2.0")
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"},
            "indented_qty", for_update=True), 2)

    def test_native_mes_native_pr_progression_keeps_requested_identity_and_current_math(self):
        from redis.exceptions import ConnectionError
        from erpnext.stock.doctype.material_request.material_request import make_purchase_order
        from deeplinkerp_branding.services import purchase_native_intent as intent
        mr, row = self.native_mes_intent()
        po = make_purchase_order(mr.name)
        po.supplier = "QA Test Supplier"
        po.items[0].qty, po.items[0].rate = 1, 10
        po.insert(set_name="QA-ATOMIC-PO-" + uuid.uuid4().hex[:16]).submit()
        pr = make_purchase_receipt(po.name).insert()
        pr.submit()  # real native PR.update_prevdoc_status -> MRI.received_qty / MR.per_received
        frappe.db.commit()
        self.remember_new_names()
        current = frappe.get_doc("Material Request", mr.name, for_update=True)
        self.assertEqual((current.items[0].qty, current.items[0].stock_qty, current.items[0].ordered_qty,
            current.items[0].received_qty), (3, 3, 1, 1))
        self.assertAlmostEqual(current.per_received, 100 / 3, places=5)
        self.assertEqual(frappe.db.count("Integration Request", {"integration_request_service": intent.SERVICE}), 1)
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            intent.dispatch(row.name)
        self.remember_new_names()
        actual = frappe.db.get_value("Integration Request", row.name, ["status", "output"], as_dict=True, for_update=True)
        self.assertEqual(actual.status, "Completed")
        witness = json.loads(actual.output)["source"]
        self.assertEqual(witness["version"], intent._source_version(current))
        self.assertNotEqual(witness["version"], json.loads(row.data)["source_persisted_version"])
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"},
            ["actual_qty", "indented_qty"], for_update=True), (1, 2))

    def test_native_mes_unregistered_current_source_qty_warehouse_or_row_removal_cannot_ack_old_generation(self):
        from redis.exceptions import ConnectionError
        from deeplinkerp_branding.services import purchase_native_intent as intent
        warehouse = frappe.get_doc({"doctype": "Warehouse", "warehouse_name": "QA-ATOMIC-" + uuid.uuid4().hex,
            "company": COMPANY, "parent_warehouse": "All Warehouses - QAB"}).insert().name
        for action in ("qty", "stock_qty", "conversion", "warehouse", "rename", "remove", "stopped", "docstatus"):
            with self.subTest(unregistered_change=action):
                mr, row = self.native_mes_intent()
                child = frappe.get_doc("Material Request Item", mr.items[0].name)
                if action == "qty":
                    child.db_set({"qty": 7, "stock_qty": 7})
                elif action == "stock_qty":
                    child.db_set("stock_qty", 7)
                elif action == "conversion":
                    child.db_set("conversion_factor", 2)
                elif action == "warehouse":
                    child.db_set("warehouse", warehouse)
                elif action == "rename":
                    from frappe.model.rename_doc import rename_doc
                    rename_doc(child.doctype, child.name, "QA-ATOMIC-CHILD-" + uuid.uuid4().hex, force=True,
                        ignore_permissions=True)
                elif action in ("stopped", "docstatus"):
                    mr.db_set("status", "Stopped") if action == "stopped" else mr.db_set("docstatus", 2)
                else:
                    child.delete(ignore_permissions=True)
                frappe.db.commit()
                self.remember_new_names()
                with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
                    intent.dispatch(row.name)
                self.remember_new_names()
                actual = frappe.db.get_value("Integration Request", row.name, ["status", "output"], as_dict=True, for_update=True)
                self.assertEqual(actual.status, "Queued",
                    "An old fixed claim must not acknowledge a new unregistered source identity")
                self.assertIsNone(actual.output)
                self.assertEqual(frappe.db.count("Bin", {"item_code": self.item}), 0,
                    "Unregistered source changes must be refused before the original executor's first Bin write")

    def test_native_mes_redis_fallback_requires_same_commit_bin_ack(self):
        from redis.exceptions import ConnectionError
        from deeplinkerp_branding.services import purchase_native_intent as intent
        import mes_integration.mes_integration.material_request as mes
        mr, row = self.native_mes_intent()
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            intent.dispatch(row.name)
        self.remember_new_names()
        actual = frappe.db.get_values("Integration Request", {"name": row.name},
            ["status", "request_description", "output"], as_dict=True, for_update=True)[0]
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"},
            "indented_qty", for_update=True), mes.get_mes_indented_qty_map([(self.item, "Stores - QAB")])[(self.item, "Stores - QAB")])
        self.assertEqual(actual.status, "Completed", "Only actual fixed-claim Bin ACK may complete native intent")
        self.assertEqual(actual.request_description, "Native MES Bin sync acknowledged")
        ack = json.loads(actual.output)
        self.assertEqual(ack["generation"], row.request_id)
        self.assertEqual(ack["user"], "Administrator")

    def test_native_mes_old_rr_failure_retains_intent_for_independent_recovery(self):
        from redis.exceptions import ConnectionError
        from erpnext.stock.utils import get_bin
        from deeplinkerp_branding.services import purchase_native_intent as intent
        from deeplinkerp_branding.services import purchase_operation as audit
        import mes_integration.mes_integration.material_request as mes
        mr, row = self.native_mes_intent()
        bin_name = get_bin(self.item, "Stores - QAB").name
        frappe.db.commit()
        self.remember_new_names()
        frappe.db.rollback()  # fixture-only evidence epoch before tested RR snapshot
        self.assertEqual(frappe.db.sql("SELECT actual_qty,reserved_qty FROM `tabBin` WHERE name=%s", (bin_name,))[0], (0, 0))
        self.assertEqual(frappe.db.sql("SELECT stock_qty FROM `tabMaterial Request Item` WHERE name=%s", (mr.items[0].name,))[0][0], 3)
        peer = self.boundary_database()
        peer_state = intent.boundary.initialize(peer)  # physical authority before the peer's first business write
        peer.sql("UPDATE `tabBin` SET actual_qty=9,reserved_qty=2 WHERE name=%s", (bin_name,))
        peer.sql("UPDATE `tabMaterial Request Item` SET qty=7,stock_qty=7 WHERE name=%s", (mr.items[0].name,))
        from types import SimpleNamespace
        from frappe.utils import CallbackManager
        request = SimpleNamespace(after_response=CallbackManager())
        with patch.object(frappe.local, "db", peer), patch.object(frappe.local, "purchase_session", peer_state), \
                patch.object(frappe.local, "request", request, create=True):
            later_source = frappe.get_doc("Material Request", mr.name)
            later_source.flags.mes_integration_request = True
            later_source.update_requested_qty()
            request.after_response.reset()
            later_name = peer.sql("SELECT name FROM `tabIntegration Request` WHERE reference_docname=%s AND name<>%s",
                (mr.name, row.name))[0][0]
            peer.commit()  # source demand change + actual later producer intent in the same peer parent transaction
        self.remember_new_names()
        before_epoch = intent.boundary.initialize().epoch
        failures, original = [], audit.runtime_log
        read_points, adapter = [], intent.adapt_query
        def observe_read(query, state):
            result = adapter(query, state)
            if "tabBin" in str(query) and str(query).lstrip().lower().startswith("select"):
                read_points.append((state.depth, len(state.locks), str(query), str(result)))
            return result
        def observe_failure(context, result, error, *args):
            failures.append((type(error).__name__, str(error)))
            return original(context, result, error, *args)
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")), \
                patch.object(audit, "runtime_log", side_effect=observe_failure):
            with patch.object(intent, "adapt_query", side_effect=observe_read):
                intent.dispatch(row.name)
        self.remember_new_names()
        self.assertEqual(len(failures), 1)
        self.assertEqual(failures[0][0], "QueryDeadlockError")
        self.assertIn("1020", failures[0][1])
        self.assertGreater(intent.boundary.initialize().epoch, before_epoch,
            "Actual stale-view failure must use the shared full physical rollback")
        unfinished = frappe.db.get_values("Integration Request", {"name": row.name},
            ["status", "request_description", "output"], as_dict=True, for_update=True)[0]
        self.assertEqual((unfinished.status, unfinished.request_description, unfinished.output),
            ("Queued", intent.PENDING, None))
        self.assertEqual(frappe.db.sql("SELECT actual_qty,reserved_qty,indented_qty FROM `tabBin` "
            "WHERE name=%s FOR UPDATE", (bin_name,))[0], (9, 2, 0))
        # End the failed callback's native Error Log transaction. The previous
        # Bin work must already be rolled back; this is not a read-view refresh.
        frappe.db.commit()
        self.remember_new_names()
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            self.assertEqual(set(intent.recover()), {row.name, later_name})
        self.remember_new_names()
        receipt = frappe.db.get_values("Integration Request", {"name": row.name},
            ["status", "request_description", "output"], as_dict=True, for_update=True)[0]
        self.assertEqual((receipt.status, receipt.request_description), ("Completed", intent.ACKNOWLEDGED))
        ack = json.loads(receipt.output)
        self.assertEqual((ack["generation"], ack["user"]), (row.request_id, "Administrator"))
        self.assertEqual(ack["bins"][0]["indented_qty"], "7.0")
        self.assertEqual(ack["claim_names"], sorted([row.name, later_name]))
        self.assertEqual(frappe.db.get_value("Integration Request", later_name, "status", for_update=True), "Completed")
        actual = frappe.db.sql("SELECT actual_qty,reserved_qty,indented_qty FROM `tabBin` WHERE name=%s FOR UPDATE", (bin_name,))[0]
        self.assertEqual(actual, (9, 2, 7), "Independent recovery transaction must commit current native Bin/demand and exact ACK")

    def test_native_bin_peer_creation_rr_refusal_captured_aliases_and_unique_fallback_current_fields(self):
        from erpnext.stock import utils, stock_balance, stock_ledger
        from erpnext.buying.doctype.purchase_order import purchase_order
        from erpnext.stock.doctype.stock_entry import stock_entry
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        self.assertIs(purchase_order.get_bin, utils.get_bin)
        self.assertIs(stock_entry.get_bin, utils.get_bin)
        self.assertIs(stock_ledger.get_or_make_bin, utils.get_or_make_bin)
        frappe.db.commit()
        self.remember_new_names()
        self.assertEqual(frappe.db.sql("SELECT name FROM `tabBin` WHERE item_code=%s AND warehouse=%s",
            (self.item, "Stores - QAB")), ())  # established real RR view with absent Bin
        peer = self.boundary_database()
        state = boundary.initialize(peer)
        with patch.object(frappe.local, "db", peer), patch.object(frappe.local, "purchase_session", state):
            bin_doc = utils.get_bin(self.item, "Stores - QAB")  # actual native concurrent creator, not SQL fixture insert
            bin_doc.db_set({"actual_qty": 9, "reserved_qty": 2})
            peer.commit()
        self.remember_new_names()
        keys = (boundary.fence_key(), boundary.lock_key("pair", COMPANY, self.item, "Stores - QAB"))
        with self.assertRaises(frappe.QueryDeadlockError) as caught:
            with boundary.execution(), boundary.acquire(keys):
                purchase_order.get_bin(self.item, "Stores - QAB")
        self.assertIn("1020", str(caught.exception))
        # A separate fresh transaction uses the original imported aliases and math.
        with boundary.execution(), boundary.acquire(keys):
            for alias in (utils.get_bin, purchase_order.get_bin, stock_entry.get_bin):
                doc = alias(self.item, "Stores - QAB")
                self.assertEqual((doc.name, doc.actual_qty, doc.reserved_qty), (bin_doc.name, 9, 2))
            self.assertEqual(stock_ledger.get_or_make_bin(self.item, "Stores - QAB"), bin_doc.name)
            with patch.object(frappe, "get_last_doc", wraps=frappe.get_last_doc) as last_doc:
                fallback = utils._create_bin(self.item, "Stores - QAB")  # actual native unique-key collision path
                last_doc.assert_called_once_with("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"})
                self.assertEqual((fallback.name, fallback.actual_qty, fallback.reserved_qty), (bin_doc.name, 9, 2))
            stock_balance.update_bin_qty(self.item, "Stores - QAB", {"indented_qty": 3})
        frappe.db.commit()
        self.remember_new_names()
        self.assertEqual(frappe.db.get_value("Bin", bin_doc.name, ["actual_qty", "reserved_qty", "indented_qty"], for_update=True),
            (9, 2, 3), "Original full db_update must preserve the current native fields, not a stale snapshot")

    def test_native_mes_all_nonstock_requires_explicit_current_noop_ack(self):
        from types import SimpleNamespace
        from redis.exceptions import ConnectionError
        from deeplinkerp_branding.services import purchase_native_intent as intent
        item = frappe.get_doc("Item", self.item)
        item.is_stock_item = 0
        item.save()
        mr, row = self.native_mes_intent()
        import mes_integration.mes_integration.material_request as mes
        with patch.dict(frappe.local.flags, {"mes_integration_request": True}), \
                patch.object(frappe.local, "job", SimpleNamespace(meta={intent.CLAIM_META: {"records": [row.name]}}), create=True):
            self.assertIsNone(mes.sync_material_request_bins(mr.name, [(self.item, "Stores - QAB")]))
        self.assertEqual(frappe.db.get_value("Integration Request", row.name, ["status", "output"], for_update=True), ("Queued", None),
            "Ordinary flags/local job metadata/empty native return must not grant a private ACK claim")
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            intent.dispatch(row.name)
        self.remember_new_names()
        actual = frappe.db.get_values("Integration Request", {"name": row.name},
            ["status", "request_description", "output"], as_dict=True, for_update=True)[0]
        self.assertEqual((actual.status, actual.request_description), ("Completed", intent.ACKNOWLEDGED),
            "Empty native stock result alone is not an ACK; fixed current no-op proof is required")
        ack = json.loads(actual.output)
        self.assertEqual(ack["bins"], [])
        self.assertEqual([(pair["item_code"], pair["warehouse"], pair["is_stock_item"])
            for pair in ack["nonstock_pairs"]], [(self.item, "Stores - QAB", 0)])
        self.assertEqual(ack["nonstock_pairs"][0]["item_modified"], str(item.modified))
        self.assertFalse(frappe.db.exists("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}))

    def test_native_mes_mixed_nonstock_proof_and_stock_to_nonstock_residue_are_distinct(self):
        from redis.exceptions import ConnectionError
        from types import SimpleNamespace
        from frappe.utils import CallbackManager
        from erpnext.stock.utils import get_bin
        from deeplinkerp_branding.services import purchase_native_intent as intent
        nonstock = "QA-ATOMIC-" + uuid.uuid4().hex[:10]
        frappe.get_doc({"doctype": "Item", "item_code": nonstock, "item_name": nonstock, "item_group": "All Item Groups",
            "stock_uom": "Nos", "is_stock_item": 0}).insert()
        request = SimpleNamespace(after_response=CallbackManager())
        with patch.object(frappe.local, "request", request, create=True):
            mr = frappe.get_doc({"doctype": "Material Request", "company": COMPANY, "material_request_type": "Purchase",
                "transaction_date": nowdate(), "schedule_date": add_days(nowdate(), 1), "items": [
                    {"item_code": item, "qty": 3, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}
                    for item in (self.item, nonstock)]})
            mr.flags.mes_integration_request = True
            mr.insert().submit()
            request.after_response.reset()
            frappe.db.commit()
            self.remember_new_names()
        row = frappe.db.get_value("Integration Request", {"integration_request_service": intent.SERVICE, "reference_docname": mr.name},
            ["name", "request_id"], as_dict=True, for_update=True)
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            intent.dispatch(row.name)
        self.remember_new_names()
        ack = json.loads(frappe.db.get_value("Integration Request", row.name, "output", for_update=True))
        self.assertEqual([entry["item_code"] for entry in ack["bins"]], [self.item])
        self.assertEqual([entry["item_code"] for entry in ack["nonstock_pairs"]], [nonstock])
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item}, "indented_qty", for_update=True), 3)
        self.assertFalse(frappe.db.exists("Bin", {"item_code": nonstock}))
        # A separately produced genuine stock claim must not silently exclude its residual Bin.
        other, pending = self.native_mes_intent()
        self.assertTrue(frappe.db.exists("Bin", {"item_code": self.item}))
        frappe.get_doc("Item", self.item).db_set("is_stock_item", 0)
        frappe.db.commit()
        self.remember_new_names()
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            intent.dispatch(pending.name)
        self.remember_new_names()
        self.assertEqual(frappe.db.get_value("Integration Request", pending.name, ["status", "output"], for_update=True), ("Queued", None))
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item}, "indented_qty", for_update=True), 3)

    def test_native_mes_late_ack_dispatch_is_verified_no_write_replay(self):
        from redis.exceptions import ConnectionError
        from deeplinkerp_branding.services import purchase_native_intent as intent
        mr, row = self.native_mes_intent()
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            intent.dispatch(row.name)
        self.remember_new_names()
        with patch.object(frappe, "enqueue", side_effect=AssertionError("ACK replay must not enqueue")) as enqueue, \
                patch.object(frappe.db, "set_value", wraps=frappe.db.set_value) as writes:
            intent.dispatch(row.name)
            enqueue.assert_not_called()
            writes.assert_not_called()

    def test_native_mes_partial_fallback_preserves_first_error_when_runtime_logger_fails(self):
        from redis.exceptions import ConnectionError
        from deeplinkerp_branding.services import purchase_native_intent as intent, purchase_operation as audit
        from erpnext.stock import stock_balance
        warehouse = frappe.get_doc({"doctype": "Warehouse", "warehouse_name": "QA-ATOMIC-" + uuid.uuid4().hex,
            "company": COMPANY, "parent_warehouse": "All Warehouses - QAB"}).insert().name
        mr, row = self.native_mes_intent((warehouse, "Stores - QAB"))
        original, calls = stock_balance.update_bin_qty, []
        def fail_second(item_code, warehouse, qty_dict=None):
            calls.append((item_code, warehouse))
            if len(calls) == 2:
                raise RuntimeError("QA first native second-pair failure")
            return original(item_code, warehouse, qty_dict)
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")), \
                patch.object(stock_balance, "update_bin_qty", side_effect=fail_second), \
                patch.object(audit, "runtime_log", side_effect=OSError("QA diagnostic logger unavailable")), \
                patch.object(frappe, "log_error", wraps=frappe.log_error) as native_errors:
            intent.dispatch(row.name)
        self.remember_new_names()
        self.assertEqual(len(calls), 2)
        message = native_errors.call_args.kwargs["message"]
        self.assertIn("QA first native second-pair failure", message)
        self.assertNotIn("QA diagnostic logger unavailable", message,
            "Best-effort diagnostic failure must not replace the first native exception")
        self.assertEqual(frappe.db.get_value("Integration Request", row.name, "status", for_update=True), "Queued")
        self.assertFalse(frappe.db.get_values("Bin", {"item_code": self.item, "warehouse": ["in", [warehouse, "Stores - QAB"]]}, "name"))
        frappe.db.commit()  # the native Error Log only; no failed Bin/ACK effect
        self.remember_new_names()
        with patch.object(frappe, "enqueue", side_effect=ConnectionError("QA exact Redis unavailable")):
            self.assertEqual(intent.recover(), (row.name,))
        self.remember_new_names()
        self.assertEqual(frappe.db.get_value("Integration Request", row.name, "status", for_update=True), "Completed")

    def test_native_mes_successful_get_locks_survive_callback_reset_until_physical_cleanup(self):
        import mes_integration.mes_integration.material_request as mes
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        names, native_sql = [], frappe.db.sql
        def observe_sql(query, values=(), **kwargs):
            result = native_sql(query, values, **kwargs)
            if str(query).strip() == "SELECT GET_LOCK(%s, %s)" and result[0][0] == 1:
                names.append(values[0])
            return result
        try:
            with patch.object(frappe.db, "sql", side_effect=observe_sql):
                with mes.lock_mes_material_request_bins({"company": COMPANY,
                        "items": [{"item_code": self.item, "warehouse": "Stores - QAB"}]}):
                    self.assertEqual(len(names), 1)
            self.assertIn(names[0], boundary.initialize().locks,
                "Actual original successful MES GET_LOCK must belong to the same physical Session")
            frappe.db.after_commit.reset()
            frappe.db.after_rollback.reset()
            frappe.db.commit()
            self.remember_new_names()
            self.assertIsNone(frappe.db.sql("SELECT IS_USED_LOCK(%s)", (names[0],))[0][0])
        finally:
            for name in names:
                if frappe.db.sql("SELECT IS_USED_LOCK(%s)", (name,))[0][0] == frappe.db.sql("SELECT CONNECTION_ID()")[0][0]:
                    frappe.db.sql("SELECT RELEASE_LOCK(%s)", (name,))  # exact synthetic lock from original acquisition

    def test_native_mes_tracked_lock_savepoint_nested_rollback_callback_close_and_kill_matrix(self):
        import mes_integration.mes_integration.material_request as mes
        from erpnext.stock.utils import get_bin
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        get_bin(self.item, "Stores - QAB")
        frappe.db.commit()
        self.remember_new_names()
        data = {"company": COMPANY, "items": [{"item_code": self.item, "warehouse": "Stores - QAB"}]}
        for phase in ("savepoint", "nested", "rollback", "early-callback-failure", "close", "kill"):
            with self.subTest(phase=phase):
                left, right = self.boundary_database(), self.boundary_database()
                state = boundary.initialize(left)
                with patch.object(frappe.local, "db", left), patch.object(frappe.local, "purchase_session", state):
                    with mes.lock_mes_material_request_bins(data):
                        if phase == "nested":
                            with mes.lock_mes_material_request_bins(data):
                                self.assertEqual(len(state.native_acquisitions), 2)
                    names = {name for name, epoch in state.native_acquisitions.values()}
                    self.assertEqual(len(names), 1)
                    for name in names:
                        self.assertEqual(right.sql("SELECT IS_USED_LOCK(%s)", (name,))[0][0], state.connection_id)
                    if phase == "savepoint":
                        left.savepoint("native_owned_lock")
                        left.rollback(save_point="native_owned_lock")
                        self.assertEqual(len(state.native_acquisitions), 1)
                        self.assertEqual(right.sql("SELECT IS_USED_LOCK(%s)", (next(iter(names)),))[0][0], state.connection_id)
                        left.commit()
                    elif phase == "early-callback-failure":
                        first = RuntimeError("QA first before physical commit")
                        left.before_commit.add(lambda: (_ for _ in ()).throw(first))
                        with self.assertRaises(RuntimeError) as caught:
                            left.commit()
                        self.assertIs(caught.exception, first)
                    elif phase == "close":
                        left.close()
                        self.assertTrue(state.poisoned)
                    elif phase == "kill":
                        right.sql("KILL CONNECTION %s", (state.connection_id,))
                        with self.assertRaises(Exception):
                            left.sql("SELECT CONNECTION_ID()")
                        self.assertTrue(state.poisoned)
                        with self.assertRaises(frappe.ValidationError):
                            boundary.initialize(left)
                    else:
                        left.rollback()
                    for name in names:
                        self.assertIsNone(right.sql("SELECT IS_USED_LOCK(%s)", (name,))[0][0])
                    self.assertEqual(state.native_acquisitions, {})

    def test_native_mes_late_original_callback_cannot_release_new_same_name_acquisition(self):
        import mes_integration.mes_integration.material_request as mes
        data = {"company": COMPANY, "items": [{"item_code": self.item, "warehouse": "Stores - QAB"}]}
        with mes.lock_mes_material_request_bins(data):
            pass
        old_callback = frappe.db.after_commit._functions[-1]
        frappe.db.commit()
        self.remember_new_names()
        names, native_sql = [], frappe.db.sql
        def observe_sql(query, values=(), **kwargs):
            result = native_sql(query, values, **kwargs)
            if str(query).strip() == "SELECT GET_LOCK(%s, %s)" and result[0][0] == 1:
                names.append(values[0])
            return result
        with patch.object(frappe.db, "sql", side_effect=observe_sql):
            with mes.lock_mes_material_request_bins(data):
                pass
        try:
            old_callback()
            self.assertEqual(frappe.db.sql("SELECT IS_USED_LOCK(%s)", (names[0],))[0][0],
                frappe.db.sql("SELECT CONNECTION_ID()")[0][0], "Old epoch callback must never consume new original native lock")
        finally:
            frappe.db.rollback()

    def test_native_mes_body_cannot_forge_successful_acquisition_or_finally_callback_witness(self):
        from types import FunctionType
        import mes_integration.mes_integration.material_request as mes
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        frappe.db.commit()
        self.remember_new_names()
        data = {"company": COMPANY, "items": [{"item_code": self.item, "warehouse": "Stores - QAB"}]}
        unknown = "mes_mr_bin_" + uuid.uuid4().hex * 2
        state = boundary.initialize()
        reached, original = [], state.native_sql
        def observe_native(query, values=(), **kwargs):
            if isinstance(values, (list, tuple)) and values and values[0] == unknown:
                reached.append(str(query))
            return original(query, values, **kwargs)
        with patch.object(state, "native_sql", side_effect=observe_native):
            for action in ("acquire", "callback"):
                with self.subTest(action=action):
                    with mes.lock_mes_material_request_bins(data):
                        if action == "acquire":
                            with self.assertRaisesRegex(frappe.ValidationError, "调用点"):
                                frappe.db.sql("SELECT GET_LOCK(%s, %s)", (unknown, 180))
                            self.assertEqual(reached, [], "Unknown body GET_LOCK must be rejected before actual acquisition")
                        else:
                            names = tuple(name for name, _ in reversed(tuple(state.native_acquisitions.values())))
                            cell = (lambda value: lambda: value)(names).__closure__
                            forged = FunctionType(mes._dlp_native_intent_lock_callback, mes.__dict__, closure=cell)
                            with self.assertRaisesRegex(frappe.ValidationError, "调用点"):
                                frappe.db.after_commit.add(forged)

    def test_native_mes_first_exception_survives_partial_or_body_error_and_callback_add_failure(self):
        import mes_integration.mes_integration.material_request as mes
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        warehouse = frappe.get_doc({"doctype": "Warehouse", "warehouse_name": "QA-ATOMIC-" + uuid.uuid4().hex,
            "company": COMPANY, "parent_warehouse": "All Warehouses - QAB"}).insert().name
        from erpnext.stock.utils import get_bin
        for name in (warehouse, "Stores - QAB"):
            get_bin(self.item, name)
        frappe.db.commit()
        self.remember_new_names()
        data = {"company": COMPANY, "items": [{"item_code": self.item, "warehouse": name}
            for name in (warehouse, "Stores - QAB")]}
        for failure in ("partial", "body"):
            for manager_name in ("after_commit", "after_rollback"):
                with self.subTest(first_failure=failure, callback=manager_name):
                    primary = RuntimeError("QA first native " + failure + " failure")
                    reached_body = []
                    calls, native_sql = [], frappe.db.sql
                    def fail_second(query, values=(), **kwargs):
                        if str(query).strip() == "SELECT GET_LOCK(%s, %s)":
                            calls.append(values[0])
                            if failure == "partial" and len(calls) == 2:
                                raise primary
                        return native_sql(query, values, **kwargs)
                    caught = None
                    try:
                        with patch.object(frappe.db, "sql", side_effect=fail_second), \
                                patch.object(getattr(frappe.db, manager_name), "add",
                                    side_effect=RuntimeError("QA callback add failure")):
                            with mes.lock_mes_material_request_bins(data):
                                reached_body.append(True)
                                raise primary
                    except Exception as error:
                        caught = error
                    self.assertEqual(bool(reached_body), failure == "body", "The intended native failure point must be exercised")
                    self.assertIs(caught, primary, "Native finally registration must not replace the first exception object")
                    self.assertFalse(boundary.initialize().native_acquisitions)

    def test_native_mes_unknown_release_all_or_commented_release_is_rejected_before_execution(self):
        import mes_integration.mes_integration.material_request as mes
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        frappe.db.commit()
        self.remember_new_names()
        captured_release = mes.release_mes_material_request_locks
        peer = self.boundary_database()
        state = boundary.initialize()
        for shape in ("generator", "list", "tuple"):
            with self.subTest(ordinary_release=shape):
                names = tuple("QA-ATOMIC-ORDINARY-RELEASE-" + uuid.uuid4().hex for _ in range(2))
                values = (name for name in names) if shape == "generator" else list(names) if shape == "list" else names
                try:
                    for name in names:
                        self.assertEqual(frappe.db.sql("SELECT GET_LOCK(%s, %s)", (name, 0))[0][0], 1)
                        self.assertNotIn(name, state.locks)
                        self.assertEqual(peer.sql("SELECT IS_USED_LOCK(%s)", (name,))[0][0], state.connection_id)
                    captured_release(values)
                    held = [peer.sql("SELECT IS_USED_LOCK(%s)", (name,))[0][0] for name in names]
                    print("C1_QUALITY_ORDINARY_RELEASE=" + json.dumps({"shape": shape, "remaining_physical_owners": held}), flush=True)
                    self.assertEqual(held, [None, None], "Ordinary one-shot iterables must reach the original native release loop")
                finally:
                    for name in names:
                        if peer.sql("SELECT IS_USED_LOCK(%s)", (name,))[0][0] == state.connection_id:
                            state.raw("SELECT RELEASE_LOCK(%s)", (name,))  # exact own untracked fixture locks, including RED leftovers
        with mes.lock_mes_material_request_bins({"company": COMPANY,
                "items": [{"item_code": self.item, "warehouse": "Stores - QAB"}]}):
            names = tuple(name for name, _ in state.native_acquisitions.values())
            for shape in ("generator", "list", "tuple"):
                with self.subTest(tracked_release=shape):
                    values = (name for name in names) if shape == "generator" else list(names) if shape == "list" else names
                    with self.assertRaisesRegex(frappe.ValidationError, "释放"):
                        captured_release(values)
                    self.assertTrue(all(peer.sql("SELECT IS_USED_LOCK(%s)", (name,))[0][0] == state.connection_id for name in names))
        frappe.db.rollback()  # native frozen callbacks / physical Session release, before the next separate lock fixture
        for query, shape in (("SELECT RELEASE_ALL_LOCKS()", "canonical"),
                ("SELECT RELEASE_LOCK /* ordinary comment */ (%s)", "canonical"),
                ("SELECT RELEASE_LOCK(%s)", "bytes"), ("SELECT RELEASE_LOCK(%s)", "upper")):
            with self.subTest(query=query, bound_shape=shape):
                peer = self.boundary_database()
                peer_state = boundary.initialize(peer)
                with patch.object(frappe.local, "db", peer), \
                        patch.object(frappe.local, "purchase_session", peer_state):
                    state = boundary.initialize()
                    with mes.lock_mes_material_request_bins({"company": COMPANY,
                            "items": [{"item_code": self.item, "warehouse": "Stores - QAB"}]}):
                        names = tuple(name for name, _ in state.native_acquisitions.values())
                        value = names[0].encode("ascii") if shape == "bytes" else names[0].upper() if shape == "upper" else names[0]
                        used = peer.sql("SELECT IS_USED_LOCK(%s)", (value,))[0][0]
                        print("NATIVE_LOCK_BOUND_EQUIVALENCE=" + json.dumps({"shape": shape,
                            "same_physical_owner": used == state.connection_id, "operation": "release"}), flush=True)
                        if used != state.connection_id:
                            self.assertIsNone(peer.sql(query, (value,))[0][0])  # native distinct/untracked name stays native
                        else:
                            try:
                                with self.assertRaisesRegex(frappe.ValidationError, "释放"):
                                    peer.sql(query, (value,) if "%s" in query else ())
                                self.assertEqual(peer.sql("SELECT IS_USED_LOCK(%s)", (names[0],))[0][0], state.connection_id)
                            finally:
                                # Exact fixture lock only: retain one original count if RED released it.
                                if state.raw("SELECT IS_USED_LOCK(%s)", (names[0],))[0][0] != state.connection_id:
                                    self.assertEqual(state.raw("SELECT GET_LOCK(%s, %s)", (names[0], 0))[0][0], 1)
                    peer.rollback()

    def test_native_mes_dormant_tracked_name_reentry_is_rejected_but_unrelated_lock_remains_native(self):
        import mes_integration.mes_integration.material_request as mes
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        with mes.lock_mes_material_request_bins({"company": COMPANY,
                "items": [{"item_code": self.item, "warehouse": "Stores - QAB"}]}):
            pass
        state = boundary.initialize()
        name = next(iter(state.native_acquisitions.values()))[0]
        reached, native_sql = [], state.native_sql
        def observe(query, values=(), **kwargs):
            if str(query).strip() == "SELECT GET_LOCK(%s, %s)":
                reached.append(values[0])
            return native_sql(query, values, **kwargs)
        with patch.object(state, "native_sql", side_effect=observe):
            for shape, value in (("canonical", name), ("bytes", name.encode("ascii")), ("upper", name.upper())):
                with self.subTest(bound_shape=shape):
                    reached.clear()
                    used = frappe.db.sql("SELECT IS_USED_LOCK(%s)", (value,))[0][0]
                    print("NATIVE_LOCK_BOUND_EQUIVALENCE=" + json.dumps({"shape": shape,
                        "same_physical_owner": used == state.connection_id, "operation": "reentry"}), flush=True)
                    if used != state.connection_id:
                        self.assertEqual(frappe.db.sql("SELECT GET_LOCK(%s, %s)", (value, 0))[0][0], 1)
                        self.assertEqual(frappe.db.sql("SELECT RELEASE_LOCK(%s)", (value,))[0][0], 1)
                    else:
                        try:
                            with self.assertRaisesRegex(frappe.ValidationError, "调用点"):
                                frappe.db.sql("SELECT GET_LOCK(%s, %s)", (value, 0))
                            self.assertEqual(reached, [])
                        finally:
                            if reached:  # release only the exact extra native count actually acquired by RED
                                self.assertEqual(state.raw("SELECT RELEASE_LOCK(%s)", (value,))[0][0], 1)
            ordinary = "QA-UNRELATED-" + uuid.uuid4().hex
            self.assertEqual(frappe.db.sql("SELECT GET_LOCK(%s, %s)", (ordinary, 0))[0][0], 1)
            self.assertEqual(frappe.db.sql("SELECT RELEASE_LOCK(%s)", (ordinary,))[0][0], 1)

    def test_native_bundle_late_packed_pairs_and_real_mes_source_use_fence_absence_gate(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        from erpnext.stock.doctype.material_request.material_request import make_purchase_order
        parent = self.scope_item()
        item = frappe.get_doc("Item", parent)
        item.is_stock_item = 0
        item.save()
        frappe.get_doc({"doctype": "Product Bundle", "new_item_code": parent,
            "items": [{"item_code": self.item, "qty": 1, "uom": "Nos"}]}).insert()
        mr = frappe.get_doc({"doctype": "Material Request", "company": COMPANY, "material_request_type": "Purchase",
            "transaction_date": nowdate(), "schedule_date": add_days(nowdate(), 1), "items": [
                {"item_code": self.item, "qty": 4, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]}).insert()
        mr.submit()
        frappe.db.set_value(mr.doctype, mr.name, "custom_request_source", "MES", update_modified=False)
        self.commit_fixture(primary_doctype="Item", name_prefix="QA-ATOMIC-")
        mr.reload()
        po = make_purchase_order(mr.name)
        po.supplier = "QA Test Supplier"
        po.items[0].rate = 10
        po.insert(set_name="QA-ATOMIC-PO-" + uuid.uuid4().hex[:16])
        self.assertEqual(po.items[0].material_request, mr.name)
        self.assertIn(boundary.fence_key(), boundary.initialize().locks)
        # Parent Item is non-stock and packed_items is initially EMPTY. The
        # untouched native calculator really creates the late stock child.
        dn = frappe.get_doc({"doctype": "Delivery Note", "company": COMPANY, "customer": "QA Purchase Payments Customer",
            "currency": "CNY", "conversion_rate": 1, "selling_price_list": "Standard Selling",
            "items": [{"item_code": parent, "qty": 1, "rate": 10, "warehouse": "Stores - QAB"}]}).insert()
        self.assertEqual([(row.item_code, row.warehouse) for row in dn.packed_items], [(self.item, "Stores - QAB")])
        owner = self.pending_owner()
        name = frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "name")
        frappe.db.set_value("Bin", name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        for doc in (po, dn):
            with self.subTest(doctype=doc.doctype), self.assertRaisesRegex(frappe.ValidationError, "未核查的原生关联路径"):
                doc.reload().save(ignore_permissions=True)
        # Fixture-only final publication/unlock also borrows the same fence.
        with boundary.execution(), boundary.acquire((boundary.fence_key(),)):
            frappe.db.set_value("Bin", name, boundary.POINTER, None, update_modified=False)
        self.commit_fixture()
        po.reload().save(ignore_permissions=True)
        dn.reload().save(ignore_permissions=True)

    def test_empty_incoming_bundle_packing_cannot_late_create_pending_child_pair(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        parent = self.scope_item()
        item = frappe.get_doc("Item", parent)
        item.is_stock_item = 0
        item.save()
        frappe.get_doc({"doctype": "Product Bundle", "new_item_code": parent,
            "items": [{"item_code": self.item, "qty": 1, "uom": "Nos"}]}).insert()
        self.order()  # real native ordered Bin for the late stock child
        def draft():
            return frappe.get_doc({"doctype": "Delivery Note", "company": COMPANY, "customer": "QA Purchase Payments Customer",
                "currency": "CNY", "conversion_rate": 1, "selling_price_list": "Standard Selling",
                "items": [{"item_code": parent, "qty": 1, "rate": 10, "warehouse": "Stores - QAB"}]})
        normal = draft().insert()
        self.assertEqual([(row.item_code, row.warehouse) for row in normal.packed_items], [(self.item, "Stores - QAB")])
        owner = self.pending_owner()
        name = frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "name")
        frappe.db.set_value("Bin", name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        incoming = draft()
        self.assertFalse(incoming.packed_items)
        with self.assertRaisesRegex(frappe.ValidationError, "未核查的原生关联路径"):
            incoming.insert(ignore_permissions=True)
        with boundary.execution(), boundary.acquire((boundary.fence_key(),)):
            frappe.db.set_value("Bin", name, boundary.POINTER, None, update_modified=False)
        self.commit_fixture()
        self.assertTrue(draft().insert().packed_items)

    def test_union_budget_real_old_new_native_pairs_and_sources_refuse_before_audit_sql(self):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        first = self.order()
        second = self.order(items=[{"item_code": self.item, "qty": 10, "rate": 12.345,
            "warehouse": "Work In Progress - QAB", "schedule_date": add_days(nowdate(), 1)}])
        receipt = make_purchase_receipt(first.name).insert()
        self.commit_fixture()
        for kind in ("pairs", "identities"):
            with self.subTest(kind=kind):
                doc = frappe.get_doc(first.doctype, first.name) if kind == "pairs" else frappe.get_doc(receipt.doctype, receipt.name)
                doc.items[0].warehouse = "Work In Progress - QAB"
                if kind == "identities":
                    doc.items[0].purchase_order, doc.items[0].purchase_order_item = second.name, second.items[0].name
                statements, sql = [], frappe.db.sql
                def observe(query, *args, **kwargs):
                    statements.append(str(query))
                    return sql(query, *args, **kwargs)
                with patch.object(scope, "MAX_PAIRS" if kind == "pairs" else "MAX_VOUCHERS", 1 if kind == "pairs" else 2), \
                        patch.object(frappe.db, "sql", side_effect=observe):
                    with self.assertRaisesRegex(frappe.ValidationError, "范围超过安全上限"):
                        doc.save(ignore_permissions=True)
                self.assertFalse(any("INSERT INTO `tabIntegration Request`" in query for query in statements))

    def test_header_only_stock_entry_native_warehouse_fill_respects_pending_absence_fence(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        self.order()
        def draft():
            return frappe.get_doc({"doctype": "Stock Entry", "company": COMPANY, "stock_entry_type": "Material Receipt",
                "to_warehouse": "Stores - QAB", "items": [{"item_code": self.item, "qty": 1, "allow_zero_valuation_rate": 1}]})
        normal = draft().insert()
        self.assertEqual(normal.items[0].t_warehouse, "Stores - QAB")
        owner = self.pending_owner()
        name = frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "name")
        frappe.db.set_value("Bin", name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        with self.assertRaisesRegex(frappe.ValidationError, "未核查的原生关联路径"):
            draft().insert(ignore_permissions=True)
        with boundary.execution(), boundary.acquire((boundary.fence_key(),)):
            frappe.db.set_value("Bin", name, boundary.POINTER, None, update_modified=False)
        self.commit_fixture()
        self.assertEqual(draft().insert().items[0].t_warehouse, "Stores - QAB")

    def test_real_putaway_rule_late_stock_entry_and_receipt_warehouse_use_absence_fence(self):
        from deeplinkerp_branding.services import purchase_repost_boundary as boundary
        source = self.order()
        self.order(items=[{"item_code": self.item, "qty": 10, "rate": 12.345,
            "warehouse": "Work In Progress - QAB", "schedule_date": add_days(nowdate(), 1)}])
        frappe.get_doc({"doctype": "Putaway Rule", "company": COMPANY, "item_code": self.item,
            "warehouse": "Work In Progress - QAB", "capacity": 100, "priority": 1, "uom": "Nos", "conversion_factor": 1}).insert()
        def draft(doctype):
            if doctype == "Purchase Receipt":
                doc = make_purchase_receipt(source.name)
                doc.apply_putaway_rule = 1
                return doc
            return frappe.get_doc({"doctype": "Stock Entry", "company": COMPANY, "stock_entry_type": "Material Receipt",
                "apply_putaway_rule": 1, "items": [{"item_code": self.item, "qty": 1, "transfer_qty": 1,
                    "uom": "Nos", "stock_uom": "Nos", "conversion_factor": 1, "t_warehouse": "Stores - QAB",
                    "allow_zero_valuation_rate": 1}]})
        for doctype in ("Stock Entry", "Purchase Receipt"):
            doc = draft(doctype).insert()
            self.assertEqual(doc.items[0].get("t_warehouse" if doctype == "Stock Entry" else "warehouse"), "Work In Progress - QAB")
        owner = self.pending_owner()
        name = frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Work In Progress - QAB"}, "name")
        frappe.db.set_value("Bin", name, boundary.POINTER, owner.name, update_modified=False)
        self.commit_fixture()
        for doctype in ("Stock Entry", "Purchase Receipt"):
            with self.subTest(doctype=doctype), self.assertRaisesRegex(frappe.ValidationError, "未核查的原生关联路径"):
                draft(doctype).insert(ignore_permissions=True)
        with boundary.execution(), boundary.acquire((boundary.fence_key(),)):
            frappe.db.set_value("Bin", name, boundary.POINTER, None, update_modified=False)
        self.commit_fixture()
        for doctype in ("Stock Entry", "Purchase Receipt"):
            draft(doctype).insert()

    def receipt(self, po):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        key = po.items[0].name
        native_draft = frappe.db.get_value("Purchase Receipt Item", {"purchase_order": po.name, "docstatus": 0}, "parent")
        args = {"target_name": native_draft, "expected_modified": frappe.db.get_value("Purchase Receipt", native_draft, "modified")} if native_draft else {}
        result = actions.record_receipt(po.name, {"items": [{"key": key, "qty": 2, "warehouse": "Stores - QAB"}]}, str(uuid.uuid4()), **args)
        self.assertFalse(result.get("failed"), result)
        self.assertEqual(result["document"]["docstatus"], 1)
        return frappe.get_doc("Purchase Receipt", result["document"]["name"])

    def test_stocked_receipt_invoice_payment_native_form_and_precommit(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        from deeplinkerp_branding.services import purchase_consistency as guard
        po = self.order()
        pr = self.receipt(po)
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "actual_qty"), 2)
        pi = make_purchase_invoice(pr.name).insert()
        from frappe.client import submit
        submit(pi.as_dict())
        pi.reload()
        guard.check_registered()
        args = dict(source_doctype="Purchase Receipt", source_name=pr.name, purchase_invoice=pi.name,
            amount_to_pay=10, bank_account="Cash - QAB", request_id=str(uuid.uuid4()))
        paid = actions.record_payment(**args)
        self.assertFalse(paid.get("failed"), paid)
        pe = frappe.get_doc("Payment Entry", paid["document"]["name"])
        self.assertEqual(pe.docstatus, 1)
        guard.check_registered()
        frappe.db.before_commit.run()  # exercise final form hooks, without committing test data
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "actual_qty"), 2)

    def test_zero_amount_native_invoice_without_financial_movement_is_legal(self):
        for automatic in (0, 1):
            with self.subTest(automatic=automatic):
                if not frappe.db.exists("Item", self.item):
                    frappe.get_doc({"doctype": "Item", "item_code": self.item, "item_name": self.item,
                        "item_group": "All Item Groups", "stock_uom": "Nos", "is_stock_item": 0}).insert()
                else:
                    frappe.db.set_value("Item", self.item, "is_stock_item", 0)
                    frappe.clear_document_cache("Item", self.item)
                frappe.db.set_value("China Finance Settings", COMPANY, "auto_submit_purchase_invoice", automatic)
                try:
                    po = self.order(qty=2, rate=0)
                    pr = make_purchase_receipt(po.name).insert()
                    pr.submit()
                    if automatic:
                        name = frappe.db.get_value("Purchase Invoice Item", {"purchase_receipt": pr.name}, "parent")
                        self.assertTrue(name)
                        pi = frappe.get_doc("Purchase Invoice", name)
                    else:
                        pi = make_purchase_invoice(pr.name).insert()
                        pi.submit()
                    self.assertEqual((pi.docstatus, pi.grand_total, pi.outstanding_amount), (1, 0, 0))
                    self.assertEqual(pi.items[0].qty, 2)
                    self.assertEqual(pi.items[0].rate, 0)
                    for doctype in ("GL Entry", "Payment Ledger Entry"):
                        self.assertFalse(frappe.db.count(doctype, {"voucher_type": pi.doctype, "voucher_no": pi.name}))
                    self.assertFalse(frappe.db.count("China Accounting Voucher", {"source_doctype": pi.doctype, "source_name": pi.name}))
                finally:
                    frappe.db.rollback()
                    frappe.clear_document_cache("Item", self.item)

    def test_order_inventory_bin_changes_are_bound_to_native_replay(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from deeplinkerp_branding.services import purchase_operation as kernel
        po = self.order()
        payload = json.dumps(po.as_dict(), default=str)
        key = str(uuid.uuid4())
        guard.savedocs(payload, "Update", request_id=key)
        self.commit_fixture()
        audit = frappe.get_doc("Integration Request", kernel.identity("Administrator", key))
        self.assertTrue(any(row["doctype"] == "Purchase Order" for row in json.loads(audit.output)["artifacts"]))
        self.assertFalse(frappe.db.count("Stock Ledger Entry", {"voucher_type": po.doctype, "voucher_no": po.name}))
        for field in ("ordered_qty", "indented_qty"):
            with self.subTest(field=field):
                quantity = frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, field)
                frappe.db.set_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, field, quantity + 1, update_modified=False)
                try:
                    with self.assertRaises(frappe.ValidationError) as caught:
                        guard.savedocs(payload, "Update", request_id=key)
                    self.assertEqual(caught.exception.purchase_error_id, "replay_evidence_changed")
                finally:
                    frappe.db.set_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, field, quantity, update_modified=False)

    def test_nonzero_invoice_requires_native_supplier_payment_ledger(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        native = guard.check_invoice_balance
        def lose_ledger(doc):
            if (doc.doctype, doc.name) == (pi.doctype, pi.name):
                self.assertTrue(frappe.db.count("Payment Ledger Entry", {"voucher_type": pi.doctype, "voucher_no": pi.name}))
                frappe.db.delete("Payment Ledger Entry", {"voucher_type": pi.doctype, "voucher_no": pi.name})
            return native(doc)
        with patch.object(guard, "check_invoice_balance", side_effect=lose_ledger), self.assertRaises(frappe.ValidationError) as caught:
            pi.submit()
        self.assertEqual(caught.exception.purchase_error_id, "purchase_invoice_ple_missing")
        self.assertFalse(frappe.db.exists("Purchase Invoice", pi.name))

    def material_request_order(self):
        from erpnext.stock.doctype.material_request.material_request import make_purchase_order
        mr = frappe.get_doc({"doctype": "Material Request", "company": COMPANY, "material_request_type": "Purchase",
            "transaction_date": nowdate(), "schedule_date": add_days(nowdate(), 1),
            "items": [{"item_code": self.item, "qty": 10, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]}).insert()
        mr.submit()
        po = make_purchase_order(mr.name)
        po.supplier = "QA Test Supplier"
        po.items[0].rate = 12.345
        po.insert(set_name="QA-ATOMIC-PO-" + uuid.uuid4().hex[:16])
        return mr, po

    def test_native_order_material_request_updates_are_replay_artifacts(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from deeplinkerp_branding.services import purchase_operation as kernel
        mr, po = self.material_request_order()
        key = str(uuid.uuid4())
        payload = json.dumps(po.as_dict(), default=str)
        guard.savedocs(payload, "Submit", key)
        self.commit_fixture()
        artifacts = json.loads(frappe.get_doc("Integration Request", kernel.identity("Administrator", key)).output)["artifacts"]
        self.assertTrue(any(row["doctype"] == mr.doctype and row["name"] == mr.name for row in artifacts))
        self.assertEqual(frappe.db.get_value("Material Request Item", mr.items[0].name, "ordered_qty"), 10)
        self.assertEqual(frappe.db.get_value(mr.doctype, mr.name, "per_ordered"), 100)
        for doctype, name, field in (("Material Request Item", mr.items[0].name, "ordered_qty"), (mr.doctype, mr.name, "per_ordered")):
            with self.subTest(field=field):
                frappe.db.set_value(doctype, name, field, frappe.db.get_value(doctype, name, field) + 1, update_modified=False)
                with self.assertRaises(frappe.ValidationError) as caught:
                    guard.savedocs(payload, "Submit", key)
                self.assertEqual(caught.exception.purchase_error_id, "replay_evidence_changed")

    def test_native_receipt_material_request_updates_are_replay_artifacts(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from deeplinkerp_branding.services import purchase_operation as kernel
        mr, po = self.material_request_order()
        po.submit()
        self.commit_fixture()
        pr = make_purchase_receipt(po.name).insert()
        key = str(uuid.uuid4())
        payload = json.dumps(pr.as_dict(), default=str)
        guard.savedocs(payload, "Submit", key)
        self.commit_fixture()
        artifact = json.loads(frappe.get_doc("Integration Request", kernel.identity("Administrator", key)).output)["artifacts"]
        self.assertTrue(any(row["doctype"] == mr.doctype and row["name"] == mr.name for row in artifact))
        for doctype, name, field in (("Material Request Item", mr.items[0].name, "received_qty"),
                (mr.doctype, mr.name, "per_received")):
            with self.subTest(field=field):
                frappe.db.set_value(doctype, name, field, frappe.db.get_value(doctype, name, field) + 1, update_modified=False)
                with self.assertRaises(frappe.ValidationError) as caught:
                    guard.savedocs(payload, "Submit", key)
                self.assertEqual(caught.exception.purchase_error_id, "replay_evidence_changed")

    def test_receipt_native_source_invoice_quantity_is_a_replay_artifact(self):
        from erpnext.buying.doctype.purchase_order.purchase_order import make_purchase_invoice as invoice_from_order
        from erpnext.accounts.doctype.purchase_invoice.purchase_invoice import make_purchase_receipt as receipt_from_invoice
        from deeplinkerp_branding.services import purchase_consistency as guard
        from deeplinkerp_branding.services import purchase_operation as kernel
        frappe.db.set_value("Item", self.item, "is_stock_item", 0)
        frappe.clear_document_cache("Item", self.item)
        po = self.order()
        pi = invoice_from_order(po.name).insert()
        pi.submit()
        self.commit_fixture()
        pr = receipt_from_invoice(pi.name).insert()
        self.assertEqual(pr.items[0].purchase_invoice, pi.name)
        key = str(uuid.uuid4())
        payload = json.dumps(pr.as_dict(), default=str)
        guard.savedocs(payload, "Submit", key)
        self.commit_fixture()
        artifacts = json.loads(frappe.get_doc("Integration Request", kernel.identity("Administrator", key)).output)["artifacts"]
        self.assertTrue(any(row["doctype"] == pi.doctype and row["name"] == pi.name for row in artifacts))
        for doctype, name, field in (("Purchase Invoice Item", pi.items[0].name, "received_qty"), (pi.doctype, pi.name, "per_received")):
            with self.subTest(field=field):
                frappe.db.set_value(doctype, name, field, frappe.db.get_value(doctype, name, field) + 1, update_modified=False)
                with self.assertRaises(frappe.ValidationError) as caught:
                    guard.savedocs(payload, "Submit", key)
                self.assertEqual(caught.exception.purchase_error_id, "replay_evidence_changed")

    def test_managed_fulfilment_invalidation_is_a_replay_artifact(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from deeplinkerp_branding.services import purchase_fulfilment_service as fulfilment
        from deeplinkerp_branding.services import purchase_operation as kernel
        po = self.order()
        result = fulfilment.save_link(po.name, po.items[0].name, COMPANY, COMPANY, "trade_custody", 2,
            str(po.modified), custody_company=COMPANY, custody_warehouse="Stores - QAB")
        name = result["name"]
        key = str(uuid.uuid4())
        payload = json.dumps(po.as_dict(), default=str)
        guard.cancel(po.doctype, po.name, request_id=key, doc=payload)
        self.commit_fixture()
        artifacts = json.loads(frappe.get_doc("Integration Request", kernel.identity("Administrator", key)).output)["artifacts"]
        self.assertTrue(any(row["doctype"] == fulfilment.DOCTYPE and row["name"] == name for row in artifacts))
        frappe.db.set_value(fulfilment.DOCTYPE, name, "source_versions_json", "{}", update_modified=False)
        with self.assertRaises(frappe.ValidationError) as caught:
            guard.cancel(po.doctype, po.name, request_id=key, doc=payload)
        self.assertEqual(caught.exception.purchase_error_id, "replay_evidence_changed")

    def cash_assignment_payment(self):
        from china_finance.services import cash_flow_assignment, cash_equivalent_scope, voucher
        from deeplinkerp_branding.services import purchase_document_actions as actions
        if "Cash - QAB" not in cash_equivalent_scope.get_cash_scope_accounts(COMPANY, nowdate()):
            frappe.get_doc({"doctype": "China Cash Equivalent Scope", "company": COMPANY, "account": "Cash - QAB",
                "classification": "库存现金", "included": 1, "policy_basis": "QA synthetic cash scope",
                "effective_from": nowdate()}).insert()
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        settings = frappe.get_doc(voucher.get_company_settings(COMPANY).as_dict())
        settings.cash_flow_assignment_activation_date = nowdate()
        key = str(uuid.uuid4())
        # Local policy fixture only; native voucher/assignment writers run real.
        with patch.object(cash_flow_assignment, "get_company_settings", return_value=settings):
            result = actions.record_payment(source_doctype=pr.doctype, source_name=pr.name, purchase_invoice=pi.name,
                amount_to_pay=10, bank_account="Cash - QAB", request_id=key)
        self.assertFalse(result.get("failed"), result)
        pe = frappe.get_doc("Payment Entry", result["document"]["name"])
        assignment = frappe.get_doc("China Cash Flow Assignment", {"source_doctype": pe.doctype, "source_name": pe.name})
        self.assertEqual((assignment.docstatus, assignment.status), (0, "Draft"))
        self.commit_fixture()
        return pe, assignment, key

    def test_native_finance_posting_cash_assignment_is_a_system_artifact(self):
        from deeplinkerp_branding.services import purchase_operation as kernel
        pe, assignment, key = self.cash_assignment_payment()
        artifacts = json.loads(frappe.get_doc("Integration Request", kernel.identity("Administrator", key)).output)["artifacts"]
        found = [row for row in artifacts if (row["doctype"], row["name"]) == (assignment.doctype, assignment.name)]
        self.assertEqual(len(found), 1)
        self.assertTrue(found[0]["system_effect"])
        frappe.db.set_value(assignment.items[0].doctype, assignment.items[0].name, "cash_amount", assignment.items[0].cash_amount + 1, update_modified=False)
        with self.assertRaises(frappe.ValidationError) as caught:
            kernel.replay_artifacts({"artifacts": artifacts})
        self.assertEqual(caught.exception.purchase_error_id, "replay_evidence_changed")

    def test_native_finance_cancellation_issue_and_cash_assignment_are_system_artifacts(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from deeplinkerp_branding.services import purchase_operation as kernel
        pe, assignment, _ = self.cash_assignment_payment()
        key = str(uuid.uuid4())
        payload = json.dumps(pe.as_dict(), default=str)
        guard.cancel(pe.doctype, pe.name, request_id=key, doc=payload)
        self.commit_fixture()
        issue = frappe.get_doc("China Voucher Sync Issue", {"issue_key": f"Cancellation|{pe.doctype}|{pe.name}"})
        artifacts = json.loads(frappe.get_doc("Integration Request", kernel.identity("Administrator", key)).output)["artifacts"]
        for target, status in ((issue, "Resolved"), (assignment, "Cancelled")):
            with self.subTest(doctype=target.doctype):
                found = [row for row in artifacts if (row["doctype"], row["name"]) == (target.doctype, target.name)]
                self.assertEqual(len(found), 1)
                self.assertTrue(found[0]["system_effect"])
                self.assertEqual(frappe.db.get_value(target.doctype, target.name, "status"), status)
                frappe.db.set_value(target.doctype, target.name, "status", "Pending" if target is issue else "Draft", update_modified=False)
                with self.assertRaises(frappe.ValidationError) as caught:
                    guard.cancel(pe.doctype, pe.name, request_id=key, doc=payload)
                self.assertEqual(caught.exception.purchase_error_id, "replay_evidence_changed")

    def test_native_resource_cancel_cannot_strip_procurement_identity_or_change_business_facts(self):
        from frappe.api import v1, v2
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        frappe.db.set_value("Item", self.item, "is_stock_item", 0)
        frappe.clear_document_cache("Item", self.item)
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        pe = get_payment_entry(po.doctype, po.name, bank_account="Cash - QAB", party_amount=10)
        pe.insert()
        pe.submit()
        self.commit_fixture()
        self.remember_effects()
        for api in (v1, v2):
            for source, change in ((pi, "links"), (pi, "amount"), (pe, "references"), (pe, "party_type")):
                with self.subTest(api=api.__name__, source=source.doctype, change=change):
                    payload = frappe.get_doc(source.doctype, source.name).as_dict(convert_dates_to_str=True)
                    payload["docstatus"] = 2
                    if change == "links":
                        for row in payload["items"]:
                            for field in ("purchase_order", "po_detail", "purchase_receipt", "pr_detail"):
                                row[field] = None
                    elif change == "amount":
                        payload["items"][0]["qty"] += 1
                    elif change == "references":
                        payload["references"] = []
                    else:
                        payload["party_type"] = "Customer"
                        payload["party"] = "QA Purchase Payments Customer"
                    try:
                        if api is v1:
                            with patch.object(v1, "get_request_form_data", return_value=payload), self.assertRaises(frappe.ValidationError) as caught:
                                api.update_doc(source.doctype, source.name)
                        else:
                            with patch.object(frappe.local, "form_dict", frappe._dict(payload)), self.assertRaises(frappe.ValidationError) as caught:
                                api.update_doc(source.doctype, source.name)
                        self.assertEqual(caught.exception.purchase_error_id, "cancellation_business_facts_changed")
                    finally:
                        frappe.db.rollback()
                    self.assert_effects_unchanged()

    def test_native_resource_docstatus_only_cancellation_keeps_finance_atomic(self):
        from datetime import timedelta
        from frappe.api import v1, v2
        from deeplinkerp_branding.services import purchase_consistency as guard
        frappe.db.set_value("Item", self.item, "is_stock_item", 0)
        frappe.clear_document_cache("Item", self.item)
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name)
        pi.set_posting_time = 1
        pi.posting_time = "00:11:22.123456"
        pi.remarks = "0:11:22.123456"
        pi.insert()
        pi.submit()
        self.commit_fixture()
        self.assertEqual(pi.meta.get_field("posting_time").fieldtype, "Time")
        self.assertNotEqual(pi.meta.get_field("remarks").fieldtype, "Time")
        self.assertEqual(frappe.db.get_value(pi.doctype, pi.name, "posting_time"),
            timedelta(minutes=11, seconds=22, microseconds=123456))
        self.remember_effects()
        for api in (v1, v2):
            for change, allowed in (({}, True), ({"posting_time": "00:11:22.123456"}, True),
                    ({"posting_time": "0:11:22.123456"}, True), ({"posting_time": "00:11:22.123457"}, False),
                    ({"posting_time": "not-a-time"}, False), ({"remarks": "00:11:22.123456"}, False)):
                with self.subTest(api=api.__name__, change=change):
                    payload = {"docstatus": 2, **change}
                    def cancel():
                        if api is v1:
                            with patch.object(v1, "get_request_form_data", return_value=payload):
                                return api.update_doc(pi.doctype, pi.name)
                        with patch.object(frappe.local, "form_dict", frappe._dict(payload)):
                            return api.update_doc(pi.doctype, pi.name)
                    try:
                        if allowed:
                            cancel()
                            guard.check_registered()
                            self.assertEqual(frappe.db.get_value(pi.doctype, pi.name, "docstatus"), 2)
                            issue = frappe.get_doc("China Voucher Sync Issue", {"issue_key": f"Cancellation|{pi.doctype}|{pi.name}"})
                            self.assertEqual(issue.status, "Resolved")
                            self.assertTrue(issue.cancellation_voucher)
                        else:
                            with self.assertRaises(frappe.ValidationError) as caught:
                                cancel()
                            self.assertEqual(caught.exception.purchase_error_id, "cancellation_business_facts_changed")
                            self.assert_effects_unchanged()
                    finally:
                        frappe.db.rollback()

    def test_zero_movement_native_cancellation_cannot_hide_a_pending_finance_issue(self):
        from frappe.api import v1
        frappe.db.set_value("Item", self.item, "is_stock_item", 0)
        frappe.clear_document_cache("Item", self.item)
        po = self.order(qty=2, rate=0)
        pr = make_purchase_receipt(po.name).insert()
        pr.submit()
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        self.commit_fixture()
        self.remember_effects()
        with patch.object(v1, "get_request_form_data", return_value={"docstatus": 2}), self.assertRaises(frappe.ValidationError) as caught:
            v1.update_doc(pi.doctype, pi.name)
        self.assertEqual(caught.exception.purchase_error_id, "finance_zero_movement_cancellation_unsupported")
        self.assertFalse(frappe.db.exists("China Voucher Sync Issue", {"issue_key": f"Cancellation|{pi.doctype}|{pi.name}"}))
        self.assert_effects_unchanged()

    def test_native_payment_invoice_cancellation_then_receipt_repost_refuses_atomically(self):
        """Synchronous PE/PI cancellation; stock PR cancellation is not fake success."""
        from collections import defaultdict
        from decimal import Decimal
        from deeplinkerp_branding.services import purchase_document_actions as actions
        from deeplinkerp_branding.services import purchase_consistency as guard
        from frappe.client import cancel

        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        paid = actions.record_payment(source_doctype="Purchase Receipt", source_name=pr.name,
            purchase_invoice=pi.name, amount_to_pay=10, bank_account="Cash - QAB", request_id=str(uuid.uuid4()))
        self.assertFalse(paid.get("failed"), paid)
        pe = frappe.get_doc("Payment Entry", paid["document"]["name"])
        full_balance = pi.grand_total
        self.commit_fixture()
        self.remember_effects()

        # Native PR.on_cancel refuses a submitted PI (and its real payment chain)
        # before stock/GL/repost. No automatic downstream cancellation is added.
        pr = frappe.get_doc("Purchase Receipt", pr.name)
        with self.assertRaisesRegex(frappe.ValidationError, "Purchase Invoice.*already submitted"):
            pr.cancel()
        self.assert_effects_unchanged()
        pr = frappe.get_doc("Purchase Receipt", pr.name)

        def amounts(snapshot):
            result = defaultdict(lambda: [Decimal(0)] * 4)
            for row in snapshot.entries:
                key = tuple(row.get(field) or "" for field in ("account", "account_currency", "party_type",
                    "party", "cost_center", "project", "finance_book", "against_voucher_type", "against_voucher")) + (
                    json.dumps(json.loads(row.get("dimensions_json") or "{}"), sort_keys=True),)
                for index, field in enumerate(("debit", "credit", "debit_in_account_currency", "credit_in_account_currency")):
                    result[key][index] += Decimal(str(row.get(field) or 0))
            return dict(result)

        for source in (pe, pi):
            with self.subTest(doctype=source.doctype):
                posting_name = frappe.db.get_value("China Accounting Voucher",
                    {"source_key": f"Posting|{source.doctype}|{source.name}", "docstatus": 1}, "name")
                self.assertTrue(posting_name, "Real posting must exist before cancellation")
                before = frappe.get_doc("China Accounting Voucher", posting_name)
                before_amounts = amounts(before)
                self.assertTrue(any(any(values) for values in before_amounts.values()))
                cancel(source.doctype, source.name)
                guard.check_registered()
                # Evaluate the final controller/finance hook chain without committing QA business rows.
                frappe.db.before_commit.run()
                self.assertEqual(frappe.db.get_value(source.doctype, source.name, "docstatus"), 2)
                reversal_name = frappe.db.get_value("China Accounting Voucher",
                    {"source_key": f"Cancellation|{source.doctype}|{source.name}"}, "name")
                reversal = frappe.get_doc("China Accounting Voucher", reversal_name)
                self.assertEqual(reversal.docstatus, 1)
                self.assertEqual(reversal.reversal_of, posting_name)
                self.assertEqual(amounts(reversal), {key: [values[1], values[0], values[3], values[2]]
                    for key, values in before_amounts.items()})
                before.reload()
                self.assertEqual((before.status, before.reversed_by), ("Reversed", reversal.name))
                self.assertEqual(amounts(before), before_amounts, "Original posted entries are immutable")
                # ERPNext retains BOTH original and reverse GL as is_cancelled=1;
                # the China snapshot must reverse only the original, never their combined net zero.
                self.assertTrue(frappe.db.count("GL Entry",
                    {"voucher_type": source.doctype, "voucher_no": source.name, "is_cancelled": 1}))
                self.assertFalse(frappe.db.count("GL Entry",
                    {"voucher_type": source.doctype, "voucher_no": source.name, "is_cancelled": 0}))
                if source.doctype == "Payment Entry":
                    self.assertEqual(frappe.db.get_value("Purchase Invoice", pi.name, "outstanding_amount"), full_balance)
                    self.assertFalse(frappe.db.count("Payment Ledger Entry",
                        {"voucher_type": "Payment Entry", "voucher_no": pe.name, "delinked": 0}))
                elif source.doctype == "Purchase Invoice":
                    self.assertFalse(frappe.db.count("Payment Ledger Entry",
                        {"voucher_type": "Purchase Invoice", "voucher_no": pi.name, "delinked": 0}))
                    self.assertEqual(frappe.db.get_value("Bin",
                        {"item_code": self.item, "warehouse": "Stores - QAB"}, "actual_qty"), 2)
        with self.assertRaises(frappe.ValidationError) as caught:
            cancel(pr.doctype, pr.name)
        self.assertEqual(caught.exception.purchase_error_id, "native_stock_repost_incomplete")
        self.assertEqual(frappe.db.get_value("Purchase Receipt", pr.name, "docstatus"), 1)
        self.assertEqual(frappe.db.get_value("Purchase Invoice", pi.name, "docstatus"), 1)
        self.assertEqual(frappe.db.get_value("Payment Entry", pe.name, "docstatus"), 1)
        self.assert_effects_unchanged()  # the refused final action restores this entire uncommitted transaction

    def test_native_form_stock_corruption_rolls_back_complete_transaction(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        po = self.order()
        pr = make_purchase_receipt(po.name).insert()
        pr.submit()
        frappe.db.set_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "actual_qty", 999)
        with self.assertRaises(frappe.ValidationError):
            frappe.db.before_commit.run()
        self.assertFalse(frappe.db.exists("Purchase Receipt", pr.name))
        self.assertFalse(frappe.db.exists("Purchase Order", po.name))

    def test_native_rejected_receipt_uses_total_received_quantity(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        po = self.order()
        pr = make_purchase_receipt(po.name)
        pr.items[0].qty = 3
        pr.items[0].received_qty = 5
        pr.items[0].rejected_qty = 2
        rejected = frappe.get_doc({"doctype": "Warehouse", "warehouse_name": "QA-ATOMIC-REJECT-" + uuid.uuid4().hex[:10],
            "company": COMPANY}).insert()
        pr.items[0].rejected_warehouse = rejected.name
        pr.insert()
        pr.submit()
        guard.check_registered()
        self.assertEqual(frappe.db.get_value("Purchase Order Item", po.items[0].name, "received_qty"), 5)
        returned = make_purchase_return(pr.name).insert()
        returned.submit()
        guard.check_registered()
        self.assertEqual(frappe.db.get_value("Purchase Order Item", po.items[0].name, "received_qty"), 0)

    def test_real_native_sales_order_row_dependency_fails_closed(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        so = frappe.get_doc({"doctype": "Sales Order", "company": COMPANY,
            "customer": "QA Purchase Payments Customer", "delivery_date": add_days(nowdate(), 1),
            "items": [{"item_code": self.item, "qty": 2, "rate": 20, "warehouse": "Stores - QAB", "delivery_date": add_days(nowdate(), 1)}]}).insert()
        so.submit()
        po = frappe.get_doc({"doctype": "Purchase Order", "company": COMPANY,
            "supplier": "QA Test Supplier", "schedule_date": add_days(nowdate(), 1),
            "items": [{"item_code": self.item, "qty": 2, "rate": 12.345, "sales_order": so.name,
                "sales_order_item": so.items[0].name, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]})
        with self.assertRaises(frappe.ValidationError) as caught:
            po.insert(set_name="QA-ATOMIC-PO-" + uuid.uuid4().hex[:16])
        self.assertEqual(caught.exception.purchase_error_id, "sales_dependency_unsupported")
        self.assertFalse(frappe.db.exists("Purchase Order", po.name))

    def internal_sales_parties(self):
        """Real, same-currency cross-company parties shared by native Sales fixtures."""
        selling_company = "Yuewei"
        self.assertEqual(frappe.db.get_value("Company", selling_company, "default_currency"), "CNY")
        frappe.db.set_value("Item", self.item, "is_stock_item", 0)
        customer = frappe.get_doc({"doctype": "Customer", "customer_name": "QA-ATOMIC-CUSTOMER-" + uuid.uuid4().hex[:10],
            "customer_type": "Company", "customer_group": "Government", "territory": "Rest Of The World",
            "is_internal_customer": 1, "represents_company": COMPANY, "companies": [{"company": selling_company}]}).insert()
        supplier = frappe.get_doc({"doctype": "Supplier", "supplier_name": "QA-ATOMIC-SUPPLIER-" + uuid.uuid4().hex[:10],
            "supplier_group": "All Supplier Groups", "supplier_type": "Company", "is_internal_supplier": 1,
            "represents_company": selling_company, "companies": [{"company": COMPANY}]}).insert()
        price_list = frappe.get_doc({"doctype": "Price List", "price_list_name": "QA-ATOMIC-PRICE-" + uuid.uuid4().hex[:10],
            "enabled": 1, "buying": 1, "selling": 1, "currency": "CNY"}).insert()
        return selling_company, customer, supplier, price_list

    def test_real_native_delivery_note_receipt_dependency_fails_closed(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from erpnext.stock.doctype.delivery_note.delivery_note import make_inter_company_purchase_receipt
        selling_company, customer, supplier, price_list = self.internal_sales_parties()
        dn = frappe.get_doc({"doctype": "Delivery Note", "company": selling_company, "customer": customer.name,
            "currency": "CNY", "conversion_rate": 1, "selling_price_list": price_list.name,
            "items": [{"item_code": self.item, "qty": 2, "rate": 20}]}).insert(
                set_name="QA-ATOMIC-DN-" + uuid.uuid4().hex[:16])
        dn.submit()
        self.commit_fixture("Delivery Note", "QA-ATOMIC-DN-")
        original = guard.artifact_evidence(frappe.get_doc("Delivery Note", dn.name))
        counts = {doctype: frappe.db.count(doctype) for doctype in self.types}
        self.assertEqual(frappe.db.get_value("Delivery Note Item", dn.items[0].name, "received_qty"), 0)
        pr = make_inter_company_purchase_receipt(dn.name)
        self.assertEqual(pr.inter_company_reference, dn.name)
        self.assertEqual(pr.items[0].delivery_note_item, dn.items[0].name)
        self.assertFalse(any(row.get(field) for row in pr.items for field in (
            "purchase_order", "purchase_order_item", "material_request", "material_request_item", "purchase_invoice", "purchase_invoice_item")))
        try:
            pr.insert()
            pr.submit()
        except frappe.ValidationError as caught:
            self.assertEqual(caught.purchase_error_id, "sales_dependency_unsupported")
        else:
            self.assertEqual(frappe.db.get_value("Purchase Receipt", pr.name, "docstatus"), 1)
            self.assertEqual(frappe.db.get_value("Delivery Note Item", dn.items[0].name, "received_qty"), 2)
            self.fail("Native standalone DN -> PR submitted and changed DN received_qty without the Sales dependency guard")
        self.assertFalse(frappe.db.exists("Purchase Receipt", pr.name))
        self.assertEqual(guard.artifact_evidence(frappe.get_doc("Delivery Note", dn.name)), original)
        self.assertEqual({doctype: frappe.db.count(doctype) for doctype in self.types}, counts)

    def test_real_native_sales_invoice_with_purchase_order_dependency_fails_closed(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from erpnext.accounts.doctype.sales_invoice.sales_invoice import make_inter_company_purchase_invoice
        selling_company, customer, supplier, price_list = self.internal_sales_parties()
        po = self.order(supplier=supplier.name, qty=2, rate=20)
        si = frappe.get_doc({"doctype": "Sales Invoice", "company": selling_company, "customer": customer.name,
            "currency": "CNY", "conversion_rate": 1, "selling_price_list": price_list.name, "update_stock": 0,
            "items": [{"item_code": self.item, "qty": 2, "rate": 20}]}).insert(
                set_name="QA-ATOMIC-SI-" + uuid.uuid4().hex[:16])
        si.submit()
        self.commit_fixture()
        self.remember_effects()
        counts = {doctype: frappe.db.count(doctype) for doctype in self.types}
        pi = make_inter_company_purchase_invoice(si.name)
        self.assertEqual(pi.inter_company_invoice_reference, si.name)
        self.assertEqual(pi.items[0].sales_invoice_item, si.items[0].name)
        self.assertFalse(si.items[0].get("sales_order") or si.items[0].get("dn_detail"))
        pi.items[0].purchase_order, pi.items[0].po_detail = po.name, po.items[0].name
        try:
            pi.insert()
        except frappe.ValidationError as caught:
            self.assertEqual(caught.purchase_error_id, "sales_dependency_unsupported")
        else:
            self.assertEqual(frappe.db.get_value("Purchase Invoice", pi.name, "docstatus"), 0)
            self.assertTrue(any(row["module"] == "Sales/procurement review" and row["result"] == "N/A"
                for row in guard.invalidate_reviews(frappe.get_doc("Purchase Invoice", pi.name))))
            self.fail("Native SI -> PI with a valid PO source saved and falsely reported Sales N/A")
        self.assertFalse(frappe.db.exists("Purchase Invoice", pi.name))
        self.assert_effects_unchanged()
        self.assertEqual({doctype: frappe.db.count(doctype) for doctype in self.types}, counts)

    def test_real_native_parent_only_intercompany_order_cancellation_fails_closed(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from erpnext.selling.doctype.sales_order.sales_order import make_inter_company_purchase_order
        selling_company, customer, supplier, price_list = self.internal_sales_parties()
        # Create a synthetic already-existing unsupported association. Only our
        # app guard is paused during setup; all native mapping/validation/effects run.
        with patch.object(guard, "check_sales_dependencies"):
            so = frappe.get_doc({"doctype": "Sales Order", "company": selling_company, "customer": customer.name,
                "currency": "CNY", "conversion_rate": 1, "selling_price_list": price_list.name,
                "delivery_date": add_days(nowdate(), 1), "items": [{"item_code": self.item, "qty": 2, "rate": 20,
                    "delivery_date": add_days(nowdate(), 1)}]}).insert(set_name="QA-ATOMIC-SO-" + uuid.uuid4().hex[:16])
            po = make_inter_company_purchase_order(so.name)
            self.assertEqual(po.inter_company_order_reference, so.name)
            po.schedule_date = add_days(nowdate(), 1)
            for row in po.items:
                row.schedule_date = po.schedule_date
                row.sales_order = row.sales_order_item = row.sales_order_packed_item = None
            po.insert(set_name="QA-ATOMIC-PO-" + uuid.uuid4().hex[:16])
            so.inter_company_order_reference = po.name
            so.save()
            so.submit()
            po.reload().submit()
            self.commit_fixture()
        self.remember_effects()
        counts = {doctype: frappe.db.count(doctype) for doctype in self.types}
        try:
            po.cancel()
        except frappe.ValidationError as caught:
            self.assertEqual(caught.purchase_error_id, "sales_dependency_unsupported")
        else:
            self.assertEqual(frappe.db.get_value("Purchase Order", po.name, "docstatus"), 2)
            self.assertFalse(frappe.db.get_value("Sales Order", so.name, "inter_company_order_reference"))
            self.assertFalse(frappe.db.get_value("Purchase Order", po.name, "inter_company_order_reference"))
            self.fail("Native parent-only PO cancellation cleared the real Sales Order header without a Sales guard")
        self.assert_effects_unchanged()
        self.assertEqual({doctype: frappe.db.count(doctype) for doctype in self.types}, counts)

    def test_real_native_parent_only_intercompany_invoice_cancellation_fails_closed(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from erpnext.accounts.doctype.sales_invoice.sales_invoice import make_inter_company_purchase_invoice
        selling_company, customer, supplier, price_list = self.internal_sales_parties()
        with patch.object(guard, "check_sales_dependencies"):
            po = self.order(supplier=supplier.name, qty=2, rate=20)
            si = frappe.get_doc({"doctype": "Sales Invoice", "company": selling_company, "customer": customer.name,
                "currency": "CNY", "conversion_rate": 1, "selling_price_list": price_list.name, "update_stock": 0,
                "items": [{"item_code": self.item, "qty": 2, "rate": 20}]}).insert(
                    set_name="QA-ATOMIC-SI-" + uuid.uuid4().hex[:16])
            pi = make_inter_company_purchase_invoice(si.name)
            self.assertEqual(pi.inter_company_invoice_reference, si.name)
            pi.items[0].purchase_order, pi.items[0].po_detail = po.name, po.items[0].name
            pi.items[0].sales_invoice_item = None
            pi.insert()
            si.inter_company_invoice_reference = pi.name
            si.save()
            si.submit()
            pi.submit()
            self.commit_fixture()
        self.remember_effects()
        counts = {doctype: frappe.db.count(doctype) for doctype in self.types}
        try:
            pi.cancel()
        except frappe.ValidationError as caught:
            self.assertEqual(caught.purchase_error_id, "sales_dependency_unsupported")
        else:
            self.assertEqual(frappe.db.get_value("Purchase Invoice", pi.name, "docstatus"), 2)
            self.assertFalse(frappe.db.get_value("Sales Invoice", si.name, "inter_company_invoice_reference"))
            self.assertFalse(frappe.db.get_value("Purchase Invoice", pi.name, "inter_company_invoice_reference"))
            self.fail("Native parent-only PI cancellation cleared the real Sales Invoice header without a Sales guard")
        self.assert_effects_unchanged()
        self.assertEqual({doctype: frappe.db.count(doctype) for doctype in self.types}, counts)

    def test_native_savedocs_queue_is_stopped_before_dispatch(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        pe = get_payment_entry(pi.doctype, pi.name, bank_account="Cash - QAB", bank_amount=10).insert()
        self.commit_fixture()
        for document in (po, pi, pe):
            payload = document.as_dict()
            if document.doctype == "Purchase Invoice":
                for row in payload["items"]:
                    for field in ("purchase_order", "po_detail", "purchase_receipt", "pr_detail"):
                        row[field] = None
            elif document.doctype == "Payment Entry":
                payload["references"] = []
            with patch.object(frappe.get_meta(document.doctype), "queue_in_background", True), \
                    patch("frappe.utils.scheduler.is_scheduler_inactive", return_value=False), \
                    patch("frappe.desk.form.save.queue_submission", side_effect=AssertionError("No queue dispatch")):
                for key in (None, str(uuid.uuid4())):
                    with self.subTest(doctype=document.doctype, request_id=key), self.assertRaises(frappe.ValidationError) as caught:
                        guard.savedocs(json.dumps(payload, default=str), "Submit", request_id=key)
                    self.assertEqual(caught.exception.purchase_error_id, "native_background_submission_unsupported")
        self.assertEqual(frappe.db.get_value("Purchase Order", po.name, "docstatus"), 1)
        self.assertEqual(self.before["Submission Queue"], frappe.db.count("Submission Queue"))

    def unrelated_financial_documents(self):
        frappe.db.set_value("Item", self.item, "is_stock_item", 0)
        frappe.clear_document_cache("Item", self.item)
        pi = frappe.get_doc({"doctype": "Purchase Invoice", "company": COMPANY,
            "supplier": "QA Test Supplier", "currency": "CNY", "conversion_rate": 1, "update_stock": 0,
            "items": [{"item_code": self.item, "qty": 2, "rate": 20}]}).insert(
                set_name="QA-ATOMIC-PI-" + uuid.uuid4().hex[:16])
        pe = frappe.get_doc({"doctype": "Payment Entry", "company": COMPANY, "payment_type": "Receive",
            "party_type": "Customer", "party": "QA Purchase Payments Customer",
            "paid_from": frappe.get_cached_value("Company", COMPANY, "default_receivable_account"),
            "paid_to": "Cash - QAB", "paid_amount": 10, "received_amount": 10,
            "source_exchange_rate": 1, "target_exchange_rate": 1,
            "reference_no": "QA-UNRELATED", "reference_date": nowdate()}).insert()
        return pi, pe

    def test_unrelated_native_savedocs_preserves_real_submission_queue_with_or_without_key(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from frappe.core.doctype.submission_queue.submission_queue import SubmissionQueue
        pi, pe = self.unrelated_financial_documents()
        self.commit_fixture("Purchase Invoice", "QA-ATOMIC-PI-")
        counts = {doctype: frappe.db.count(doctype) for doctype in self.types}
        for document in (pi, pe):
            self.assertFalse(guard.is_procurement(document))
            for key in (None, str(uuid.uuid4())):
                with self.subTest(doctype=document.doctype, request_id=key):
                    try:
                        with patch.object(frappe.get_meta(document.doctype), "queue_in_background", True), \
                                patch("frappe.utils.scheduler.is_scheduler_inactive", return_value=False), \
                                patch("frappe.desk.form.save.is_scheduler_inactive", return_value=False), \
                                patch.object(SubmissionQueue, "queue_action"):  # isolate external dispatch, not native queue insertion
                            guard.savedocs(json.dumps(document.as_dict(), default=str), "Submit", request_id=key)
                        queues = frappe.get_all("Submission Queue", filters={"ref_doctype": document.doctype,
                            "ref_docname": document.name}, fields=["name", "status"])
                        self.assertEqual(len(queues), 1)
                        self.assertEqual(queues[0].status, "Queued")
                        self.assertEqual(frappe.db.get_value(document.doctype, document.name, "docstatus"), 0)
                        self.assertEqual(frappe.db.count("Integration Request"), counts["Integration Request"])
                    finally:
                        frappe.db.rollback()
        self.assertEqual(counts, {doctype: frappe.db.count(doctype) for doctype in self.types})

    def test_unrelated_native_savedocs_and_cancel_keep_native_results_and_acl(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from frappe import permissions
        native_permission = permissions.has_permission
        def denied_permission(doctype, ptype="read", *args, **kwargs):
            if doctype in ("Purchase Invoice", "Payment Entry") and ptype in ("write", "cancel"):
                return False
            return native_permission(doctype, ptype, *args, **kwargs)
        pi, pe = self.unrelated_financial_documents()
        self.commit_fixture("Purchase Invoice", "QA-ATOMIC-PI-")
        for document in (pi, pe):
            for key in (None, str(uuid.uuid4())):
                with self.subTest(doctype=document.doctype, action="Save", request_id=key):
                    payload = frappe.get_doc(document.doctype, document.name).as_dict()
                    payload["remarks"] = "QA unrelated native save " + str(key)
                    before = frappe.db.count("Integration Request")
                    with patch.object(permissions, "has_permission", side_effect=denied_permission), self.assertRaises(frappe.PermissionError):
                        guard.savedocs(json.dumps(payload, default=str), "Save", request_id=key)
                    self.assertEqual(frappe.db.count("Integration Request"), before)
                    guard.savedocs(json.dumps(payload, default=str), "Save", request_id=key)
                    self.assertEqual(frappe.response.docs[-1]["name"], document.name)
                    self.assertEqual(frappe.db.count("Integration Request"), before)
                    # Distinct real requests each commit their own successful
                    # native save. A later denied request's full rollback must
                    # not erase a previous uncommitted fixture mutation.
                    self.commit_fixture("Purchase Invoice", "QA-ATOMIC-PI-")
            document.reload().submit()
            self.commit_fixture("Purchase Invoice", "QA-ATOMIC-PI-")
        self.commit_fixture("Purchase Invoice", "QA-ATOMIC-PI-")
        for document in (pi, pe):
            for key in (None, str(uuid.uuid4())):
                with self.subTest(doctype=document.doctype, action="Cancel", request_id=key):
                    try:
                        before = frappe.db.count("Integration Request")
                        with patch.object(permissions, "has_permission", side_effect=denied_permission), self.assertRaises(frappe.PermissionError):
                            guard.cancel(document.doctype, document.name, request_id=key)
                        self.assertEqual(frappe.db.get_value(document.doctype, document.name, "docstatus"), 1)
                        self.assertEqual(frappe.db.count("Integration Request"), before)
                        guard.cancel(document.doctype, document.name, request_id=key)  # native cancel needs no supplied snapshot
                        self.assertEqual(frappe.response.docs[-1]["docstatus"], 2)
                        self.assertEqual(frappe.db.get_value(document.doctype, document.name, "docstatus"), 2)
                        self.assertEqual(frappe.db.count("Integration Request"), before)
                    finally:
                        frappe.db.rollback()

    def test_auto_invoice_rejects_real_purchase_user_without_invoice_create(self):
        po = self.order()
        user = "qa-atomic-" + uuid.uuid4().hex[:12] + "@example.invalid"
        frappe.get_doc({"doctype": "User", "email": user, "first_name": "Atomic QA", "send_welcome_email": 0,
            "roles": [{"role": "Purchase User"}]}).insert()
        self.commit_fixture()
        frappe.db.set_value("China Finance Settings", COMPANY, "auto_submit_purchase_invoice", 1)
        pr = make_purchase_receipt(po.name).insert()
        frappe.set_user(user)
        self.assertTrue(frappe.has_permission("Purchase Receipt", "submit", doc=pr))
        self.assertFalse(frappe.has_permission("Purchase Invoice", "create"))
        try:
            with self.assertRaises(frappe.PermissionError):
                pr.submit()
        finally:
            frappe.set_user("Administrator")
            frappe.clear_cache(user=user)
        self.assertFalse(frappe.db.exists("Purchase Receipt", pr.name))
        self.assertFalse(frappe.db.exists("Purchase Invoice Item", {"purchase_receipt": pr.name}))

    def test_auto_invoice_noop_cannot_acknowledge_receipt(self):
        from china_finance.services import auto_invoice
        po = self.order()
        self.commit_fixture()
        frappe.db.set_value("China Finance Settings", COMPANY, "auto_submit_purchase_invoice", 1)
        pr = make_purchase_receipt(po.name).insert()
        with patch.object(auto_invoice, "_create_and_submit", return_value=None):
            with self.assertRaisesRegex(frappe.ValidationError, "完整覆盖"):
                pr.submit()
        self.assertFalse(frappe.db.exists("Purchase Receipt", pr.name))

    def test_new_auto_invoice_requires_actual_document_submit_permission(self):
        from frappe import permissions
        po = self.order()
        user = "qa-atomic-" + uuid.uuid4().hex[:12] + "@example.invalid"
        frappe.get_doc({"doctype": "User", "email": user, "first_name": "Atomic QA", "send_welcome_email": 0,
            "roles": [{"role": "Purchase User"}, {"role": "Accounts User"}]}).insert()
        self.commit_fixture()
        frappe.db.set_value("China Finance Settings", COMPANY, "auto_submit_purchase_invoice", 1)
        pr = make_purchase_receipt(po.name).insert()
        original = permissions.has_controller_permissions
        def deny_mapped_invoice_submit(doc, ptype, **kwargs):
            if doc.doctype == "Purchase Invoice" and ptype == "submit" and any(
                    row.purchase_receipt == pr.name for row in doc.items):
                return False
            return original(doc, ptype, **kwargs)
        frappe.set_user(user)
        try:
            self.assertTrue(frappe.has_permission("Purchase Invoice", "create"))
            self.assertTrue(frappe.has_permission("Purchase Invoice", "submit"))
            with patch.object(permissions, "has_controller_permissions", side_effect=deny_mapped_invoice_submit):
                with self.assertRaises(frappe.PermissionError) as caught:
                    pr.submit()
            self.assertEqual(caught.exception.purchase_error_id, "auto_pi_permission_missing")
        finally:
            frappe.set_user("Administrator")
            frappe.clear_cache(user=user)
        self.assertFalse(frappe.db.exists("Purchase Receipt", pr.name))
        self.assertFalse(frappe.db.exists("Purchase Invoice Item", {"purchase_receipt": pr.name}))

    def test_partial_existing_auto_invoice_draft_is_preserved_on_failure(self):
        from frappe.client import save
        po = self.order()
        pr = make_purchase_receipt(po.name).insert()
        pi = frappe.get_doc({"doctype": "Purchase Invoice", "company": COMPANY, "supplier": pr.supplier,
            "currency": pr.currency, "conversion_rate": pr.conversion_rate, "bill_no": "QA-ATOMIC-MANUAL",
            "bill_date": add_days(nowdate(), -2), "items": [{"item_code": self.item, "qty": 1, "rate": po.items[0].rate,
                "purchase_order": po.name, "po_detail": po.items[0].name, "purchase_receipt": pr.name,
                "pr_detail": pr.items[0].name, "warehouse": "Stores - QAB"}]}).insert()
        self.commit_fixture()
        pi.remarks = "QA-ATOMIC-native draft PR reference"
        save(json.dumps(pi.as_dict(), default=str))  # real ordinary RPC draft succeeds
        self.commit_fixture()  # that successful request is committed independently
        with self.assertRaises(frappe.ValidationError):
            pi.reload().submit()  # no draft-source submit capability is granted
        self.assertEqual(frappe.db.get_value(pi.doctype, pi.name, "docstatus"), 0)
        self.assertEqual(frappe.db.get_value(pr.doctype, pr.name, "docstatus"), 0)
        before = pi.as_dict()
        frappe.db.set_value("China Finance Settings", COMPANY, "auto_submit_purchase_invoice", 1)
        with self.assertRaisesRegex(frappe.ValidationError, "完整覆盖"):
            pr.submit()
        pi.reload()
        self.assertEqual(pi.docstatus, 0)
        self.assertEqual((pi.items[0].qty, pi.items[0].rate, pi.bill_no, str(pi.bill_date), pi.taxes),
            (before["items"][0].qty, before["items"][0].rate, before["bill_no"], str(before["bill_date"]), []))
        self.assertEqual(frappe.db.get_value("Purchase Receipt", pr.name, "docstatus"), 0)
        self.assertEqual(frappe.db.count("Purchase Invoice Item", {"purchase_receipt": pr.name}), 1)

    def test_auto_invoice_success_records_all_native_and_finance_effects(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        from deeplinkerp_branding.services import purchase_operation as kernel
        po = self.order()
        self.commit_fixture()
        frappe.db.set_value("China Finance Settings", COMPANY, "auto_submit_purchase_invoice", 1)
        key = str(uuid.uuid4())
        changes = {"items": [{"key": po.items[0].name, "qty": 2, "warehouse": "Stores - QAB"}]}
        target = frappe.db.get_value("Purchase Receipt Item", {"purchase_order": po.name, "docstatus": 0}, "parent")
        args = {"target_name": target, "expected_modified": frappe.db.get_value("Purchase Receipt", target, "modified")} if target else {}
        result = actions.record_receipt(po.name, changes, key, **args)
        self.assertFalse(result.get("failed"), result)
        pr = frappe.get_doc("Purchase Receipt", result["document"]["name"])
        pi = frappe.get_doc("Purchase Invoice", frappe.db.get_value("Purchase Invoice Item", {"purchase_receipt": pr.name}, "parent"))
        self.assertEqual(pi.docstatus, 1)
        self.assertEqual(pi.items[0].qty, 2)
        audit = frappe.get_doc("Integration Request", kernel.identity("Administrator", key))
        evidence = json.loads(audit.output)["artifacts"]
        self.assertTrue(any(row["doctype"] == "Purchase Invoice" and row["name"] == pi.name for row in evidence))
        finance = [row for row in evidence if row["doctype"] == "China Accounting Voucher"]
        self.assertEqual(len(finance), 2)
        self.assertTrue(all(row["system_effect"] for row in finance))
        replay = actions.record_receipt(po.name, changes, key, **args)
        self.assertTrue(replay["reused"])
        frappe.db.set_value("China Accounting Voucher", finance[0]["name"], "source_name", "QA-ATOMIC-TAMPER")
        with self.assertRaises(frappe.ValidationError):
            actions.record_receipt(po.name, changes, key, **args)

    def test_auto_invoice_respects_an_installed_native_workflow(self):
        from frappe.model.workflow import get_workflow_name
        po = self.order()
        self.commit_fixture()
        # Existing status column avoids fixture DDL/implicit commits.
        for state in ("Draft", "Submitted"):
            if not frappe.db.exists("Workflow State", state):
                frappe.get_doc({"doctype": "Workflow State", "workflow_state_name": state}).insert()
        action = "QA-ATOMIC-APPROVE-" + uuid.uuid4().hex[:10]
        frappe.get_doc({"doctype": "Workflow Action Master", "workflow_action_name": action}).insert()
        workflow = frappe.get_doc({"doctype": "Workflow", "workflow_name": "QA-ATOMIC-WORKFLOW-" + uuid.uuid4().hex[:10],
            "document_type": "Purchase Invoice", "is_active": 1, "workflow_state_field": "status", "states": [
                {"state": "Draft", "doc_status": "0", "allow_edit": "Accounts User"},
                {"state": "Submitted", "doc_status": "1", "allow_edit": "Accounts User"}],
            "transitions": [{"state": "Draft", "action": action, "next_state": "Submitted", "allowed": "Accounts User"}]}).insert()
        self.assertEqual(get_workflow_name("Purchase Invoice"), workflow.name)
        frappe.db.set_value("China Finance Settings", COMPANY, "auto_submit_purchase_invoice", 1)
        pr = make_purchase_receipt(po.name).insert()
        try:
            with self.assertRaisesRegex(frappe.ValidationError, "原生工作流"):
                pr.submit()
        finally:
            frappe.clear_cache(doctype="Purchase Invoice")
        self.assertFalse(frappe.db.exists("Purchase Receipt", pr.name))

    def test_replay_rejects_actual_revoked_automatic_invoice_submit_permission(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        po = self.order()
        user = "qa-atomic-" + uuid.uuid4().hex[:12] + "@example.invalid"
        frappe.get_doc({"doctype": "User", "email": user, "first_name": "Atomic QA", "send_welcome_email": 0,
            "roles": [{"role": "Purchase User"}, {"role": "Accounts User"}]}).insert()
        self.commit_fixture()
        frappe.db.set_value("China Finance Settings", COMPANY, "auto_submit_purchase_invoice", 1)
        target = frappe.db.get_value("Purchase Receipt Item", {"purchase_order": po.name, "docstatus": 0}, "parent")
        args = {"target_name": target, "expected_modified": frappe.db.get_value("Purchase Receipt", target, "modified")} if target else {}
        key = str(uuid.uuid4())
        changes = {"items": [{"key": po.items[0].name, "qty": 2, "warehouse": "Stores - QAB"}]}
        frappe.set_user(user)
        try:
            result = actions.record_receipt(po.name, changes, key, **args)
            self.assertFalse(result.get("failed"), result)
            frappe.db.delete("Has Role", {"parent": user, "role": "Accounts User"})
            frappe.clear_cache(user=user)
            frappe.set_user("Administrator"); frappe.set_user(user)
            self.assertTrue(frappe.has_permission("Purchase Invoice", "read"))
            self.assertFalse(frappe.has_permission("Purchase Invoice", "submit"))
            with self.assertRaises(frappe.PermissionError):
                actions.record_receipt(po.name, changes, key, **args)
        finally:
            frappe.set_user("Administrator")
            frappe.clear_cache(user=user)

    def test_native_direct_stock_invoice_and_return_preserve_received_and_billed(self):
        from erpnext.buying.doctype.purchase_order.purchase_order import make_purchase_invoice as from_order
        from erpnext.controllers.sales_and_purchase_return import make_return_doc
        from deeplinkerp_branding.services import purchase_consistency as guard
        po = self.order()
        pi = from_order(po.name)
        pi.update_stock = 1
        pi.items[0].qty = 2
        pi.items[0].warehouse = "Stores - QAB"
        pi.insert(); pi.submit()
        guard.check_registered()
        self.assertEqual(frappe.db.get_value("Purchase Order Item", po.items[0].name, "received_qty"), 2)
        returned = make_return_doc("Purchase Invoice", pi.name)
        returned.update_billed_amount_in_purchase_order = 1
        returned.insert(); returned.submit()
        guard.check_registered()
        self.assertEqual(frappe.db.get_value("Purchase Order Item", po.items[0].name, "received_qty"), 0)
        self.assertEqual(frappe.db.get_value("Purchase Order Item", po.items[0].name, "billed_amt"), 0)

    def test_native_fx_invoice_and_payment_match_both_currency_ledgers(self):
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        from deeplinkerp_branding.services import purchase_consistency as guard
        account = frappe.get_doc({"doctype": "Account", "account_name": "QA-ATOMIC-USD-" + uuid.uuid4().hex[:10],
            "company": COMPANY, "parent_account": "Accounts Payable - QAB", "account_type": "Payable", "account_currency": "USD"}).insert()
        supplier = frappe.get_doc({"doctype": "Supplier", "supplier_name": "QA-ATOMIC-USD-" + uuid.uuid4().hex[:10],
            "supplier_type": "Company", "supplier_group": "All Supplier Groups", "default_currency": "USD",
            "accounts": [{"company": COMPANY, "account": account.name}]}).insert()
        po = self.order(currency="USD", conversion_rate=7.15, supplier=supplier.name)
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name)
        pi.currency, pi.conversion_rate, pi.credit_to = "USD", 7.15, account.name
        pi.items[0].rate = 12.345
        pi.insert(); pi.submit()
        pe = get_payment_entry("Purchase Invoice", pi.name, bank_account="Cash - QAB")
        pe.reference_no, pe.reference_date = "QA-FX", nowdate()
        pe.insert(); pe.submit()
        guard.check_registered()
        self.assertEqual(pe.paid_from_account_currency, "CNY")
        self.assertEqual(pe.paid_to_account_currency, "USD")
        self.assertEqual(frappe.db.get_value("Purchase Invoice", pi.name, "outstanding_amount"), 0)

    def test_existing_no_update_return_then_new_invoice_uses_native_submitted_rows(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from erpnext.accounts.doctype.purchase_invoice.purchase_invoice import make_debit_note
        po = self.order()
        pr = make_purchase_receipt(po.name).insert()
        pr.submit()
        first = make_purchase_invoice(pr.name)
        first.items[0].qty = 2
        first.insert().submit()
        original = first.grand_total
        returned = make_debit_note(first.name)
        returned.items[0].qty = -1
        returned.update_billed_amount_in_purchase_order = 0
        returned.update_billed_amount_in_purchase_receipt = 0
        returned.insert().submit()
        # These flags skip this return's updater, not the return in future sums.
        self.assertEqual(frappe.db.get_value("Purchase Order Item", po.items[0].name, "billed_amt"), original)
        self.assertEqual(frappe.db.get_value("Purchase Receipt Item", pr.items[0].name, "billed_amt"), original)
        self.commit_fixture()
        next_invoice = make_purchase_invoice(pr.name)
        next_invoice.items[0].qty = 1
        next_invoice.insert().submit()
        guard.check_registered()
        expected = first.items[0].amount + returned.items[0].amount + next_invoice.items[0].amount
        self.assertEqual(frappe.db.get_value("Purchase Order Item", po.items[0].name, "billed_amt"), expected)
        self.assertEqual(frappe.db.get_value("Purchase Receipt Item", pr.items[0].name, "billed_amt"), expected)
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        frappe.db.set_single_value("Buying Settings", "set_landed_cost_based_on_purchase_invoice_rate", 1)
        frappe.clear_document_cache("Buying Settings", "Buying Settings")
        try:
            result = self.read_reversal_scope(returned)
            self.assertEqual(result.vouchers, (scope.DocumentIdentity(returned.doctype, returned.name),))
            self.assertIn(scope.DocumentIdentity(pr.doctype, pr.name), result.sources)
        finally:
            frappe.db.set_single_value("Buying Settings", "set_landed_cost_based_on_purchase_invoice_rate", 0)
            frappe.clear_document_cache("Buying Settings", "Buying Settings")

    def test_native_invoice_rate_calculator_on_copy_cannot_write_or_repost(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from erpnext.stock.doctype.purchase_receipt.purchase_receipt import PurchaseReceipt
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()  # normal native rate path; no queued cost adjustment
        self.commit_fixture()
        frappe.db.set_single_value("Buying Settings", "set_landed_cost_based_on_purchase_invoice_rate", 1)
        frappe.clear_document_cache("Buying Settings", "Buying Settings")
        self.assertEqual(frappe.db.get_value("Purchase Receipt", pr.name, "per_billed"), 100)
        sql = frappe.db.sql
        def read_only(query, *args, **kwargs):
            self.assertNotIn(str(query).strip().split()[0].lower(), {"insert", "update", "delete", "replace", "alter"})
            return sql(query, *args, **kwargs)
        with patch.object(frappe.db, "sql", side_effect=read_only), \
                patch.object(frappe.db, "set_value", side_effect=AssertionError("No validation DB write")), \
                patch.object(frappe, "enqueue", side_effect=AssertionError("No validation queue")), \
                patch.object(PurchaseReceipt, "update_valuation_rate", side_effect=AssertionError("No validation valuation")), \
                patch.object(PurchaseReceipt, "repost_future_sle_and_gle", side_effect=AssertionError("No validation repost")):
            guard.check_billing(frappe.get_doc("Purchase Invoice", pi.name))
        frappe.db.rollback()
        frappe.clear_document_cache("Buying Settings", "Buying Settings")

    def future_receipts(self):
        po = self.order(transaction_date=add_days(nowdate(), -3))
        receipts = []
        for days in (-2, -1):
            pr = make_purchase_receipt(po.name)
            pr.items[0].qty = 2
            pr.set_posting_time = 1
            pr.posting_date, pr.posting_time = add_days(nowdate(), days), "12:00:00"
            pr.insert().submit()
            receipts.append(pr)
        self.commit_fixture()
        self.remember_effects()
        return po, *receipts

    def scope_item(self):
        code = "QA-ATOMIC-" + uuid.uuid4().hex[:10]
        return frappe.get_doc({"doctype": "Item", "item_code": code, "item_name": code,
            "item_group": "All Item Groups", "stock_uom": "Nos", "is_stock_item": 1}).insert().name

    def scope_stock_entry(self, purpose, date, rows, **headers):
        return frappe.get_doc({"doctype": "Stock Entry", "company": COMPANY, "stock_entry_type": purpose,
            "set_posting_time": 1, "posting_date": date, "posting_time": "12:00:00", "items": rows, **headers}).insert().submit()

    def read_reversal_scope(self, doc, *, cancellation=True):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        sql = frappe.db.sql
        def read_only(query, *args, **kwargs):
            text = str(query).strip().lower()
            self.assertNotIn(text.split()[0], {"insert", "update", "delete", "replace", "alter"})
            self.assertNotIn("for update", text)
            return sql(query, *args, **kwargs)
        with patch.object(frappe.db, "sql", side_effect=read_only), \
                patch.object(frappe.db, "set_value", side_effect=AssertionError("scope cannot write")), \
                patch.object(frappe.db, "commit", side_effect=AssertionError("scope cannot commit")), \
                patch.object(frappe, "enqueue", side_effect=AssertionError("scope cannot enqueue")):
            return (scope.collect_cancellation_scope if cancellation else scope.collect_document_scope)(doc)

    def test_reversal_scope_native_transfer_repack_new_item_warehouse_full_set_is_read_only(self):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        po = self.order(transaction_date=add_days(nowdate(), -4))
        root = make_purchase_receipt(po.name)
        root.items[0].qty = 2
        root.set_posting_time = 1
        root.posting_date, root.posting_time = add_days(nowdate(), -3), "12:00:00"
        root.insert().submit()
        warehouse2, warehouse3 = "Finished Goods - QAB", "Work In Progress - QAB"
        output, extra = self.scope_item(), self.scope_item()
        transfer = self.scope_stock_entry("Material Transfer", add_days(nowdate(), -2), [
            {"item_code": self.item, "qty": 1, "s_warehouse": "Stores - QAB", "t_warehouse": warehouse2}])
        repack = self.scope_stock_entry("Repack", add_days(nowdate(), -1), [
            {"item_code": self.item, "qty": 1, "s_warehouse": warehouse2},
            {"item_code": output, "qty": 1, "t_warehouse": warehouse2, "is_finished_item": 1,
                "set_basic_rate_manually": 1, "basic_rate": 6},
            {"item_code": extra, "qty": 1, "t_warehouse": warehouse3, "set_basic_rate_manually": 1, "basic_rate": 6}])
        repack_sles = frappe.db.get_values("Stock Ledger Entry", {"voucher_type": repack.doctype, "voucher_no": repack.name},
            ["item_code", "warehouse", "dependant_sle_voucher_detail_no", "actual_qty"], as_dict=True)
        self.assertTrue(any(row.dependant_sle_voucher_detail_no for row in repack_sles), repack_sles)
        future = self.scope_stock_entry("Material Issue", nowdate(), [
            {"item_code": extra, "qty": 1, "s_warehouse": warehouse3}])
        self.commit_fixture()
        self.remember_effects()
        before = {dt: frappe.db.count(dt) for dt in self.types}
        result = self.read_reversal_scope(root)
        self.assertEqual(result.company, COMPANY)
        self.assertEqual(set(result.sources), {scope.DocumentIdentity("Purchase Order", po.name)})
        self.assertEqual(set(result.vouchers), {scope.DocumentIdentity(doc.doctype, doc.name)
            for doc in (root, transfer, repack, future)})
        self.assertEqual(set(result.pairs), {scope.StockPair(self.item, "Stores - QAB"),
            scope.StockPair(self.item, warehouse2), scope.StockPair(output, warehouse2), scope.StockPair(extra, warehouse3)})
        self.assertEqual(before, {dt: frappe.db.count(dt) for dt in self.types})
        self.assert_effects_unchanged()

    def test_reversal_document_first_stock_draft_needs_no_bin_but_freeze_refuses_missing_bin(self):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        draft = frappe.get_doc({"doctype": "Purchase Receipt", "company": COMPANY, "supplier": "QA Test Supplier",
            "posting_date": nowdate(), "posting_time": "12:00:00", "items": [{"item_code": self.item,
                "stock_uom": "Nos", "qty": 1, "warehouse": "Stores - QAB"}]})
        self.assertFalse(frappe.db.exists("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}))
        result = scope.collect_document_scope(draft)
        self.assertEqual(result.pairs, (scope.StockPair(self.item, "Stores - QAB"),))
        po = self.order()
        root = self.receipt(po)
        frappe.db.delete("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"})
        with self.assertRaises(frappe.ValidationError):
            scope.collect_cancellation_scope(root)

    def test_reversal_scope_native_same_item_warehouse_is_bound_to_exact_receipt_detail(self):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        warehouses = ("Stores - QAB", "Work In Progress - QAB")
        po = self.order(items=[{"item_code": self.item, "qty": 1, "rate": 10, "warehouse": warehouse,
            "schedule_date": add_days(nowdate(), 1)} for warehouse in warehouses])
        root = make_purchase_receipt(po.name).insert().submit()
        self.commit_fixture()
        self.remember_effects()
        result = self.read_reversal_scope(root)
        self.assertEqual(set(result.pairs), {scope.StockPair(self.item, warehouse) for warehouse in warehouses})
        entry = frappe.db.get_value("Stock Ledger Entry", {"voucher_type": root.doctype, "voucher_no": root.name,
            "voucher_detail_no": root.items[0].name}, "name")
        self.assertTrue(entry)
        original = frappe.db.get_value("Stock Ledger Entry", entry, "warehouse")
        try:
            other = next(warehouse for warehouse in warehouses if warehouse != original)
            frappe.db.set_value("Stock Ledger Entry", entry, "warehouse", other, update_modified=False)
            with self.assertRaises(frappe.ValidationError): self.read_reversal_scope(root)
        finally:
            frappe.db.set_value("Stock Ledger Entry", entry, "warehouse", original, update_modified=False)
        self.assert_effects_unchanged()

    def test_reversal_scope_native_gl_cartesian_whole_voucher_does_not_seed_extra_stock_frontier(self):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        second, extra = self.scope_item(), self.scope_item()
        warehouse2, warehouse3 = "Finished Goods - QAB", "Work In Progress - QAB"
        po = self.order(transaction_date=add_days(nowdate(), -4), items=[
            {"item_code": code, "qty": 2, "rate": 10, "warehouse": warehouse, "schedule_date": add_days(nowdate(), 1)}
            for code, warehouse in ((self.item, "Stores - QAB"), (second, warehouse2))])
        root = make_purchase_receipt(po.name)
        root.set_posting_time = 1
        root.posting_date, root.posting_time = add_days(nowdate(), -3), "12:00:00"
        root.insert().submit()
        cross = self.scope_stock_entry("Material Receipt", add_days(nowdate(), -2), [
            {"item_code": self.item, "qty": 1, "t_warehouse": warehouse2, "basic_rate": 10},
            {"item_code": extra, "qty": 1, "t_warehouse": warehouse3, "basic_rate": 10}])
        future = self.scope_stock_entry("Material Issue", add_days(nowdate(), -1), [
            {"item_code": extra, "qty": 1, "s_warehouse": warehouse3}])
        self.commit_fixture()
        self.remember_effects()
        result = scope.collect_cancellation_scope(root)
        self.assertEqual(set(result.vouchers), {scope.DocumentIdentity(root.doctype, root.name),
            scope.DocumentIdentity(cross.doctype, cross.name)})
        self.assertNotIn(scope.DocumentIdentity(future.doctype, future.name), result.vouchers)
        self.assertEqual(set(result.pairs), {scope.StockPair(self.item, "Stores - QAB"),
            scope.StockPair(second, warehouse2), scope.StockPair(self.item, warehouse2), scope.StockPair(extra, warehouse3)})
        self.assertNotIn(scope.StockPair(second, "Stores - QAB"), result.pairs)
        self.assert_effects_unchanged()

    def test_reversal_scope_native_billing_receipts_and_cancelled_raw_seed(self):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        po, prior, future = self.future_receipts()
        invoice = frappe.get_doc({"doctype": "Purchase Invoice", "company": COMPANY, "supplier": po.supplier,
            "posting_date": nowdate(), "items": [{"item_code": self.item, "qty": 1, "rate": 10, "stock_uom": "Nos",
                "purchase_order": po.name, "po_detail": po.items[0].name} ]})
        result = scope.collect_document_scope(invoice)
        self.assertEqual(set(result.sources), {scope.DocumentIdentity(doc.doctype, doc.name) for doc in (po, prior, future)})
        frappe.db.set_value("Stock Ledger Entry", {"voucher_type": prior.doctype, "voucher_no": prior.name}, "is_cancelled", 1)
        result = scope.collect_cancellation_scope(prior)
        self.assertEqual(set(result.vouchers), {scope.DocumentIdentity(doc.doctype, doc.name) for doc in (prior, future)})
        self.assertEqual(result.pairs, (scope.StockPair(self.item, "Stores - QAB"),))

    def test_reversal_scope_resolves_native_material_request_and_persists_no_business_effect(self):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        from erpnext.stock.doctype.material_request.material_request import make_purchase_order
        request = frappe.get_doc({"doctype": "Material Request", "company": COMPANY, "material_request_type": "Purchase",
            "transaction_date": nowdate(), "schedule_date": add_days(nowdate(), 1), "items": [{"item_code": self.item,
                "qty": 2, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]}).insert().submit()
        order = make_purchase_order(request.name)
        order.supplier = "QA Test Supplier"
        order.insert(set_name="QA-ATOMIC-PO-" + uuid.uuid4().hex[:16]).submit()
        result = scope.collect_document_scope(order)
        self.assertEqual(result.sources, (scope.DocumentIdentity("Material Request", request.name),))
        self.assertEqual(result.pairs, (scope.StockPair(self.item, "Stores - QAB"),))
        self.assertEqual(result.vouchers, ())
        frappe.db.set_value("Material Request", request.name, "custom_request_source", "MES")
        with self.assertRaises(frappe.ValidationError): scope.collect_document_scope(order)

    def test_reversal_scope_native_stock_entry_direct_transit_mr_and_old_supply_sources(self):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        po, prior, future = self.future_receipts()
        request = frappe.get_doc({"doctype": "Material Request", "company": COMPANY, "material_request_type": "Material Transfer",
            "transaction_date": nowdate(), "schedule_date": nowdate(), "items": [{"item_code": self.item, "qty": 1,
                "from_warehouse": "Stores - QAB", "warehouse": "Work In Progress - QAB", "schedule_date": nowdate()}]}).insert().submit()
        outgoing = self.scope_stock_entry("Material Transfer", nowdate(), [{"item_code": self.item, "qty": 1,
            "s_warehouse": "Stores - QAB", "t_warehouse": "Goods In Transit - QAB",
            "material_request": request.name, "material_request_item": request.items[0].name}], add_to_transit=1)
        self.assertEqual(self.read_reversal_scope(outgoing, cancellation=False).sources,
            (scope.DocumentIdentity(request.doctype, request.name),))
        incoming = self.scope_stock_entry("Material Transfer", nowdate(), [{"item_code": self.item, "qty": 1,
            "s_warehouse": "Goods In Transit - QAB", "t_warehouse": "Work In Progress - QAB",
            "against_stock_entry": outgoing.name, "ste_detail": outgoing.items[0].name}], outgoing_stock_entry=outgoing.name)
        result = self.read_reversal_scope(incoming, cancellation=False)
        self.assertEqual(set(result.sources), {scope.DocumentIdentity(doc.doctype, doc.name) for doc in (request, outgoing)})
        self.assertIn(scope.StockPair(self.item, "Work In Progress - QAB"), result.pairs)
        frappe.db.set_value("Material Request", request.name, "custom_request_source", "MES")
        with self.assertRaises(frappe.ValidationError): self.read_reversal_scope(incoming, cancellation=False)
        frappe.db.set_value("Material Request", request.name, "custom_request_source", "手动创建")
        finished = self.scope_item()
        order = self.order(items=[{"item_code": finished, "qty": 10, "rate": 10, "warehouse": "Stores - QAB", "schedule_date": nowdate()}])
        # Exact new synthetic old-PO source records; no historical PO/BOM repair.
        frappe.db.set_value("Purchase Order", order.name, {"is_subcontracted": 1, "is_old_subcontracting_flow": 1})
        supplied = frappe.get_doc({"doctype": "Purchase Order Item Supplied", "parent": order.name, "parenttype": "Purchase Order",
            "parentfield": "supplied_items", "rm_item_code": self.item, "main_item_code": finished,
            "docstatus": 1, "stock_uom": "Nos", "reserve_warehouse": "Work In Progress - QAB", "required_qty": 10, "idx": 1})
        self.assertTrue(frappe.get_meta(supplied.doctype).has_field("stock_uom"))
        self.assertFalse(frappe.get_meta(supplied.doctype).has_field("uom"))
        supplied.db_insert()
        sent = self.scope_stock_entry("Send to Subcontractor", nowdate(), [{"item_code": self.item, "qty": 1,
            "s_warehouse": "Stores - QAB", "t_warehouse": "Finished Goods - QAB", "po_detail": supplied.name,
            "subcontracted_item": finished}], purchase_order=order.name, supplier=order.supplier)
        result = self.read_reversal_scope(sent, cancellation=False)
        self.assertEqual(result.sources, (scope.DocumentIdentity(order.doctype, order.name),))
        self.assertIn(scope.StockPair(self.item, "Work In Progress - QAB"), result.pairs)
        frappe.db.set_value(supplied.doctype, supplied.name, "stock_uom", "Kg", update_modified=False)
        with self.assertRaises(frappe.ValidationError): self.read_reversal_scope(sent, cancellation=False)
        frappe.db.set_value(supplied.doctype, supplied.name, "stock_uom", "Nos", update_modified=False)

    def test_reversal_scope_native_invoice_cost_setting_direct_and_fifo_potential_seeds(self):
        from deeplinkerp_branding.services import purchase_reversal_scope as scope
        from erpnext.buying.doctype.purchase_order.purchase_order import make_purchase_invoice as from_order
        po, prior, future = self.future_receipts()
        invoice = make_purchase_invoice(prior.name)
        invoice.items[0].qty = 1
        invoice.insert().submit()
        fifo = from_order(po.name)
        fifo.items[0].qty = 1
        fifo.insert().submit()
        stock = from_order(po.name)
        stock.update_stock = 1; stock.items[0].qty = 1; stock.items[0].warehouse = "Stores - QAB"
        stock.insert().submit()
        self.commit_fixture()
        self.remember_effects()
        for enabled in (0, 1):
            frappe.db.set_single_value("Buying Settings", "set_landed_cost_based_on_purchase_invoice_rate", enabled)
            frappe.clear_document_cache("Buying Settings", "Buying Settings")
            for root in (invoice, fifo, stock):
                with self.subTest(enabled=enabled, update_stock=root.update_stock, direct=bool(root.items[0].pr_detail)):
                    result = self.read_reversal_scope(root)
                    expected = {scope.DocumentIdentity(root.doctype, root.name)}
                    if enabled:
                        expected |= {scope.DocumentIdentity(doc.doctype, doc.name) for doc in (prior, future, stock)}
                        self.assertEqual(result.posting_datetime, str(prior.posting_date) + " " + str(prior.posting_time))
                    self.assertEqual(set(result.vouchers), expected)
        # Real native cancellation invokes the PR force-repost path and the
        # strict floor rejects it atomically; never run its native executor.
        from erpnext.stock.doctype.purchase_receipt.purchase_receipt import PurchaseReceipt
        native = PurchaseReceipt.repost_future_sle_and_gle
        with patch.object(PurchaseReceipt, "repost_future_sle_and_gle", autospec=True, side_effect=native) as repost:
            with self.assertRaises(frappe.ValidationError) as caught: invoice.reload().cancel()
        self.assertEqual(caught.exception.purchase_error_id, "native_stock_repost_incomplete")
        self.assertTrue(any(call.kwargs.get("force") for call in repost.call_args_list))
        self.assert_effects_unchanged()

    def remember_effects(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        self.fixture_effects = {(doctype, name): guard.artifact_evidence(frappe.get_doc(doctype, name))
            for doctype in ("Purchase Order", "Purchase Receipt", "Purchase Invoice", "Payment Entry", "China Accounting Voucher", "Integration Request", "Repost Item Valuation", "Delivery Note", "Sales Order", "Sales Invoice", "Stock Entry")
            for name in self.committed_names[doctype]}

    def assert_effects_unchanged(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        for identity, evidence in self.fixture_effects.items():
            self.assertEqual(guard.artifact_evidence(frappe.get_doc(*identity)), evidence, str(identity))

    def test_pending_native_repost_with_future_sle_refuses_receipt_cancellation(self):
        po, prior, future = self.future_receipts()
        frappe.flags.dont_execute_stock_reposts = True  # native production-equivalent RIV.on_submit branch
        try:
            with self.assertRaises(frappe.ValidationError) as caught:
                prior.cancel()
            self.assertEqual(caught.exception.purchase_error_id, "native_stock_repost_incomplete")
        finally:
            frappe.flags.dont_execute_stock_reposts = False
        self.assertEqual(frappe.db.get_value("Purchase Receipt", prior.name, "docstatus"), 1)
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "actual_qty"), 4)
        self.assertEqual(self.before["Repost Item Valuation"], frappe.db.count("Repost Item Valuation"))
        self.assert_effects_unchanged()

    def test_pending_native_repost_refuses_invoice_rate_cost_adjustment(self):
        po, prior, future = self.future_receipts()
        frappe.db.set_single_value("Buying Settings", "set_landed_cost_based_on_purchase_invoice_rate", 1)
        frappe.db.set_single_value("Buying Settings", "maintain_same_rate", 0)
        frappe.clear_document_cache("Buying Settings", "Buying Settings")
        pi = make_purchase_invoice(prior.name)
        pi.items[0].rate -= 1
        pi.insert()
        frappe.flags.dont_execute_stock_reposts = True
        try:
            with self.assertRaises(frappe.ValidationError) as caught:
                pi.submit()
            self.assertEqual(caught.exception.purchase_error_id, "native_stock_repost_incomplete")
        finally:
            frappe.flags.dont_execute_stock_reposts = False
            frappe.clear_document_cache("Buying Settings", "Buying Settings")
        self.assertFalse(frappe.db.exists("Purchase Invoice", pi.name))
        self.assertEqual(frappe.db.get_value("Purchase Receipt", prior.name, "per_billed"), 0)
        self.assertEqual(self.before["Repost Item Valuation"], frappe.db.count("Repost Item Valuation"))
        self.assert_effects_unchanged()

    def test_existing_item_based_repost_blocks_forward_invoice_without_history_changes(self):
        po, prior, future = self.future_receipts()
        repost = frappe.get_doc({"doctype": "Repost Item Valuation", "based_on": "Item and Warehouse",
            "company": COMPANY, "item_code": self.item, "warehouse": "Stores - QAB",
            "posting_date": prior.posting_date, "posting_time": prior.posting_time})
        repost.flags.dont_run_in_test = True
        repost.insert().submit()  # synthetic historical pending task, outside a procurement operation
        self.assertEqual(repost.status, "Queued")
        for status in ("Queued", "Skipped"):
            with self.subTest(status=status):
                frappe.db.set_value(repost.doctype, repost.name, "status", status)
                self.commit_fixture()
                self.remember_effects()
                pi = make_purchase_invoice(prior.name).insert()
                with self.assertRaises(frappe.ValidationError) as caught:
                    pi.submit()
                self.assertEqual(caught.exception.purchase_error_id, "native_stock_repost_pending")
                self.assertFalse(frappe.db.exists("Purchase Invoice", pi.name))
                self.assertEqual(frappe.db.get_value(repost.doctype, repost.name, "status"), status)
                self.assert_effects_unchanged()

    def test_native_po_advance_ledger_submit_cancel_and_corruption(self):
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        from deeplinkerp_branding.services import purchase_consistency as guard
        from frappe.client import cancel
        po = self.order()
        self.commit_fixture()
        account = frappe.get_doc({"doctype": "Account", "account_name": "QA-ATOMIC-ADVANCE-" + uuid.uuid4().hex[:10],
            "company": COMPANY, "parent_account": "Current Assets - QAB", "account_type": "Payable", "account_currency": "CNY"}).insert()
        frappe.db.set_value("Company", COMPANY, {"book_advance_payments_in_separate_party_account": 1,
            "default_advance_paid_account": account.name})
        frappe.clear_document_cache("Company", COMPANY)
        pe = get_payment_entry("Purchase Order", po.name, bank_account="Cash - QAB", bank_amount=10)
        pe.paid_amount = pe.received_amount = 10
        pe.references[0].allocated_amount = 10
        pe.reference_no, pe.reference_date = "QA-ADVANCE", nowdate()
        pe.insert(); pe.submit()
        guard.check_registered()
        self.assertEqual(frappe.db.get_value("Purchase Order", po.name, "advance_paid"), 10)
        cancel(pe.doctype, pe.name)
        guard.check_registered()
        self.assertEqual(frappe.db.get_value("Purchase Order", po.name, "advance_paid"), 0)
        other = get_payment_entry("Purchase Order", po.name, bank_account="Cash - QAB", bank_amount=10)
        other.paid_amount = other.received_amount = 10
        other.references[0].allocated_amount = 10
        other.reference_no, other.reference_date = "QA-ADVANCE-CORRUPT", nowdate()
        other.insert()
        native = guard.check_document
        def corrupt(doc):
            if doc.doctype == "Payment Entry" and doc.docstatus == 1:
                frappe.db.set_value("Purchase Order", po.name, "advance_paid", 999)
            return native(doc)
        try:
            with patch.object(guard, "check_document", side_effect=corrupt), self.assertRaises(frappe.ValidationError):
                other.submit()
        finally:
            frappe.clear_document_cache("Company", COMPANY)
        self.assertEqual(frappe.db.get_value("Purchase Order", po.name, "advance_paid"), 0)

    def test_mutable_cancellation_preserves_original_and_reverse_payment_ledgers(self):
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        from deeplinkerp_branding.services import purchase_consistency as guard
        from frappe.client import cancel
        for ledger in ("Payment Ledger Entry", "Advance Payment Ledger Entry"):
            for target in ("original", "reverse"):
                with self.subTest(ledger=ledger, target=target):
                    if not frappe.db.exists("Item", self.item):
                        frappe.get_doc({"doctype": "Item", "item_code": self.item, "item_name": self.item,
                            "item_group": "All Item Groups", "stock_uom": "Nos", "is_stock_item": 1}).insert()
                    po = self.order()
                    account = frappe.get_doc({"doctype": "Account", "account_name": "QA-ATOMIC-ADVANCE-" + uuid.uuid4().hex[:10],
                        "company": COMPANY, "parent_account": "Current Assets - QAB", "account_type": "Payable", "account_currency": "CNY"}).insert()
                    frappe.db.set_value("Company", COMPANY, {"book_advance_payments_in_separate_party_account": 1,
                        "default_advance_paid_account": account.name})
                    frappe.clear_document_cache("Company", COMPANY)
                    pe = get_payment_entry("Purchase Order", po.name, bank_account="Cash - QAB", bank_amount=10)
                    pe.paid_amount = pe.received_amount = pe.references[0].allocated_amount = 10
                    pe.reference_no, pe.reference_date = "QA-ADVANCE-INTEGRITY", nowdate()
                    pe.insert(); pe.submit()
                    originals = set(frappe.get_all(ledger, filters={"voucher_type": pe.doctype, "voucher_no": pe.name}, pluck="name"))
                    self.assertTrue(originals)
                    native = guard.check_cancelled_ledgers
                    def corrupt(doc):
                        if (doc.doctype, doc.name) == (pe.doctype, pe.name):
                            rows = frappe.get_all(ledger, filters={"voucher_type": pe.doctype, "voucher_no": pe.name}, pluck="name")
                            selected = next(name for name in rows if (name in originals) == (target == "original"))
                            frappe.db.set_value(ledger, selected, "amount", 999, update_modified=False)
                        return native(doc)
                    try:
                        with patch.object(guard, "check_cancelled_ledgers", side_effect=corrupt), self.assertRaises(frappe.ValidationError):
                            cancel(pe.doctype, pe.name)
                    finally:
                        frappe.db.rollback()
                        frappe.clear_document_cache("Company", COMPANY)
                    self.assertEqual(self.before, {doctype: frappe.db.count(doctype) for doctype in self.types})
    def test_keyed_native_cancel_retransmission_keeps_one_reversal(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        payload, key = json.dumps(pi.as_dict(), default=str), str(uuid.uuid4())
        frappe.response.docs = []
        guard.cancel(pi.doctype, pi.name, request_id=key, doc=payload)
        frappe.response.docs = []
        guard.cancel(pi.doctype, pi.name, request_id=key, doc=payload)
        self.assertEqual(frappe.response.docs[-1]["docstatus"], 2)
        self.assertEqual(frappe.db.count("China Accounting Voucher", {"source_doctype": pi.doctype,
            "source_name": pi.name, "source_event": "Cancellation"}), 1)

    def test_native_savedocs_retransmission_binds_complete_payload(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from frappe.handler import execute_cmd
        from werkzeug.test import EnvironBuilder
        from werkzeug.wrappers import Request
        payload = {"doctype": "Purchase Order", "name": "new-purchase-order-" + uuid.uuid4().hex,
            "__islocal": 1, "company": COMPANY, "supplier": "QA Test Supplier", "currency": "CNY",
            "schedule_date": add_days(nowdate(), 1), "items": [{"doctype": "Purchase Order Item", "item_code": self.item,
                "qty": 2, "rate": 12.345, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]}
        self.assertTrue(callable(getattr(guard, "savedocs", None)), "Native savedocs stable request bridge missing")
        key = str(uuid.uuid4())
        self.assertEqual(frappe.override_whitelisted_method("frappe.desk.form.save.savedocs"),
            "deeplinkerp_branding.services.purchase_consistency.savedocs")
        def endpoint(value):
            with patch.object(frappe, "form_dict", frappe._dict(doc=json.dumps(value), action="Save", request_id=key)), \
                    patch.object(frappe.local, "request", Request(EnvironBuilder(method="POST",
                        path="/api/method/frappe.desk.form.save.savedocs", base_url="http://" + SITE).get_environ()), create=True):
                return execute_cmd("frappe.desk.form.save.savedocs")
        frappe.response.docs = []
        self.assertIsNone(endpoint(payload))
        first = frappe.response.docs[-1]
        frappe.response.docs = []  # discard the first response, as a client network loss would
        self.assertIsNone(endpoint(payload))
        self.assertEqual(frappe.response.docs[-1]["name"], first["name"])
        self.assertEqual(frappe.response.docs[-1]["localname"], payload["name"])
        self.assertEqual(frappe.db.count("Purchase Order Item", {"item_code": self.item}), 1)
        changed = json.loads(json.dumps(payload))
        changed["items"][0]["rate"] += 1
        with self.assertRaises(frappe.ValidationError):
            endpoint(changed)

    def test_actual_system_rounding_policy_is_used_by_quantity_guard(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from frappe.utils import flt
        from frappe.core.doctype.system_settings.system_settings import clear_system_settings_cache
        try:
            frappe.db.set_single_value("System Settings", {"rounding_method": "Banker's Rounding", "float_precision": "3"})
            clear_system_settings_cache()
            frappe.clear_document_cache("System Settings", "System Settings")
            self.assertEqual(frappe.get_system_settings("rounding_method"), "Banker's Rounding")
            # Real client-local RLock holds background Redis invalidations. A
            # document-cache-only fixture can read the previous rounding policy.
            with frappe.client_cache.lock:
                for policy, expected in (("Banker's Rounding", .002), ("Commercial Rounding", .003)):
                    with self.subTest(policy=policy):
                        frappe.db.set_single_value("System Settings", {"rounding_method": policy, "float_precision": "3"})
                        clear_system_settings_cache()
                        frappe.clear_document_cache("System Settings", "System Settings")
                        row = frappe.new_doc("Purchase Order Item")
                        self.assertEqual(row.precision("qty"), 3)
                        self.assertEqual(flt(.0025, 3), expected)
                        guard.equal(row, "qty", .0025, expected, "实际系统舍入")
        finally:
            frappe.db.rollback()
            clear_system_settings_cache()
            frappe.clear_document_cache("System Settings", "System Settings")

    def test_identical_new_native_docs_with_distinct_requests_both_insert(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        self.assertTrue(callable(getattr(guard, "savedocs", None)), "Native savedocs stable request bridge missing")
        payload = {"doctype": "Purchase Order", "name": "new-purchase-order-" + uuid.uuid4().hex,
            "__islocal": 1, "company": COMPANY, "supplier": "QA Test Supplier", "currency": "CNY",
            "schedule_date": add_days(nowdate(), 1), "items": [{"doctype": "Purchase Order Item", "item_code": self.item,
                "qty": 2, "rate": 12.345, "warehouse": "Stores - QAB", "schedule_date": add_days(nowdate(), 1)}]}
        for _ in range(2):
            frappe.response.docs = []
            guard.savedocs(json.dumps(payload), "Save", request_id=str(uuid.uuid4()))
        self.assertEqual(frappe.db.count("Purchase Order Item", {"item_code": self.item}), 2)

    def test_missing_own_payment_ledger_cannot_hide_behind_restored_invoice_balance(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        from deeplinkerp_branding.services import purchase_consistency as guard
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        self.commit_fixture()
        original = guard.check_document

        def lose_payment_ledger(doc):
            if doc.doctype == "Payment Entry" and doc.docstatus == 1:
                frappe.db.delete("Payment Ledger Entry", {"voucher_type": doc.doctype, "voucher_no": doc.name})
                frappe.db.set_value("Purchase Invoice", pi.name, "outstanding_amount", pi.grand_total)
            return original(doc)

        with patch.object(guard, "check_document", side_effect=lose_payment_ledger):
            result = actions.record_payment(source_doctype="Purchase Receipt", source_name=pr.name,
                purchase_invoice=pi.name, amount_to_pay=10, bank_account="Cash - QAB", request_id=str(uuid.uuid4()))
        self.assertTrue(result.get("failed"), "Missing PE's own PLE must roll back even when PI and its remaining PLE agree")
        self.assertEqual(frappe.db.get_value("Purchase Invoice", pi.name, "outstanding_amount"), pi.grand_total)

    def test_native_billed_row_and_percent_corruption_roll_back_invoice(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        for doctype, field in (("Purchase Order Item", "billed_amt"), ("Purchase Receipt Item", "billed_amt"),
                ("Purchase Order", "per_billed"), ("Purchase Receipt", "per_billed")):
            with self.subTest(doctype=doctype, field=field):
                po = self.order()
                pr = self.receipt(po)
                original = guard.check_document
                def corrupt(doc):
                    if doc.doctype == "Purchase Invoice" and doc.docstatus == 1:
                        target = po.items[0].name if doctype == "Purchase Order Item" else pr.items[0].name if doctype == "Purchase Receipt Item" else po.name if doctype == "Purchase Order" else pr.name
                        frappe.db.set_value(doctype, target, field, 999)
                    return original(doc)
                with patch.object(guard, "check_document", side_effect=corrupt):
                    pi = make_purchase_invoice(pr.name).insert()
                    with self.assertRaises(frappe.ValidationError):
                        pi.submit()
                self.assertFalse(frappe.db.exists("Purchase Invoice", pi.name))
                # A failed uncommitted synthetic operation rolls back its fixture too.
                frappe.get_doc({"doctype": "Item", "item_code": self.item, "item_name": self.item,
                    "item_group": "All Item Groups", "stock_uom": "Nos", "is_stock_item": 1}).insert()

    def test_immutable_native_cancellation_reverses_gl_and_finance(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from frappe.client import cancel
        po = self.order(transaction_date=add_days(nowdate(), -1))
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.set_posting_time = 1
        pi.posting_date = add_days(nowdate(), -1)
        pi.save()
        pi.submit()
        self.commit_fixture()
        frappe.db.set_single_value("Accounts Settings", "enable_immutable_ledger", 1)
        frappe.clear_document_cache("Accounts Settings", "Accounts Settings")
        cancel(pi.doctype, pi.name)
        guard.check_registered()
        self.assertEqual(frappe.db.get_value(pi.doctype, pi.name, "docstatus"), 2)
        self.assertTrue(frappe.db.count("GL Entry", {"voucher_type": pi.doctype, "voucher_no": pi.name, "is_cancelled": 0}))
        self.assertFalse(frappe.db.count("GL Entry", {"voucher_type": pi.doctype, "voucher_no": pi.name, "is_cancelled": 1}))
        self.assertTrue(frappe.db.exists("China Accounting Voucher", {"source_doctype": pi.doctype,
            "source_name": pi.name, "source_event": "Cancellation", "docstatus": 1}))
        self.assertEqual(str(frappe.db.get_value("China Accounting Voucher", {"source_doctype": pi.doctype,
            "source_name": pi.name, "source_event": "Cancellation"}, "posting_date")), pi.posting_date)
        dates = {str(row.posting_date) for row in frappe.db.get_values("GL Entry", {"voucher_type": pi.doctype,
            "voucher_no": pi.name}, ["posting_date"], as_dict=True)}
        self.assertEqual(dates, {pi.posting_date, nowdate()}, "Immutable GL reversal date differs from native Finance source date")

    def test_native_controller_failure_rolls_back_audit_and_stock(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        from erpnext.stock.doctype.purchase_receipt.purchase_receipt import PurchaseReceipt
        po = self.order()
        original = PurchaseReceipt.on_submit
        def fail_after_native(doc):
            original(doc)
            raise RuntimeError("QA fault after real stock and GL")
        with patch.object(PurchaseReceipt, "on_submit", fail_after_native):
            with self.assertRaises(RuntimeError):
                self.receipt(po)
        self.assertEqual(self.before, {doctype: frappe.db.count(doctype) for doctype in self.types})

    def test_database_receipt_recovers_without_redis_and_without_second_native_write(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        from deeplinkerp_branding.services import purchase_operation as kernel
        po = self.order()
        key = str(uuid.uuid4())
        native_draft = frappe.db.get_value("Purchase Receipt Item", {"purchase_order": po.name, "docstatus": 0}, "parent")
        args = dict(source_name=po.name, changes={"items": [{"key": po.items[0].name, "qty": 2}]}, request_id=key,
            target_name=native_draft, expected_modified=frappe.db.get_value("Purchase Receipt", native_draft, "modified"))
        first = actions.record_receipt(**args)
        self.assertFalse(first.get("failed"), first)
        audit = frappe.get_doc("Integration Request", kernel.identity("Administrator", key))
        self.assertEqual(audit.status, "Completed")
        self.commit_fixture()  # a new request transaction must recover the committed receipt
        with patch.object(frappe.cache, "lock", side_effect=ConnectionError("QA Redis lock unavailable")):
            second = actions.record_receipt(**args)
        self.assertTrue(second["reused"])
        self.assertEqual(first["document"]["name"], second["document"]["name"])
        self.assertEqual(frappe.db.count("Purchase Receipt", {"name": first["document"]["name"]}), 1)
        self.assertEqual(frappe.db.get_value("Purchase Order Item", po.items[0].name, "received_qty"), 2)

    def test_native_savedocs_deadlock_retry_discards_rolled_back_artifacts(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from deeplinkerp_branding.services import purchase_operation as kernel
        from frappe.desk.form.save import savedocs as native
        po = self.order()
        self.commit_fixture()
        payload = frappe.copy_doc(po).as_dict()
        payload.update(name="new-purchase-order-" + uuid.uuid4().hex, __islocal=1)
        response_docs_before = len(frappe.response.get("docs") or [])
        names = []
        def retry_once(*args, **kwargs):
            native(*args, **kwargs)
            names.append(frappe.response.docs[-1]["name"])
            if len(names) == 1:
                raise frappe.QueryDeadlockError(Exception(1213, "synthetic native deadlock after insert"))
        key = str(uuid.uuid4())
        from erpnext.buying.doctype.purchase_order.purchase_order import PurchaseOrder
        def random_name(doc):
            doc.name = "QA-ATOMIC-PO-" + uuid.uuid4().hex[:16]
        with patch("frappe.desk.form.save.savedocs", side_effect=retry_once), \
             patch.object(PurchaseOrder, "autoname", random_name, create=True):
            guard.savedocs(json.dumps(payload, default=str), "Save", key)
        self.assertEqual(len(names), 2)
        self.assertNotEqual(*names)
        self.assertFalse(frappe.db.exists("Purchase Order", names[0]))
        self.assertEqual([row["name"] for row in frappe.response.docs[response_docs_before:]], [names[1]])
        audit = frappe.get_doc("Integration Request", kernel.identity("Administrator", key))
        artifacts = json.loads(audit.output)["artifacts"]
        self.assertEqual({row["name"] for row in artifacts}, {names[1]})
        self.assertNotIn(names[0], audit.data)

    def test_payment_receipt_records_invoice_effect_and_refuses_changed_source_replay(self):
        from deeplinkerp_branding.services import purchase_document_actions as actions
        from deeplinkerp_branding.services import purchase_operation as kernel
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        self.commit_fixture()
        key = str(uuid.uuid4())
        args = dict(source_doctype="Purchase Receipt", source_name=pr.name, purchase_invoice=pi.name,
            amount_to_pay=10, bank_account="Cash - QAB", request_id=key)
        paid = actions.record_payment(**args)
        self.assertFalse(paid.get("failed"), paid)
        audit = frappe.get_doc("Integration Request", kernel.identity("Administrator", key))
        artifacts = json.loads(audit.output)["artifacts"]
        invoice_effects = [row for row in artifacts if (row["doctype"], row["name"]) == (pi.doctype, pi.name)]
        self.assertEqual(len(invoice_effects), 1)
        self.assertEqual(invoice_effects[0]["permission"], "read")
        replay = actions.record_payment(**args)
        self.assertTrue(replay["reused"])
        outstanding = frappe.db.get_value(pi.doctype, pi.name, "outstanding_amount")
        frappe.db.set_value(pi.doctype, pi.name, "outstanding_amount", outstanding + 1, update_modified=False)
        with self.assertRaises(frappe.ValidationError) as caught:
            actions.record_payment(**args)
        self.assertEqual(caught.exception.purchase_error_id, "replay_evidence_changed")

    def test_direct_payment_draft_replay_rechecks_acknowledged_evidence_and_write_acl(self):
        from deeplinkerp_branding.services import purchase_consistency as guard
        from deeplinkerp_branding.services import purchase_payment_service as service
        from deeplinkerp_branding.services import purchase_document_actions as actions
        from deeplinkerp_branding.services import purchase_operation as kernel
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        self.commit_fixture()
        for change in ({"remarks": "QA legitimate changed draft"}, {"amount": 11}):
            with self.subTest(change=change):
                args = dict(source_doctype="Purchase Receipt", source_name=pr.name, purchase_invoice=pi.name,
                    amount_to_pay=10, bank_account="Cash - QAB", request_id=str(uuid.uuid4()))
                original = service.create_payment_draft(**args)
                self.commit_fixture()
                receipt = json.loads(frappe.get_doc("Integration Request",
                    kernel.identity("Administrator", args["request_id"])).output)
                self.assertTrue(any(row["name"] == original["name"] for row in receipt["artifacts"]))
                self.assertEqual(service.create_payment_draft(**args), {**original, "reused": True})
                from frappe import permissions
                native_permission = permissions.has_permission
                def deny_write(doctype, ptype="read", *args, **kwargs):
                    if doctype == "Payment Entry" and ptype == "write":
                        return False
                    return native_permission(doctype, ptype, *args, **kwargs)
                with patch.object(permissions, "has_permission", side_effect=deny_write), self.assertRaises(frappe.PermissionError):
                    service.create_payment_draft(**args)
                current = actions.preview_payment(original["name"])["document"]
                actions.update_payment_draft(original["name"], change, current["modified"])
                self.commit_fixture()
                before = guard.artifact_evidence(frappe.get_doc("Payment Entry", original["name"]))
                counts = {doctype: frappe.db.count(doctype) for doctype in self.types}
                with self.assertRaises(frappe.ValidationError) as caught:
                    service.create_payment_draft(**args)
                self.assertEqual(caught.exception.purchase_error_id, "replay_evidence_changed")
                self.assertEqual(before, guard.artifact_evidence(frappe.get_doc("Payment Entry", original["name"])))
                self.assertEqual(counts, {doctype: frappe.db.count(doctype) for doctype in self.types})

    def test_existing_finance_posting_cannot_be_stranded_when_settings_turn_inactive(self):
        from china_finance.services import voucher
        from frappe.client import cancel
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        self.commit_fixture()
        self.remember_effects()
        settings = voucher.get_company_settings(COMPANY)
        frappe.db.set_value(settings.doctype, settings.name, "enabled", 0)
        frappe.clear_document_cache(settings.doctype, settings.name)
        with self.assertRaises(frappe.ValidationError) as caught:
            cancel(pi.doctype, pi.name)
        self.assertEqual(caught.exception.purchase_error_id, "finance_cancellation_settings_inactive")
        self.assert_effects_unchanged()

    def test_posting_snapshot_header_must_match_native_finance_effect(self):
        from china_finance.services import voucher
        from deeplinkerp_branding.services import purchase_consistency as guard
        po = self.order()
        self.commit_fixture()
        self.remember_effects()
        native = voucher.create_voucher_from_source
        for altered in ({"currency": "USD"}, {"status": "Reversed"}, {"reversal_of": "QA-ATOMIC-INVALID"},
                {"source_key": "QA-ATOMIC-INVALID"}, {"total_debit": 999, "total_credit": 999},
                {"posting_date": add_days(nowdate(), -1)}, {"fiscal_year": "1900"},
                {"accounting_period": "1900-01"}, {"reversed_by": "QA-ATOMIC-INVALID"}):
            with self.subTest(fields=tuple(altered)):
                pr = make_purchase_receipt(po.name).insert()
                def corrupt(source, event):
                    name = native(source, event)
                    if name and (source.doctype, source.name, event) == (pr.doctype, pr.name, "Posting"):
                        frappe.db.set_value("China Accounting Voucher", name, altered, update_modified=False)
                    return name
                try:
                    with patch.object(guard, "finance_service", return_value=frappe._dict(create_voucher_from_source=corrupt)), \
                            self.assertRaises(frappe.ValidationError):
                        pr.submit()
                finally:
                    frappe.db.rollback()
                self.assertFalse(frappe.db.exists("Purchase Receipt", pr.name))
                self.assert_effects_unchanged()

    def test_china_voucher_failure_after_real_stock_and_gl_rolls_back_everything(self):
        from china_finance.services import voucher
        po = self.order()
        with patch.object(voucher, "create_voucher_from_source", side_effect=RuntimeError("QA finance fault after native stock/GL")):
            with self.assertRaises(RuntimeError):
                self.receipt(po)
        self.assertEqual(self.before, {doctype: frappe.db.count(doctype) for doctype in self.types})

    def test_pending_or_exception_cancellation_restores_committed_native_state(self):
        from china_finance.services import voucher
        from frappe.client import cancel
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        self.commit_fixture()
        stock = frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "actual_qty")
        original = frappe.db.get_value("China Accounting Voucher", {"source_doctype": pi.doctype,
            "source_name": pi.name, "source_event": "Posting"}, "name")
        for outcome in ({"status": "pending", "error": "private QA detail"}, RuntimeError("private QA detail")):
            with self.subTest(outcome=type(outcome).__name__), patch.object(voucher, "process_cancellation_snapshot",
                **({"side_effect": outcome} if isinstance(outcome, Exception) else {"return_value": outcome})) as process:
                with self.assertRaises((frappe.ValidationError, RuntimeError)):
                    cancel(pi.doctype, pi.name)
                self.assertTrue(process.called, "The Finance failure must be reached, not masked by the stock-repost guard")
            self.assertEqual(frappe.db.get_value(pi.doctype, pi.name, "docstatus"), 1)
            self.assertEqual(frappe.db.get_value("Purchase Receipt", pr.name, "docstatus"), 1)
            self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "actual_qty"), stock)
            self.assertFalse(frappe.db.get_value("China Accounting Voucher", original, "reversed_by"))
            self.assertFalse(frappe.db.exists("China Accounting Voucher", {"source_doctype": pi.doctype,
                "source_name": pi.name, "source_event": "Cancellation"}))

    def test_balanced_wrong_reversal_restores_committed_stock_and_posting(self):
        from china_finance.services import voucher
        from frappe.client import cancel
        po = self.order()
        pr = self.receipt(po)
        pi = make_purchase_invoice(pr.name).insert()
        pi.submit()
        self.commit_fixture()
        process = voucher.process_cancellation_snapshot

        def corrupt_reversal(*args, **kwargs):
            result = process(*args, **kwargs)
            self.assertEqual(result["status"], "resolved")
            reversal = frappe.get_doc("China Accounting Voucher", result["voucher"])
            for row in reversal.entries:
                frappe.db.set_value(row.doctype, row.name, {field: (row.get(field) or 0) * 2 for field in
                    ("debit", "credit", "debit_in_account_currency", "credit_in_account_currency")})
            frappe.db.set_value(reversal.doctype, reversal.name,
                {"total_debit": reversal.total_debit * 2, "total_credit": reversal.total_credit * 2})
            return result

        with patch.object(voucher, "process_cancellation_snapshot", side_effect=corrupt_reversal):
            with self.assertRaisesRegex(frappe.ValidationError, "反向金额"):
                cancel(pi.doctype, pi.name)
        self.assertEqual(frappe.db.get_value(pi.doctype, pi.name, "docstatus"), 1)
        self.assertEqual(frappe.db.get_value("Purchase Receipt", pr.name, "docstatus"), 1)
        self.assertEqual(frappe.db.get_value("Bin", {"item_code": self.item, "warehouse": "Stores - QAB"}, "actual_qty"), 2)
        self.assertFalse(frappe.db.exists("China Accounting Voucher", {"source_doctype": pi.doctype,
            "source_name": pi.name, "source_event": "Cancellation"}))

    def test_audit_retention_and_normal_delete_protection(self):
        from deeplinkerp_branding.services import purchase_operation as kernel
        self.order()
        name = frappe.db.get_value("Integration Request", {"integration_request_service": kernel.SERVICE}, "name")
        audit = frappe.get_doc("Integration Request", name)
        with patch.object(frappe.db, "delete") as delete:
            type(audit).clear_old_logs(30)
        predicate = str(delete.call_args.kwargs["filters"])
        self.assertIn(kernel.SERVICE, predicate)
        self.assertIn("integration_request_service", predicate)
        with self.assertRaises(frappe.PermissionError):
            audit.delete()


def boundary_peer(payload):
    """One native peer in the SAME guarded synthetic QA harness, no executor."""
    from deeplinkerp_branding.services import purchase_repost_boundary as boundary
    if frappe.local.site != SITE or frappe.conf.db_name != "qa_procurement_5" or frappe.conf.db_host != "db":
        raise RuntimeError("Isolated synthetic procurement database only")
    if not str(payload.get("item") or "").startswith("QA-ATOMIC-"):
        raise RuntimeError("Exact new synthetic item only")
    frappe.in_test = False
    frappe.flags.in_test = True
    frappe.flags.mute_emails = True
    if not payload["action"].startswith("entry-"):
        boundary.initialize()
    frappe.db.after_commit.reset()  # before any peer document acquires leases
    # Reuse the same harness's cleanup inventory without executing setUp or
    # creating another cancellation fixture/test class.
    import ast
    import inspect
    import textwrap
    method = ast.parse(textwrap.dedent(inspect.getsource(NativeAtomicPurchaseTests.setUp))).body[0]
    types = ()
    for node in method.body:
        target = node.targets[0] if isinstance(node, ast.Assign) else node.target if isinstance(node, ast.AugAssign) else None
        if isinstance(target, ast.Attribute) and target.attr == "types":
            value = ast.literal_eval(node.value)
            types = types + value if isinstance(node, ast.AugAssign) else value
    before = {doctype: set(frappe.get_all(doctype, pluck="name", limit_page_length=0)) for doctype in types}
    action = payload["action"]
    if action.startswith("entry-"):
        from deeplinkerp_branding.services import purchase_document_actions as actions, purchase_source_service as sources
        queries = []
        native_sql = frappe.local.db.sql
        @wraps(native_sql)  # observer must retain the real keyword-only SQL contract
        def observe(query, *args, **kwargs):
            queries.append(str(query).upper())
            return native_sql(query, *args, **kwargs)
        frappe.local.db.sql = observe
        physical = frappe.local.db._conn
        if action == "entry-payment-update":
            actions.update_payment_draft(payload["name"], {"remarks": "QA-ATOMIC-EARLY"}, payload["modified"])
        elif action == "entry-payment-submit":
            actions.submit_document("Payment Entry", payload["name"], payload["modified"])
        elif action == "entry-source-create":
            try:
                sources.create_purchase_order_from_source("QA-ATOMIC-SOURCE-MISSING-" + uuid.uuid4().hex,
                    "", COMPANY, "QA Test Supplier", "CNY", add_days(nowdate(), 1), [])
            except frappe.DoesNotExistError:
                pass  # exact nonexistent source: still exercises real earliest locked public entry
            else:
                raise AssertionError("Missing source unexpectedly accepted")
        else:
            raise RuntimeError("Unknown public entry")
        state = boundary.initialize()
        result = {"identity_query": next(i for i, query in enumerate(queries) if "CONNECTION_ID()" in query),
            "first_lock": next(i for i, query in enumerate(queries) if "FOR UPDATE" in query),
            "same_physical": state.connection is physical, "execution_id": state.execution_id}
    elif action == "defaults":
        entry_type = frappe.get_doc("Stock Entry Type", payload["entry_type"])
        if not entry_type.name.startswith("QA-ATOMIC-TYPE-"):
            raise RuntimeError("Exact new synthetic Stock Entry Type only")
        # Purpose is native set_only_once. This positive peer changes only the
        # legally mutable Item stock_uom, through its real save/cache hooks.
        item = frappe.get_doc("Item", payload["item"])
        item.stock_uom = "Kg"
        item.save()
        result = {"updated": True}
    elif action == "source":
        receipt = frappe.get_doc("Purchase Receipt", payload["name"])
        receipt.items[0].purchase_order = payload["source"]
        receipt.items[0].purchase_order_item = payload["detail"]
        receipt.items[0].warehouse = "Work In Progress - QAB"
        receipt.save()
        result = {"name": receipt.name, "detail": receipt.items[0].name}
    elif action in ("transfer", "draft"):
        draft = frappe.get_doc({"doctype": "Stock Entry", "company": COMPANY,
            "stock_entry_type": payload.get("entry_type") or "Material Transfer",
            "set_posting_time": 1, "posting_date": add_days(nowdate(), 2), "posting_time": "12:00:00",
            "items": [{"item_code": payload["item"], "qty": 1, "s_warehouse": "Stores - QAB",
                "t_warehouse": "Work In Progress - QAB"}]}).insert()
        if action == "transfer":
            draft.submit()
        result = {"name": draft.name, "detail": draft.items[0].name, "purpose": draft.purpose,
            "stock_uom": draft.items[0].stock_uom}
    else:
        raise RuntimeError("Unknown synthetic peer action")
    frappe.db.commit()
    result["created"] = {doctype: sorted(set(frappe.get_all(doctype, pluck="name", limit_page_length=0)) - before[doctype])
        for doctype in types}
    print(json.dumps(result))


if __name__ == "__main__":
    frappe.init(site=SITE, sites_path="/home/frappe/frappe-bench/sites")
    if frappe.local.site != SITE or frappe.conf.db_name != "qa_procurement_5" or frappe.conf.db_host != "db":
        raise RuntimeError("Isolated synthetic procurement database only, before connection")
    frappe.connect()
    try:
        # The legacy single-transaction harness predates whole-request rollback.
        # Run explicit native transaction scenarios on the guarded isolated clone.
        if len(sys.argv) == 3 and sys.argv[1] == "--boundary-peer":
            boundary_peer(json.loads(sys.argv[2]))
        else:
            result = unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(NativeAtomicPurchaseTests))
            raise SystemExit(not result.wasSuccessful())
    finally:
        frappe.destroy()
