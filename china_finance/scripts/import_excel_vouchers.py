"""Import the reviewed May/June 2026 Yuewei accounting voucher workbook.

The workbook is an accounting voucher register rather than an ERPNext import
template.  This module parses and validates it first, then creates submitted
Journal Entries so the China Finance hooks create the immutable China
Accounting Voucher snapshots and monthly statutory numbers.
"""

from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import frappe
from frappe import _
from frappe.utils import getdate

try:
	from openpyxl import load_workbook
except ImportError:  # pragma: no cover - the dependency is provided by Frappe
	load_workbook = None


COMPANY = "悦为智能技术(东莞)有限公司"
VOUCHER_SHEET = "26年凭证"
EXPECTED_PERIODS = {"2026-05", "2026-06"}
ACCOUNT_NUMBERS = {
	"银行存款": "100201",
	"其他应收款-备用金": "1221",
	"实收资本": "4001",
	"本年利润": "4103",
	"营业外收入": "6301",
	"管理费用-工资": "660209",
	"管理费用-办公费": "660201",
	"管理费用-房租": "660202",
	"管理费用-物业费": "660203",
	"管理费用-水电费": "660204",
	"管理费用-社保": "660208",
	"管理费用-公积金": "660229",
	"财务费用-利息收入": "660302",
	"财务费用-手续费": "660303",
	"应付职工薪酬-工资": "221101",
	"应付职工薪酬-社保": "221103",
	"应付职工薪酬-公积金": "221104",
	"应交税费-应交个人所得税": "222112",
	"预收账款-东莞市奥领数控": "2203",
	"其他应付款-社保": "224101",
	"其他应付款-公积金": "224102",
}


def _text(value):
	return str(value).strip() if value is not None else ""


def _decimal(value, row_number, fieldname):
	if value in (None, ""):
		return Decimal("0.00")
	try:
		return Decimal(str(value)).quantize(Decimal("0.01"))
	except (InvalidOperation, ValueError) as exc:
		frappe.throw(_("Excel 第 {0} 行的{1}不是有效金额：{2}").format(row_number, fieldname, value))
		raise exc


def _posting_date(value, row_number):
	if isinstance(value, datetime):
		return value.date()
	if isinstance(value, date):
		return value
	try:
		return getdate(value)
	except Exception:
		frappe.throw(_("Excel 第 {0} 行的日期无效：{1}").format(row_number, value))


def parse_workbook(file_path):
	"""Parse and validate the voucher sheet without writing anything."""
	if load_workbook is None:
		frappe.throw("当前环境未安装 openpyxl，无法读取 Excel")

	path = Path(file_path)
	if not path.is_file():
		frappe.throw(_("找不到 Excel 文件：{0}").format(file_path))

	workbook = load_workbook(path, data_only=True, read_only=True)
	if VOUCHER_SHEET not in workbook.sheetnames:
		frappe.throw(_("Excel 缺少工作表：{0}").format(VOUCHER_SHEET))

	worksheet = workbook[VOUCHER_SHEET]
	rows = worksheet.iter_rows(values_only=True)
	headers = [_text(value) for value in next(rows, ())]
	expected_headers = ["日期", "凭证号", "摘要", "借方科目", "贷方科目", "借方金额", "贷方金额"]
	if headers[: len(expected_headers)] != expected_headers:
		frappe.throw(_("凭证工作表表头不符合预期：{0}").format(headers))

	vouchers = []
	current = None
	active_summary = ""
	for row_number, raw_row in enumerate(rows, start=2):
		row = list(raw_row[:7]) + [None] * max(0, 7 - len(raw_row))
		posting_date_value, voucher_no_value, summary_value, debit_label, credit_label, debit_value, credit_value = row[:7]
		if not any(value not in (None, "") for value in row):
			continue

		if posting_date_value not in (None, "") or voucher_no_value not in (None, ""):
			if posting_date_value in (None, "") or voucher_no_value in (None, ""):
				frappe.throw(_("Excel 第 {0} 行的日期和凭证号必须同时存在").format(row_number))
			try:
				voucher_no = int(voucher_no_value)
			except (TypeError, ValueError):
				frappe.throw(_("Excel 第 {0} 行的凭证号无效：{1}").format(row_number, voucher_no_value))
			current = {
				"posting_date": _posting_date(posting_date_value, row_number),
				"original_number": voucher_no,
				"lines": [],
				"source_rows": [],
			}
			vouchers.append(current)
			active_summary = ""

		if current is None:
			frappe.throw(_("Excel 第 {0} 行在任何凭证之前出现了分录").format(row_number))

		line_summary = _text(summary_value) or active_summary
		if summary_value not in (None, ""):
			active_summary = _text(summary_value)
		current["source_rows"].append(row_number)

		if not debit_label and not credit_label:
			continue
		if debit_label:
			debit_label = _text(debit_label)
			current["lines"].append(
				{
					"label": debit_label,
					"debit": _decimal(debit_value, row_number, "借方金额"),
					"credit": Decimal("0.00"),
					"summary": line_summary,
					"source_row": row_number,
				}
			)
		if credit_label:
			credit_label = _text(credit_label)
			# The workbook intentionally records interest income as a negative
			# debit on the credit-side account. Keep that representation so the
			# Journal Entry/GL logic can perform the requested offset.
			if credit_value in (None, "") and debit_value not in (None, ""):
				negative_debit = _decimal(debit_value, row_number, "借方金额")
				if negative_debit >= 0:
					frappe.throw(_("Excel 第 {0} 行贷方科目使用借方金额时必须为负数").format(row_number))
				current["lines"].append(
					{
						"label": credit_label,
						"debit": negative_debit,
						"credit": Decimal("0.00"),
						"summary": line_summary,
						"source_row": row_number,
					}
				)
				continue
			current["lines"].append(
				{
					"label": credit_label,
					"debit": Decimal("0.00"),
					"credit": _decimal(credit_value, row_number, "贷方金额"),
					"summary": line_summary,
					"source_row": row_number,
				}
			)

	for voucher in vouchers:
		if not voucher["lines"]:
			frappe.throw(_("原凭证 {0} 没有有效分录").format(voucher["original_number"]))
		debit = sum((line["debit"] for line in voucher["lines"]), Decimal("0.00"))
		credit = sum((line["credit"] for line in voucher["lines"]), Decimal("0.00"))
		if abs(debit - credit) > Decimal("0.005"):
			frappe.throw(
				_("原凭证 {0} 借贷不平：借方 {1}，贷方 {2}").format(
					voucher["original_number"], debit, credit
				)
			)
		voucher["total_debit"] = debit
		voucher["total_credit"] = credit
		voucher["period"] = voucher["posting_date"].strftime("%Y-%m")
		if voucher["period"] not in EXPECTED_PERIODS:
			frappe.throw(_("发现不在 2026 年 5、6 月范围内的凭证：{0}").format(voucher["posting_date"]))

	by_period = defaultdict(list)
	for voucher in vouchers:
		by_period[voucher["period"]].append(voucher)
	for period, period_vouchers in by_period.items():
		numbers = [voucher["original_number"] for voucher in period_vouchers]
		if numbers != list(range(1, len(numbers) + 1)):
			frappe.throw(_("{0} 的凭证号不是从 1 连续排列：{1}").format(period, numbers))
		ordered_keys = [(voucher["posting_date"], voucher["original_number"]) for voucher in period_vouchers]
		if ordered_keys != sorted(ordered_keys):
			frappe.throw(_("{0} 的凭证日期/字号不是升序排列：{1}").format(period, ordered_keys))

	return vouchers


def _resolve_accounts(company, vouchers):
	labels = sorted({line["label"] for voucher in vouchers for line in voucher["lines"]})
	unknown = sorted(set(labels) - set(ACCOUNT_NUMBERS))
	if unknown:
		frappe.throw(_("Excel 中存在未配置的科目：{0}").format("、".join(unknown)))

	resolved = {}
	for label in labels:
		number = ACCOUNT_NUMBERS[label]
		rows = frappe.get_all(
			"Account",
			filters={"company": company, "account_number": number},
			fields=["name", "account_name", "is_group", "disabled", "account_type"],
		)
		if len(rows) != 1:
			frappe.throw(_("科目 {0} 应唯一匹配 {1}，实际匹配 {2} 条").format(label, number, len(rows)))
		account = rows[0]
		if account.is_group:
			frappe.throw(_("科目 {0} 匹配到了分组科目 {1}，不能直接记账").format(label, account.name))
		if account.disabled:
			frappe.throw(_("科目 {0} 已停用：{1}").format(label, account.name))
		resolved[label] = account
	return resolved


def preview(file_path, company=COMPANY):
	"""Return the validated import summary without changing the site."""
	vouchers = parse_workbook(file_path)
	accounts = _resolve_accounts(company, vouchers)
	by_period = defaultdict(list)
	for voucher in vouchers:
		by_period[voucher["period"]].append(voucher)
	return {
		"status": "validated",
		"company": company,
		"file": str(file_path),
		"voucher_count": len(vouchers),
		"periods": {
			period: {
				"count": len(period_vouchers),
				"debit": str(sum((v["total_debit"] for v in period_vouchers), Decimal("0.00"))),
				"credit": str(sum((v["total_credit"] for v in period_vouchers), Decimal("0.00"))),
			}
			for period, period_vouchers in sorted(by_period.items())
		},
		"account_mapping": {
			label: {"number": ACCOUNT_NUMBERS[label], "name": account.name, "account_name": account.account_name}
			for label, account in sorted(accounts.items())
		},
	}


def import_from_file(file_path, company=COMPANY):
	"""Validate and import the workbook into the requested company."""
	if company != COMPANY:
		frappe.throw("此 Excel 导入脚本仅允许处理 aaa 的悦为智能技术(东莞)有限公司")

	vouchers = parse_workbook(file_path)
	accounts = _resolve_accounts(company, vouchers)
	period_dates = [voucher["posting_date"] for voucher in vouchers]
	existing = frappe.get_all(
		"Journal Entry",
		filters={"company": company, "posting_date": ["between", [min(period_dates), max(period_dates)]]},
		fields=["name", "posting_date", "docstatus"],
		limit=1,
	)
	if existing:
		frappe.throw(
			_("aaa 在导入日期范围内已有 Journal Entry（{0}），为避免重复导入已停止").format(existing[0].name)
		)

	created = []
	try:
		for voucher in vouchers:
			period = voucher["period"]
			original_number = voucher["original_number"]
			tag = f"Excel导入：悦为智能(1).xlsx；{period}；原凭证号：{original_number}"
			voucher_summary = next(
				(line["summary"] for line in voucher["lines"] if line["summary"]), ""
			)
			journal_entry = frappe.new_doc("Journal Entry")
			journal_entry.posting_date = voucher["posting_date"]
			journal_entry.company = company
			journal_entry.voucher_type = "Journal Entry"
			journal_entry.custom_remark = 1
			# Keep the business summary in the visible remark fields. The import
			# marker is stored in Title after submit so this batch can still be
			# removed precisely without polluting the voucher summary shown to users;
			# ERPNext derives Title while inserting a new Journal Entry.
			journal_entry.remark = voucher_summary
			journal_entry.user_remark = voucher_summary
			needs_party_bypass = False
			for line in voucher["lines"]:
				account = accounts[line["label"]]
				if account.account_type in {"Receivable", "Payable"}:
					needs_party_bypass = True
				journal_entry.append(
					"accounts",
					{
						"account": account.name,
						"debit_in_account_currency": float(line["debit"]),
						"credit_in_account_currency": float(line["credit"]),
						"user_remark": line["summary"] or tag,
					},
				)
			if needs_party_bypass:
				journal_entry.party_not_required = 1
			journal_entry.insert(ignore_permissions=True)
			journal_entry.submit()
			frappe.db.set_value("Journal Entry", journal_entry.name, "title", tag, update_modified=False)
			created.append(journal_entry.name)
		frappe.db.commit()
	except Exception:
		frappe.db.rollback()
		raise

	return {
		"status": "imported",
		"company": company,
		"file": str(file_path),
		"journal_entries": created,
		"voucher_count": len(created),
	}


def remove_imported_vouchers(file_name="悦为智能(1).xlsx", company=COMPANY, apply=False):
	"""Preview or remove only vouchers imported from the named workbook."""
	if company != COMPANY:
		frappe.throw("此清理脚本仅允许处理 aaa 的悦为智能技术(东莞)有限公司")

	tag_pattern = f"Excel导入：{file_name}；%"
	sources = frappe.db.sql(
		"""
		SELECT name, docstatus, posting_date
		FROM `tabJournal Entry`
		WHERE company=%s
			AND posting_date BETWEEN %s AND %s
			AND (user_remark LIKE %s OR title LIKE %s)
		ORDER BY posting_date DESC, creation DESC, name DESC
		""",
		(company, "2026-05-01", "2026-06-30", tag_pattern, tag_pattern),
		as_dict=True,
	)
	source_names = [row.name for row in sources]
	snapshots = (
		frappe.get_all(
			"China Accounting Voucher",
			filters={"source_doctype": "Journal Entry", "source_name": ["in", source_names]},
			fields=["name", "docstatus", "source_name"],
		)
		if source_names
		else []
	)
	assignments = (
		frappe.get_all(
			"China Cash Flow Assignment",
			filters={"china_accounting_voucher": ["in", [row.name for row in snapshots]]},
			pluck="name",
		)
		if snapshots
		else []
	)
	return_data = {
		"company": company,
		"source_journal_entries": [row.name for row in sources],
		"source_count": len(sources),
		"snapshot_count": len(snapshots),
		"assignment_count": len(assignments),
	}
	if not apply:
		return return_data
	if not source_names:
		return {**return_data, "status": "nothing_to_remove"}

	try:
		# Cancel submitted source entries first so ERPNext creates the standard
		# cancelling GL rows. They are removed below because this is test data.
		for row in sources:
			if row.docstatus == 1:
				doc = frappe.get_doc("Journal Entry", row.name)
				doc.flags.ignore_permissions = True
				doc.cancel()

		# Remove audit snapshots and their auxiliary assignments. Snapshot cancel
		# hooks may enqueue deferred jobs, but those jobs see deleted sources and
		# safely skip; their sync-issue rows are removed below as well.
		for row in snapshots:
			if row.docstatus == 1:
				doc = frappe.get_doc("China Accounting Voucher", row.name)
				doc.flags.ignore_permissions = True
				doc.cancel()
		for assignment_name in assignments:
			frappe.delete_doc(
				"China Cash Flow Assignment",
				assignment_name,
				force=True,
				ignore_permissions=True,
			)
		for row in snapshots:
			frappe.delete_doc(
				"China Accounting Voucher",
				row.name,
				force=True,
				ignore_permissions=True,
			)

		# GL and payment-ledger rows are linked records, not Journal Entry child
		# rows, so remove the posting and cancellation rows explicitly.
		for doctype in ("GL Entry", "Payment Ledger Entry"):
			if frappe.db.exists("DocType", doctype):
				frappe.db.delete(doctype, {"voucher_type": "Journal Entry", "voucher_no": ["in", source_names]})

		for row in sources:
			frappe.delete_doc(
				"Journal Entry",
				row.name,
				force=True,
				ignore_permissions=True,
			)

		# Remove sync issues created by cancellation hooks, including any issue
		# created while cancelling the audit snapshots themselves.
		if frappe.db.exists("DocType", "China Voucher Sync Issue"):
			frappe.db.delete(
				"China Voucher Sync Issue",
				{"source_doctype": "Journal Entry", "source_name": ["in", source_names]},
			)
			if snapshots:
				frappe.db.delete(
					"China Voucher Sync Issue",
					{
						"source_doctype": "China Accounting Voucher",
						"source_name": ["in", [row.name for row in snapshots]],
					},
				)

		# Reset only the imported months. Other months' numbering state is not
		# touched.
		for period in ("2026-05", "2026-06"):
			sequence_name = frappe.db.get_value(
				"China Voucher Sequence",
				{
					"company": company,
					"fiscal_year": "2026",
					"accounting_period": period,
					"voucher_word": "记",
				},
				"name",
			)
			if sequence_name:
				frappe.db.set_value("China Voucher Sequence", sequence_name, "current_value", 0, update_modified=False)

		frappe.db.commit()
	except Exception:
		frappe.db.rollback()
		raise

	return {**return_data, "status": "removed"}
